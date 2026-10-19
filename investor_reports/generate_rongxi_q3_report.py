# -*- coding: utf-8 -*-
"""荣熙共赢 2026年第三季度投资报告。致投资人写入共赢FOF三季度回顾，业绩页与追风季报同一套周净值口径。"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
matplotlib.rcParams["pdf.fonttype"] = 42
matplotlib.rcParams["ps.fonttype"] = 42
matplotlib.rcParams["axes.unicode_minus"] = False

import matplotlib.pyplot as plt
import pandas as pd
import psycopg2
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Rectangle
from PIL import Image

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from generate_q3_investor_report import (
    C_BG,
    C_BORDER,
    C_MUTED,
    C_NAVY,
    C_NOTE_BG,
    C_TEXT,
    DISCLAIMER,
    DPI,
    MUTED,
    NAVY,
    PAGE_IN,
    SLATE,
    TEXT,
    WHITE,
    add_bottom_border,
    add_page_field,
    add_text,
    draw_wrapped,
    load_fonts,
    section_bar,
    wrap_text,
    set_cell_margins,
    set_cell_width,
    set_run_font,
    set_table_fixed,
    shade_cell,
    stitch_long,
)
from q3_strategy_history import load_strategy_history, render_history_page

ENV_FILE = Path(r"d:\coding\market_website\.env.local")
NAV_FILE = Path(r"d:\coding\market_dashboard_storage\custom_funds\nav\380004.json")
BEIAN = "SBNX55"
CUTOFF = "2026-09-30"
PERIOD_BEGIN = "2026-07-01"
INCEPTION = "2026-01-20"
PRODUCT = "荣熙共赢私募证券投资基金"
SHORT = "荣熙共赢"
BRAND = "上海荣熙私募基金"
MANAGER = "上海荣熙私募基金管理有限公司"
ADVISOR = "陈家本"
CUSTODIAN = "招商证券股份有限公司"

HIGHLIGHT = {"区间收益", "最大回撤", "相对沪深300"}
FOF_ASSET = Path(__file__).resolve().parent / "assets" / "rongxi_2026q3_fof"

LETTER_LEAD = "共赢FOF三季度的配置整体变化不大，核心持仓为50%的打板策略，叠加宏观交易与主观期货策略各15%，剩余20%为卫星持仓。"
REVIEW_STOCK = "股票市场，7月份科技抱团瓦解，市场大幅缩量下跌，同期沪深300回调近8%；8月市场在短暂弱反弹后继续向下，9月大盘最低下探到3800点附近，市场情绪极为悲观。打板策略在7月份市场急跌时相对抗跌，但在8-9月份震荡下跌过程中持续磨损；而多数打板策略在上半年高流动性环境中又超募了不少规模，导致了此次超预期的回撤。从市场成交量来看，今年6月22日创下3.74万亿的天量后Q3持续缩量，短线打板策略缺少承接资金，热点轮动快缺乏持续性，不利于打板策略表现。"
REVIEW_FUTURES = "期货市场，三季度商品期货市场整体走强，南华商品指数季末收于186.57点，较6月末上涨8%；其中能化受美伊战争影响持续领涨，贵金属与有色震荡上涨，黑色板块焦煤上演逼空行情，农产品涨跌不一。三季度行情非常利于量化CTA策略表现，主观基本面交易难度极大，组合中期货策略整体是正向贡献，宏观交易策略回撤明显。"
REVIEW_PORTFOLIO = "从组合配置上来看，打板策略只依赖于股票市场短线情绪，与A股大盘指数以及其他策略相关性都较低；主观期货策略因为自身多空双向交易，品种非常分散，也并非跟随南华商品指数波动；只有宏观交易型策略，在今年高波动的宏观环境中出现极度不适，与黄金+股指的beta波动相关性较强，交易alpha为负，对组合拖累较大。组合前三大持仓（30%+15%+15%=60%仓位占比）基金表现如下："
REVIEW_OUTLOOK = "展望后市，美伊战争终将结束，高油价带来的通胀缓解，加息预期减弱，宏观扰动将会告一段落，但商品市场大概率会继续维持高波动，组合也将择机提升量化CTA配比；A股市场随着三季报陆续披露，市场有望迎来新的热点题材，整个指数的跨年行情同样值得期待，组合计划适当降低打板策略权重，回归均衡配置，加入中性对冲、红利低波等低相关策略，提升组合抗风险能力。在接下里的几个月内，我们团队将会更加专注更加紧密地跟踪市场环境变化，积极调整组合配置，争取尽快修复回撤。最后感谢各位投资人一直以来的理解与包容！"

_REVIEW_TOP = 0.892
_REVIEW_BOTTOM = 0.078
_REVIEW_LINE = 0.0172
_REVIEW_CHARS = 54
_REVIEW_IMG_W = 0.908


def _review_blocks() -> list[tuple[str, object]]:
    return [
        ("p", REVIEW_STOCK),
        ("img", FOF_ASSET / "limitup_temperature.png"),
        ("img", FOF_ASSET / "market_turnover.png"),
        ("p", REVIEW_FUTURES),
        ("img", FOF_ASSET / "nanhua_sectors.png"),
        ("p", REVIEW_PORTFOLIO),
        ("img", FOF_ASSET / "top_holdings_monthly.png"),
        ("img", FOF_ASSET / "top_holdings_corr.png"),
        ("p", REVIEW_OUTLOOK),
    ]


def _review_image_height(path: Path) -> float:
    with Image.open(path) as im:
        return _REVIEW_IMG_W * (im.size[1] / im.size[0]) * (PAGE_IN[0] / PAGE_IN[1])


def _review_block_height(kind: str, payload) -> float:
    if kind == "p":
        return max(1, len(wrap_text(str(payload), _REVIEW_CHARS))) * _REVIEW_LINE + 0.009
    return _review_image_height(Path(payload)) + 0.012


def _review_groups() -> list[list[tuple[str, object]]]:
    """Keep a paragraph with the charts that follow it."""
    groups: list[list[tuple[str, object]]] = []
    pending: list[tuple[str, object]] = []
    for block in _review_blocks():
        if block[0] == "p" and pending:
            groups.append(pending)
            pending = [block]
        else:
            pending.append(block)
    if pending:
        groups.append(pending)
    return groups


def _review_plan() -> list[list[tuple[str, object]]]:
    pages: list[list[tuple[str, object]]] = []
    current: list[tuple[str, object]] = []
    y = _REVIEW_TOP
    for group in _review_groups():
        height = sum(_review_block_height(*block) for block in group)
        if current and y - height < _REVIEW_BOTTOM:
            pages.append(current)
            current = []
            y = _REVIEW_TOP
        if y - height < _REVIEW_BOTTOM:
            raise RuntimeError("review group does not fit on a page")
        current.extend(group)
        y -= height
    if current:
        pages.append(current)
    return pages


def review_page_count() -> int:
    return len(_review_plan())


def _review_footer(ax, font, page: int, total_pages: int):
    ax.add_patch(Rectangle((0.046, 0.052), 0.908, 0.0008, facecolor=C_BORDER, edgecolor="none", transform=ax.transAxes, zorder=2))
    ax.text(0.5, 0.040, DISCLAIMER, transform=ax.transAxes, ha="center", va="top", fontsize=6.45, color=C_MUTED, fontproperties=font, zorder=4)
    ax.text(
        0.5, 0.024,
        f"数据截止日  2026年9月30日      ·      第 {page} 页 / 共 {total_pages} 页",
        transform=ax.transAxes, ha="center", va="top", fontsize=6.3, color="#A0AEC0", fontproperties=font, zorder=4,
    )


def _draw_review_image(fig, ax, path: Path, y: float) -> float:
    height = _review_image_height(path)
    x = (1 - _REVIEW_IMG_W) / 2
    image_ax = fig.add_axes([x, y - height, _REVIEW_IMG_W, height])
    image_ax.imshow(Image.open(path), aspect="auto")
    image_ax.set_xticks([])
    image_ax.set_yticks([])
    for spine in image_ax.spines.values():
        spine.set_color(C_BORDER)
        spine.set_linewidth(0.6)
    ax.add_patch(Rectangle((x, y - height), _REVIEW_IMG_W, height, facecolor="none", edgecolor="none", transform=ax.transAxes, zorder=1))
    return y - height


def render_review_pages(font, font_bold, first_page: int, total_pages: int):
    figs = []
    for index, blocks in enumerate(_review_plan()):
        fig = plt.figure(figsize=PAGE_IN, dpi=DPI, facecolor=C_BG)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis("off")
        ax.text(0.62, 0.42, "荣熙", transform=ax.transAxes, ha="center", va="center", fontsize=58, color="#E7EDF4", rotation=28, fontproperties=font_bold, zorder=0)
        ax.add_patch(Rectangle((0, 0.958), 1, 0.042, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=1))
        ax.text(0.046, 0.979, "CONFIDENTIAL    ·    仅供指定机构投资者", transform=ax.transAxes, ha="left", va="center", fontsize=7.4, color="#E2E8F0", fontproperties=font, zorder=2)
        ax.text(0.954, 0.979, BRAND, transform=ax.transAxes, ha="right", va="center", fontsize=8.6, color="white", fontproperties=font_bold, zorder=2)
        ax.add_patch(Rectangle((0, 0.954), 1, 0.004, facecolor="#C53030", edgecolor="none", transform=ax.transAxes, zorder=1))
        ax.text(0.5, 0.938, PRODUCT, transform=ax.transAxes, ha="center", va="top", fontsize=13.2, color=C_TEXT, fontproperties=font_bold, zorder=2)
        title = "共赢FOF季度回顾与展望" + ("（续）" if index else "")
        ax.text(0.5, 0.916, title, transform=ax.transAxes, ha="center", va="top", fontsize=11.2, color=C_NAVY, fontproperties=font_bold, zorder=2)
        y = _REVIEW_TOP
        for kind, payload in blocks:
            if kind == "p":
                y = draw_wrapped(ax, str(payload), 0.058, y, font, 8.85, _REVIEW_CHARS, C_TEXT, _REVIEW_LINE)
                y -= 0.009
            else:
                y = _draw_review_image(fig, ax, Path(payload), y)
                y -= 0.012
        if y < _REVIEW_BOTTOM:
            raise RuntimeError(f"review page {index + 1} overflow, bottom y={y:.3f}")
        _review_footer(ax, font, first_page + index, total_pages)
        figs.append(fig)
    return figs


def _db_url() -> str:
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        if line.startswith("DATABASE_URL="):
            return line.split("=", 1)[1].strip().strip('"')
    raise RuntimeError("DATABASE_URL missing")


def _cn(iso: str) -> str:
    day = pd.Timestamp(iso)
    return f"{day.year}年{day.month}月{day.day}日"


def load_product():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "product_quarterly"))
    import generate_product_quarterly as pq

    conn = psycopg2.connect(_db_url())
    cur = conn.cursor()
    cur.execute("SELECT nav_series::text FROM ops_private_fund_detail_nav_cache WHERE beian_hao = %s", (BEIAN,))
    fetched = cur.fetchone()
    if not fetched:
        raise RuntimeError(f"no nav cache for {BEIAN}")
    series = json.loads(fetched[0])
    rows = sorted(
        (
            row["price_date"][:10],
            float(row["nav"]),
            float(row.get("cumulative_nav") or row.get("cum_nav_withdrawal") or row["nav"]),
        )
        for row in series
        if row["price_date"][:10] <= CUTOFF
    )
    if len(rows) < 2:
        raise RuntimeError("荣熙共赢 nav series is empty")

    def index_rows(ts_code: str):
        cur.execute(
            """
            SELECT trade_date::text, close::float
            FROM raw_ashare_index_daily
            WHERE ts_code = %s AND trade_date >= %s AND trade_date <= %s AND close > 0
            ORDER BY trade_date
            """,
            (ts_code, rows[0][0], CUTOFF),
        )
        return cur.fetchall()

    sse = index_rows("000001.SH")
    cur.execute(
        """
        SELECT DISTINCT ON (trade_date) trade_date::text, close::float
        FROM raw_spot_daily
        WHERE symbol = 'IF'
          AND trade_date >= %s
          AND trade_date <= %s
          AND close IS NOT NULL
          AND close > 0
        ORDER BY trade_date ASC, fetched_at DESC
        """,
        (rows[0][0], CUTOFF),
    )
    csi = cur.fetchall()
    conn.close()

    def align(index):
        last = None
        i = 0
        out = []
        for day, _unit, _adj in rows:
            while i < len(index) and index[i][0] <= day:
                last = index[i][1]
                i += 1
            out.append(last)
        return out

    frame = pd.DataFrame({
        "date": pd.to_datetime([row[0] for row in rows]),
        "unit_nav": [row[1] for row in rows],
        "nav": [row[2] for row in rows],
        "bench1": align(sse),
        "bench2": align(csi),
    })
    work = frame[frame["date"] <= pd.Timestamp(CUTOFF)].copy()
    work = pq.apply_nav_frequency(work, "weekly")
    metrics = pq.calc_metrics(work, pd.Timestamp(PERIOD_BEGIN), pd.Timestamp(CUTOFF), "window", "division")
    return frame, pq, metrics


def _sheets(pq, metrics) -> dict[str, str]:
    fund = metrics["fund"]
    return {
        "期末单位净值": f"{metrics['unit_nav']:.4f}",
        "区间收益": pq._fmt_pct(fund["ret"]),
        "年化波动": pq._fmt_pct(fund["ann_vol"], signed=False),
        "最大回撤": pq._fmt_pct(fund["mdd"]),
        "同期沪深300": pq._fmt_pct(metrics["bench2"]["ret"]),
        "相对沪深300": pq._fmt_pct(metrics["excess2"]),
    }


def render_letter_page(font, font_bold, sheets: dict[str, str], *, perf_page: int, hist_page: int, total_pages: int):
    fig = plt.figure(figsize=PAGE_IN, dpi=DPI, facecolor=C_BG)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.62, 0.42, "荣熙", transform=ax.transAxes, ha="center", va="center", fontsize=58, color="#E7EDF4", rotation=28, fontproperties=font_bold, zorder=0)

    ax.add_patch(Rectangle((0, 0.958), 1, 0.042, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=1))
    ax.text(0.046, 0.979, "CONFIDENTIAL    ·    仅供指定机构投资者", transform=ax.transAxes, ha="left", va="center", fontsize=7.4, color="#E2E8F0", fontproperties=font, zorder=2)
    ax.text(0.954, 0.979, BRAND, transform=ax.transAxes, ha="right", va="center", fontsize=8.6, color="white", fontproperties=font_bold, zorder=2)
    ax.add_patch(Rectangle((0, 0.954), 1, 0.004, facecolor="#C53030", edgecolor="none", transform=ax.transAxes, zorder=1))

    ax.text(0.5, 0.936, PRODUCT, transform=ax.transAxes, ha="center", va="top", fontsize=15.2, color=C_TEXT, fontproperties=font_bold, zorder=2)
    ax.text(0.5, 0.907, "2026年第三季度投资报告", transform=ax.transAxes, ha="center", va="top", fontsize=12.6, color=C_NAVY, fontproperties=font_bold, zorder=2)
    ax.text(0.5, 0.882, f"报告期    {_cn(PERIOD_BEGIN)}    至    {_cn(CUTOFF)}", transform=ax.transAxes, ha="center", va="top", fontsize=8.5, color=C_MUTED, fontproperties=font, zorder=2)
    ax.text(0.5, 0.862, f"管理人  {MANAGER}      投资经理  {ADVISOR}", transform=ax.transAxes, ha="center", va="top", fontsize=7.5, color="#4A5568", fontproperties=font, zorder=2)
    ax.text(0.5, 0.846, f"托管人  {CUSTODIAN}      成立时间  {_cn(INCEPTION)}", transform=ax.transAxes, ha="center", va="top", fontsize=7.5, color="#4A5568", fontproperties=font, zorder=2)

    y = section_bar(ax, 0.824, "一、致投资人", font_bold)
    y = draw_wrapped(ax, LETTER_LEAD, 0.058, y, font, 8.85, 54, C_TEXT, 0.0172)
    y -= 0.006
    y = draw_wrapped(ax, "股票与期货市场回顾、前三大持仓表现及后市展望见第2页。", 0.058, y, font, 8.85, 54, C_TEXT, 0.0172)
    y -= 0.012

    y = section_bar(ax, y, "二、本产品本季度", font_bold)
    ax.text(
        0.058, y,
        f"以下只列荣熙共赢本季度业绩，对应投资者所持份额。详表见第{perf_page}页。",
        transform=ax.transAxes, ha="left", va="top", fontsize=8.2, color=C_TEXT, fontproperties=font, zorder=4,
    )
    y -= 0.026
    kpis = [
        ("区间收益", sheets["区间收益"], "本季度"),
        ("年化波动", sheets["年化波动"], "按周净值"),
        ("最大回撤", sheets["最大回撤"], "季度内"),
        ("期末单位净值", sheets["期末单位净值"], CUTOFF),
    ]
    box_h = 0.078
    gap = 0.012
    width = (0.908 - gap * 3) / 4
    x0 = 0.046
    for i, (label, value, hint) in enumerate(kpis):
        x = x0 + i * (width + gap)
        ax.add_patch(Rectangle((x, y - box_h), width, box_h, facecolor="#F8FAFC", edgecolor=C_BORDER, lw=0.7, transform=ax.transAxes, zorder=2))
        ax.text(x + width / 2, y - 0.016, label, transform=ax.transAxes, ha="center", va="top", fontsize=7.3, color=C_MUTED, fontproperties=font, zorder=3)
        ax.text(x + width / 2, y - 0.042, value, transform=ax.transAxes, ha="center", va="center", fontsize=13, color=C_NAVY, fontproperties=font_bold, zorder=3)
        ax.text(x + width / 2, y - 0.062, hint, transform=ax.transAxes, ha="center", va="top", fontsize=6.5, color=C_MUTED, fontproperties=font, zorder=3)
    y = y - box_h - 0.014
    ax.text(
        0.058, y,
        f"同期沪深300 {sheets['同期沪深300']}，相对沪深300 {sheets['相对沪深300']}。",
        transform=ax.transAxes, ha="left", va="top", fontsize=8.0, color=C_TEXT, fontproperties=font, zorder=4,
    )
    y -= 0.026
    bridge = [
        "本产品于2026年1月20日成立，运作时间较短。",
        f"第{hist_page}页补充同一策略自2020年4月3日起的历史业绩，供了解策略过往表现。",
    ]
    note_h = 0.014 + len(bridge) * 0.017
    ax.add_patch(Rectangle((0.046, y - note_h), 0.908, note_h, facecolor=C_NOTE_BG, edgecolor=C_BORDER, lw=0.6, transform=ax.transAxes, zorder=2))
    ax.add_patch(Rectangle((0.046, y - note_h), 0.005, note_h, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=3))
    ny = y - 0.009
    for line in bridge:
        ax.text(0.062, ny, line, transform=ax.transAxes, ha="left", va="top", fontsize=7.5, color="#2D3748", fontproperties=font, zorder=4)
        ny -= 0.017

    ax.add_patch(Rectangle((0.046, 0.052), 0.908, 0.0008, facecolor=C_BORDER, edgecolor="none", transform=ax.transAxes, zorder=2))
    ax.text(0.5, 0.040, DISCLAIMER, transform=ax.transAxes, ha="center", va="top", fontsize=6.45, color=C_MUTED, fontproperties=font, zorder=4)
    ax.text(
        0.5, 0.024,
        f"数据截止日  2026年9月30日      ·      净值已复权      ·      基准为沪深300      ·      第 1 页 / 共 {total_pages} 页",
        transform=ax.transAxes, ha="center", va="top", fontsize=6.3, color="#A0AEC0", fontproperties=font, zorder=4,
    )
    return fig


def holding_config(page_label: str = "") -> dict:
    return {
        "product_name": PRODUCT,
        "short_name": SHORT,
        "brand_name": BRAND,
        "watermark": "荣熙",
        "end_date": CUTOFF,
        "period_begin": PERIOD_BEGIN,
        "chart_start": INCEPTION,
        "inception_date": INCEPTION,
        "nav_frequency": "weekly",
        "ann_vol_instead_of_return": True,
        "metrics_match_product_page": True,
        "show_bench1": False,
        "page_label": page_label,
        "bench1_label": "上证指数",
        "bench2_label": "沪深300",
        "overview": {
            "manager": MANAGER,
            "investment_manager": ADVISOR,
            "custodian": CUSTODIAN,
            "inception_date": "2026.1.20",
            "product_type": "私募证券投资基金",
        },
    }


def build_docx(out_dir: Path, sheets: dict[str, str], perf_image: Path, hist_image: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.50)
    section.bottom_margin = Cm(1.55)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    rpr = normal.element.get_or_add_rPr()
    rfonts = qn_fonts(rpr)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "微软雅黑")

    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.paragraph_format.space_before = Pt(0)
    header_p.paragraph_format.space_after = Pt(1)
    set_run_font(header_p.add_run("CONFIDENTIAL  ·  仅供指定机构投资者"), 8, color=MUTED)
    set_run_font(header_p.add_run(f"          {BRAND}"), 8, bold=True, color=NAVY)
    add_bottom_border(header_p)

    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.paragraph_format.space_before = Pt(1)
    foot.paragraph_format.space_after = Pt(0)
    set_run_font(foot.add_run(DISCLAIMER), 7.5, color=MUTED)
    foot2 = footer.add_paragraph()
    foot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot2.paragraph_format.space_before = Pt(0)
    foot2.paragraph_format.space_after = Pt(0)
    set_run_font(foot2.add_run("数据截止日 2026年9月30日  ·  第 "), 8, color=MUTED)
    add_page_field(foot2, "PAGE")
    set_run_font(foot2.add_run(" 页 / 共 "), 8, color=MUTED)
    add_page_field(foot2, "NUMPAGES")
    set_run_font(foot2.add_run(" 页"), 8, color=MUTED)

    add_text(doc, PRODUCT, 18, bold=True, align="center", before=0, after=1)
    add_text(doc, "2026年第三季度投资报告", 14, bold=True, color=NAVY, align="center", before=0, after=1)
    add_text(doc, f"报告期  {_cn(PERIOD_BEGIN)}  至  {_cn(CUTOFF)}", 10.5, color=MUTED, align="center", before=0, after=1)
    add_text(doc, f"管理人  {MANAGER}          投资经理  {ADVISOR}", 9, color=RGBColor(0x4A, 0x55, 0x68), align="center", before=0, after=0)
    meta = add_text(doc, f"托管人  {CUSTODIAN}          成立时间  {_cn(INCEPTION)}", 9, color=RGBColor(0x4A, 0x55, 0x68), align="center", before=0, after=6)
    add_bottom_border(meta, "E2E8F0", "8")

    add_text(doc, "一、致投资人", 13, bold=True, color=NAVY, before=2, after=2)
    add_text(doc, LETTER_LEAD, 10.5, before=2, after=4)
    add_text(doc, "股票与期货市场回顾、前三大持仓表现及后市展望见后文。", 10.5, before=0, after=6)

    add_text(doc, "二、本产品本季度", 13, bold=True, color=NAVY, before=4, after=2)
    add_text(doc, "以下只列荣熙共赢本季度业绩，对应投资者所持份额。详表见后文本季度业绩。", 9, color=MUTED, before=0, after=3)
    rows = list(sheets.items())
    widths = [3600, 6370]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, widths)
    for j, (title, width) in enumerate(zip(("指标", "荣熙共赢  7月1日–9月30日"), widths)):
        cell = table.rows[0].cells[j]
        shade_cell(cell, "1A365D")
        set_cell_width(cell, width)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(title), 9, bold=True, color=WHITE)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, (label, val) in enumerate(rows):
        row = table.rows[i + 1]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, (cell_text, width) in enumerate(zip((label, val), widths)):
            cell = row.cells[j]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if label in HIGHLIGHT:
                fill = "E8EEF6"
            elif j == 0:
                fill = "EDF2F7"
            else:
                fill = "FFFFFF" if i % 2 == 0 else "F7FAFC"
            shade_cell(cell, fill)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(
                paragraph.add_run(cell_text), 9,
                bold=(j == 0 or label in HIGHLIGHT),
                color=NAVY if j == 1 and label in HIGHLIGHT else TEXT,
            )
    add_text(doc, "本产品于2026年1月20日成立，运作时间较短。文末补充同一策略自2020年4月3日起的历史业绩，供了解策略过往表现。", 9, color=SLATE, before=6, after=2)

    review_heading = add_text(doc, "三、共赢FOF季度回顾与展望", 13, bold=True, color=NAVY, before=0, after=4)
    review_heading.paragraph_format.page_break_before = True
    add_text(doc, REVIEW_STOCK, 10.5, before=2, after=6)
    _add_review_picture(doc, FOF_ASSET / "limitup_temperature.png")
    _add_review_picture(doc, FOF_ASSET / "market_turnover.png")
    add_text(doc, REVIEW_FUTURES, 10.5, before=2, after=6)
    _add_review_picture(doc, FOF_ASSET / "nanhua_sectors.png")
    add_text(doc, REVIEW_PORTFOLIO, 10.5, before=2, after=6)
    _add_review_picture(doc, FOF_ASSET / "top_holdings_monthly.png")
    _add_review_picture(doc, FOF_ASSET / "top_holdings_corr.png")
    add_text(doc, REVIEW_OUTLOOK, 10.5, before=2, after=6)

    sec3 = add_text(doc, "四、荣熙共赢本季度业绩", 13, bold=True, color=NAVY, before=0, after=1)
    sec3.paragraph_format.page_break_before = True
    add_text(doc, f"下图为本季度持有期业绩，对应投资者所持份额。报告期 {_cn(PERIOD_BEGIN)}至{_cn(CUTOFF)}，期末单位净值 {sheets['期末单位净值']}。", 9, color=MUTED, before=0, after=3)
    doc.add_picture(str(perf_image), width=Cm(15.2))
    picture = doc.paragraphs[-1]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(0)

    sec4 = add_text(doc, "五、同策略历史业绩", 13, bold=True, color=NAVY, before=0, after=1)
    sec4.paragraph_format.page_break_before = True
    add_text(doc, "下图为同一策略自2020年4月3日至2026年9月30日的历史业绩，补充产品成立以来较短的记录，供了解策略过往表现。", 9, color=MUTED, before=0, after=3)
    doc.add_picture(str(hist_image), width=Cm(16.0))
    picture = doc.paragraphs[-1]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(0)

    core = doc.core_properties
    core.title = "荣熙共赢私募证券投资基金 2026年第三季度投资报告"
    core.subject = "共赢FOF三季度回顾、本季度业绩与同策略历史净值"
    core.author = BRAND
    core.category = "季度投资报告"
    path = out_dir / "荣熙共赢_2026年第三季度投资报告.docx"
    doc.save(path)
    return path


def qn(name: str):
    from docx.oxml.ns import qn as _qn
    return _qn(name)


def qn_fonts(rpr):
    found = rpr.find(qn("w:rFonts"))
    if found is None:
        found = OxmlElement("w:rFonts")
        rpr.append(found)
    return found


def _add_review_picture(doc: Document, path: Path):
    doc.add_picture(str(path), width=Cm(16.4))
    picture = doc.paragraphs[-1]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(2)
    picture.paragraph_format.space_after = Pt(8)


def main():
    out_dir = Path(__file__).resolve().parent / "荣熙共赢_2026Q3"
    out_dir.mkdir(parents=True, exist_ok=True)
    font, font_bold = load_fonts()
    n_review = review_page_count()
    total_pages = 1 + n_review + 2
    perf_page = 2 + n_review
    hist_page = perf_page + 1
    frame, pq, metrics = load_product()
    sheets = _sheets(pq, metrics)
    history = load_strategy_history(
        NAV_FILE,
        product_name=SHORT,
        inception=_cn(INCEPTION),
        brand=BRAND,
        watermark="荣熙",
    )
    fig1 = render_letter_page(font, font_bold, sheets, perf_page=perf_page, hist_page=hist_page, total_pages=total_pages)
    review_figs = render_review_pages(font, font_bold, first_page=2, total_pages=total_pages)
    fig_perf = pq.build_report_figure(frame, holding_config(page_label=f"第 {perf_page} 页 / 共 {total_pages} 页"))
    fig_hist = render_history_page(font, font_bold, history, page_label=f"第 {hist_page} 页 / 共 {total_pages} 页")

    page1_path = out_dir / "第1页_致投资人.png"
    perf_path = out_dir / f"第{perf_page}页_产品持有期业绩.png"
    hist_path = out_dir / f"第{hist_page}页_同策略历史业绩.png"
    review_paths = [out_dir / f"第{index}页_季度回顾.png" for index in range(2, 2 + n_review)]
    long_path = out_dir / "荣熙共赢_2026年第三季度投资报告.png"
    pdf_path = out_dir / "荣熙共赢_2026年第三季度投资报告.pdf"
    fig1.savefig(page1_path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    for path, fig in zip(review_paths, review_figs):
        fig.savefig(path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    fig_perf.savefig(perf_path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    fig_hist.savefig(hist_path, dpi=DPI, facecolor=C_BG, pad_inches=0)

    pages = [Image.open(page1_path).convert("RGB")]
    pages.extend(Image.open(path).convert("RGB") for path in review_paths)
    pages.append(Image.open(perf_path).convert("RGB"))
    pages.append(Image.open(hist_path).convert("RGB"))
    stitch_long(pages).save(long_path, dpi=(DPI, DPI), optimize=True)
    with PdfPages(pdf_path) as pdf:
        pdf.savefig(fig1, facecolor=C_BG)
        info = pdf.infodict()
        info["Title"] = "荣熙共赢私募证券投资基金 2026年第三季度投资报告"
        info["Author"] = BRAND
        info["Subject"] = "共赢FOF三季度回顾、本季度业绩与同策略更长历史"
        for fig in review_figs:
            pdf.savefig(fig, facecolor=C_BG)
        pdf.savefig(fig_perf, facecolor=C_BG)
        pdf.savefig(fig_hist, facecolor=C_BG)
    plt.close(fig1)
    for fig in review_figs:
        plt.close(fig)
    plt.close(fig_perf)
    plt.close(fig_hist)
    for page in pages:
        page.close()
    for stale in (
        out_dir / "第2页_产品持有期业绩.png",
        out_dir / "第3页_同策略历史业绩.png",
    ):
        if stale.exists() and stale not in {perf_path, hist_path}:
            stale.unlink()
    docx_path = build_docx(out_dir, sheets, perf_path, hist_path)
    for key, value in sheets.items():
        print(key.encode("unicode_escape").decode(), value)
    print("HISTORY", history["start"], history["end"], history["n"], f"{history['total']*100:.2f}%")
    print("OUT", out_dir)
    print("DOCX", docx_path)


if __name__ == "__main__":
    main()
