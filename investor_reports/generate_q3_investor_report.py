# -*- coding: utf-8 -*-
"""金舆追风1号 2026年第三季度投资报告：投资说明 + 产品业绩 + 同策略历史净值。"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
matplotlib.rcParams["pdf.fonttype"] = 42
matplotlib.rcParams["ps.fonttype"] = 42
matplotlib.rcParams["axes.unicode_minus"] = False

import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.font_manager import FontProperties, fontManager
from matplotlib.patches import Rectangle
from PIL import Image

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from q3_strategy_history import load_strategy_history, render_history_page

ASSETS = Path(r"C:\Users\Administrator\.cursor\projects\d-coding-market-website\assets")
PRODUCT_IMG = ASSETS / (
    "c__Users_Administrator_AppData_Roaming_Cursor_User_workspaceStorage_"
    "3fc4f3e308e1a15b7408c15bdb7ec093_images_____1__2026_7_15__2026_9_30______"
    "-d3e39699-4d32-4211-81e8-313420cdd908.png"
)
STRATEGY_IMG = ASSETS / (
    "c__Users_Administrator_AppData_Roaming_Cursor_User_workspaceStorage_"
    "3fc4f3e308e1a15b7408c15bdb7ec093_images_____________2026_7_1__2026_9_30______"
    "-389a9b86-a6d2-4919-81e0-e9cf93afe524.png"
)

FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"

C_NAVY = "#1A365D"
C_NAVY_MID = "#2C5282"
C_RED = "#C53030"
C_TEXT = "#1A202C"
C_MUTED = "#718096"
C_BORDER = "#CBD5E1"
C_BG = "#FFFFFF"
C_ROW = "#F7FAFC"
C_NOTE_BG = "#F8FAFC"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
TEXT = RGBColor(0x1A, 0x20, 0x2C)
MUTED = RGBColor(0x71, 0x80, 0x96)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x2D, 0x37, 0x48)

PAGE_IN = (8.27, 11.69)
DPI = 220
BRAND = "上海金舆资产管理"

LETTER = [
    ("salute", "各位投资人："),
    ("body", "今年三季度，我们FOF下投的强势股打板策略经历了较长的不适应期，导致母基金出现了4%左右的回撤。"),
    ("body", "主要原因是当前市场资金缺乏主线，连续下跌使得交易量大幅萎缩，强势股接力资金匮乏，导致打板赛道需要的连板投机情绪退潮，次日低开率显著飙升，这也是打板策略最艰难的阶段。"),
    ("body", "此外，沪深两市近期要求整改托管服务器，对打板策略的报单时效也产生了一些干扰。"),
    ("body", "面对以上情况，底层子基金正在努力尝试以下几点应对措施："),
    ("item", "优化选股模型，提高对封单量、超大单排撤、异动指数等因子的筛选阈值，努力过滤掉“量化陷阱”，提高封板率。"),
    ("item", "控制仓位，灵活调整打板时机：面对成交量显著萎缩的市场行情，充分发挥我们自研“打板启停开关”的避险作用，尽量少亏钱，保住现有净值。"),
    ("item", "耐心等待，择机用奇：历史经验告诉我们，没有永远有效的策略，也没有永远退潮的市场。打板策略本身挣的就是“市场情绪的溢价收益”，每一次情绪的冰点，都是下一轮行情的开始。一旦市场情绪回暖、打板生态恢复，我们将即刻把握机会，尽最大努力收复失地。"),
    ("body", "艰难的三季度已经过去，希望投资人再多给我们一些时间去修复净值。最后，感谢大家的理解与陪伴，如有需要，欢迎联系！"),
]

DISCLAIMER = "风险提示：本报告仅供指定机构投资者参考，不构成投资建议、要约或承诺。私募基金有投资风险，管理人过往业绩不预示未来表现。"

_WRAP_BREAK_AFTER = set("。！？；，、： ")
_NO_SPLIT_CHARS = set("0123456789+-.%")
_LINE_START_FORBIDDEN = set("。！？；，、：）)」』")
_LINE_END_FORBIDDEN = set("「『（(")


def prepare_output_dir() -> Path:
    parent = Path(__file__).resolve().parent
    out = parent / "金舆追风1号_2026Q3"
    out.mkdir(parents=True, exist_ok=True)
    for child in parent.iterdir():
        if child.is_dir() and child.resolve() != out.resolve() and child.name.endswith("2026Q3"):
            try:
                child.rmdir()
            except OSError:
                pass
    return out


def load_fonts() -> tuple[FontProperties, FontProperties]:
    fontManager.addfont(FONT_REG)
    fontManager.addfont(FONT_BOLD)
    regular = FontProperties(fname=FONT_REG)
    bold = FontProperties(fname=FONT_BOLD)
    matplotlib.rcParams["font.family"] = "sans-serif"
    matplotlib.rcParams["font.sans-serif"] = [regular.get_name(), "Microsoft YaHei", "SimHei"]
    return regular, bold


def wrap_text(text: str, max_chars: int) -> list[str]:
    lines: list[str] = []
    for para in text.split("\n"):
        para = para.strip()
        if not para:
            continue
        start = 0
        while start < len(para):
            rest = para[start:]
            if len(rest) <= max_chars:
                lines.append(rest)
                break
            end = max_chars
            while end > 1 and end < len(rest) and rest[end - 1] in _NO_SPLIT_CHARS and rest[end] in _NO_SPLIT_CHARS:
                end -= 1
            while end > 1 and rest[end - 1] in _LINE_END_FORBIDDEN:
                end -= 1
            while end < len(rest) and rest[end] in _LINE_START_FORBIDDEN:
                end += 1
            if end < 1:
                end = 1
            lines.append(rest[:end])
            start += end
    return lines

def draw_wrapped(ax, text, x, y, font, fontsize, max_chars, color, line_h) -> float:
    lines = wrap_text(text, max_chars)
    for i, line in enumerate(lines):
        ax.text(
            x, y - i * line_h, line,
            transform=ax.transAxes, ha="left", va="top",
            fontsize=fontsize, color=color, fontproperties=font, zorder=4,
        )
    return y - len(lines) * line_h


def section_bar(ax, y, title, font_bold) -> float:
    bar_h = 0.0245
    ax.add_patch(Rectangle((0.046, y - bar_h), 0.908, bar_h, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=2))
    ax.plot([0.059, 0.059], [y - bar_h + 0.005, y - 0.005], color="white", lw=1.6, transform=ax.transAxes, zorder=3, solid_capstyle="round")
    ax.text(0.074, y - bar_h / 2, title, transform=ax.transAxes, va="center", ha="left", color="white", fontsize=9.3, fontproperties=font_bold, zorder=3)
    return y - bar_h - 0.013


def render_letter_page(font, font_bold, market: dict[str, str]):
    fig = plt.figure(figsize=PAGE_IN, dpi=DPI, facecolor=C_BG)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.text(0.62, 0.42, "金舆", transform=ax.transAxes, ha="center", va="center", fontsize=58, color="#E7EDF4", rotation=28, fontproperties=font_bold, zorder=0)

    ax.add_patch(Rectangle((0, 0.958), 1, 0.042, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=1))
    ax.text(0.046, 0.979, "CONFIDENTIAL    ·    仅供指定机构投资者", transform=ax.transAxes, ha="left", va="center", fontsize=7.4, color="#E2E8F0", fontproperties=font, zorder=2)
    ax.text(0.954, 0.979, BRAND, transform=ax.transAxes, ha="right", va="center", fontsize=8.6, color="white", fontproperties=font_bold, zorder=2)
    ax.add_patch(Rectangle((0, 0.954), 1, 0.004, facecolor=C_RED, edgecolor="none", transform=ax.transAxes, zorder=1))

    ax.text(0.5, 0.936, "金舆追风1号私募证券投资基金", transform=ax.transAxes, ha="center", va="top", fontsize=15.2, color=C_TEXT, fontproperties=font_bold, zorder=2)
    ax.text(0.5, 0.907, "2026年第三季度投资报告", transform=ax.transAxes, ha="center", va="top", fontsize=12.6, color=C_NAVY, fontproperties=font_bold, zorder=2)
    ax.text(0.5, 0.882, "持有期    2026年7月15日    至    2026年9月30日", transform=ax.transAxes, ha="center", va="top", fontsize=8.5, color=C_MUTED, fontproperties=font, zorder=2)
    ax.text(0.5, 0.862, "管理人  上海金舆资产管理有限公司      投资经理  陈家本", transform=ax.transAxes, ha="center", va="top", fontsize=7.5, color="#4A5568", fontproperties=font, zorder=2)
    ax.text(0.5, 0.846, "托管人  招商证券股份有限公司      成立时间  2026年7月15日", transform=ax.transAxes, ha="center", va="top", fontsize=7.5, color="#4A5568", fontproperties=font, zorder=2)

    y = 0.824
    y = section_bar(ax, y, "一、致投资人", font_bold)
    item_no = 0
    for kind, text in LETTER:
        if kind == "salute":
            ax.text(0.058, y, text, transform=ax.transAxes, ha="left", va="top", fontsize=9.4, color=C_TEXT, fontproperties=font_bold, zorder=4)
            y -= 0.022
        elif kind == "body":
            y = draw_wrapped(ax, text, 0.058, y, font, 8.85, 54, C_TEXT, 0.0172)
            y -= 0.0075
        else:
            item_no += 1
            ax.text(0.060, y, f"{item_no}.", transform=ax.transAxes, ha="left", va="top", fontsize=8.85, color=C_NAVY, fontproperties=font_bold, zorder=4)
            y = draw_wrapped(ax, text, 0.086, y, font, 8.85, 50, C_TEXT, 0.0172)
            y -= 0.0065

    y -= 0.008
    y = section_bar(ax, y, "\u4e8c\u3001\u672c\u4ea7\u54c1\u672c\u5b63\u5ea6", font_bold)
    ax.text(
        0.058, y,
        "\u4ee5\u4e0b\u53ea\u5217\u91d1\u8206\u8ffd\u98ce1\u53f7\u6301\u6709\u671f\u4e1a\u7ee9\uff0c\u5bf9\u5e94\u6295\u8d44\u8005\u6240\u6301\u4efd\u989d\u3002\u8be6\u8868\u89c1\u7b2c2\u9875\u3002",
        transform=ax.transAxes, ha="left", va="top", fontsize=8.2, color=C_TEXT, fontproperties=font, zorder=4,
    )
    y -= 0.026
    kpis = [
        ("\u533a\u95f4\u6536\u76ca", "-3.92%", "\u672c\u6301\u6709\u671f"),
        ("\u5e74\u5316\u6ce2\u52a8", "5.54%", "\u6309\u5468\u51c0\u503c"),
        ("\u6700\u5927\u56de\u64a4", "-3.92%", "\u6301\u6709\u671f\u5185"),
        ("\u671f\u672b\u5355\u4f4d\u51c0\u503c", "0.9607", "2026-09-30"),
    ]
    box_h = 0.078
    gap = 0.012
    width = (0.908 - gap * 3) / 4
    x0 = 0.046
    for i, (label, value, hint) in enumerate(kpis):
        x = x0 + i * (width + gap)
        ax.add_patch(Rectangle((x, y - box_h), width, box_h, facecolor="#F8FAFC", edgecolor=C_BORDER, lw=0.7, transform=ax.transAxes, zorder=2))
        ax.text(x + width / 2, y - 0.016, label, transform=ax.transAxes, ha="center", va="top", fontsize=7.3, color=C_MUTED, fontproperties=font, zorder=3)
        ax.text(x + width / 2, y - 0.042, value, transform=ax.transAxes, ha="center", va="center", fontsize=13, color=C_NAVY, fontproperties=font_bold, zorder=3)
        ax.text(x + width / 2, y - 0.062, hint, transform=ax.transAxes, ha="center", va="top", fontsize=6.5, color=C_MUTED, fontproperties=font, zorder=3)
    y = y - box_h - 0.014
    ax.text(
        0.058, y,
        (
            f"\u540c\u671f\u6caa\u6df1300 {market['csi']}\uff0c\u76f8\u5bf9\u6caa\u6df1300 {market['xcsi']}\u3002"
        ),
        transform=ax.transAxes, ha="left", va="top", fontsize=8.0, color=C_TEXT, fontproperties=font, zorder=4,
    )
    y -= 0.026
    bridge = [
        "\u672c\u4ea7\u54c1\u4e8e2026\u5e747\u670815\u65e5\u6210\u7acb\uff0c\u8fd0\u4f5c\u65f6\u95f4\u8f83\u77ed\u3002",
        "\u7b2c3\u9875\u8865\u5145\u540c\u4e00\u7b56\u7565\u81ea2023\u5e7412\u670829\u65e5\u8d77\u7684\u5386\u53f2\u4e1a\u7ee9\uff0c\u4f9b\u4e86\u89e3\u7b56\u7565\u8fc7\u5f80\u8868\u73b0\u3002",
    ]
    note_h = 0.014 + len(bridge) * 0.017
    ax.add_patch(Rectangle((0.046, y - note_h), 0.908, note_h, facecolor=C_NOTE_BG, edgecolor=C_BORDER, lw=0.6, transform=ax.transAxes, zorder=2))
    ax.add_patch(Rectangle((0.046, y - note_h), 0.005, note_h, facecolor=C_NAVY, edgecolor="none", transform=ax.transAxes, zorder=3))
    ny = y - 0.009
    for line in bridge:
        ax.text(0.062, ny, line, transform=ax.transAxes, ha="left", va="top", fontsize=7.5, color="#2D3748", fontproperties=font, zorder=4)
        ny -= 0.017
    y = y - note_h

    print(f"content_bottom={y:.3f}")
    if y < 0.062:
        raise RuntimeError(f"page 1 overflow, content bottom y={y:.3f}")

    ax.add_patch(Rectangle((0.046, 0.052), 0.908, 0.0008, facecolor=C_BORDER, edgecolor="none", transform=ax.transAxes, zorder=2))
    ax.text(0.5, 0.040, DISCLAIMER, transform=ax.transAxes, ha="center", va="top", fontsize=6.45, color=C_MUTED, fontproperties=font, zorder=4)
    ax.text(
        0.5, 0.024,
        "数据截止日  2026年9月30日      ·      净值已复权      ·      基准为沪深300      ·      第 1 页 / 共 3 页",
        transform=ax.transAxes, ha="center", va="top", fontsize=6.3, color="#A0AEC0", fontproperties=font, zorder=4,
    )
    return fig


def fit_exhibit(src_path: Path, size: tuple[int, int]) -> Image.Image:
    src = Image.open(src_path).convert("RGB")
    canvas = Image.new("RGB", size, "white")
    scale = min(size[0] / src.width, size[1] / src.height)
    resized = src.resize((max(1, int(round(src.width * scale))), max(1, int(round(src.height * scale)))), Image.Resampling.LANCZOS)
    canvas.paste(resized, ((size[0] - resized.width) // 2, (size[1] - resized.height) // 2))
    return canvas


def stitch_long(pages: list[Image.Image]) -> Image.Image:
    gap = 16
    height = sum(page.height for page in pages) + gap * (len(pages) - 1)
    long_im = Image.new("RGB", (pages[0].width, height), "#E8EEF4")
    y = 0
    for page in pages:
        long_im.paste(page, (0, y))
        y += page.height + gap
    return long_im


def set_run_font(run, size, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def add_text(doc, text, size, bold=False, color=TEXT, align="left", before=0, after=6):
    paragraph = doc.add_paragraph()
    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    set_run_font(paragraph.add_run(text), size, bold=bold, color=color)
    return paragraph


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(old)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcMar")):
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "10"), ("left", "46"), ("bottom", "10"), ("right", "46")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def set_table_fixed(table, widths: list[int]):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    total = sum(widths)
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E1")
        borders.append(el)
    tbl_pr.append(borders)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr.addnext(grid)


def add_bottom_border(paragraph, color="C53030", sz="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_page_field(paragraph, instr: str):
    run = paragraph.add_run()
    set_run_font(run, 8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instr} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def build_docx(out_dir: Path, market: dict[str, str]):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.50)
    section.bottom_margin = Cm(1.55)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "微软雅黑")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")

    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.paragraph_format.space_before = Pt(0)
    header_p.paragraph_format.space_after = Pt(1)
    set_run_font(header_p.add_run("CONFIDENTIAL  ·  仅供指定机构投资者"), 8, color=MUTED)
    set_run_font(header_p.add_run("          上海金舆资产管理"), 8, bold=True, color=NAVY)
    add_bottom_border(header_p)

    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.paragraph_format.space_before = Pt(1)
    foot.paragraph_format.space_after = Pt(0)
    set_run_font(foot.add_run(DISCLAIMER), 7.5, color=MUTED)
    foot2 = footer.add_paragraph()
    foot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot2.paragraph_format.space_before = Pt(0)
    foot2.paragraph_format.space_after = Pt(0)
    set_run_font(foot2.add_run("数据截止日 2026年9月30日  ·  第 "), 8, color=MUTED)
    add_page_field(foot2, "PAGE")
    set_run_font(foot2.add_run(" 页 / 共 "), 8, color=MUTED)
    add_page_field(foot2, "NUMPAGES")
    set_run_font(foot2.add_run(" 页"), 8, color=MUTED)

    add_text(doc, "金舆追风1号私募证券投资基金", 18, bold=True, align="center", before=0, after=1)
    add_text(doc, "2026年第三季度投资报告", 14, bold=True, color=NAVY, align="center", before=0, after=1)
    add_text(doc, "持有期  2026年7月15日  至  2026年9月30日", 10.5, color=MUTED, align="center", before=0, after=1)
    add_text(doc, "管理人  上海金舆资产管理有限公司          投资经理  陈家本", 9, color=RGBColor(0x4A, 0x55, 0x68), align="center", before=0, after=0)
    meta = add_text(doc, "托管人  招商证券股份有限公司          成立时间  2026年7月15日", 9, color=RGBColor(0x4A, 0x55, 0x68), align="center", before=0, after=6)
    add_bottom_border(meta, "E2E8F0", "8")

    add_text(doc, "一、致投资人", 13, bold=True, color=NAVY, before=2, after=2)
    item_no = 0
    for kind, text in LETTER:
        if kind == "salute":
            add_text(doc, text, 10.5, bold=True, before=0, after=3)
        elif kind == "item":
            item_no += 1
            paragraph = add_text(doc, f"{item_no}.  {text}", 10.5, before=1, after=2)
            paragraph.paragraph_format.left_indent = Cm(0.25)
        else:
            add_text(doc, text, 10.5, before=0, after=2)

    add_text(doc, "\u4e8c\u3001\u672c\u4ea7\u54c1\u672c\u5b63\u5ea6", 13, bold=True, color=NAVY, before=4, after=2)
    add_text(doc, "\u4ee5\u4e0b\u53ea\u5217\u91d1\u8206\u8ffd\u98ce1\u53f7\u6301\u6709\u671f\u4e1a\u7ee9\uff0c\u5bf9\u5e94\u6295\u8d44\u8005\u6240\u6301\u4efd\u989d\u3002\u8be6\u8868\u89c1\u540e\u9875\u3002", 9, color=MUTED, before=0, after=3)
    prod_rows = [
        ("\u671f\u672b\u5355\u4f4d\u51c0\u503c", "0.9607"),
        ("\u533a\u95f4\u6536\u76ca", "-3.92%"),
        ("\u5e74\u5316\u6ce2\u52a8", "5.54%"),
        ("\u6700\u5927\u56de\u64a4", "-3.92%"),
        ("\u540c\u671f\u6caa\u6df1300", market["csi"]),
        ("\u76f8\u5bf9\u6caa\u6df1300", market["xcsi"]),
    ]
    widths = [3600, 6370]
    table = doc.add_table(rows=1 + len(prod_rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, widths)
    for j, (title, fill, width) in enumerate(zip(("\u6307\u6807", "\u91d1\u8206\u8ffd\u98ce1\u53f7  7\u670815\u65e5\u20139\u670830\u65e5"), ("1A365D", "1A365D"), widths)):
        cell = table.rows[0].cells[j]
        shade_cell(cell, fill)
        set_cell_width(cell, width)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(title), 9, bold=True, color=WHITE)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, (label, val) in enumerate(prod_rows):
        row = table.rows[i + 1]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, (cell_text, width) in enumerate(zip((label, val), widths)):
            cell = row.cells[j]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "E8EEF6" if label in {"\u533a\u95f4\u6536\u76ca", "\u6700\u5927\u56de\u64a4", "\u76f8\u5bf9\u4e0a\u8bc1\u6307\u6570", "\u76f8\u5bf9\u6caa\u6df1300"} else ("EDF2F7" if j == 0 else ("FFFFFF" if i % 2 == 0 else "F7FAFC")))
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(cell_text), 9, bold=(j == 0 or label in {"\u533a\u95f4\u6536\u76ca", "\u6700\u5927\u56de\u64a4"}), color=NAVY if j == 1 and label in {"\u533a\u95f4\u6536\u76ca", "\u6700\u5927\u56de\u64a4", "\u76f8\u5bf9\u4e0a\u8bc1\u6307\u6570", "\u76f8\u5bf9\u6caa\u6df1300"} else TEXT)
    add_text(doc, "\u672c\u4ea7\u54c1\u4e8e2026\u5e747\u670815\u65e5\u6210\u7acb\uff0c\u8fd0\u4f5c\u65f6\u95f4\u8f83\u77ed\u3002\u540e\u9875\u8865\u5145\u540c\u4e00\u7b56\u7565\u81ea2023\u5e7412\u670829\u65e5\u8d77\u7684\u5386\u53f2\u4e1a\u7ee9\uff0c\u4f9b\u4e86\u89e3\u7b56\u7565\u8fc7\u5f80\u8868\u73b0\u3002", 9, color=SLATE, before=6, after=2)

    sec3 = add_text(doc, "三、金舆追风1号持有期业绩", 13, bold=True, color=NAVY, before=0, after=1)
    sec3.paragraph_format.page_break_before = True
    add_text(doc, "下图为已出具的产品持有期报告，对应投资者所持份额。持有期 2026年7月15日至2026年9月30日，期末单位净值 0.9607。", 9, color=MUTED, before=0, after=3)
    doc.add_picture(str(out_dir / "第2页_产品持有期业绩.png"), width=Cm(15.2))
    picture = doc.paragraphs[-1]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(0)

    sec4 = add_text(doc, "\u56db\u3001\u540c\u7b56\u7565\u5386\u53f2\u4e1a\u7ee9", 13, bold=True, color=NAVY, before=0, after=1)
    sec4.paragraph_format.page_break_before = True
    add_text(doc, "\u4e0b\u56fe\u4e3a\u540c\u4e00\u7b56\u7565\u81ea2023\u5e7412\u670829\u65e5\u81f32026\u5e749\u670830\u65e5\u7684\u5386\u53f2\u4e1a\u7ee9\uff0c\u8865\u5145\u4ea7\u54c1\u6210\u7acb\u4ee5\u6765\u8f83\u77ed\u7684\u8bb0\u5f55\uff0c\u4f9b\u4e86\u89e3\u7b56\u7565\u8fc7\u5f80\u8868\u73b0\u3002", 9, color=MUTED, before=0, after=3)
    doc.add_picture(str(out_dir / "\u7b2c3\u9875_\u540c\u7b56\u7565\u5386\u53f2\u4e1a\u7ee9.png"), width=Cm(16.0))
    picture = doc.paragraphs[-1]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(0)
    picture.paragraph_format.space_after = Pt(0)

    core = doc.core_properties
    core.title = "金舆追风1号私募证券投资基金 2026年第三季度投资报告"
    core.subject = "投资说明、产品持有期业绩与同策略历史净值"
    core.author = BRAND
    core.category = "季度投资报告"
    path = out_dir / "金舆追风1号_2026年第三季度投资报告.docx"
    doc.save(path)
    return path



def load_holding_frame():
    import json
    import sys
    import pandas as pd
    import psycopg2

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "product_quarterly"))
    import generate_product_quarterly as pq

    env = Path(r"d:\coding\market_website\.env.local").read_text(encoding="utf-8")
    url = next(line.split("=", 1)[1].strip().strip('"') for line in env.splitlines() if line.startswith("DATABASE_URL="))
    conn = psycopg2.connect(url)
    cur = conn.cursor()
    cur.execute("SELECT nav_series::text FROM ops_private_fund_detail_nav_cache WHERE beian_hao = 'SCJ536'")
    series = json.loads(cur.fetchone()[0])
    rows = sorted(
        (
            row["price_date"][:10],
            float(row["nav"]),
            float(row.get("cumulative_nav") or row.get("cum_nav_withdrawal") or row["nav"]),
        )
        for row in series
        if row["price_date"][:10] <= "2026-09-30"
    )
    cur.execute(
        """
        SELECT trade_date::text, close::float
        FROM raw_ashare_index_daily
        WHERE ts_code = %s AND trade_date >= %s AND trade_date <= %s AND close > 0
        ORDER BY trade_date
        """,
        ("000001.SH", rows[0][0], "2026-09-30"),
    )
    sse = cur.fetchall()
    cur.execute(
        """
        SELECT DISTINCT ON (trade_date) trade_date::text, close::float
        FROM raw_spot_daily
        WHERE symbol = 'IF'
          AND trade_date >= %s
          AND trade_date <= %s
          AND close IS NOT NULL
          AND close > 0
        ORDER BY trade_date ASC, fetched_at DESC
        """,
        (rows[0][0], "2026-09-30"),
    )
    csi = cur.fetchall()
    conn.close()

    def align(index_rows):
        last = None
        i = 0
        out = []
        for day, _unit, _adj in rows:
            while i < len(index_rows) and index_rows[i][0] <= day:
                last = index_rows[i][1]
                i += 1
            out.append(last)
        return out

    frame = pd.DataFrame({
        "date": pd.to_datetime([row[0] for row in rows]),
        "unit_nav": [row[1] for row in rows],
        "nav": [row[2] for row in rows],
        "bench1": align(sse),
        "bench2": align(csi),
    })
    return frame, pq


def holding_market(nav_df, pq) -> dict[str, str]:
    import pandas as pd

    period_end = pd.Timestamp(HOLDING_CONFIG["end_date"])
    period_begin = pd.Timestamp(HOLDING_CONFIG["period_begin"])
    work = nav_df[nav_df["date"] <= period_end].copy()
    work = pq.apply_nav_frequency(work, str(HOLDING_CONFIG.get("nav_frequency") or "weekly"))
    metrics = pq.calc_metrics(work, period_begin, period_end, "window", "division")
    return {
        "sse": pq._fmt_pct(metrics["bench1"]["ret"]),
        "csi": pq._fmt_pct(metrics["bench2"]["ret"]),
        "xsse": pq._fmt_pct(metrics["excess1"]),
        "xcsi": pq._fmt_pct(metrics["excess2"]),
    }


HOLDING_CONFIG = {
    "product_name": "金舆追风1号私募证券投资基金",
    "short_name": "金舆追风1号",
    "brand_name": "上海金舆资产管理",
    "watermark": "金舆",
    "end_date": "2026-09-30",
    "period_begin": "2026-07-15",
    "inception_date": "2026-07-15",
    "nav_frequency": "weekly",
    "ann_vol_instead_of_return": True,
    "metrics_match_product_page": True,
    "show_bench1": False,
    "bench1_label": "上证指数",
    "bench2_label": "沪深300",
    "overview": {
        "manager": "上海金舆资产管理有限公司",
        "investment_manager": "陈家本",
        "custodian": "招商证券股份有限公司",
        "inception_date": "2026.7.15",
        "product_type": "私募证券投资基金",
    },
}


def build_visuals(out_dir: Path):
    font, font_bold = load_fonts()
    history = load_strategy_history()
    nav_df, pq = load_holding_frame()
    market = holding_market(nav_df, pq)
    fig = render_letter_page(font, font_bold, market)
    fig3 = render_history_page(font, font_bold, history)
    page1_path = out_dir / "\u7b2c1\u9875_\u81f4\u6295\u8d44\u4eba.png"
    page2_path = out_dir / "\u7b2c2\u9875_\u4ea7\u54c1\u6301\u6709\u671f\u4e1a\u7ee9.png"
    page3_path = out_dir / "\u7b2c3\u9875_\u540c\u7b56\u7565\u5386\u53f2\u4e1a\u7ee9.png"
    long_path = out_dir / "\u91d1\u8206\u8ffd\u98ce1\u53f7_2026\u5e74\u7b2c\u4e09\u5b63\u5ea6\u6295\u8d44\u62a5\u544a.png"
    pdf_path = out_dir / "\u91d1\u8206\u8ffd\u98ce1\u53f7_2026\u5e74\u7b2c\u4e09\u5b63\u5ea6\u6295\u8d44\u62a5\u544a.pdf"

    fig2 = pq.build_report_figure(nav_df, HOLDING_CONFIG)
    fig.savefig(page1_path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    fig2.savefig(page2_path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    fig3.savefig(page3_path, dpi=DPI, facecolor=C_BG, pad_inches=0)
    page1 = Image.open(page1_path).convert("RGB")
    page2 = Image.open(page2_path).convert("RGB")
    page3 = Image.open(page3_path).convert("RGB")
    long_im = stitch_long([page1, page2, page3])
    long_im.save(long_path, dpi=(DPI, DPI), optimize=True)

    with PdfPages(pdf_path) as pdf:
        pdf.savefig(fig, facecolor=C_BG)
        info = pdf.infodict()
        info["Title"] = "金舆追风1号私募证券投资基金 2026年第三季度投资报告"
        info["Author"] = BRAND
        info["Subject"] = "投资说明、本产品持有期业绩与同策略更长历史"
        pdf.savefig(fig2, facecolor=C_BG)
        pdf.savefig(fig3, facecolor=C_BG)
    plt.close(fig)
    plt.close(fig2)
    plt.close(fig3)
    page1.close()
    page2.close()
    page3.close()
    return page1_path, pdf_path, long_path, market


def main():
    out_dir = prepare_output_dir()
    page1, pdf_path, long_path, market = build_visuals(out_dir)
    docx_path = build_docx(out_dir, market)
    print("REL", market["xsse"], market["xcsi"])
    for path in sorted(out_dir.iterdir()):
        print(f"{path.name}\t{path.stat().st_size}")
    print("OUT", out_dir)
    print("PAGE1", page1)
    print("PDF", pdf_path)
    print("LONG", long_path)
    print("DOCX", docx_path)


if __name__ == "__main__":
    main()
