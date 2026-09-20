#!/usr/bin/env python3
"""Export NAV series for funds listed in find_data/*.xlsx, matching the
private-funds-nav6m zip layout, plus a found/missing status file.
"""
from __future__ import annotations

import csv
import os
import re
import socket
import subprocess
import sys
import time
import zipfile
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import psycopg2

ROOT = Path(r"C:\coding\market_dashboard_website")
FIND_DIR = ROOT / "find_data"
STAMP = date.today().isoformat()
OUT_NAME = f"private-funds-nav-find-data-{STAMP}"
OUT_DIR = ROOT / "data" / "exports" / OUT_NAME
DESKTOP_ZIP = Path(os.environ.get("USERPROFILE", str(ROOT))) / "Desktop" / f"{OUT_NAME}.zip"
PREV_CSV = ROOT / "outside_data" / "团队跟踪_对照私募基金清单.csv"

LOCAL_PORT = 5433
SSH_HOST = "root@8.154.33.143"
DEFAULT_DB_URL = f"postgresql://market_user:2026SmartDashboard%21@127.0.0.1:{LOCAL_PORT}/market_data"

EMPTY = {"", "-", "—", "–", "--", "nan", "none", "null", "n/a", "na", "-%"}
NAME_SUFFIXES = (
    "私募证券投资基金",
    "私募投资基金",
    "私募基金",
    "集合资产管理计划",
    "集合资金信托计划",
    "证券投资基金",
    "资产管理计划",
)
MANAGER_STRIP = (
    "私募证券基金管理有限公司",
    "私募基金管理有限公司",
    "基金管理有限公司",
    "投资管理有限公司",
    "资产管理有限公司",
    "投资咨询有限公司",
    "管理有限公司",
    "有限责任公司",
    "股份有限公司",
    "有限合伙",
    "有限公司",
    "投资管理",
    "资产管理",
    "投资基金",
    "基金管理",
    "私募基金",
    "私募",
    "资管",
    "投资",
    "基金",
    "公司",
)

FILE_BUCKET = {
    "FOF底层汇总（私募）": "FOF底层",
    "团队跟踪（私募）_多资产策略": "多资产",
    "团队跟踪（私募）_套利策略": "套利",
    "团队跟踪（私募）_期权策略": "期权",
    "团队跟踪（私募）_期货策略_主观期货": "商品/主观期货",
    "团队跟踪（私募）_期货策略_量化期货": "商品/CTA",
    "团队跟踪（私募）_期货策略": "商品/CTA",
    "团队跟踪（私募）_股票多头": "股票/股票多头",
    "团队跟踪（私募）_股票对冲": "股票/股票对冲",
}

BUCKET_RANK = {
    "商品/主观期货": 10,
    "商品/CTA": 20,
    "股票/股票多头": 20,
    "股票/股票对冲": 20,
    "套利": 20,
    "期权": 30,
    "多资产": 30,
    "FOF底层": 40,
    "未分类": 90,
}


def log(msg: str) -> None:
    print(msg, flush=True)


def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    for path in (ROOT / ".env.local", ROOT / ".env"):
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            s = line.strip()
            if not s or s.startswith("#") or "=" not in s:
                continue
            k, v = s.split("=", 1)
            env.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    return env


def port_open(port: int, timeout: float = 0.4) -> bool:
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=timeout):
            return True
    except OSError:
        return False


def start_tunnel() -> subprocess.Popen | None:
    if port_open(LOCAL_PORT):
        log(f"Using existing listener on localhost:{LOCAL_PORT}")
        return None
    key = Path(os.environ.get("USERPROFILE", "")) / ".ssh" / "id_ed25519_server"
    if not key.exists():
        raise SystemExit(f"SSH key not found: {key}")
    child = subprocess.Popen(
        [
            "ssh",
            "-i",
            str(key),
            "-L",
            f"{LOCAL_PORT}:127.0.0.1:5432",
            "-N",
            "-o",
            "StrictHostKeyChecking=accept-new",
            "-o",
            "ExitOnForwardFailure=yes",
            SSH_HOST,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    deadline = time.time() + 25
    while time.time() < deadline:
        if child.poll() is not None:
            raise SystemExit("SSH tunnel exited before the port opened")
        if port_open(LOCAL_PORT):
            log(f"SSH tunnel ready on localhost:{LOCAL_PORT}")
            return child
        time.sleep(0.4)
    child.kill()
    raise SystemExit("SSH tunnel did not open localhost:5433 within 25s")


def trim(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in EMPTY else s


def iso(v: object) -> str:
    s = trim(v)
    return s[:10] if s else ""


def norm_name(s: str) -> str:
    s = (s or "").strip().replace("\u3000", "").replace("\xa0", "")
    s = re.sub(r"\s+", "", s)
    s = s.replace("（", "(").replace("）", ")").replace("【", "[").replace("】", "]")
    return s.lower()


def core_name(s: str) -> str:
    n = norm_name(s)
    for suf in NAME_SUFFIXES:
        suf_n = suf.lower()
        if n.endswith(suf_n):
            n = n[: -len(suf_n)]
    n = re.sub(r"[()\[\]·•\-_/]", "", n)
    return n


def manager_key(s: str) -> str:
    n = norm_name(s)
    for suf in MANAGER_STRIP:
        n = n.replace(suf.lower(), "")
    return n[:8]


def managers_overlap(a: str, b: str) -> bool:
    ka, kb = manager_key(a), manager_key(b)
    if not ka or not kb:
        return False
    return ka == kb or ka.startswith(kb) or kb.startswith(ka)


def csv_escape(v: object) -> str:
    s = "" if v is None else str(v)
    return f"\"{s.replace('\"', '\"\"')}\"" if re.search(r'["\n\r,]', s) else s


def write_utf8_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [",".join(csv_escape(h) for h in header)]
    lines.extend(",".join(csv_escape(c) for c in row) for row in rows)
    path.write_text("\ufeff" + "\n".join(lines) + "\n", encoding="utf-8")


def safe_file_part(name: str, max_len: int = 60) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]', "_", name or "")
    cleaned = re.sub(r"\s+", "", cleaned).rstrip(".")
    return (cleaned or "unnamed")[:max_len]


def bucket_from_file(filename: str, team_l1: str, team_l2: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"_\d{8,}$", "", stem)
    for prefix, bucket in FILE_BUCKET.items():
        if stem.startswith(prefix):
            if prefix == "团队跟踪（私募）_期货策略" and "主观" in (team_l2 + team_l1):
                return "商品/主观期货"
            return bucket
    labeled = f"{team_l1} {team_l2}"
    if "套利" in labeled:
        return "套利"
    if "主观期货" in labeled:
        return "商品/主观期货"
    if "股票对冲" in labeled:
        return "股票/股票对冲"
    if "股票多头" in labeled or "指数增强" in labeled:
        return "股票/股票多头"
    if "期权" in labeled:
        return "期权"
    if "多资产" in labeled:
        return "多资产"
    if "期货" in labeled or "CTA" in labeled.upper():
        return "商品/CTA"
    return "未分类"


def better_bucket(cur: str, nxt: str) -> str:
    return nxt if BUCKET_RANK.get(nxt, 99) < BUCKET_RANK.get(cur, 99) else cur


def read_find_data() -> list[dict]:
    rows: list[dict] = []
    for path in sorted(FIND_DIR.glob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [str(x or "").strip() for x in next(ws.iter_rows(max_row=1, values_only=True))]
        idx = {h: i for i, h in enumerate(header)}
        for raw in ws.iter_rows(min_row=2, values_only=True):
            name = trim(raw[idx["产品名称"]]) if "产品名称" in idx else ""
            if not name or name in {"合计", "总计", "汇总"}:
                continue
            l1 = trim(raw[idx["团队一级策略"]]) if "团队一级策略" in idx else trim(raw[idx["平台一级策略"]]) if "平台一级策略" in idx else ""
            l2 = trim(raw[idx["团队二级策略"]]) if "团队二级策略" in idx else trim(raw[idx["平台二级策略"]]) if "平台二级策略" in idx else ""
            l3 = trim(raw[idx["团队三级策略"]]) if "团队三级策略" in idx else trim(raw[idx["平台三级策略"]]) if "平台三级策略" in idx else ""
            rows.append(
                {
                    "source_file": path.name,
                    "seq": trim(raw[idx["序号"]]) if "序号" in idx else "",
                    "product_name": name,
                    "manager": trim(raw[idx["基金管理人"]]) if "基金管理人" in idx else "",
                    "sheet_nav": trim(raw[idx["最新单位净值"]]) if "最新单位净值" in idx else "",
                    "sheet_nav_date": iso(raw[idx["最新净值日期"]]) if "最新净值日期" in idx else "",
                    "sheet_cum": trim(raw[idx["最新累计净值"]]) if "最新累计净值" in idx else "",
                    "team_l1": l1,
                    "team_l2": l2,
                    "team_l3": l3,
                    "bucket": bucket_from_file(path.name, l1, l2),
                }
            )
        wb.close()
    return rows


def read_prev_hint() -> dict[str, str]:
    if not PREV_CSV.exists():
        return {}
    out: dict[str, str] = {}
    with PREV_CSV.open(encoding="utf-8-sig", newline="") as fh:
        for r in csv.DictReader(fh):
            name = trim(r.get("product_name"))
            code = trim(r.get("beian_hao")).upper()
            if name and code:
                out.setdefault(norm_name(name), code)
                out.setdefault(core_name(name), code)
    return out


def table_exists(cur, name: str) -> bool:
    cur.execute(
        """
        SELECT 1 FROM information_schema.tables
        WHERE table_schema='public' AND table_name=%s
        """,
        (name,),
    )
    return cur.fetchone() is not None


def table_cols(cur, name: str) -> set[str]:
    cur.execute(
        """
        SELECT column_name FROM information_schema.columns
        WHERE table_schema='public' AND table_name=%s
        """,
        (name,),
    )
    return {r[0] for r in cur.fetchall()}


def connect():
    env = load_env()
    url = env.get("DATABASE_URL") or DEFAULT_DB_URL
    if "127.0.0.1:5433" not in url and "localhost:5433" not in url:
        url = DEFAULT_DB_URL
    return psycopg2.connect(url)


class DbFund:
    __slots__ = ("code", "name", "short", "manager", "nav", "nav_date", "src")

    def __init__(self, code, name, short="", manager="", nav="", nav_date="", src=""):
        self.code = (code or "").strip().upper()
        self.name = name or ""
        self.short = short or ""
        self.manager = manager or ""
        self.nav = nav or ""
        self.nav_date = (nav_date or "")[:10]
        self.src = src


def load_db_catalog(cur):
    funds: list[DbFund] = []
    by_code: dict[str, DbFund] = {}

    def add(item: DbFund) -> None:
        if not item.code:
            return
        prev = by_code.get(item.code)
        if prev is None:
            by_code[item.code] = item
            funds.append(item)
            return
        if not prev.manager and item.manager:
            prev.manager = item.manager
        if not prev.nav_date and item.nav_date:
            prev.nav, prev.nav_date = item.nav, item.nav_date
        if not prev.short and item.short:
            prev.short = item.short

    if table_exists(cur, "private_fund_info"):
        log("Loading private_fund_info ...")
        cur.execute(
            """
            SELECT UPPER(BTRIM(beian_hao)), COALESCE(product_name,''), COALESCE(manager,''),
                   latest_nav::text, latest_nav_date::text
            FROM private_fund_info
            WHERE beian_hao IS NOT NULL AND BTRIM(beian_hao) <> ''
            """
        )
        for code, name, manager, nav, nav_date in cur.fetchall():
            add(DbFund(code, name, "", manager, nav or "", nav_date or "", "info"))
        log(f"  info codes={len(by_code)}")

    if table_exists(cur, "private_fund_info_bfl"):
        log("Loading private_fund_info_bfl ...")
        cur.execute(
            """
            SELECT UPPER(BTRIM(beian_hao)), COALESCE(product_name,''), COALESCE(short_name,'')
            FROM private_fund_info_bfl
            WHERE beian_hao IS NOT NULL AND BTRIM(beian_hao) <> ''
            """
        )
        for code, name, short in cur.fetchall():
            add(DbFund(code, name, short, "", "", "", "bfl"))

    if table_exists(cur, "type6_ops_team_full"):
        log("Loading type6_ops_team_full ...")
        cols = table_cols(cur, "type6_ops_team_full")
        name_col = "product_name" if "product_name" in cols else None
        mgr_col = next((c for c in ("manager", "fund_manager", "company_name") if c in cols), None)
        if name_col:
            extra = f", COALESCE({mgr_col},'')" if mgr_col else ", ''"
            cur.execute(
                f"""
                SELECT UPPER(BTRIM(register_number)), COALESCE({name_col},''){extra}
                FROM type6_ops_team_full
                WHERE register_number IS NOT NULL AND BTRIM(register_number) <> ''
                """
            )
            for code, name, manager in cur.fetchall():
                add(DbFund(code, name, "", manager, "", "", "team"))

    if table_exists(cur, "ops_tracking_funds_list_cache"):
        log("Loading ops_tracking_funds_list_cache ...")
        cols = table_cols(cur, "ops_tracking_funds_list_cache")
        nav_col = "unit_nav" if "unit_nav" in cols else None
        date_col = "nav_date" if "nav_date" in cols else None
        short_col = "short_name" if "short_name" in cols else None
        sel_nav = f", {nav_col}::text" if nav_col else ", NULL"
        sel_date = f", {date_col}::text" if date_col else ", NULL"
        sel_short = f", COALESCE({short_col},'')" if short_col else ", ''"
        cur.execute(
            f"""
            SELECT UPPER(BTRIM(beian_hao)), COALESCE(product_name,''){sel_short}{sel_nav}{sel_date}
            FROM ops_tracking_funds_list_cache
            WHERE beian_hao IS NOT NULL AND BTRIM(beian_hao) <> ''
            """
        )
        for code, name, short, nav, nav_date in cur.fetchall():
            add(DbFund(code, name, short, "", nav or "", nav_date or "", "tracking"))

    log(f"Catalog size: {len(funds)}")
    indexes = build_indexes(funds)
    return funds, indexes, by_code


def build_indexes(funds: list[DbFund]) -> dict[str, dict[str, list[DbFund]]]:
    exact: dict[str, list[DbFund]] = defaultdict(list)
    core: dict[str, list[DbFund]] = defaultdict(list)
    prefix4: dict[str, list[DbFund]] = defaultdict(list)
    seen_exact: set[tuple[str, str]] = set()
    seen_core: set[tuple[str, str]] = set()
    for f in funds:
        for raw in (f.name, f.short):
            if not raw:
                continue
            n = norm_name(raw)
            c = core_name(raw)
            if n and (n, f.code) not in seen_exact:
                exact[n].append(f)
                seen_exact.add((n, f.code))
            if c and (c, f.code) not in seen_core:
                core[c].append(f)
                seen_core.add((c, f.code))
                if len(c) >= 4:
                    prefix4[c[:4]].append(f)
    return {"exact": exact, "core": core, "prefix4": prefix4}


def pick_best(cands: list[DbFund], query_name: str, query_mgr: str) -> DbFund | None:
    if not cands:
        return None
    qn = core_name(query_name)
    uniq: dict[str, DbFund] = {}
    for f in cands:
        uniq.setdefault(f.code, f)
    items = list(uniq.values())

    def score(f: DbFund) -> tuple:
        fn = core_name(f.name)
        fs = core_name(f.short)
        exact = int(fn == qn or fs == qn)
        mgr = int(managers_overlap(query_mgr, f.manager))
        prefix = int(fn.startswith(qn) or qn.startswith(fn) or (fs and (fs.startswith(qn) or qn.startswith(fs))))
        has_nav = int(bool(f.nav_date))
        return (-exact, -mgr, -prefix, -has_nav, -len(fn))

    items.sort(key=score)
    best = items[0]
    # reject weak prefix-only matches when many candidates and no manager/exact
    fn = core_name(best.name)
    fs = core_name(best.short)
    exact = fn == qn or fs == qn
    mgr = managers_overlap(query_mgr, best.manager)
    if not exact and not mgr and len(items) > 8:
        return None
    if not exact and qn and fn and not (fn.startswith(qn) or qn.startswith(fn) or (fs and (fs.startswith(qn) or qn.startswith(fs)))):
        return None
    if not exact and qn and min(len(qn), max(len(fn), len(fs))) < 4:
        return None
    return best


def match_one(name: str, manager: str, indexes, prev_hint: dict[str, str], by_code: dict[str, DbFund]) -> tuple[DbFund | None, str]:
    n = norm_name(name)
    c = core_name(name)
    hits = indexes["exact"].get(n, [])
    if hits:
        picked = pick_best(hits, name, manager)
        if picked:
            return picked, "exact_name"
    hits = indexes["core"].get(c, [])
    if hits:
        picked = pick_best(hits, name, manager)
        if picked:
            return picked, "core_name"
    hint = prev_hint.get(n) or prev_hint.get(c)
    if hint and hint in by_code:
        return by_code[hint], "prev_csv"
    if len(c) >= 4:
        pool = indexes["prefix4"].get(c[:4], [])
        cands = []
        for f in pool:
            fn = core_name(f.name)
            fs = core_name(f.short)
            if fn.startswith(c) or (c.startswith(fn) and len(fn) >= 4) or (fs and (fs.startswith(c) or (c.startswith(fs) and len(fs) >= 4))):
                cands.append(f)
        picked = pick_best(cands, name, manager)
        if picked:
            return picked, "prefix_name"
    if len(c) >= 6:
        pool = []
        for key, items in indexes["core"].items():
            if c in key or (len(key) >= 6 and key in c):
                pool.extend(items)
        picked = pick_best(pool, name, manager)
        if picked:
            fn = core_name(picked.name)
            if c in fn or fn in c:
                return picked, "contains_name"
    return None, "unmatched"


def fetch_nav(cur, codes: list[str]) -> dict[str, list[tuple[str, str, str, str]]]:
    """code -> list of (date, nav, cum, source) already merged by date."""
    out: dict[str, dict[str, tuple[str, str, str, str]]] = defaultdict(dict)
    if not codes:
        return {}

    def ingest(rows, source: str, pri: int) -> None:
        for code, d, nav, cum in rows:
            code = (code or "").upper()
            d = (d or "")[:10]
            if not code or not d or not nav:
                continue
            prev = out[code].get(d)
            if prev is None or pri < prev[3]:  # store pri in hidden slot then rewrite
                out[code][d] = (nav, cum or "", source, pri)

    sources = [
        ("private_fund_nav", "vendor", 0),
        ("private_fund_nav_group_type6", "email_type6", 1),
        ("private_fund_nav_group", "email_group", 2),
        ("private_fund_nav_group_hy", "email_hy", 3),
    ]
    BATCH = 120
    for table, source, pri in sources:
        if not table_exists(cur, table):
            continue
        cols = table_cols(cur, table)
        date_col = next((c for c in ("price_date", "nav_date", "trade_date") if c in cols), None)
        nav_col = next((c for c in ("nav", "unit_nav", "nav_value") if c in cols), None)
        cum_col = next((c for c in ("cumulative_nav", "cum_nav", "acc_nav") if c in cols), None)
        if not date_col or not nav_col or "beian_hao" not in cols:
            continue
        cum_sql = f"{cum_col}::text" if cum_col else "NULL"
        log(f"  NAV {table}: date={date_col} nav={nav_col}")
        for i in range(0, len(codes), BATCH):
            batch = codes[i : i + BATCH]
            cur.execute(
                f"""
                SELECT UPPER(BTRIM(beian_hao)), {date_col}::text, {nav_col}::text, {cum_sql}
                FROM {table}
                WHERE UPPER(BTRIM(beian_hao)) = ANY(%s)
                  AND {nav_col} IS NOT NULL
                """,
                (batch,),
            )
            ingest([(r[0], r[1], r[2], r[3]) for r in cur.fetchall()], source, pri)

    merged: dict[str, list[tuple[str, str, str, str]]] = {}
    for code, by_date in out.items():
        series = []
        for d in sorted(by_date):
            nav, cum, source, _pri = by_date[d]
            series.append((d, nav, cum, source))
        merged[code] = series
    return merged


def fetch_nav_by_names(cur, names: list[str]) -> dict[str, list[tuple[str, str, str, str]]]:
    """product_name -> merged series, for funds whose 备案号 has no history."""
    if not names:
        return {}
    out: dict[str, dict[str, tuple[str, str, str, str]]] = defaultdict(dict)
    sources = [
        ("private_fund_nav", "vendor_name", 0),
        ("private_fund_nav_group_type6", "email_type6_name", 1),
        ("private_fund_nav_group", "email_group_name", 2),
        ("private_fund_nav_group_hy", "email_hy_name", 3),
    ]
    for table, source, pri in sources:
        if not table_exists(cur, table):
            continue
        cols = table_cols(cur, table)
        if "product_name" not in cols:
            continue
        date_col = next((c for c in ("price_date", "nav_date", "trade_date") if c in cols), None)
        nav_col = next((c for c in ("nav", "unit_nav", "nav_value") if c in cols), None)
        cum_col = next((c for c in ("cumulative_nav", "cum_nav", "acc_nav") if c in cols), None)
        if not date_col or not nav_col:
            continue
        cum_sql = f"{cum_col}::text" if cum_col else "NULL"
        cur.execute(
            f"""
            SELECT BTRIM(product_name), {date_col}::text, {nav_col}::text, {cum_sql}
            FROM {table}
            WHERE BTRIM(product_name) = ANY(%s)
              AND {nav_col} IS NOT NULL
            """,
            (names,),
        )
        for pname, d, nav, cum in cur.fetchall():
            d = (d or "")[:10]
            if not pname or not d or not nav:
                continue
            prev = out[pname].get(d)
            if prev is None or pri < prev[3]:
                out[pname][d] = (nav, cum or "", source, pri)
    merged: dict[str, list[tuple[str, str, str, str]]] = {}
    for pname, by_date in out.items():
        merged[pname] = [(d, *by_date[d][:3]) for d in sorted(by_date)]
    return merged


def fetch_fof99(cur, codes: list[str]) -> dict[str, tuple[str, str]]:
    if not codes or not table_exists(cur, "fof99_nav_universe"):
        return {}
    cur.execute(
        """
        SELECT UPPER(BTRIM(reg_code)), COALESCE(policy,''), COALESCE(reason,'')
        FROM fof99_nav_universe
        WHERE UPPER(BTRIM(reg_code)) = ANY(%s)
        """,
        (codes,),
    )
    return {r[0]: (r[1], r[2]) for r in cur.fetchall()}


def classify_status(matched: bool, points: int, info_nav: bool) -> str:
    if points > 0:
        return "已导出净值序列"
    if matched and info_nav:
        return "已匹配但仅有最新净值点"
    if matched:
        return "已匹配备案号但库中无净值"
    return "未能匹配备案号"


def main() -> int:
    tunnel = start_tunnel()
    try:
        sheet_rows = read_find_data()
        log(f"find_data rows={len(sheet_rows)} unique_names={len({r['product_name'] for r in sheet_rows})}")
        prev_hint = read_prev_hint()
        log(f"prev csv hints={len(prev_hint)}")

        conn = connect()
        cur = conn.cursor()
        _funds, indexes, by_code = load_db_catalog(cur)

        # unique export keys: name + manager
        unique: dict[tuple[str, str], dict] = {}
        for r in sheet_rows:
            key = (r["product_name"], r["manager"])
            if key not in unique:
                unique[key] = {
                    **r,
                    "sources": [r["source_file"]],
                    "buckets": [r["bucket"]],
                }
            else:
                unique[key]["sources"].append(r["source_file"])
                unique[key]["buckets"].append(r["bucket"])
                unique[key]["bucket"] = better_bucket(unique[key]["bucket"], r["bucket"])
                if (not unique[key]["sheet_nav_date"]) and r["sheet_nav_date"]:
                    unique[key]["sheet_nav"] = r["sheet_nav"]
                    unique[key]["sheet_nav_date"] = r["sheet_nav_date"]
                    unique[key]["sheet_cum"] = r["sheet_cum"]
                if not unique[key]["team_l1"] and r["team_l1"]:
                    unique[key]["team_l1"] = r["team_l1"]
                    unique[key]["team_l2"] = r["team_l2"]
                    unique[key]["team_l3"] = r["team_l3"]

        units = list(unique.values())
        log(f"unique products={len(units)}")

        matches: list[tuple[dict, DbFund | None, str]] = []
        method_counts = Counter()
        for u in units:
            fund, method = match_one(u["product_name"], u["manager"], indexes, prev_hint, by_code)
            matches.append((u, fund, method))
            method_counts[method] += 1
        log(f"match methods: {dict(method_counts)}")

        codes = sorted({f.code for _, f, _ in matches if f})
        log(f"resolved codes={len(codes)}; fetching NAV ...")
        nav_map = fetch_nav(cur, codes)
        missing_names = [
            u["product_name"]
            for u, f, _m in matches
            if not ((f and nav_map.get(f.code)) if f else False)
        ]
        log(f"codes with NAV series={sum(1 for c in codes if nav_map.get(c))}; name fallback {len(missing_names)}")
        nav_by_name = fetch_nav_by_names(cur, missing_names)
        log(f"name-based NAV hits={len(nav_by_name)}")
        fof = fetch_fof99(cur, codes)
        conn.close()

        series_of: dict[tuple[str, str], list] = {}
        for u, f, _m in matches:
            key = (u["product_name"], u["manager"])
            s = nav_map.get(f.code, []) if f else []
            if not s:
                s = nav_by_name.get(u["product_name"], [])
            series_of[key] = s

        if OUT_DIR.exists():
            import shutil

            shutil.rmtree(OUT_DIR)
        OUT_DIR.mkdir(parents=True, exist_ok=True)

        buckets = [
            "股票/股票多头",
            "股票/股票对冲",
            "商品/主观期货",
            "商品/CTA",
            "套利",
            "期权",
            "多资产",
            "FOF底层",
            "未分类",
        ]
        for b in buckets:
            (OUT_DIR / b).mkdir(parents=True, exist_ok=True)

        product_rows = []
        status_rows = []
        written = 0
        empty_hist = 0
        used_names: set[str] = set()

        for u, fund, method in matches:
            code = fund.code if fund else ""
            series = list(series_of.get((u["product_name"], u["manager"]), []))
            info_nav = bool(fund and (fund.nav_date or fund.nav))
            if not series and fund and fund.nav and fund.nav_date:
                series = [(fund.nav_date, fund.nav, "", "info_latest")]
                tip_only = True
            else:
                tip_only = False
            if not series:
                empty_hist += 1

            file_rel = ""
            if series:
                raw_name = f"{safe_file_part(code or 'NOCODE', 20)}_{safe_file_part(u['product_name'])}.csv"
                if raw_name in used_names:
                    raw_name = f"{safe_file_part(code or 'NOCODE', 20)}_{safe_file_part(u['product_name'], 40)}_{safe_file_part(u['manager'], 12)}.csv"
                used_names.add(raw_name)
                dest = OUT_DIR / u["bucket"] / raw_name
                csv_rows = []
                prev = None
                for d, nav, cum, _src in series:
                    try:
                        nv = float(nav)
                    except (TypeError, ValueError):
                        nv = None
                    chg = ""
                    if nv is not None and prev not in (None, 0):
                        chg = f"{(nv / prev) - 1:.8f}"
                    if nv is not None:
                        prev = nv
                    csv_rows.append([d, nav, cum, chg])
                write_utf8_csv(dest, ["日期", "单位净值", "累计净值", "日涨跌"], csv_rows)
                file_rel = f"{u['bucket']}/{raw_name}".replace("\\", "/")
                written += 1

            sources = sorted(set(u["sources"]))
            first_d = series[0][0] if series else ""
            last_d = series[-1][0] if series else ""
            last_nav = series[-1][1] if series else (fund.nav if fund else "")
            nav_src = ",".join(sorted({s[3] for s in series})) if series else ""
            fof_pol, fof_reason = fof.get(code, ("", ""))
            hist_pts = 0 if tip_only else len(series)
            status = classify_status(bool(fund), hist_pts, info_nav)
            if tip_only and status == "已导出净值序列":
                status = "已匹配但仅有最新净值点（已写入单点）"

            product_rows.append(
                [
                    u["bucket"].split("/")[0],
                    u["bucket"].split("/")[-1],
                    u["bucket"],
                    method,
                    code,
                    u["product_name"],
                    u["manager"],
                    fund.name if fund else "",
                    fund.manager if fund else "",
                    u["team_l1"],
                    u["team_l2"],
                    u["team_l3"],
                    last_nav,
                    last_d,
                    len(series),
                    first_d,
                    nav_src,
                    ";".join(sources),
                    file_rel,
                    status,
                ]
            )
            status_rows.append(
                [
                    ";".join(sources),
                    u["seq"],
                    u["product_name"],
                    u["manager"],
                    u["team_l1"],
                    u["team_l2"],
                    u["team_l3"],
                    u["sheet_nav"],
                    u["sheet_nav_date"],
                    method,
                    code,
                    fund.name if fund else "",
                    "Y" if fund else "N",
                    "Y" if hist_pts else "N",
                    hist_pts,
                    last_nav,
                    last_d,
                    first_d,
                    nav_src,
                    "Y" if code in fof else "N",
                    fof_pol,
                    fof_reason,
                    file_rel,
                    status,
                ]
            )

        write_utf8_csv(
            OUT_DIR / "_产品清单.csv",
            [
                "分类一级",
                "分类二级",
                "分类路径",
                "匹配方式",
                "备案号",
                "产品名称",
                "管理人",
                "数据库产品名称",
                "数据库管理人",
                "团队一级",
                "团队二级",
                "团队三级",
                "最新净值",
                "净值日期",
                "净值点数",
                "最早净值日期",
                "净值来源",
                "来源文件",
                "净值文件",
                "状态",
            ],
            product_rows,
        )
        write_utf8_csv(
            OUT_DIR / "_查找状态.csv",
            [
                "来源文件",
                "原表序号",
                "产品名称",
                "管理人",
                "团队一级",
                "团队二级",
                "团队三级",
                "表格最新净值",
                "表格净值日期",
                "匹配方式",
                "备案号",
                "数据库产品名称",
                "是否匹配备案号",
                "是否有净值序列",
                "净值点数",
                "库内最新净值",
                "库内最新净值日期",
                "最早净值日期",
                "净值来源",
                "是否在火富牛宇宙",
                "火富牛策略",
                "火富牛原因",
                "净值文件",
                "状态",
            ],
            status_rows,
        )

        status_counts = Counter(r[-1] for r in status_rows)
        bucket_counts = Counter(u["bucket"] for u, _, _ in matches)
        missing_only = [r for r in status_rows if r[-1] != "已导出净值序列"]
        write_utf8_csv(
            OUT_DIR / "_缺失清单.csv",
            [
                "产品名称",
                "管理人",
                "团队一级",
                "备案号",
                "数据库产品名称",
                "表格最新净值",
                "表格净值日期",
                "匹配方式",
                "状态",
                "来源文件",
            ],
            [
                [r[2], r[3], r[4], r[10], r[11], r[7], r[8], r[9], r[-1], r[0]]
                for r in missing_only
            ],
        )

        write_utf8_csv(
            OUT_DIR / "_分类汇总.csv",
            ["分类路径", "产品数", "已导出净值文件"],
            [
                [
                    b,
                    str(bucket_counts.get(b, 0)),
                    str(
                        sum(
                            1
                            for u, f, _m in matches
                            if u["bucket"] == b and series_of.get((u["product_name"], u["manager"]))
                        )
                    ),
                ]
                for b in buckets
                if bucket_counts.get(b, 0)
            ],
        )

        have_series = sum(1 for u, _f, _m in matches if series_of.get((u["product_name"], u["manager"])))
        tip_only_n = sum(
            1
            for u, f, _m in matches
            if f
            and not series_of.get((u["product_name"], u["manager"]))
            and f.nav
            and f.nav_date
        )
        matched_no_nav = sum(
            1
            for u, f, _m in matches
            if f
            and not series_of.get((u["product_name"], u["manager"]))
            and not (f.nav and f.nav_date)
        )
        unmatched_n = sum(1 for _u, f, _m in matches if not f)

        readme = [
            "find_data 私募基金净值导出（对照 private-funds-nav6m 格式）",
            f"导出日期：{STAMP}",
            f"源文件目录：find_data（{len(list(FIND_DIR.glob('*.xlsx')))} 个 xlsx）",
            f"源表行数：{len(sheet_rows)}",
            f"去重后产品数：{len(units)}",
            "",
            "每个已找到净值的产品一个 CSV：日期, 单位净值, 累计净值, 日涨跌。",
            "分类文件夹按源表策略归入：股票/股票多头、股票/股票对冲、商品/主观期货、商品/CTA、套利、期权、多资产、FOF底层。",
            "",
            "净值来源优先级：平台 vendor(private_fund_nav) > 邮箱 type6 > 邮箱 group > 邮箱 hy；若都没有则写入 private_fund_info 最新净值单点。",
            "未调用火富牛付费接口，只用库里已有数据。",
            "",
            "状态说明：",
            "  已导出净值序列 = 库中有历史净值，已写入 CSV",
            "  已匹配但仅有最新净值点 = 对上备案号，但只有最新净值、没有序列",
            "  已匹配备案号但库中无净值 = 对上备案号，库中完全没有净值",
            "  未能匹配备案号 = 按产品名称/简称/前缀+管理人未能对上",
            "",
            "详见 _查找状态.csv / _产品清单.csv。",
        ]
        (OUT_DIR / "_说明.txt").write_text("\n".join(readme) + "\n", encoding="utf-8")

        summary = [
            f"源表行数: {len(sheet_rows)}",
            f"去重后产品: {len(units)}",
            f"已匹配备案号: {len(units) - unmatched_n}",
            f"已导出净值序列: {have_series}",
            f"仅最新净值单点: {tip_only_n}",
            f"已匹配无净值: {matched_no_nav}",
            f"未能匹配备案号: {unmatched_n}",
            f"已写净值文件: {written}",
            f"匹配方式: {dict(method_counts)}",
            f"状态: {dict(status_counts)}",
            *[f"  {k}: {v}" for k, v in sorted(bucket_counts.items())],
            f"输出目录: {OUT_DIR}",
        ]
        (OUT_DIR / "_导出日志.txt").write_text("\n".join(summary) + "\n", encoding="utf-8")
        log("\n".join(summary))

        DESKTOP_ZIP.parent.mkdir(parents=True, exist_ok=True)
        if DESKTOP_ZIP.exists():
            DESKTOP_ZIP.unlink()
        log(f"Zipping {DESKTOP_ZIP} ...")
        with zipfile.ZipFile(DESKTOP_ZIP, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for path in OUT_DIR.rglob("*"):
                if path.is_file():
                    zf.write(path, f"{OUT_NAME}/{path.relative_to(OUT_DIR).as_posix()}")
        log(f"ZIP ready: {DESKTOP_ZIP} ({DESKTOP_ZIP.stat().st_size} bytes)")
        return 0
    finally:
        if tunnel is not None:
            tunnel.kill()


if __name__ == "__main__":
    raise SystemExit(main())
