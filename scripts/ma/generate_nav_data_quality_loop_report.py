# -*- coding: utf-8 -*-
"""私募净值数据质量闭环方案（中文 Word）。

把「持续发现、实习生核定、归因、通则与例外」收成一条可运行的流程。
案例与阈值对齐现行代码：净值异常、异常成因、数据概览、跟踪池规则。
"""
from __future__ import annotations

import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.font_manager import FontProperties, fontManager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "data-quality"
CHART_DIR = OUT_DIR / "charts"
REPORT_PATH = OUT_DIR / "私募净值数据质量闭环方案.docx"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
TEXT = RGBColor(0x2D, 0x37, 0x48)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

C_NAVY = "#1A365D"
C_BLUE = "#2B6CB0"
C_TEAL = "#2C7A7B"
C_GOLD = "#B7791F"
C_GREEN = "#276749"
C_RED = "#C53030"
C_ORANGE = "#C05621"
C_PURPLE = "#6B46C1"
C_GRAY = "#4A5568"
C_LINE = "#A0AEC0"
C_BG = "#F7FAFC"
C_SOFT = "#EDF2F7"

_CN: FontProperties | None = None


def configure_matplotlib() -> None:
    global _CN
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.facecolor"] = "white"
    plt.rcParams["savefig.facecolor"] = "white"
    for path in [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyhbd.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
    ]:
        if os.path.isfile(path):
            try:
                fontManager.addfont(path)
                _CN = FontProperties(fname=path)
                plt.rcParams["font.family"] = "sans-serif"
                plt.rcParams["font.sans-serif"] = [_CN.get_name(), "Microsoft YaHei", "SimHei"]
                return
            except Exception:
                continue
    _CN = FontProperties(family="Microsoft YaHei")


def fp(size: float | None = None) -> FontProperties:
    if _CN is None:
        font = FontProperties()
    else:
        font = _CN.copy()
    if size is not None:
        font.set_size(size)
    return font


def set_run_font(run, *, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_text(p, text, *, size=11, bold=False, color=None):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def para(doc, text="", *, size=11, bold=False, color=None, align=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.2
    if align:
        p.alignment = align
    if text:
        add_text(p, text, size=size, bold=bold, color=color or TEXT)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    size = 16 if level == 1 else 13
    for run in p.runs:
        set_run_font(run, size=size, bold=True, color=NAVY)
    return p


def caption(doc, text):
    para(doc, text, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=12)


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="E2E8F0"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margin(cell, *, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_text(cell, text, *, size=9, bold=False, color=None, align="left", fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.08
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, "" if text is None else str(text), size=size, bold=bold, color=color or TEXT)


def add_table(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=9, bold=True, color=WHITE, align="center", fill="1A365D")
    for r_i, row in enumerate(rows):
        fill = "F7FAFC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, size=8.5, fill=fill, align="left")
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Cm(w)
    para(doc, "", size=6, space_after=4)
    return table


def add_picture(doc, path: Path, width_cm=16.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.4)
    shade(cell, "F7FAFC")
    set_cell_border(cell, "CBD5E0")
    set_cell_margin(cell, top=80, bottom=80, left=120, right=120)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    add_text(p1, title, size=11, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    add_text(p2, body, size=10.5, color=TEXT)
    para(doc, "", size=4, space_after=6)


def prevent_row_split(table):
    for row in table.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(hp, "内部方案", size=9, color=NAVY, bold=True)
    add_text(hp, "    私募净值数据质量闭环    ", size=9, color=MUTED)
    add_text(hp, "2026-10-15", size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp_ = footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(fp_, "跟踪池 · 产品页 · 邮件净值 · 火富牛    ", size=9, color=MUTED)
    add_page_number(fp_)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.2


def new_ax(w, h):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    return fig, ax


def rbox(ax, x, y, w, h, text, fc, tc="#FFFFFF", fs=9, ec=None, lw=0.8, weight="regular"):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.2,rounding_size=1.2",
        linewidth=lw,
        edgecolor=ec or fc,
        facecolor=fc,
        mutation_aspect=0.6,
    )
    ax.add_patch(patch)
    font = fp(fs)
    ax.text(
        x + w / 2,
        y + h / 2,
        text,
        ha="center",
        va="center",
        color=tc,
        fontproperties=font,
        linespacing=1.35,
        zorder=3,
    )


def arrow(ax, x1, y1, x2, y2, color=C_LINE, text=None, fs=8, text_color=C_GRAY, rad=0.0):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=11,
            linewidth=1.15,
            color=color,
            connectionstyle=f"arc3,rad={rad}",
            shrinkA=0,
            shrinkB=0,
        )
    )
    if text:
        ax.text((x1 + x2) / 2, (y1 + y2) / 2 + 2.2, text, ha="center", va="bottom", color=text_color, fontproperties=fp(fs))


def save(fig, name: str) -> Path:
    path = CHART_DIR / name
    fig.savefig(path, dpi=160, bbox_inches="tight", pad_inches=0.18)
    plt.close(fig)
    return path


def chart_loop() -> Path:
    fig, ax = new_ax(11.4, 3.55)
    labels = [
        ("发现", "检测器每日跑\n写入待核定"),
        ("核定", "实习生判断\n真问题或误报"),
        ("归因", "系统缺陷\n或源数据四分"),
        ("处置", "改规则、转运营\n或登记例外"),
        ("验收", "同一检测器\n复跑通过"),
        ("回流", "误报与漏报\n改检测器"),
    ]
    colors = [C_NAVY, C_TEAL, C_BLUE, C_ORANGE, C_GREEN, C_PURPLE]
    n = len(labels)
    gap = 1.15
    w = (100 - gap * (n - 1) - 4) / n
    y, h = 28, 36
    xs = []
    for i, ((title, sub), color) in enumerate(zip(labels, colors)):
        x = 2 + i * (w + gap)
        xs.append(x)
        rbox(ax, x, y, w, h, f"{i + 1}  {title}\n{sub}", color, fs=8.2)
        if i < n - 1:
            arrow(ax, x + w + 0.15, y + h / 2, x + w + gap - 0.35, y + h / 2)
    left_x = xs[0] + w / 2
    right_x = xs[-1] + w / 2
    top = y + h
    rail = 82
    ax.plot([right_x, right_x, left_x], [top + 1.2, rail, rail], color=C_PURPLE, lw=1.35, solid_capstyle="round")
    arrow(ax, left_x, rail - 0.2, left_x, top + 1.5, color=C_PURPLE)
    ax.text(50, 90, "回流把新规则、新阈值、新样例送回发现", ha="center", va="center", color=C_PURPLE, fontproperties=fp(9))
    ax.text(50, 12, "验收未通过则退回处置。关闭只发生在同一检测器复跑通过之后。", ha="center", va="center", color=C_GRAY, fontproperties=fp(8.5))
    ax.set_ylim(4, 98)
    return save(fig, "01_loop.png")


def chart_coverage() -> Path:
    fig, ax = plt.subplots(figsize=(11.2, 5.4))
    items = [
        ("净值跳变、尖刺、安静序列尖刺", 1, "已有规则，数据概览可筛"),
        ("异常成因：源数据 / 计算规则 / 未能判断", 1, "已有 diagnoseNavAnomaly"),
        ("净值断档（缺失超过应有净值一成）", 1, "已有，与列表筛选同一规则"),
        ("新鲜度分桶、成立对齐、更新渠道", 1, "数据概览已分桶，尚未自动开单"),
        ("列表日期、涨跌幅与产品页对账", 0, "2026-08、2026-10 两次人工发现"),
        ("同一产品多个备案号 / 估值表科目", 0, "诚奇浦江四行，规则已补，检测未日常化"),
        ("复权 ≥ 累计 ≥ 单位，逐行报警", 0, "计算规则已写明，未进日报"),
        ("邮件「未授权」与托管应到未到", 0, "正文能见到未授权，未单独成单"),
        ("解析失败（邮件入库但抽不出净值）", 0, "格式问题目前混在缺失里"),
    ]
    items = list(reversed(items))
    y = list(range(len(items)))
    colors = [C_TEAL if flag else C_ORANGE for _, flag, _ in items]
    ax.barh(y, [1] * len(items), color=colors, height=0.72, zorder=2)
    ax.set_yticks(y)
    ax.set_yticklabels([name for name, _, _ in items], fontproperties=fp(9))
    ax.set_xlim(0, 1.02)
    ax.set_xticks([])
    for spine in ax.spines.values():
        spine.set_visible(False)
    for yi, (_, flag, note) in zip(y, items):
        tag = "已有素材" if flag else "闭环里还没有"
        ax.text(0.02, yi, f"{tag}    {note}", va="center", ha="left", color="white", fontproperties=fp(8.2))
    ax.set_title("检测覆盖：深青是已经能算的，橙色是还要收成日常工单的", fontproperties=fp(12), color=C_NAVY, pad=12, loc="left")
    fig.tight_layout()
    return save(fig, "02_coverage.png")


def chart_symptom_owner() -> Path:
    fig, ax = new_ax(11.2, 5.6)
    ax.text(50, 96, "先看用户碰到的现象，再决定最可能的责任", ha="center", va="center", color=C_NAVY, fontproperties=fp(12))
    headers = [(8, "现象"), (38, "先问的一句话"), (74, "常见责任")]
    for x, text in [(6, "现象"), (34, "先问的一句话"), (70, "常见责任")]:
        ax.text(x, 86, text, ha="left", va="center", color=C_GRAY, fontproperties=fp(8.5))
    rows = [
        ("净值突然大跌或尖刺", "这根跳变在合并之前的原始列里吗", "两边都有：源数据再四分\n只有合并后才有：系统计算"),
        ("中间缺了一段", "邮件到了没有，解析成功没有", "未授权 / 托管未更 / 格式 / 序列被拆开"),
        ("整段不再更新", "产品页和列表是一起停，还是只有列表停", "一起停：源端或授权\n只有列表停：系统缓存"),
        ("列表和产品页不一致", "两边是否应该读同一条已解析序列", "展示分叉即系统缺陷"),
        ("同一只产品出现多行", "备案号、份额、估值表科目是否同一只", "身份归并，系统缺陷"),
    ]
    colors = [C_RED, C_GOLD, C_ORANGE, C_BLUE, C_PURPLE]
    for i, ((a, b, c), color) in enumerate(zip(rows, colors)):
        y = 72 - i * 14
        rbox(ax, 3, y, 28, 11, a, color, fs=8.4)
        rbox(ax, 34, y, 32, 11, b, C_SOFT, tc=C_NAVY, fs=8, ec="#CBD5E0", lw=0.6)
        rbox(ax, 69, y, 28, 11, c, "#EBF4FF" if i >= 3 else "#FFFAF0", tc=C_NAVY, fs=7.6, ec="#CBD5E0", lw=0.6)
    ax.set_ylim(0, 100)
    return save(fig, "03_symptom.png")


def chart_cause_l1() -> Path:
    fig, ax = new_ax(11.2, 6.3)
    ax.text(50, 96, "归因第一层：这根坏数是算出来的，还是带进来的", ha="center", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 32, 82, 36, 9, "检测器报出一条具体的点或一段缺失", C_NAVY, fs=9)
    arrow(ax, 50, 82, 50, 76)
    rbox(ax, 22, 64, 56, 11, "列表、产品页、原始列，读到的是不是同一条序列？", C_BLUE, fs=9.5)
    arrow(ax, 28, 64, 16, 52, text="展示已经分叉")
    arrow(ax, 72, 64, 78, 52, text="序列一致，再看跳变从哪来")
    rbox(ax, 3, 36, 26, 15, "系统缺陷\n身份归并、列表缓存、\n复权拼接、份额互填", C_BLUE, fs=8.2)
    rbox(ax, 62, 48, 35, 10, "单一原始列里，跳变或错误值是否已经存在？", "#2A4365", fs=8.2)
    arrow(ax, 68, 48, 55, 30)
    arrow(ax, 79, 48, 79, 30)
    arrow(ax, 92, 48, 96, 30)
    rbox(ax, 36, 16, 26, 13, "原始列里没有\n合并、复权、修正之后才出现\n系统·计算规则", C_BLUE, fs=7.8)
    rbox(ax, 64, 16, 22, 13, "原始列里已有\n进入源数据四分诊\n见下一张图", C_ORANGE, fs=7.8)
    rbox(ax, 88, 16, 11, 13, "读不到\n原始净值\n待查", C_GRAY, fs=7.6)
    ax.set_ylim(8, 100)
    return save(fig, "04_cause_l1.png")


def chart_cause_l2() -> Path:
    fig, ax = new_ax(11.2, 6.5)
    ax.text(50, 96, "归因第二层：源数据里再分四种，出口不同", ha="center", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 28, 82, 44, 9, "跳变或缺失已经写在某一路原始数据里", C_ORANGE, fs=9)
    tests = [
        (3, C_PURPLE, "待授权", "邮件正文或字段出现「未授权」\n或跟踪池产品从未收到该托管净值邮件\n且运营侧没有授权记录", "转运营授权\n不改净值、不记成数据错误"),
        (27, C_GOLD, "格式", "邮件到了，解析失败、列错位\n或累计净值被写入单位净值", "已知模板：改解析器\n新模板：留样例再适配"),
        (51, C_ORANGE, "托管未更新", "授权和格式都正常\n按该产品自己的披露间隔\n已经超过一档还没新日期", "催托管补发\n最新净值保持上一个可信点"),
        (75, C_RED, "数据本身错误", "格式能解析，多源互相矛盾\n或单源自相矛盾\n且对不上分红、份额切换", "隔离该点\n联系管理人或数据源\n不用猜测值覆盖"),
    ]
    for x, color, title, test, action in tests:
        rbox(ax, x, 62, 22, 8, title, color, fs=10)
        arrow(ax, x + 11, 62, x + 11, 56)
        rbox(ax, x, 36, 22, 19, test, C_SOFT, tc=C_NAVY, fs=7.3, ec="#E2E8F0", lw=0.7)
        arrow(ax, x + 11, 36, x + 11, 30)
        rbox(ax, x, 12, 22, 17, action, "#1A365D", fs=7.4)
    ax.set_ylim(6, 100)
    return save(fig, "05_cause_l2.png")


def chart_states() -> Path:
    fig, ax = new_ax(11.2, 4.6)
    ax.text(50, 94, "一张工单只走这些状态，关闭条件写在验收上", ha="center", color=C_NAVY, fontproperties=fp(12))
    nodes = [
        (3, 62, "待核定", C_TEAL),
        (23, 62, "确认", C_NAVY),
        (43, 62, "处置中", C_ORANGE),
        (63, 62, "待验收", C_BLUE),
        (83, 62, "已关闭", C_GREEN),
    ]
    for x, y, text, color in nodes:
        rbox(ax, x, y, 14, 13, text, color, fs=9)
    for x in (17, 37, 57, 77):
        arrow(ax, x, 68.5, x + 6, 68.5)
    rbox(ax, 3, 26, 14, 13, "误报", C_GRAY, fs=9)
    arrow(ax, 8, 62, 8, 40)
    ax.text(10.2, 50, "不成立", color=C_GRAY, fontproperties=fp(8), ha="left")
    rbox(ax, 23, 26, 14, 13, "待授权", C_PURPLE, fs=9)
    ax.plot([14, 14, 30], [62, 50, 50], color=C_PURPLE, lw=1.15, solid_capstyle="round")
    arrow(ax, 30, 49.6, 30, 40, color=C_PURPLE)
    ax.text(22, 53, "转运营", color=C_PURPLE, fontproperties=fp(8), ha="center")
    rbox(ax, 63, 24, 14, 15, "验收失败\n退回处置", C_RED, fs=8)
    arrow(ax, 72, 62, 72, 40, color=C_RED)
    ax.annotate(
        "",
        xy=(57, 66),
        xytext=(63, 32),
        arrowprops=dict(arrowstyle="-|>", color=C_RED, lw=1.15, connectionstyle="arc3,rad=0.28"),
    )
    ax.text(54, 44, "退回", color=C_RED, fontproperties=fp(8), ha="center")
    ax.text(50, 12, "误报关闭时写下检测器哪里过宽。待授权关闭时写上授权完成日期。", ha="center", color=C_GRAY, fontproperties=fp(8))
    ax.set_ylim(4, 100)
    return save(fig, "06_states.png")


def chart_intern_week() -> Path:
    fig, ax = new_ax(11.2, 5.2)
    ax.text(4, 94, "实习生一周做两件不同的事", ha="left", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 4, 48, 42, 38, "", "#E6FFFA", tc=C_NAVY, ec=C_TEAL, lw=1.2)
    rbox(ax, 54, 48, 42, 38, "", "#FFFAF0", tc=C_NAVY, ec=C_ORANGE, lw=1.2)
    ax.text(25, 78, "核定自动报警", ha="center", color=C_TEAL, fontproperties=fp(12))
    ax.text(25, 68, "回答：报出来的是不是真问题\n全量看 S1，看完当日新增 S2\n判误报时写下哪条规则过宽", ha="center", color=C_NAVY, fontproperties=fp(8.5))
    ax.text(75, 78, "覆盖抽检", ha="center", color=C_ORANGE, fontproperties=fp(12))
    ax.text(75, 68, "回答：没报出来的里面还有没有问题\n每周五抽 20 只检测器全绿的跟踪池产品\n对照列表、产品页、最近一封原始邮件", ha="center", color=C_NAVY, fontproperties=fp(8.5))
    rbox(ax, 18, 14, 64, 22, "两路都汇成规则变更单\n误报 → 收紧或加排除条件          漏报 → 新检测器先影子运行 14 天\n待授权 → 转运营，不进工程队列", C_NAVY, fs=8.5)
    arrow(ax, 25, 48, 40, 37)
    arrow(ax, 75, 48, 60, 37)
    ax.set_ylim(4, 100)
    return save(fig, "07_intern.png")


def chart_rules() -> Path:
    fig, ax = new_ax(11.2, 4.8)
    ax.text(50, 94, "能写成通则的进代码和回归，只能个案放过的进登记表", ha="center", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 4, 18, 44, 66, "", "#EBF8FF", tc=C_NAVY, ec=C_BLUE, lw=1.1)
    rbox(ax, 52, 18, 44, 66, "", "#FFFAF0", tc=C_NAVY, ec=C_GOLD, lw=1.1)
    ax.text(26, 74, "通则", ha="center", color=C_BLUE, fontproperties=fp(13))
    ax.text(
        26,
        48,
        "对所有产品默认生效\n改动必须带回归用例\n并写入「不得再出现」\n\n例：列表日期跟上产品页\n复权不低于累计\n同一列表一个产品一行\n跳变阈值与现有检测器一致",
        ha="center",
        va="center",
        color=C_NAVY,
        fontproperties=fp(8.5),
    )
    ax.text(74, 74, "例外", ha="center", color=C_GOLD, fontproperties=fp(13))
    ax.text(
        74,
        48,
        "只对写明的备案号和日期区间生效\n有理由、批准人、到期日\n到期不复审就自动重新打开\n\n例：清盘产品豁免新鲜度\n某托管固定周五披露\n已确认的高净值尺度产品\n不走平台日期覆盖",
        ha="center",
        va="center",
        color=C_NAVY,
        fontproperties=fp(8.5),
    )
    ax.set_ylim(8, 100)
    return save(fig, "08_rules.png")


def chart_roadmap() -> Path:
    fig, ax = plt.subplots(figsize=(11.2, 4.2))
    phases = [
        ("工单与严重度\n接入已有异常、断档、新鲜度", 1, 2, C_TEAL),
        ("列表与产品页对账\n身份重复检测", 3, 5, C_BLUE),
        ("核定台与每周覆盖抽检", 3, 6, C_NAVY),
        ("源数据四分\n待授权、格式、停滞、错误", 6, 9, C_ORANGE),
        ("例外登记、回归包、月度成绩单", 9, 12, C_GREEN),
    ]
    for i, (label, start, end, color) in enumerate(phases):
        y = len(phases) - 1 - i
        width = end - start + 1
        left = start - 0.4
        ax.barh(y, width, left=left, height=0.62, color=color, zorder=2)
        shown = label.replace("\n", "  ")
        if width >= 4:
            ax.text(left + 0.25, y, shown, va="center", ha="left", color="white", fontproperties=fp(8.5))
        else:
            ax.text(left + width + 0.25, y, shown, va="center", ha="left", color=color, fontproperties=fp(8.5))
    ax.set_yticks([])
    ax.set_xlim(0.4, 12.8)
    ax.set_xticks(range(1, 13))
    ax.set_xticklabels([f"第{i}周" for i in range(1, 13)], fontproperties=fp(8), color=C_GRAY)
    ax.set_title("九十天：先把已有检测接进工单，再补对账和源数据四分", fontproperties=fp(12), color=C_NAVY, loc="left", pad=10)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("#E2E8F0")
    ax.tick_params(length=0)
    ax.set_axisbelow(True)
    fig.tight_layout()
    return save(fig, "09_roadmap.png")


def build_doc(charts: dict[str, Path]) -> None:
    doc = Document()
    setup_section(doc)
    core = doc.core_properties
    core.title = "私募净值数据质量闭环方案"
    core.subject = "持续发现、核定、归因、通则与例外"
    core.category = "数据质量"

    para(doc, "内部方案", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "私募净值数据质量闭环", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(
        doc,
        "持续发现 · 人工核定 · 归因分诊 · 通则与例外",
        size=12,
        color=C_TEAL and RGBColor(0x2C, 0x7A, 0x7B),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    para(
        doc,
        "适用对象：跟踪池列表、产品页、邮件净值、火富牛与手动净值　　日期：2026年10月15日",
        size=10,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    callout(
        doc,
        "一句话",
        "把已经会算的净值异常、断档和新鲜度，加上还没日常化的列表对账、身份重复和托管授权，收成同一条工单。"
        "检测器负责每天把问题捞出来。实习生负责判断捞得对不对、还有没有没捞到的。工程只修系统缺陷。运营只处理托管授权和催更。坏点先隔离，再用同一条检测器验收。",
    )

    heading(doc, "1. 要解决的事", 1)
    para(
        doc,
        "使用里反复出现的是五类现象：净值异常大跌或尖刺，中间缺数，整段不再更新，列表和产品页对不上，同一只产品拆成多行。"
        "这些现象会反复出现，因为新邮件、火富牛更新和夜间缓存每天都在写新数。一次修好历史，挡不住下一封格式变了的托管邮件。",
    )
    para(
        doc,
        "因此这套方案管的是一条流水线，而不是一次清洗。每条问题都要能回答四件事：谁发现的，实习生核定后还成不成立，责任在系统和源数据的哪一支，修完之后哪条检测器会再次变绿。",
    )

    heading(doc, "2. 把原先的三步收成闭环", 1)
    para(
        doc,
        "原先的骨架是对的：先搜索，再分类，再修复；搜索要持续自动跑；实习生判断自动结果对不对、覆盖全不全；源数据还要再看出错在数本身、托管没更新、格式，还是运营没做托管授权；修复要有通则，也要允许个案；新数据进来后同一套办法继续转。",
    )
    para(doc, "按这个骨架直接开工，会在六个地方断掉。方案把它们补进流程，而不是另起一套口径。", space_after=6)

    add_table(
        doc,
        ["原先的一步", "直接做会断在哪里", "闭环里怎么补"],
        [
            ["持续搜索", "只加新阈值，和现有跳变、断档各说各话", "检测器目录。已有规则原样接入，缺口单独列项目"],
            ["实习生判断对不对", "只审报警，漏报永远看不见", "核定和覆盖抽检分成两件周工作"],
            ["系统还是源数据", "靠经验贴标签，同一类案例前后判法不同", "两层决策树，每层只有一个问题"],
            ["源数据再细分", "格式问题里混着解析器缺陷和托管换模板", "格式再拆一次：已知模板进工程，新模板留样例"],
            ["通则加个案", "个案写成代码里的永久特殊判断，半年后没人敢动", "例外登记表，必有到期日，到期重新打开"],
            ["修完即结束", "下次缓存或邮件一来，同一症状再出现", "关闭条件是同一检测器复跑通过，并留下回归"],
        ],
        [3.4, 6.4, 6.6],
    )

    para(
        doc,
        "另外补一条使用上的优先级。跟踪池和在管产品上的最新净值、最新涨跌幅先处理。平台全库里的同类问题降级排队。数据质量的目标是投资画面可信，不是把全库每一行都审完。",
    )

    heading(doc, "3. 现有系统已经算到哪一步", 1)
    para(
        doc,
        "数据概览和列表筛选里，已经有一套可以当种子的检测。净值异常看图表序列：相邻涨跌达到 30% 算跳变；大约一个月内翻倍或腰斩也算；相对前后都偏了 15% 以上的尖刺算；平时一步很小、前后又回到原水平的 5% 尖刺也算。图表取值顺序是复权、累计、单位，所以单纯的分红下跌不会先被当成跳变。",
    )
    para(
        doc,
        "成因判断也已经有了。某一路原始列自己就能解释这根跳变，记为源数据异常。合并、复权或修正之后才出现，记为计算或修正导致。原始净值没读到，记为未能判断。原始列包括虚拟净值、邮件净值表、邮件估值表、其他邮件净值、手动净值、火富牛、Type6、分组净值和 HY 净值。",
    )
    para(
        doc,
        "断档、新鲜度、成立对齐、更新渠道同样已经按产品汇总。新鲜度取产品页上团队日期和平台日期里较新的一天，按上海日期分桶。这些结果今天要人打开数据概览才看得见，还不会自己变成待办。",
    )
    add_picture(doc, charts["coverage"], 16.2)
    caption(doc, "图 1    深青：规则已经在系统里。橙色：还停在人工发现或文档记录，没有每日工单。")

    para(
        doc,
        "橙色五项是这九十天要补的检测，不是另一套阈值。列表和产品页对账、身份归并，用的就是跟踪池规则里已经写过的不变量：列表上的最新净值日期必须等于产品页，同一列表里同一只产品只留一行。复权不低于累计、累计不低于单位，用的是净值计算规则里已经写明的不变量。未授权和解析失败，用的是邮件正文里已经出现的字段。",
    )

    heading(doc, "4. 现象和责任分开记", 1)
    para(
        doc,
        "实习生和工程容易把「看起来跌了很多」直接当成「源数据错了」。现象是用户看见的，责任是修的时候要打开的门。同一现象可以走出两扇门。列表涨跌幅突然变成正百分之三十五，可能是源数据，也可能是把单位净值接到了旧的复权上。两条的修法完全不同。",
    )
    add_picture(doc, charts["symptom"], 16.2)
    caption(doc, "图 2    五类现象各自先问一句。展示层分叉和身份重复，责任在系统。跳变和缺失才进入源数据四分。")

    para(doc, "严重度按这只产品会不会出现在投资画面上，不按跌幅绝对值。", space_after=6)
    add_table(
        doc,
        ["级别", "范围", "实习生何时必须看"],
        [
            ["S1", "跟踪池或在管产品。列表与产品页的日期或涨跌幅不一致；最新净值被错误点带偏", "当天清完"],
            ["S2", "跟踪池内的跳变、断档、超过自身披露间隔仍未更新、待授权", "当天处理新增"],
            ["S3", "未进跟踪池的平台数据，同类问题", "每周抽样"],
            ["S4", "要素缺失、成立日偏差，不改变当期净值展示", "有余量再看"],
        ],
        [2.2, 9.6, 4.6],
    )

    heading(doc, "5. 六步闭环", 1)
    para(
        doc,
        "每条工单按六个状态往前走。检测器可以重复报警，工单按「产品 + 检测器 + 问题日期」合并，避免同一根尖刺每天开一张新单。",
    )
    add_picture(doc, charts["loop"], 16.2)
    caption(doc, "图 3    六步闭合。回流箭头把误报、漏报和新样例送回检测器，下一封邮件进来时用的已经是更新后的规则。")

    add_table(
        doc,
        ["步骤", "谁做", "产出", "完成的样子"],
        [
            ["发现", "检测器，净值入库后跑", "待核定工单，附片段和原始列", "S1、S2 出现在当日队列"],
            ["核定", "实习生", "成立、误报，或转授权", "误报写明过宽的条件"],
            ["归因", "实习生给出建议，工程确认系统类", "系统，或源数据四类之一", "决策树的每一问都有答案"],
            ["处置", "工程、运营，或数据负责人批例外", "代码与回归、授权、催更、隔离、例外行", "动作落到唯一出口"],
            ["验收", "同一检测器再跑，实习生抽看画面", "通过或退回", "列表与产品页重新一致"],
            ["回流", "实习生写规则变更单，工程合入", "阈值、样例、不得再出现", "下一周检测器成绩单能看到变化"],
        ],
        [2.0, 4.4, 5.2, 4.8],
    )

    add_picture(doc, charts["states"], 16.2)
    caption(doc, "图 4    状态机。验收失败回到处置。误报和待授权是两条旁路，关闭时都要留下原因。")

    heading(doc, "6. 持续发现", 1)
    heading(doc, "6.1 什么时候跑", 2)
    para(
        doc,
        "检测器挂在写数之后，不挂在人打开页面的时候。三条触发就够：夜间缓存和火富牛写入之后跑全量跟踪池；新的托管邮件或手动净值落库后，只跑涉及的备案号；列表缓存键或归并规则发布后，对跟踪池再跑一遍对账。平台全库每天最多跑一次，结果进 S3，不占用当天的 S1 队列。",
    )
    para(
        doc,
        "交易日历用中国开市日。周末和法定假日不因为「昨天没有新净值」开停滞单。披露频率也不用一个全局的「超过七天就算滞后」。每只产品用自己最近十二段间隔的中位数当作正常步长，超过一个步长再加两个交易日仍没有新日期，才开停滞。周频产品在周四不会被当成断更。数据概览里的「滞后一周」分桶继续保留，它描述的是库的年龄，停滞检测器描述的是这只产品该不该已经有下一笔。",
    )

    heading(doc, "6.2 检测器目录", 2)
    para(doc, "第一批只上这些。每条都写清输入、报警条件、以及它对应哪类现象。", space_after=6)
    add_table(
        doc,
        ["检测器", "报警条件", "现象"],
        [
            ["图表跳变与尖刺", "沿用现有 30%、翻倍、15% 尖刺、安静序列 5% 尖刺", "大跌或尖刺"],
            ["成因标记", "沿用源数据 / 计算规则 / 未能判断，写入工单，不另算阈值", "为归因预备"],
            ["复权链", "任一行复权 < 累计，或累计 < 单位", "计算被分红或拼接打穿"],
            ["断档", "沿用现有规则：运作日到最近净值，交易日缺失超过一成", "中间缺数"],
            ["停滞", "超过该产品自身步长加两个交易日", "失去更新"],
            ["列表对账", "跟踪池列表的日期或涨跌幅与产品页不一致", "页与表不符"],
            ["身份重复", "同一列表里，名称相同且份额或别名可归并，仍出现多行", "一个产品多行"],
            ["待授权", "正文或字段含「未授权」，或应有托管邮件而从未入库", "失去更新的一种"],
            ["解析失败", "邮件已入库，按该托管模板应抽出净值但没有", "格式"],
        ],
        [3.2, 8.6, 4.6],
    )
    para(
        doc,
        "国信一类托管邮件里，「未授权」常常出现在日期和后面的单位、累计之间。现有解析会跳过这两个词，继续读取后面的数字。所以「邮件里有未授权」和「完全没有净值」是两种工单：前者通知运营补授权，已抽出的数字仍要走跳变检测；后者在没有数字时走待授权或停滞。",
    )

    heading(doc, "6.3 新检测器先影子运行", 2)
    para(
        doc,
        "新写的检测器前 14 天只记结果，不进当日待办。实习生每天标一小批。满 14 天并且抽检准确率达到七成、同时至少抓住一个旧目录没报过的真问题，才转正。连续两周准确率低于五成的正式检测器退回影子，避免队列被误报淹没。准确率用实习生核定结果计算：成立的张数除以核定过的张数。",
    )

    heading(doc, "7. 实习生的两件工作", 1)
    para(
        doc,
        "核定回答「报出来的对不对」。覆盖抽检回答「没报的里面还有没有」。两件事情用同一张规则变更单回流，但队列不能混在一起。只做核定，检测器会越来越安静，漏掉的问题一直留在产品页上。",
    )
    add_picture(doc, charts["intern"], 16.2)
    caption(doc, "图 5    左路消耗自动报警，右路主动找漏报。两边都不直接改数据库。")

    heading(doc, "7.1 每天的核定", 2)
    para(
        doc,
        "打开工单时，画面上同时给四样东西：产品页图表在问题日期前后各五个点、列表上的日期和涨跌幅、该日期每一路原始列的单位 / 累计 / 复权、邮件主题或「未授权」片段。实习生只做三选一。",
    )
    add_table(
        doc,
        ["选择", "含义", "接着发生的事"],
        [
            ["成立", "用户画面确实错了，或该有的数确实没有", "进入归因"],
            ["误报", "分红、份额切换、正常周频空档，检测器过宽", "关闭，并写一句过宽条件"],
            ["转运营", "缺的是托管授权，数本身还无法判断", "待授权，通知运营"],
        ],
        [2.6, 7.4, 6.4],
    )
    para(
        doc,
        "第一月的时间按每天一到一个半小时排。只保证 S1 清零和当日新增 S2 被看过。看不完的 S3 留到周五，不从 S1 里挤时间。",
    )

    heading(doc, "7.2 每周的覆盖抽检", 2)
    para(
        doc,
        "周五另留两小时。从跟踪池里抽 20 只「所有正式检测器都是绿的」产品，不是从报警里抽。对照三处：列表行、产品页曲线、最近一封相关邮件或火富牛原始点。发现漏报就按现象补一张工单，来源记为抽检，并写规则变更单。",
    )
    para(
        doc,
        "20 只是第一月的工作量，不是统计意义上的样本量。某一周只要抽到两只同类漏报，当周就开影子检测器，不等凑够一个季度。第二月按实际耗时把 20 调高或调低。抽检名单按托管人轮换，避免每周都抽到同一家、同一套邮件模板。",
    )

    heading(doc, "7.3 实习生不改数", 2)
    para(
        doc,
        "核定、归因建议、规则变更单、转运营，这四件事是这个岗位的全部写权限。改解析、改复权、改列表缓存、改别名表，进工程。给托管发授权、向托管催邮件，进运营。确认某一天的原始净值错误后，隔离这个点的动作由工程按通则执行，实习生在工单上点名要隔离的日期。这样每一笔改动能追溯到一张工单，也不会在核定台上直接改出一条无法回归的手工净值。",
    )

    heading(doc, "8. 归因怎么问", 1)
    para(
        doc,
        "归因只有两层，每层一个问题。第一层决定进工程还是进源数据。第二层只在源数据上问，决定是运营、解析器、催托管，还是隔离坏点。",
    )
    add_picture(doc, charts["cause_l1"], 16.2)
    caption(doc, "图 6    第一层。展示已经分叉，或原始列里没有这根跳变，责任在系统。原始列里已经有，才进入下一张图。")
    add_picture(doc, charts["cause_l2"], 16.2)
    caption(doc, "图 7    第二层四个出口。待授权不改数。格式看模板是新是旧。托管未更新保持上一个可信点。数据本身错误只隔离、不猜测。")

    para(doc, "几个容易判反的边界，写成固定句，避免每次重新讨论。", space_after=6)
    add_table(
        doc,
        ["看到的情况", "判定", "原因"],
        [
            ["列表停在旧日期，产品页已经是新净值", "系统·缓存或列表拼接", "源如果没更新，两边会一起停"],
            ["单位下跌，累计和复权平滑", "规则内分红，误报", "图表本来就优先画复权"],
            ["单位下跌，复权被算到累计下面", "系统·计算规则", "不变量被分红公式打穿"],
            ["累计被写进单位，虚拟净值里单位更低", "格式，常常还夹着邮件优先级", "解析和选源各记一条，两处都修"],
            ["两路原始列同一天差一个数量级", "数据本身错误，先隔离", "在搞清哪一路可信之前不覆盖"],
            ["正文是未授权，后面仍有单位和累计", "待授权，同时保留跳变检测", "授权未完成，已有数字仍可能错"],
            ["未能判断，因为没有原始列", "待查，先补证", "不把「没读到」记成源数据错误"],
        ],
        [6.4, 4.6, 5.4],
    )

    heading(doc, "9. 通则、例外和隔离", 1)
    add_picture(doc, charts["rules"], 16.2)
    caption(doc, "图 8    左边对全部产品生效，右边只对登记过的备案号和日期生效。")

    heading(doc, "9.1 通则", 2)
    para(doc, "下面八条是默认行为。改任何一条，都要带一只已经出过事的产品作回归，并在对应规则文档里留一句「不得再出现」。", space_after=6)
    add_table(
        doc,
        ["编号", "通则"],
        [
            ["G1", "每一行复权净值 ≥ 累计净值 ≥ 单位净值"],
            ["G2", "跟踪池列表的最新净值日期、单位净值和涨跌幅，与产品页同一套解析结果"],
            ["G3", "同一列表中，父代码、份额代码、多写的 S、估值表科目路径，归成一行"],
            ["G4", "图表异常沿用现有阈值，不另设一套百分比"],
            ["G5", "断档沿用「缺失超过一成应有净值」"],
            ["G6", "停滞用产品自身步长；交易日历排除周末和法定假日"],
            ["G7", "同一天的更新渠道优先级与数据概览一致：手动，然后邮箱，然后火富牛。托管净值邮件在合并时优先于火富牛"],
            ["G8", "已确认错误的原始点从图表和最新净值中隔离，原值保留可查"],
        ],
        [1.6, 14.8],
    )
    para(
        doc,
        "G3 依赖已经存在的份额归并和估值表别名表。新的托管科目若没有登记别名，身份检测器会报警。报警本身就是提醒去补别名，而不是在列表里暂时隐藏。",
    )

    heading(doc, "9.2 例外登记", 2)
    para(
        doc,
        "例外是一张表，不是代码里的永久分支。每一行至少有：备案号、检测器编号、生效起止日期、理由、批准人、到期日。到期日最长九十天。到期当天检测器重新打开这只产品，实习生决定续期或取消。没有到期日的行不允许保存。",
    )
    para(
        doc,
        "允许登记的典型例外：已经清盘、不再披露的产品，豁免停滞；某家托管固定在周五发送，停滞步长改成周频；计算规则里已经标明要保留高净值尺度的产品，不走「用平台单位净值去接复权」那条覆盖。这些例外每次续期都要重新写理由，避免一份旧备忘变成永久沉默。",
    )

    heading(doc, "9.3 隔离，而不是改掉原值", 2)
    para(
        doc,
        "源数据被核定为错误时，原始邮件和火富牛行保留。图表、最新净值、跟踪池涨跌幅改读「隔离之后的序列」。隔离记录写明日期、字段、工单号。以后源端把这一天更正了，检测器看到原值变化，工单重新打开，由人决定解除隔离。系统不自动用相邻两天的平均数填上一个新净值。",
    )

    heading(doc, "10. 新数据怎样进入同一条流水线", 1)
    para(
        doc,
        "新的一天的净值不走特殊通道。邮件、火富牛、手动净值一旦写入，就触发第 6 节的增量检测。新出现的问题类型也不走特殊通道：抽检或使用中发现的漏报，先成为规则变更单，再成为影子检测器，14 天后才进入正式队列。这样「修已经存在的问题」和「防明天的新问题」用的是同一套工单，差别只在检测器当时是正式还是影子。",
    )
    para(doc, "一条新问题从被人看见到变成日常检测，步骤固定为下面六拍。", space_after=6)
    add_table(
        doc,
        ["拍", "动作"],
        [
            ["1", "使用或抽检发现一个正式检测器没报的问题，开漏报工单"],
            ["2", "用第 8 节的决策树归因。系统问题直接进工程，不必等新检测器"],
            ["3", "若同类问题还会再来，写检测规格：输入、阈值、反例"],
            ["4", "影子运行 14 天，实习生标注"],
            ["5", "准确率够了就转正，从此进入每日队列"],
            ["6", "对应的回归样例进「不得再出现」。下一次代码改动若让这只产品再次变红，发布不过"],
        ],
        [1.6, 14.8],
    )

    heading(doc, "11. 案例按这条流水线重走一遍", 1)
    para(
        doc,
        "下面六则都是系统里已经发生过、并且规则文档里写了原因的事。用它们检验流程：现象归哪一类，责任走哪一扇门，通则落在哪一条，以后靠哪一个检测器拦住。",
    )

    heading(doc, "11.1 卓尚天道1号：列表涨跌幅 +35%，产品页是 −0.25%", 2)
    para(
        doc,
        "2026-10-13，跟踪池列出 SGT288 单位净值 1.0498，最新涨跌幅 +35.01%。产品页同一天是 −0.25%，由 1.0498 除以 1.0524 得到。平台信息已经写到 10 月 13 日的单位净值，火富牛净值序列停在 9 月 22 日，那天的复权是 0.777548。列表把新的单位净值当成复权，去接旧的复权，于是 1.0498 / 0.777548 − 1。",
    )
    para(
        doc,
        "现象是列表和产品页不一致，同时用户会把它看成异常大涨。第一层问题：两边没有读到同一条序列。责任是系统拼接，不是托管写错了 10 月 13 日。处置落在通则 G2：缺日期时用单位涨跌把上一日复权带过去，禁止把单位净值直接推进复权序列。验收是列表对账检测器变绿，并且浏览器列表缓存升版，否则旧的 +35% 还会在本地留几天。回归样例就是 SGT288 这一对日期。",
    )

    heading(doc, "11.2 大约五千只产品的列表停在 2026-08-28", 2)
    para(
        doc,
        "跟踪产品上一大片最新净值日期是 2026-08-28，少量是 8 月 21 日。打开产品页，常见的是 9 月 30 日或 9 月 24 日。约 5,075 行缓存停在 8 月 28 日，其中约 4,990 行的平台信息日期已经更新。SABF94 是其中一只：缓存 8 月 28 日、净值 1.7762，刷新时间却是 9 月 10 日；净值表和产品页都已是 9 月 30 日、1.7557。",
    )
    para(
        doc,
        "产品页和解析器当时已经能读到新日期。列表在后台重算被关掉之后，一直在发 9 月 10 日写下的缓存，而火富牛仍在更新平台信息。现象是失去更新，但只发生在列表上。判定句用得上：只有列表停、产品页没停，就是系统缓存。修法是每次列表响应都采用较新的平台交易日，并把前进了的日期写回缓存。检测器是列表对账，严重度 S1。它不该开成「托管未更新」，否则运营会去催一家其实已经披露过的托管。",
    )

    heading(doc, "11.3 诚奇浦江CTA1号：一个名字四行", 2)
    para(
        doc,
        "搜索诚奇浦江CTA1号，列表给出四条：SB595A 是 10 月 13 日、约 −0.90%，和产品页一致；SSB595A 多写了一个 S，停在 10 月 9 日、+0.34%；父代码 SSB595 又是另一套数；估值表科目「场外_已上市_开放式_私募_成本.诚奇浦江CTA1号」再占一行，同样停在 +0.34%。",
    )
    para(
        doc,
        "现象是同一产品多行。责任是身份归并，通则 G3。份额代码、多出来的 S、父代码、估值表科目，只要能对上名称和已知别名，就只留一行。别名表里没有的新科目会再次报警，这是预期行为：报警推动补登记，而不是把科目路径显示给投资看。关键字搜索走的是另一条读路径，归并必须在那条路径上同样执行，否则主列表修好之后，搜索框仍会把四行拿回来。",
    )

    heading(doc, "11.4 钜融添宝20号：累计净值被当成单位净值", 2)
    para(
        doc,
        "投资之后，同一只底层有两路邮件。2026-06-23 的 TA 虚拟净值，单位净值 1.3358，这是对的。2026-06-25 的资产净值公告把 1.7600 写进单位净值，这个数实际是累计。系统当时更信任附件、排除虚拟净值，FOF 底层列表和产品页都显示 1.7600，复权还留着旧的 1.3496，落到累计下面。",
    )
    para(
        doc,
        "这是格式问题和系统选源叠在一起，要开两张相关联的工单，不要合成一句「源数据错了」。格式工单：附件里单位等于累计，而虚拟净值历史里两者有缺口，单位应按最近一次虚拟净值的单位/累计比例回推。系统工单：投后虚拟净值的优先级高于资产净值公告。验收时两张都要绿。只改优先级，下一次公告仍会把累计写入单位；只改这一个数，下一只基金还会选错邮件。",
    )

    heading(doc, "11.5 荣熙恒盈2号：分红把复权算穿，又被错误修复放大", 2)
    para(
        doc,
        "SBAH99 在 2026 年劳动节后分红，每单位大约 0.21。第一次计算用累计之比去接复权，除息日当天复权被拉到和累计一样高，之后复权落到累计下面。随后一次修复改成用单位净值的倒数去乘，复权先放大，再在之后每一天继续放大，直到 2.11 左右。",
    )
    para(
        doc,
        "现象像尖刺和失真。第一层：原始分红是真的，失真出现在复权计算之后，责任是系统。同时要有一条误报规则：单位下跌而累计、复权保持平滑时，跳变检测器不应开 S2。复权链检测器才是这只产品的正式报警，它直接检查 G1。这个案例也说明修复本身要被检测器看着。错误的「修复」会制造比原问题更大的跳变，所以计算规则的改动必须先在这只产品的除息窗口上跑绿，才能发布。",
    )

    heading(doc, "11.6 待授权：通知运营，不补一条假净值", 2)
    para(
        doc,
        "托管邮件可以在净值数字旁边带「未授权」。另一种情况是跟踪池里的产品按合同应由某家托管发送净值，库里却从来没有这路邮件。两种都归到待授权。实习生在核定台上转给运营，工单离开工程队列。运营完成授权后，把完成日期写回工单，检测器从下一次邮件开始验收。在授权完成之前，不允许为了让曲线不断，去复制上一周净值或用火富牛暂代却不标明来源。火富牛若已有独立序列，产品页可以继续显示火富牛，更新渠道按现有优先级标明，待授权单仍然开着，直到托管这路真正进来。",
    )

    heading(doc, "12. 谁在什么节奏上工作", 1)
    add_table(
        doc,
        ["角色", "日常", "不承担"],
        [
            ["检测器", "入库后写工单，验收时复跑", "不判定业务对错"],
            ["实习生", "每天核定 S1/S2；周五抽检 20 只；写规则变更单；转运营", "不改库、不改代码、不对外发送授权"],
            ["工程", "系统类工单、解析器、回归、隔离实现、检测器上下线", "不代替运营催托管"],
            ["运营", "托管授权、按工单催更、回写完成日期", "不改复权公式"],
            ["复核", "批准例外，看每月成绩单", "不逐条重审 S2"],
        ],
        [2.4, 8.4, 5.6],
    )
    para(doc, "节奏固定成四档，避免质量工作变成「有空再看」。", space_after=6)
    add_table(
        doc,
        ["节奏", "内容"],
        [
            ["每次入库", "增量检测。新的 S1 当天出现在队列"],
            ["每个工作日", "实习生清 S1，处理新增 S2。工程处理已归因的系统单"],
            ["每周五", "覆盖抽检、规则变更单、影子检测器标注、误报退回"],
            ["每月", "检测器成绩单、到期例外复审、S3 积压是否要升优先级"],
        ],
        [3.2, 13.2],
    )

    heading(doc, "13. 成绩单只看四个数", 1)
    para(doc, "每月一页就够。四个数分别回答发现得准不准、漏得多不多、积压是否在变老、同一类问题有没有杀回来。", space_after=6)
    add_table(
        doc,
        ["指标", "算法", "第一月用来观察的线"],
        [
            ["核定准确率", "成立张数 / 已核定张数，按检测器分开", "低于 50% 退回影子，高于 70% 才允许新检测器转正"],
            ["抽检漏报", "周五 20 只里发现的漏报只数，及是否同类", "同一类出现 2 只就开影子检测器"],
            ["S1 账龄", "从开单到验收通过的工作日", "超过 2 个工作日在成绩单上点名"],
            ["复发", "已关闭工单的回归样例再次变红的次数", "目标是 0。出现 1 次就阻断相关发布"],
        ],
        [3.2, 7.2, 6.0],
    )
    para(
        doc,
        "待授权另计一张小表：未完成张数、最老一张的天数、本月完成张数。它不进「源数据错误率」。授权积压是运营事项，混进数据错误率会让解析问题看起来变多。",
    )

    heading(doc, "14. 九十天怎么落地", 1)
    para(
        doc,
        "不先做完整的新平台。第一月先让已有检测结果变成工单，并让实习生按第 7 节的方式核定。对账和源数据四分排在后面，因为那两项要新写检测器，适合走影子期。",
    )
    add_picture(doc, charts["roadmap"], 16.2)
    caption(doc, "图 9    前六周让人和已有规则先跑起来。后六周补源数据四分、例外表和成绩单。")

    add_table(
        doc,
        ["周", "交付", "怎样算做完"],
        [
            ["1–2", "工单字段、严重度、把跳变、断档、新鲜度导入待核定", "跟踪池当天的 S1/S2 能列出来，并附上现有成因"],
            ["3–5", "列表对账、身份重复。先影子，再转正", "用卓尚天道、八月缓存、诚奇浦江三则样例能报出"],
            ["3–6", "核定记录和周五抽检记录", "每一张误报都有过宽原因，每一只漏报都有规则变更单"],
            ["6–9", "未授权、解析失败、自身步长停滞、复权链", "钜融添宝这类格式问题能和「托管没更新」分开"],
            ["9–12", "例外表、隔离、月度成绩单、回归清单", "例外没有到期日就无法保存；复发数出现在成绩单上"],
        ],
        [2.0, 8.0, 6.4],
    )
    para(
        doc,
        "前两周可以用一张结构化的队列表加现有数据概览链接跑起来，不必等新的页面。核定动作（成立、误报、转运营）要先有地方记。页面可以后做，记录不能后补，否则第一月的准确率无法计算，影子期也没有标注。",
    )

    heading(doc, "15. 范围", 1)
    para(
        doc,
        "这套工单覆盖私募净值，从原始邮件、火富牛、手动净值，到跟踪池列表和产品页。MOM 每日风控和单账户结算单继续用各自的文件格式、导入目录和数据表。净值工单不写入那两套表，也不为了修净值去改它们的解析。份额互填、高净值尺度、估值表别名这些已经写进净值规则的行为，作为通则或例外登记进本方案，避免再长出第三套口径。",
    )

    heading(doc, "附录 A    工单字段", 1)
    add_table(
        doc,
        ["字段", "填写"],
        [
            ["单号 / 合并键", "备案号 + 检测器 + 问题日期。重复报警合并到原单"],
            ["严重度", "S1–S4"],
            ["现象", "跳变、断档、停滞、列表不一致、身份重复、待授权、解析失败、复权链"],
            ["检测器证据", "问题日期前后片段、列表值、产品页值、各原始列"],
            ["核定", "成立 / 误报 / 转运营，加上一句原因"],
            ["归因", "系统·缓存、系统·计算、系统·身份、系统·解析；或待授权、格式·新模板、托管未更新、数据本身错误、待查"],
            ["处置与关联", "代码变更、运营单号、例外行、隔离日期"],
            ["验收", "复跑结果、抽看人、日期"],
            ["回流", "规则变更单编号。误报和漏报必填"],
        ],
        [3.6, 12.8],
    )

    heading(doc, "附录 B    例外登记字段", 1)
    add_table(
        doc,
        ["字段", "约束"],
        [
            ["备案号、检测器", "必填。检测器必须是目录中的一个"],
            ["起止日期", "只覆盖这个区间内的点"],
            ["理由", "写事实，例如清盘日、托管发送日历、高净值尺度"],
            ["批准人、批准日期", "复核人"],
            ["到期日", "距批准日不超过 90 天。空值拒绝保存"],
            ["关联工单", "哪一次核定产生了这次例外"],
        ],
        [4.2, 12.2],
    )

    heading(doc, "附录 C    和现有实现的对应", 1)
    add_table(
        doc,
        ["方案中的名字", "现在的位置"],
        [
            ["图表跳变、尖刺、安静序列尖刺", "lib/server/nav-anomaly.ts 中的 1.3 倍、2 倍、15%、5%"],
            ["源数据 / 计算规则 / 未能判断", "lib/server/nav-anomaly-cause.ts"],
            ["断档、新鲜度、成因汇总", "数据概览，lib/server/ops-data-overview.ts"],
            ["列表必须等于产品页", "docs/tracking-pool-rules.md，2026-10-09 与 2026-10-15"],
            ["一个产品一行", "同一文档 2026-10-14，诚奇浦江CTA1号"],
            ["复权 ≥ 累计 ≥ 单位", "docs/nav-calculation-rules.md"],
            ["虚拟净值与公告优先级、单位回推", "同一文档，钜融添宝20号 SNF018"],
            ["分红后的复权", "同一文档，荣熙恒盈2号 SBAH99"],
            ["原始列种类", "虚拟净值、邮件净值表、邮件估值表、其他邮件、手动、火富牛、Type6、分组、HY"],
        ],
        [6.2, 10.2],
    )

    para(
        doc,
        "方案的使用方式：从第 5 节的六步和第 8 节的两张决策图开始执行。案例用来培训核定，不用来代替检测器。第一张月度成绩单出来之后，再决定哪些阈值要收紧，而不是在开工前再设计一套新的百分比。",
        space_before=8,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)


def main() -> None:
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    charts = {
        "loop": chart_loop(),
        "coverage": chart_coverage(),
        "symptom": chart_symptom_owner(),
        "cause_l1": chart_cause_l1(),
        "cause_l2": chart_cause_l2(),
        "states": chart_states(),
        "intern": chart_intern_week(),
        "rules": chart_rules(),
        "roadmap": chart_roadmap(),
    }
    build_doc(charts)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
