# -*- coding: utf-8 -*-
"""发给实习生的每周工作说明。只保留机器做不到的判断。"""
from __future__ import annotations

import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.font_manager import FontProperties, fontManager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "data-quality"
CHART_DIR = OUT_DIR / "charts" / "intern"
REPORT_PATH = OUT_DIR / "实习生每周工作说明.docx"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
TEXT = RGBColor(0x2D, 0x37, 0x48)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

C_NAVY = "#1A365D"
C_BLUE = "#2B6CB0"
C_TEAL = "#2C7A7B"
C_GOLD = "#B7791F"
C_GREEN = "#276749"
C_RED = "#C53030"
C_ORANGE = "#C05621"
C_PURPLE = "#6B46C1"
C_GRAY = "#4A5568"
C_LINE = "#A0AEC0"
C_SOFT = "#EDF2F7"

_CN: FontProperties | None = None


def configure_matplotlib() -> None:
    global _CN
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.facecolor"] = "white"
    plt.rcParams["savefig.facecolor"] = "white"
    for path in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"]:
        if os.path.isfile(path):
            try:
                fontManager.addfont(path)
                _CN = FontProperties(fname=path)
                plt.rcParams["font.family"] = "sans-serif"
                plt.rcParams["font.sans-serif"] = [_CN.get_name(), "Microsoft YaHei", "SimHei"]
                return
            except Exception:
                continue
    _CN = FontProperties(family="Microsoft YaHei")


def fp(size: float | None = None) -> FontProperties:
    font = FontProperties(family="Microsoft YaHei") if _CN is None else _CN.copy()
    if size is not None:
        font.set_size(size)
    return font


def set_run_font(run, *, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_text(p, text, *, size=11, bold=False, color=None):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def para(doc, text="", *, size=11, bold=False, color=None, align=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.18
    if align:
        p.alignment = align
    if text:
        add_text(p, text, size=size, bold=bold, color=color or TEXT)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    return p


def caption(doc, text):
    para(doc, text, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=10)


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="E2E8F0"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margin(cell, *, top=40, bottom=40, left=70, right=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_text(cell, text, *, size=8.5, bold=False, color=None, align="left", fill=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    add_text(p, "" if text is None else str(text), size=size, bold=bold, color=color or TEXT)


def add_table(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=8.5, bold=True, color=WHITE, align="center", fill="1A365D")
    for r_i, row in enumerate(rows):
        fill = "F7FAFC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, size=8, fill=fill)
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Cm(w)
    para(doc, "", size=4, space_after=2)
    return table


def add_picture(doc, path: Path, width_cm=16.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.width = Cm(16.4)
    shade(cell, "F7FAFC")
    set_cell_border(cell, "CBD5E0")
    set_cell_margin(cell, top=70, bottom=70, left=120, right=120)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    add_text(p1, title, size=11, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    add_text(p2, body, size=10.5, color=TEXT)
    para(doc, "", size=4, space_after=4)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.45)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    add_text(hp, "实习生岗位说明", size=9, color=NAVY, bold=True)
    add_text(hp, "    数据质量 · 每周例行    ", size=9, color=MUTED)
    add_text(hp, "2026-10-15", size=9, color=MUTED)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp_ = footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(fp_, "只做机器下不了的判断    ", size=9, color=MUTED)
    add_page_number(fp_)
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18


def new_ax(w, h):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    return fig, ax


def rbox(ax, x, y, w, h, text, fc, tc="#FFFFFF", fs=9, ec=None, lw=0.8):
    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.2,rounding_size=1.1",
            linewidth=lw,
            edgecolor=ec or fc,
            facecolor=fc,
            mutation_aspect=0.55,
        )
    )
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", color=tc, fontproperties=fp(fs), linespacing=1.35)


def arrow(ax, x1, y1, x2, y2, color=C_LINE):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=11,
            linewidth=1.15,
            color=color,
            shrinkA=0,
            shrinkB=0,
        )
    )


def save(fig, name: str) -> Path:
    path = CHART_DIR / name
    fig.savefig(path, dpi=160, bbox_inches="tight", pad_inches=0.16)
    plt.close(fig)
    return path


def chart_split() -> Path:
    fig, ax = new_ax(11.2, 5.15)
    ax.text(50, 95, "找问题和算数归机器，你只留下机器分不清的那一层", ha="center", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 3, 14, 46, 74, "", "#EBF8FF", tc=C_NAVY, ec=C_BLUE, lw=1.1)
    rbox(ax, 52, 14, 45, 74, "", "#FFFAF0", tc=C_NAVY, ec=C_ORANGE, lw=1.2)
    ax.text(26, 80, "机器做", ha="center", color=C_BLUE, fontproperties=fp(13))
    ax.text(
        26,
        46,
        "扫全库的跳变、尖刺、断档\n给每条异常标上成因\n计算或修正导致的，直接给工程\n列表日期和涨跌幅对不上产品页的，直接给工程\n复权低于累计的，直接给工程\n邮件里写明未授权的，直接给运营\n披露节奏之内没有新净值的，不开单",
        ha="center",
        va="center",
        color=C_NAVY,
        fontproperties=fp(8.3),
    )
    ax.text(74.5, 80, "你做", ha="center", color=C_ORANGE, fontproperties=fp(13))
    ax.text(
        74.5,
        46,
        "每天最多 6 张灰单\n源数据异常再分成四类\n未能判断的，看材料够不够\n周五 8 只盲区，专找没被报出的\n把重复出现的判断收成一句规则\n规则写完，这一类就离开你的队列",
        ha="center",
        va="center",
        color=C_NAVY,
        fontproperties=fp(8.3),
    )
    return save(fig, "01_split.png")


def chart_week() -> Path:
    fig, ax = new_ax(11.2, 3.9)
    ax.text(4, 90, "一周固定这些时段，做完就停", ha="left", color=C_NAVY, fontproperties=fp(12))
    days = ["周一", "周二", "周三", "周四", "周五"]
    for i, day in enumerate(days):
        x = 4 + i * 19
        rbox(ax, x, 70, 16, 10, day, C_NAVY, fs=10)
        if i < 4:
            rbox(ax, x, 28, 16, 36, "灰单\n≤ 40 分钟\n最多 6 张\n队列空了就结束", C_TEAL, fs=8)
        else:
            rbox(ax, x, 52, 16, 12, "盲区 8 只\n40 分钟", C_ORANGE, fs=7.6)
            rbox(ax, x, 36, 16, 12, "规则 ≤ 5 条\n20 分钟", C_BLUE, fs=7.6)
            rbox(ax, x, 20, 16, 12, "授权三行\n10 分钟", C_PURPLE, fs=7.6)
    ax.text(50, 8, "周四若灰单已空，不要提前做周五的 8 只。周五的价值是换一种抽法。", ha="center", color=C_GRAY, fontproperties=fp(8))
    ax.set_ylim(2, 100)
    return save(fig, "02_week.png")


def chart_stop() -> Path:
    fig, ax = new_ax(11.2, 4.3)
    ax.text(4, 92, "四条停手线，碰到就停，不算没做完", ha="left", color=C_NAVY, fontproperties=fp(12))
    rows = [
        (C_TEAL, "时间到了", "40 分钟写完手头这一张就关页面"),
        (C_ORANGE, "张数到了", "当天 6 张灰单写完就停，剩下的留到明天"),
        (C_RED, "同一句话出现第 3 次", "前 3 张照常写。第 4 张起不要再打开，改写一条规则"),
        (C_PURPLE, "材料不够", "句子里写明缺哪一封邮件或哪一列，标待查，当天不要再翻"),
    ]
    for i, (color, title, body) in enumerate(rows):
        y = 74 - i * 18
        rbox(ax, 4, y, 22, 14, title, color, fs=9)
        rbox(ax, 29, y, 67, 14, body, C_SOFT, tc=C_NAVY, fs=9, ec="#E2E8F0", lw=0.6)
    ax.set_ylim(0, 100)
    return save(fig, "03_stop.png")


def chart_friday() -> Path:
    fig, ax = new_ax(11.2, 4.6)
    ax.text(50, 93, "周五的 8 只按格子填，不自己挑眼熟的产品", ha="center", color=C_NAVY, fontproperties=fp(12))
    slots = [
        (3, C_TEAL, "1–3", "全绿\n三个不同托管"),
        (27, C_GOLD, "4–5", "单位跌\n累计仍平"),
        (51, C_ORANGE, "6–7", "源数据异常\n隔行抽 2 只"),
        (75, C_PURPLE, "8", "上周那条规则\n还在不在报"),
    ]
    for x, color, no, text in slots:
        rbox(ax, x, 58, 22, 22, f"{no}\n{text}", color, fs=9)
        arrow(ax, x + 11, 58, x + 11, 50)
    rbox(ax, 14, 18, 72, 30, "每只只回答两句\n列表、产品页、原始列，是不是同一件事？\n有没有一种坏法，现有筛选项里没有名字？\n两句都是否定的，结论写「全绿」，不再加备注。", C_NAVY, fs=9)
    ax.set_ylim(8, 100)
    return save(fig, "04_friday.png")


def chart_doors() -> Path:
    fig, ax = new_ax(11.2, 5.0)
    ax.text(50, 94, "每张单最后只能进一扇门", ha="center", color=C_NAVY, fontproperties=fp(12))
    rbox(ax, 30, 78, 40, 10, "看完规定的三处证据", C_NAVY, fs=10)
    doors = [
        (3, C_GRAY, "误报", "分红或正常空档\n检测器判宽了"),
        (27, C_BLUE, "成立 · 系统", "拼接、缓存、复权\n身份重复"),
        (51, C_ORANGE, "成立 · 源", "格式、托管没更新\n或数本身是错的"),
        (75, C_PURPLE, "转运营", "缺的是托管授权\n你不判断净值对错"),
    ]
    for x, color, title, sub in doors:
        arrow(ax, 50, 78, x + 11, 66)
        rbox(ax, x, 40, 22, 24, f"{title}\n{sub}", color, fs=8)
    rbox(ax, 30, 12, 40, 16, "材料缺一件就写待查\n待查不是第五类业务结论", C_SOFT, tc=C_NAVY, fs=9, ec=C_GRAY, lw=0.8)
    ax.set_ylim(4, 100)
    return save(fig, "05_doors.png")


def build_doc(charts: dict[str, Path]) -> None:
    doc = Document()
    setup_section(doc)
    doc.core_properties.title = "数据质量实习生每周工作说明"
    doc.core_properties.subject = "只做机器下不了的判断"
    doc.core_properties.category = "岗位说明"

    para(doc, "岗位说明", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "数据质量实习生每周工作说明", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "只做机器下不了的判断", size=12, color=RGBColor(0xC0, 0x56, 0x21), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, "生效：2026年10月15日起　　对象：数据质量实习生　　阅读时间约 15 分钟", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    callout(
        doc,
        "你这一周交什么",
        "灰单结论、周五 8 只盲区结论、最多 5 条规则、需要运营动手的授权三行。没有别的周报。处理张数多，不算做得好。同一句话连续三周还要你重写，算没做好。",
    )

    heading(doc, "1. 岗位", 1)
    para(
        doc,
        "系统负责把跟踪池里的净值扫一遍：跳变、尖刺、断档、列表和产品页对不上、复权被算穿。这些计算每天都在做，不需要你逐只打开。你的价值是机器分不清的那一层，以及机器根本没有名字的那一种坏法。",
    )
    para(
        doc,
        "分不清的例子：原始净值里确实有一根大跌，它到底是分红、托管把累计写成了单位、托管这周没发、授权没做，还是数字本身印错了。没有名字的例子：筛选全是绿的，列表和产品页却不是同一个数，或者一种新的邮件格式把列读歪了。",
    )
    add_picture(doc, charts["split"], 16.2)
    caption(doc, "图 1    左边的单不经过你。你每周的时间只放在右边。")

    para(doc, "三个词先记住，后面判断用得上。", space_after=4)
    add_table(
        doc,
        ["词", "你需要知道的"],
        [
            ["单位净值", "分红当天会下跌"],
            ["累计净值", "把已经分掉的钱加回去，分红当天通常不跌"],
            ["复权净值", "产品页曲线默认画它。分红当天应保持平滑，并且不低于累计、累计不低于单位"],
        ],
        [3.2, 13.2],
    )

    heading(doc, "2. 一周怎么过", 1)
    para(
        doc,
        "每周大约四个小时。周一到周四各一次，每次最多 40 分钟。周五 70 分钟，分成三截，到点就停。灰单队列是空的，当天就结束，不要另找产品来看。",
    )
    add_picture(doc, charts["week"], 16.2)
    caption(doc, "图 2    周四的空档不要拿来提前做周五的抽查。两种抽法回答的是不同的问题。")

    add_table(
        doc,
        ["时段", "上限", "做完的标志"],
        [
            ["周一至周四，各 1 次", "40 分钟，且最多 6 张灰单", "当天新开的灰单写完，或碰到停手线"],
            ["周五前 40 分钟", "8 只，按第 6 节的格子", "8 格都有结论，允许某一格写「不足 15 页，停」"],
            ["周五中 20 分钟", "规则最多 5 条", "每条有正例和反例，没有反例的不提交"],
            ["周五后 10 分钟", "只处理还没移交的授权", "每只产品三行：备案号、托管、邮件主题或「从未收到」"],
        ],
        [4.2, 5.4, 6.8],
    )

    heading(doc, "3. 这些单不由你来找", 1)
    para(
        doc,
        "机器负责把问题从全库里捞出来，并给跳变标上「计算或修正导致」或「源数据异常」。你不另开一份名单去找系统拼接、复权算穿或未授权。自动把这些单送进工程和运营的队列还没有每天跑起来。在那之前，你只在自己本来就要打开的灰单和周五 8 只里撞见它们。撞见了写一句交出去，然后停，不为这一类把列表再翻一遍。",
        space_after=4,
    )
    add_table(
        doc,
        ["情形", "谁接手", "你做什么"],
        [
            ["列表上异常成因已经是「计算或修正导致」", "工程", "不打开，不写"],
            ["你打开的这只，列表日期或涨跌幅和产品页不是同一个数", "工程", "写一句「成立·系统拼接或缓存」，不改数"],
            ["你打开的这只，复权低于累计，或累计低于单位", "工程", "写一句「成立·系统计算」，不改公式"],
            ["你打开的这只，邮件里有「未授权」，或应有的托管邮件一封都没有", "运营", "只写第 8 节的三行，不再判断这天净值对不对"],
            ["无异常，而且日期还在这只产品平时的间隔里", "无人", "不打开。周五全绿格子除外"],
            ["不在跟踪池里", "无人", "不打开。平台全库留给机器"],
        ],
        [7.4, 2.2, 6.8],
    )
    para(
        doc,
        "灰单名单的入口是：运维，数据维护，数据概览，范围选「跟踪池」，再点「源数据异常」或「未能判断」。不要从平台数据的全库列表往下翻。",
    )

    heading(doc, "4. 周一到周四：只写灰单", 1)
    heading(doc, "4.1 灰单是什么", 2)
    para(
        doc,
        "灰单只有两种。第一种，数据概览里跟踪池的「源数据异常」：机器已经确定这根跳变在某一路原始净值里，但还不知道是格式、托管没更新，还是数本身错了。第二种，「未能判断」：机器没读到原始列。计算或修正导致的不在这里。",
    )
    para(
        doc,
        "打开顺序固定。数据概览选跟踪池，点进对应数字，列表从上到下。不跳着点眼熟的管理人，不按跌幅大小重新排序。当天从昨天停下的备案号接着往下。",
    )

    heading(doc, "4.2 每张单只看三处", 2)
    para(doc, "看完这三处就要下结论。三处之外的页面，只有标成待查之后、缺的那一件到了，才允许再打开。", space_after=4)
    add_table(
        doc,
        ["处", "看什么", "时间"],
        [
            ["列表这一行", "异常日期、涨跌、小图、数据源、成因说明", "1 分钟"],
            ["产品页同一日期", "单位、累计、复权各是多少；列表上的日期和涨跌幅是否同一套", "2 分钟"],
            ["最近一封相关邮件", "只在你倾向于「格式」或「转运营」时打开。看主题、日期、正文里有没有未授权、单位和累计是不是同一列", "最多 3 分钟"],
        ],
        [3.4, 9.6, 3.4],
    )
    para(doc, "三分钟内找不到那封邮件，结论写待查，句子写「缺邮件原文」。不要到邮箱里按产品名搜索一轮。")

    heading(doc, "4.3 句子的规格", 2)
    para(
        doc,
        "每张单一句，一行写完。必须有备案号、日期、两个可以核对的数，以及下面九个结论里的一个。没有第二段背景介绍。",
    )
    add_table(
        doc,
        ["结论（只许用这些字）", "句子里必须点名的"],
        [
            ["误报", "为什么这是分红或正常间隔，复权有没有保持平滑"],
            ["成立·系统拼接或缓存", "列表用了哪个数，产品页用了哪个数"],
            ["成立·系统计算", "哪一天复权落到了累计下面，或被放大到了多少"],
            ["成立·身份重复", "哪几个备案号其实是同一只"],
            ["成立·格式", "哪一列被写进了单位净值，正确的单位大约是多少"],
            ["成立·托管未更新", "上一笔日期，这只产品平时隔几天有一笔"],
            ["成立·数据本身错误", "哪一路、哪一天、两个互相矛盾的数"],
            ["转运营", "托管名称，邮件主题，或写明从未收到这路邮件"],
            ["待查", "缺的那一件东西的名字"],
        ],
        [5.2, 11.2],
    )

    add_picture(doc, charts["doors"], 16.2)
    caption(doc, "图 3    四扇门。待查表示证据没齐，不是一种新的业务原因。")

    heading(doc, "4.4 对照着写", 2)
    para(doc, "下面两句是规格，不是要你把历史产品重新审一遍。", space_after=4)
    add_table(
        doc,
        ["如果看到", "结论", "句子"],
        [
            ["SGT288，列表 2026-10-13 涨跌幅约 +35%，产品页约 −0.25%。10-13 的单位是 1.0498，9-22 的复权是 0.7775", "成立·系统拼接或缓存", "10-13 单位 1.0498 被当成复权去接 09-22 复权 0.7775，产品页 −0.25% 是对的"],
            ["单位明显下跌，累计和复权几乎不动", "误报", "除息日单位下跌，累计和复权平滑，这张异常单应关掉"],
            ["附件里单位等于累计，都在 1.76；虚拟净值邮件里单位约 1.34、累计约 1.75", "成立·格式", "公告把累计 1.76 写进了单位，虚拟净值单位 1.34 才是单位"],
            ["正文在日期和净值之间写着未授权，后面没有更新的托管净值", "转运营", "某托管，邮件主题原文，授权未完成"],
        ],
        [5.6, 3.4, 7.4],
    )
    para(
        doc,
        "第一句这种单，若列表上的成因已经是「计算或修正导致」，你直接跳过，不用再写。句子留给机器还没标成系统问题的那些。",
    )

    heading(doc, "5. 停手", 1)
    add_picture(doc, charts["stop"], 16.0)
    caption(doc, "图 4    停手是岗位要求。把剩余名单看完，不增加判断的质量。")
    para(
        doc,
        "同一句话出现第 3 次，指的是结论相同，而且原因能用同一句规则描述。例如连续三只都是「公告把累计写进单位，虚拟净值里有更低的单位」。第 4 只开始，备案号记在规则的例子里即可，不要再打开产品页。规则周五统一写，当天在记录表的「一句」里写「同第 3 张，并入规则」。",
    )
    para(
        doc,
        "当天 6 张写完之后列表里还有灰单，在记录表最下加一行：未看完，下一只备案号是多少。明天从那一只开始。不要为了清空列表超过 40 分钟。列表很长，说明机器把太多源数据异常交给了你，周五用一条规则把其中重复的部分收走。",
    )

    heading(doc, "6. 周五：8 只盲区", 1)
    para(
        doc,
        "周一到周四回答的是：报出来的这张单，机器分的类对不对。周五回答的是：机器认为正常的里面，还有没有它没有名字的问题。所以周五不从灰单列表里挑，也不要挑你这周已经写过的产品，第 8 格除外。",
    )
    add_picture(doc, charts["friday"], 16.2)
    caption(doc, "图 5    八个格子固定。某一格按抽法抽不满，就停下，把实抽只数写上。")

    heading(doc, "6.1 格子怎么填", 2)
    para(doc, "抽的时候按备案号在列表里的现成顺序，从上到下。不按规模、不按跌幅、不按熟不熟悉。", space_after=4)
    add_table(
        doc,
        ["格", "从哪份名单", "怎么取", "打开上限"],
        [
            ["1–3 全绿", "数据概览，跟踪池，点「无异常」", "从上到下。托管人与已取的重复就跳过。取满 3 个不同托管", "最多翻 30 行"],
            ["4–5 分红外形", "同一份无异常名单，接着往下", "打开产品页。最近一段单位跌了、累计几乎没跌，就收下", "最多再打开 15 个产品页"],
            ["6–7 源数据抽查", "跟踪池，「有异常」且「源数据异常」", "跳过本周已经写过结论的。从第 1 只未写过的开始，隔 20 行再取 1 只", "只取 2 只"],
            ["8 规则回看", "上周规则变更单的正例备案号", "打开列表，看同一类异常这周还在不在。没有上周规则，这一格写「无」", "1 只"],
        ],
        [3.0, 4.6, 5.6, 3.2],
    )
    para(
        doc,
        "第 4、5 格若 15 个产品页里一只都没遇上「单位跌、累计平」，两格都写「未遇到」，不要换一个名单继续找。找不到，本身就是本周的结果。",
    )

    heading(doc, "6.2 每只只写两句", 2)
    para(doc, "第一句：列表、产品页、你看到的原始列，日期和净值是否同一套。第二句：有没有一种坏法，是「净值异常 / 异常成因 / 断档」这三个筛选项没有名字的。", space_after=4)
    add_table(
        doc,
        ["两句的结果", "结论怎么填"],
        [
            ["同一套，也没有新的坏法", "全绿。不要再写原因"],
            ["不是同一套", "成立·系统拼接或缓存，或成立·身份重复。按第 4 节的句子规格"],
            ["是同一套，但邮件列读歪了，筛选项只显示缺失或无异常", "成立·格式。这就是一条新规则的材料"],
            ["单位跌、累计平，复权也平", "全绿。分红被机器正确放过了"],
            ["单位跌、累计平，复权掉到累计下面", "成立·系统计算"],
        ],
        [8.0, 8.4],
    )

    heading(doc, "7. 周五的规则，最多五条", 1)
    para(
        doc,
        "把这一周的误报和漏报收成规则。同一句话已经能覆盖的多张单，合成一条。写不满 5 条是正常的。没有可以合并的判断，这一栏写「本周无新规则」，不要为了凑数改写句子。",
    )
    para(doc, "每条四行，少一行就不提交。", space_after=4)
    add_table(
        doc,
        ["行", "写什么"],
        [
            ["机器现在", "它把这种单放进了哪一个筛选，或为什么没报"],
            ["应该", "以后这种单应进哪一扇门，或不应再出现"],
            ["正例", "备案号和日期，来自你这周写过的句子"],
            ["反例", "什么样的看起来很像，但不要算进去。写不出反例，这条规则就还不能交给工程"],
        ],
        [2.6, 13.8],
    )
    para(doc, "例：", space_after=2)
    callout(
        doc,
        "规则示例",
        "机器现在：列表把新的单位净值去除以旧的复权，标成很大的涨跌幅，成因有时还不是「计算或修正导致」。\n"
        "应该：这种单直接进工程，实习生不再打开。\n"
        "正例：SGT288，2026-10-13，1.0498 接在 0.7775 后面。\n"
        "反例：列表和产品页用的是同一套复权，涨跌幅一致。那种不是这条规则。",
    )
    para(
        doc,
        "规则交给工程之后，下一周第 8 格就回看这条正例。机器已经不再把同类单交给你，这条规则就从你的周工作里消失。这是你的工作变少的方式。",
    )

    heading(doc, "8. 授权只写三行", 1)
    para(
        doc,
        "机器已经能从正文里认出「未授权」的，你不要再写。你只处理两种机器还没送走的：灰单被你判成转运营；周五盲区里发现跟踪池产品应该有托管邮件、库里一封都没有。",
    )
    para(doc, "每只产品三行，贴到当天记录的最下面，运营照着就能做。", space_after=4)
    add_table(
        doc,
        ["行", "例"],
        [
            ["备案号和产品名", "SAUV26　邦客鼎成精选"],
            ["托管", "国信托管"],
            ["证据", "邮件主题原文；或写「跟踪池内从未收到该托管净值邮件」"],
        ],
        [3.4, 13.0],
    )
    para(doc, "不写建议用语，不跟进对方有没有回复，不在授权完成前用别的净值把曲线补齐。授权完成日期由运营回写。")

    heading(doc, "9. 记录表", 1)
    para(
        doc,
        "每周一张表，表名用该周周一的日期。一行一只产品。结论列只能填第 4 节的九个词，或「全绿」。一句列写不下第二句时，说明这张单还没看完，回到三处证据里删到一句。",
        space_after=4,
    )
    add_table(
        doc,
        ["列", "填法"],
        [
            ["星期", "周一至周五"],
            ["格子", "灰单 / 全绿 / 分红外形 / 源数据抽查 / 规则回看 / 授权"],
            ["备案号", "列表上的代码"],
            ["结论", "九个词之一，或全绿"],
            ["一句", "按第 4 节。全绿留空"],
            ["并入规则", "是 / 否。同一句话的第 4 张起填是"],
        ],
        [3.2, 13.2],
    )
    para(doc, "一星期正常看起来像这样。", space_after=4)
    add_table(
        doc,
        ["星期", "格子", "备案号", "结论", "一句", "并入规则"],
        [
            ["周一", "灰单", "SGT288", "成立·系统拼接或缓存", "10-13 单位 1.0498 接在 09-22 复权 0.7775 上，产品页 −0.25% 是对的", "否"],
            ["周一", "灰单", "SBAH99", "误报", "除息日单位下跌，累计和复权平滑", "否"],
            ["周二", "灰单", "（第 4 只同类）", "成立·格式", "同第 3 张，并入规则", "是"],
            ["周五", "全绿", "（抽中的备案号）", "全绿", "", "否"],
            ["周五", "规则回看", "SGT288", "全绿", "同类拼接这周没有再进灰单", "否"],
        ],
        [1.6, 2.2, 2.4, 3.6, 5.2, 1.4],
    )

    heading(doc, "10. 这一周怎样算做完", 1)
    add_table(
        doc,
        ["检查", "做完"],
        [
            ["时间", "四天的灰单都有开始和结束时刻，单次不超过 40 分钟"],
            ["灰单", "每张有九个结论之一，一句里有备案号、日期和两个数；或标了并入规则"],
            ["停手", "第 4 张同类没有继续打开产品页"],
            ["周五", "8 格都有字。抽不满的格子写了实抽只数和停的原因"],
            ["规则", "0 到 5 条，每条有正例和反例"],
            ["授权", "需要移交的都是三行，没有夹带别的说明"],
            ["下周会变少", "第 8 格能指出：上周至少一条规则，机器已经不再把同类单放进灰单"],
        ],
        [3.2, 13.2],
    )
    para(
        doc,
        "带教每周五只看三样：规则是不是都有反例，待查是不是都写了缺的那一件，第 8 格有没有把上周的规则销掉。不看你打开过多少只产品。",
        space_before=4,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)


def main() -> None:
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    charts = {
        "split": chart_split(),
        "week": chart_week(),
        "stop": chart_stop(),
        "friday": chart_friday(),
        "doors": chart_doors(),
    }
    build_doc(charts)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
