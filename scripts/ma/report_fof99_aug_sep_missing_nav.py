#!/usr/bin/env python3
"""Funds in fof99_aug_sep_not_in_6m_table.csv with no private_fund_nav series,
plus 火富牛 tracking status. Writes CSV + Word report.
"""
from __future__ import annotations

import csv
import os
import socket
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
import psycopg2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "scripts" / "ma" / "fof99_aug_sep_not_in_6m_table.csv"
OUT_DIR = ROOT / "reports"
CHART_DIR = OUT_DIR / "_fof99_missing_nav_charts"
CSV_OUT = OUT_DIR / "fof99_aug_sep_no_db_nav.csv"
REPORT_PATH = OUT_DIR / "火富牛8-9月缺口产品库内无净值与跟踪状态报告.docx"
TODAY = date(2026, 9, 15)

NAVY = RGBColor(0x1A, 0x36, 0x5D)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
TEXT = RGBColor(0x2D, 0x37, 0x48)
MUTED = RGBColor(0x64, 0x74, 0x8B)
RED = RGBColor(0xC5, 0x30, 0x30)
GREEN = RGBColor(0x2F, 0x85, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

C_NAVY = "#1A365D"
C_TEAL = "#2B6CB0"
C_GOLD = "#C9A227"
C_GREEN = "#2F855A"
C_RED = "#C53030"
C_ORANGE = "#DD6B20"
C_GRAY = "#718096"
C_PURPLE = "#6B46C1"
PALETTE = [C_NAVY, C_TEAL, C_GOLD, C_GREEN, C_ORANGE, C_PURPLE, C_RED, C_GRAY]

POLICY_CN = {
    "weekly": "weekly（每周五付费拉取）",
    "weekly_plus": "weekly_plus（邮件优先，落后才拉）",
    "skip": "skip（已确认火富牛无序列，永不重试）",
    "update_slow": "update_slow（火富牛最新净值偏旧，不周更）",
    "": "未纳入跟踪宇宙",
    None: "未纳入跟踪宇宙",
}

_CN_FONT: FontProperties | None = None


def load_env() -> None:
    for p in (ROOT / ".env.local", ROOT / ".env"):
        if p.exists():
            load_dotenv(p, override=False)


def iso(raw: object) -> str:
    if raw is None:
        return ""
    if hasattr(raw, "isoformat"):
        return str(raw)[:10]
    s = str(raw).strip()[:10]
    return s if len(s) == 10 and s[0].isdigit() else ""


def port_open(port: int = 5433) -> bool:
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=1.5):
            return True
    except OSError:
        return False


def ensure_db() -> None:
    url = os.environ.get("DATABASE_URL")
    if not url:
        os.environ["DATABASE_URL"] = (
            "postgresql://market_user:2026SmartDashboard!@127.0.0.1:5433/market_data"
        )
    if port_open(5433):
        return
    key = Path.home() / ".ssh" / "id_ed25519_server"
    if not key.exists():
        raise SystemExit("DATABASE_URL host not reachable and SSH key missing")
    subprocess.Popen(
        [
            "ssh",
            "-i",
            str(key),
            "-L",
            "5433:127.0.0.1:5432",
            "-N",
            "-o",
            "ExitOnForwardFailure=yes",
            "root@8.154.33.143",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(40):
        time.sleep(0.4)
        if port_open(5433):
            return
    raise SystemExit("SSH tunnel to 5433 failed")


def configure_matplotlib() -> None:
    global _CN_FONT
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.facecolor"] = "white"
    plt.rcParams["axes.facecolor"] = "white"
    for path in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ):
        if os.path.isfile(path):
            try:
                fontManager.addfont(path)
                _CN_FONT = FontProperties(fname=path)
                plt.rcParams["font.family"] = "sans-serif"
                plt.rcParams["font.sans-serif"] = [_CN_FONT.get_name(), "Microsoft YaHei", "SimHei"]
                return
            except Exception:
                continue
    _CN_FONT = FontProperties(family="Microsoft YaHei")


def fp() -> dict:
    return {"fontproperties": _CN_FONT} if _CN_FONT is not None else {}


def set_run_font(run, *, size=11, bold=False, color=None, name="微软雅黑", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_text(p, text, *, size=11, bold=False, color=None, italic=False):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def para(doc, text="", *, size=11, bold=False, color=None, align=None, space_after=8, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.28
    if first_line and align is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align:
        p.alignment = align
    if text:
        add_text(p, text, size=size, bold=bold, color=color or TEXT)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    size = {1: 16, 2: 13, 3: 12}.get(level, 11)
    for run in p.runs:
        set_run_font(run, size=size, bold=True, color=NAVY)
    return p


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E0")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, size=9, bold=False, color=None, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, "" if text is None else str(text), size=size, bold=bold, color=color or TEXT)
    set_cell_border(cell)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=9, bold=True, color=WHITE)
        shade(table.rows[0].cells[i], "1A365D")
    for r_i, row in enumerate(rows):
        bg = "F7FAFC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, size=9, align="center" if c_i else "left")
            shade(table.rows[r_i + 1].cells[c_i], bg)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def is_share_class(name: str, code: str) -> bool:
    n = name or ""
    if any(x in n for x in ("类份额", "A类", "B类", "C类", "D类", "E类")):
        return True
    return bool(code) and code[-1].isalpha() and any(ch.isdigit() for ch in code[:-1])


def fetch_db(codes: list[str]) -> dict:
    conn = psycopg2.connect(os.environ["DATABASE_URL"])
    cur = conn.cursor()
    out: dict = {}

    cur.execute(
        """
        SELECT UPPER(BTRIM(beian_hao)), COUNT(*)::int,
               MIN(price_date), MAX(price_date)
        FROM private_fund_nav
        WHERE UPPER(BTRIM(beian_hao)) = ANY(%s)
        GROUP BY 1
        """,
        (codes,),
    )
    nav = {r[0]: {"nav_rows": r[1], "nav_min": iso(r[2]), "nav_max": iso(r[3])} for r in cur.fetchall()}

    cur.execute(
        """
        SELECT UPPER(BTRIM(beian_hao)), COALESCE(product_name, ''),
               latest_nav_date, latest_nav, COALESCE(manager, '')
        FROM private_fund_info
        WHERE UPPER(BTRIM(beian_hao)) = ANY(%s)
        """,
        (codes,),
    )
    info = {
        r[0]: {
            "in_info": True,
            "info_name": r[1],
            "list_nav_date": iso(r[2]),
            "list_nav": "" if r[3] is None else str(r[3]),
            "manager": r[4] or "",
        }
        for r in cur.fetchall()
    }

    cur.execute(
        """
        SELECT UPPER(BTRIM(reg_code)), policy, COALESCE(reason, ''),
               listed_at, updated_at
        FROM fof99_nav_universe
        WHERE UPPER(BTRIM(reg_code)) = ANY(%s)
        """,
        (codes,),
    )
    uni = {
        r[0]: {
            "in_universe": True,
            "policy": r[1] or "",
            "reason": r[2] or "",
            "listed_at": iso(r[3]),
            "updated_at": iso(r[4]),
        }
        for r in cur.fetchall()
    }

    cur.execute(
        """
        SELECT DISTINCT ON (UPPER(BTRIM(reg_code)))
               UPPER(BTRIM(reg_code)), price_date, status, nav, fetched_at
        FROM fof99_nav_fetch_log
        WHERE UPPER(BTRIM(reg_code)) = ANY(%s)
          AND price_date > DATE '1970-01-01'
        ORDER BY UPPER(BTRIM(reg_code)), fetched_at DESC NULLS LAST, price_date DESC
        """,
        (codes,),
    )
    last_fetch = {
        r[0]: {
            "last_fetch_date": iso(r[1]),
            "last_fetch_status": r[2] or "",
            "last_fetch_nav": "" if r[3] is None else str(r[3]),
            "last_fetched_at": iso(r[4]),
        }
        for r in cur.fetchall()
    }

    cur.execute(
        """
        SELECT DISTINCT ON (UPPER(BTRIM(reg_code)))
               UPPER(BTRIM(reg_code)), price_date, status, nav
        FROM fof99_nav_fetch_log
        WHERE UPPER(BTRIM(reg_code)) = ANY(%s)
          AND status = 'ok'
          AND price_date > DATE '1970-01-01'
        ORDER BY UPPER(BTRIM(reg_code)), price_date DESC
        """,
        (codes,),
    )
    last_ok = {
        r[0]: {"last_ok_date": iso(r[1]), "last_ok_nav": "" if r[3] is None else str(r[3])}
        for r in cur.fetchall()
    }

    cur.execute(
        """
        SELECT UPPER(BTRIM(reg_code)), status
        FROM fof99_nav_fetch_log
        WHERE price_date = DATE '1970-01-01'
          AND UPPER(BTRIM(reg_code)) = ANY(%s)
        """,
        (codes,),
    )
    probe = {r[0]: r[1] or "" for r in cur.fetchall()}

    cur.close()
    conn.close()
    out["nav"] = nav
    out["info"] = info
    out["uni"] = uni
    out["last_fetch"] = last_fetch
    out["last_ok"] = last_ok
    out["probe"] = probe
    return out


def make_charts(missing: list[dict], all_rows: list[dict]) -> dict[str, Path]:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # 1. source file funnel
    n_src = len(all_rows)
    n_has = sum(1 for r in all_rows if int(r["nav_rows"] or 0) > 0)
    n_miss = n_src - n_has
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    labels = ["源文件产品数", "库内已有净值序列", "库内无净值序列"]
    vals = [n_src, n_has, n_miss]
    colors = [C_NAVY, C_GREEN, C_RED]
    bars = ax.bar(labels, vals, color=colors, width=0.55)
    ax.set_title("源文件 vs 库内净值覆盖", **fp())
    ax.set_ylabel("产品数", **fp())
    for i, (lab, b) in enumerate(zip(labels, bars)):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + max(vals) * 0.02, f"{vals[i]:,}", ha="center", **fp())
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, **fp())
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    p = CHART_DIR / "01_coverage.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    paths["coverage"] = p

    # 2. policy among missing (almost all not in universe)
    pol = Counter(r["fof99_policy"] or "未纳入宇宙" for r in missing)
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    keys = ["未纳入宇宙", "weekly"]
    sizes = [pol.get(k, 0) for k in keys]
    bars = ax.bar(keys, sizes, color=[C_ORANGE, C_TEAL], width=0.5)
    ax.set_title("无净值产品的火富牛跟踪策略", **fp())
    ax.set_ylabel("产品数", **fp())
    ax.set_xticks(range(len(keys)))
    ax.set_xticklabels(keys, **fp())
    for b, v in zip(bars, sizes):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + max(sizes) * 0.02, f"{v:,}", ha="center", **fp())
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    p = CHART_DIR / "02_policy.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    paths["policy"] = p

    # 3. missing_reason x has_nav
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    reasons = ["in_table_null_latest_nav_date", "not_in_private_fund_info"]
    reason_cn = {
        "in_table_null_latest_nav_date": "在产品表但最新净值日期为空",
        "not_in_private_fund_info": "不在 private_fund_info",
    }
    has_vals = []
    miss_vals = []
    for reason in reasons:
        sub = [r for r in all_rows if r["src_missing_reason"] == reason]
        has_vals.append(sum(1 for r in sub if int(r["nav_rows"] or 0) > 0))
        miss_vals.append(sum(1 for r in sub if int(r["nav_rows"] or 0) == 0))
    x = range(len(reasons))
    ax.bar(x, has_vals, width=0.38, label="库内有净值", color=C_GREEN)
    ax.bar([i + 0.4 for i in x], miss_vals, width=0.38, label="库内无净值", color=C_RED)
    ax.set_xticks([i + 0.2 for i in x])
    ax.set_xticklabels([reason_cn[r] for r in reasons], **fp())
    ax.set_ylabel("产品数", **fp())
    ax.set_title("源文件缺口原因 vs 库内是否有净值序列", **fp())
    ax.legend(prop=_CN_FONT)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    p = CHART_DIR / "03_reason.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    paths["reason"] = p

    # 4. strategy among missing
    strat = Counter(r["fof99_strategy_one"] or "(空白)" for r in missing)
    top = strat.most_common(8)
    fig, ax = plt.subplots(figsize=(7.6, 4.4))
    labels = [k for k, _ in top][::-1]
    vals = [v for _, v in top][::-1]
    ax.barh(labels, vals, color=C_TEAL)
    ax.set_xlabel("产品数", **fp())
    ax.set_title("无净值产品 · 火富牛一级策略（前8）", **fp())
    ax.set_yticks(range(len(labels)))
    ax.set_yticklabels(labels, **fp())
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    p = CHART_DIR / "04_strategy.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    paths["strategy"] = p

    # 5. share class
    share_n = sum(1 for r in missing if r["is_share_class"] == "yes")
    main_n = len(missing) - share_n
    fig, ax = plt.subplots(figsize=(6.4, 4.0))
    ax.bar(["主份额 / 非份额类", "A/B/C 等份额类"], [main_n, share_n], color=[C_NAVY, C_GOLD], width=0.5)
    ax.set_ylabel("产品数", **fp())
    ax.set_title("无净值产品 · 是否份额类", **fp())
    ax.set_xticks([0, 1])
    ax.set_xticklabels(["主份额 / 非份额类", "A/B/C 等份额类"], **fp())
    for i, v in enumerate([main_n, share_n]):
        ax.text(i, v + max(main_n, share_n) * 0.02, str(v), ha="center", **fp())
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    p = CHART_DIR / "05_share.png"
    fig.savefig(p, dpi=160)
    plt.close(fig)
    paths["share"] = p

    return paths


def write_report(all_rows: list[dict], missing: list[dict], charts: dict[str, Path]) -> None:
    n_src = len(all_rows)
    n_has = n_src - len(missing)
    n_miss = len(missing)
    pol = Counter(r["fof99_policy"] or "未纳入宇宙" for r in missing)
    reason_m = Counter(r["src_missing_reason"] for r in missing)
    in_info_m = sum(1 for r in missing if r["in_private_fund_info"] == "yes")
    share_n = sum(1 for r in missing if r["is_share_class"] == "yes")
    in_uni = sum(1 for r in missing if r["in_fof99_universe"] == "yes")
    month = Counter((r["fof99_price_date"] or "")[:7] for r in missing)
    has_nav_rows = [r for r in all_rows if int(r["nav_rows"] or 0) > 0]
    has_but_null = [
        r
        for r in all_rows
        if int(r["nav_rows"] or 0) > 0 and r["src_missing_reason"] == "in_table_null_latest_nav_date"
    ]

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    p = para(doc, "内部数据核对", size=10, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    p.paragraph_format.first_line_indent = Cm(0)
    t = para(doc, "火富牛 8–9 月净值产品：库内无净值序列与跟踪状态", size=18, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    t.paragraph_format.first_line_indent = Cm(0)
    s = para(
        doc,
        f"核对日 {TODAY.isoformat()}  ·  源文件 {SRC.name}  ·  净值表 public.private_fund_nav  ·  跟踪表 public.fof99_nav_universe",
        size=9,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    s.paragraph_format.first_line_indent = Cm(0)

    heading(doc, "一、结论", 1)
    para(
        doc,
        f"源文件共 {n_src:,} 只产品（火富牛 advancedlist 最新净值日在 2026-08 或 2026-09，且未进入私募基金列表「净值日期 6 个月以内 / 6 个月以上」筛选）。"
        f"以库内 master 净值表 private_fund_nav 是否存在任意一行作为「有净值数据」的标准：其中仅 {n_has:,} 只已有净值序列，"
        f"{n_miss:,} 只完全没有净值行。无净值产品占源文件的 {n_miss / n_src:.1%}。",
    )
    para(
        doc,
        f"火富牛跟踪方面，源文件 {n_src:,} 只里只有 {sum(1 for r in all_rows if r['in_fof99_universe']=='yes'):,} 只在 fof99_nav_universe（全部为 weekly）。"
        f"其中 6 只新发产品已由周五任务写入 1–2 条 9 月初净值，但仍未回填产品表 latest_nav_date，所以仍出现在源文件里；"
        f"另外 1 只 SADG72（凯瑞稳健二号）虽为 weekly（operator override），库内净值表仍为空。"
        f"其余 {n_miss - in_uni:,} 只无净值产品均未纳入跟踪宇宙，周五 ETL 不会覆盖，fof99_nav_fetch_log 也没有任何抓取记录。",
    )

    add_table(
        doc,
        ["口径", "数量", "占源文件"],
        [
            ["源文件产品数", f"{n_src:,}", "100%"],
            ["库内已有净值序列", f"{n_has:,}", f"{n_has / n_src:.1%}"],
            ["库内无净值序列（本报告对象）", f"{n_miss:,}", f"{n_miss / n_src:.1%}"],
            ["无净值且已纳入跟踪宇宙", f"{in_uni:,}", f"{in_uni / n_src:.1%}"],
            ["无净值且未纳入跟踪宇宙", f"{n_miss - in_uni:,}", f"{(n_miss - in_uni) / n_src:.1%}"],
        ],
    )
    doc.add_picture(str(charts["coverage"]), width=Inches(6.1))
    cap = para(doc, "图1  源文件覆盖：有净值 vs 无净值", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.paragraph_format.first_line_indent = Cm(0)

    heading(doc, "二、核对口径", 1)
    para(
        doc,
        "源文件 scripts/ma/fof99_aug_sep_not_in_6m_table.csv 是此前对火富牛 advancedlist 的截取：最新 price_date 落在 2026 年 8 月或 9 月，"
        "但在网站「私募基金」列表按净值日期「6 个月以内」或「6 个月以上」筛选时都进不去。原因只有两类——"
        "产品在 private_fund_info 中但 latest_nav_date 为空，或备案号根本不在 private_fund_info。",
    )
    para(
        doc,
        "列表筛选用的是产品表上的 latest_nav_date，不等于库内是否有历史净值。因此本报告另行查询 private_fund_nav："
        "该备案号一行都没有，才记为「库内无净值数据」。只要存在任意一条净值，即视为已有净值（即使产品表 latest_nav_date 仍为空）。",
    )
    para(
        doc,
        "火富牛跟踪状态取自 fof99_nav_universe.policy：weekly 为每周五 FundMultiPrice 付费拉取；weekly_plus 为邮件通常已覆盖、仅当列表尖落后于周五才拉；"
        "skip 为已探测确认火富牛无序列、永不重试；update_slow 为火富牛最新净值偏旧、不参与周更。表中无行则记为「未纳入跟踪宇宙」。"
        "同时附最近一次非探测抓取日志（fof99_nav_fetch_log）状态。",
    )

    heading(doc, "三、无净值产品的火富牛跟踪状态", 1)
    pol_rows = []
    for k, v in pol.most_common():
        label = POLICY_CN.get(k if k != "未纳入宇宙" else "", k)
        if k == "未纳入宇宙":
            label = "未纳入跟踪宇宙"
        pol_rows.append([label, f"{v:,}", f"{v / n_miss:.1%}"])
    add_table(doc, ["跟踪策略", "无净值产品数", "占无净值"], pol_rows)
    doc.add_picture(str(charts["policy"]), width=Inches(6.0))
    cap = para(doc, "图2  无净值产品按火富牛跟踪策略分布", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.paragraph_format.first_line_indent = Cm(0)

    para(
        doc,
        "无净值集合里没有 weekly_plus、skip、update_slow。跟踪状态几乎是同一句话：尚未进入宇宙。"
        "唯一例外 SADG72 凯瑞稳健二号：policy=weekly，reason 为 operator override，但 private_fund_nav 0 行、抓取日志也没有该备案号的非探测记录。"
        "它会被周五任务选中，当前却还没有写成净值。",
    )
    tracked = [r for r in missing if r["in_fof99_universe"] == "yes"]
    if tracked:
        add_table(
            doc,
            ["备案号", "产品名称", "策略", "reason", "火富牛净值日"],
            [
                [
                    r["beian_hao"],
                    (r["fof99_fund_name"] or "")[:22],
                    r["fof99_policy"],
                    (r["fof99_reason"] or "")[:36],
                    r["fof99_price_date"],
                ]
                for r in tracked
            ],
        )

    heading(doc, "四、与源文件缺口原因的对照", 1)
    para(
        doc,
        f"源文件 {n_src:,} = 产品表日期为空 {sum(1 for r in all_rows if r['src_missing_reason']=='in_table_null_latest_nav_date'):,} "
        f"+ 不在 private_fund_info {sum(1 for r in all_rows if r['src_missing_reason']=='not_in_private_fund_info'):,}。"
        f"不在产品表的全部库内无净值。"
        f"产品表日期为空的子集里，{reason_m.get('in_table_null_latest_nav_date', 0):,} 只净值表也是空的，"
        f"另外 {len(has_nav_rows):,} 只其实已经有净值行（全部是 2026-09 新发、policy=weekly），只是 latest_nav_date 未回填，所以仍被挡在列表日期筛选之外。",
    )
    if has_nav_rows:
        add_table(
            doc,
            ["备案号", "产品名称", "净值行数", "净值区间", "跟踪策略"],
            [
                [
                    r["beian_hao"],
                    (r["fof99_fund_name"] or "")[:22],
                    str(r["nav_rows"]),
                    f"{r['nav_min_date']} ~ {r['nav_max_date']}",
                    r["fof99_policy"] or "未纳入",
                ]
                for r in has_nav_rows
            ],
        )
    doc.add_picture(str(charts["reason"]), width=Inches(6.2))
    cap = para(doc, "图3  源文件缺口原因与库内净值有无", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.paragraph_format.first_line_indent = Cm(0)

    heading(doc, "五、产品结构", 1)
    para(
        doc,
        f"无净值产品按火富牛最新净值月份：8 月 {month.get('2026-08', 0):,} 只，9 月 {month.get('2026-09', 0):,} 只。"
        f"份额类 {share_n:,} 只（{share_n / n_miss:.1%}），与此前观察一致：不少 A/B/C 类份额在协会主表没有独立备案行，"
        "因此既进不了列表筛选，库内也没有净值序列。",
    )
    doc.add_picture(str(charts["share"]), width=Inches(5.4))
    cap = para(doc, "图4  无净值产品是否份额类", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.paragraph_format.first_line_indent = Cm(0)
    doc.add_picture(str(charts["strategy"]), width=Inches(6.1))
    cap = para(doc, "图5  无净值产品火富牛一级策略（前8）", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.paragraph_format.first_line_indent = Cm(0)

    strat = Counter(r["fof99_strategy_one"] or "(空白)" for r in missing)
    add_table(
        doc,
        ["火富牛一级策略", "无净值产品数", "占无净值"],
        [[k, f"{v:,}", f"{v / n_miss:.1%}"] for k, v in strat.most_common(10)],
    )

    heading(doc, "六、明细表（抽样）", 1)
    para(
        doc,
        f"完整 {n_miss:,} 行见同目录 CSV：{CSV_OUT.name}。除 SADG72 外均为未纳入宇宙，下表按源文件原因各抽 5 只。",
        first_line=True,
    )
    sample = []
    for key in ("in_table_null_latest_nav_date", "not_in_private_fund_info"):
        cand = [r for r in missing if r["src_missing_reason"] == key]
        sample.extend(cand[:5])
    add_table(
        doc,
        ["备案号", "产品简称", "火富牛净值日", "跟踪策略", "源文件原因"],
        [
            [
                r["beian_hao"],
                (r["fof99_short_name"] or r["fof99_fund_name"] or "")[:18],
                r["fof99_price_date"],
                r["fof99_policy"] or "未纳入",
                "不在产品表" if r["src_missing_reason"] == "not_in_private_fund_info" else "产品表日期空",
            ]
            for r in sample
        ],
    )

    heading(doc, "七、建议", 1)
    para(
        doc,
        f"1. 这 {n_miss:,} 只里 {n_miss - in_uni:,} 只不在跟踪宇宙，周五任务不会自动补齐。若要对主份额（约 {n_miss - share_n:,} 只非份额类）建序列，"
        "需按 admit 流程写入 fof99_nav_universe（weekly），先用 FundPrice 回填历史再进周五 FundMultiPrice。不要对整表付费盲拉。",
    )
    para(
        doc,
        "2. SADG72 已是 weekly override 但仍无净值、无抓取日志，应单独排查：是尚未排进本周五批次、写入失败，还是备案号与火富牛列表不一致。",
    )
    para(
        doc,
        "3. SB8145、SB8164、SBKB21、SBPR25、SCE083、SCF808 六只新发已有 9 月净值，应立刻回填 private_fund_info.latest_nav / latest_nav_date，"
        "否则会继续从「6 个月以内 / 以上」筛选中消失——这与「没有净值」是两类问题。",
    )
    para(
        doc,
        "4. 约 450 只 A/B/C 类份额不在协会主表，列表与净值主表都不会自然出现。是否单独建档取决于产品页是否按份额展示，不宜默认全部付费拉取。",
    )

    heading(doc, "附录  字段说明", 1)
    add_table(
        doc,
        ["字段", "含义"],
        [
            ["beian_hao", "备案号"],
            ["has_db_nav", "private_fund_nav 是否有任意一行（本 CSV 均为 no）"],
            ["nav_rows", "净值行数"],
            ["in_fof99_universe", "是否在 fof99_nav_universe"],
            ["fof99_policy", "weekly / weekly_plus / skip / update_slow / 空"],
            ["fof99_reason", "宇宙行上的 reason"],
            ["last_fetch_status", "最近一次非探测抓取 ok / no_data / error"],
            ["empty_date_probe", "1970-01-01 空日期探测结果"],
            ["src_missing_reason", "源文件为何进不了 6 个月筛选"],
        ],
    )
    foot = para(
        doc,
        f"生成时间 {datetime.now().strftime('%Y-%m-%d %H:%M')}  ·  不调用火富牛付费接口，仅读库与源 CSV。",
        size=9,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    foot.paragraph_format.first_line_indent = Cm(0)
    doc.save(REPORT_PATH)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    load_env()
    ensure_db()
    configure_matplotlib()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    src_rows = list(csv.DictReader(SRC.open(encoding="utf-8-sig", newline="")))
    codes = []
    by_code: dict[str, dict] = {}
    for r in src_rows:
        code = (r.get("beian_hao") or "").strip().upper()
        if not code:
            continue
        codes.append(code)
        by_code[code] = r
    codes = list(dict.fromkeys(codes))
    print("source unique funds", len(codes))

    db = fetch_db(codes)
    fields = [
        "beian_hao",
        "fof99_fund_name",
        "fof99_short_name",
        "fof99_advisor",
        "fof99_price_date",
        "fof99_price_nav",
        "fof99_strategy_one",
        "fof99_strategy_two",
        "is_share_class",
        "src_missing_reason",
        "in_private_fund_info",
        "in_amac",
        "amac_fund_name",
        "amac_working_state",
        "amac_manager",
        "has_db_nav",
        "nav_rows",
        "nav_min_date",
        "nav_max_date",
        "info_latest_nav_date",
        "info_latest_nav",
        "in_fof99_universe",
        "fof99_policy",
        "fof99_reason",
        "universe_listed_at",
        "last_fetch_date",
        "last_fetch_status",
        "last_ok_date",
        "empty_date_probe",
    ]
    all_rows = []
    for code in codes:
        src = by_code[code]
        nav = db["nav"].get(code, {})
        info = db["info"].get(code, {})
        uni = db["uni"].get(code, {})
        lf = db["last_fetch"].get(code, {})
        lo = db["last_ok"].get(code, {})
        name = src.get("fof99_fund_name") or ""
        rows_n = int(nav.get("nav_rows") or 0)
        row = {
            "beian_hao": code,
            "fof99_fund_name": name,
            "fof99_short_name": src.get("fof99_short_name") or "",
            "fof99_advisor": src.get("fof99_advisor") or "",
            "fof99_price_date": src.get("fof99_price_date") or "",
            "fof99_price_nav": src.get("fof99_price_nav") or "",
            "fof99_strategy_one": src.get("fof99_strategy_one") or "",
            "fof99_strategy_two": src.get("fof99_strategy_two") or "",
            "is_share_class": "yes" if is_share_class(name, code) else "no",
            "src_missing_reason": src.get("missing_reason") or "",
            "in_private_fund_info": "yes" if info.get("in_info") else "no",
            "in_amac": src.get("in_amac") or "",
            "amac_fund_name": src.get("amac_fund_name") or "",
            "amac_working_state": src.get("amac_working_state") or "",
            "amac_manager": src.get("amac_manager") or "",
            "has_db_nav": "yes" if rows_n > 0 else "no",
            "nav_rows": rows_n,
            "nav_min_date": nav.get("nav_min") or "",
            "nav_max_date": nav.get("nav_max") or "",
            "info_latest_nav_date": info.get("list_nav_date") or "",
            "info_latest_nav": info.get("list_nav") or "",
            "in_fof99_universe": "yes" if uni.get("in_universe") else "no",
            "fof99_policy": uni.get("policy") or "",
            "fof99_reason": uni.get("reason") or "",
            "universe_listed_at": uni.get("listed_at") or "",
            "last_fetch_date": lf.get("last_fetch_date") or "",
            "last_fetch_status": lf.get("last_fetch_status") or "",
            "last_ok_date": lo.get("last_ok_date") or "",
            "empty_date_probe": db["probe"].get(code, ""),
        }
        all_rows.append(row)

    missing = [r for r in all_rows if r["has_db_nav"] == "no"]
    missing.sort(key=lambda r: (r["fof99_policy"] or "zzz", r["fof99_price_date"] or "", r["beian_hao"]))

    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(missing)

    print("wrote csv", CSV_OUT, "rows", len(missing))
    print("source", len(all_rows), "has_nav", len(all_rows) - len(missing), "no_nav", len(missing))
    print("policy", dict(Counter(r["fof99_policy"] or "(none)" for r in missing)))
    print("reason", dict(Counter(r["src_missing_reason"] for r in missing)))
    print("share", dict(Counter(r["is_share_class"] for r in missing)))
    print("in_universe", dict(Counter(r["in_fof99_universe"] for r in missing)))

    charts = make_charts(missing, all_rows)
    write_report(all_rows, missing, charts)
    print("wrote report", REPORT_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
