# -*- coding: utf-8 -*-
"""Chinese nginx + login Word report from reports/_week_traffic_raw."""
from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.font_manager import FontProperties, fontManager

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "reports" / "_week_traffic_raw"
OUT_DIR = ROOT / "reports"
CHART_DIR = OUT_DIR / "_week_traffic_charts"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
TEXT = RGBColor(0x2D, 0x37, 0x48)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

C_NAVY = "#1A365D"
C_TEAL = "#2B6CB0"
C_GOLD = "#C9A227"
C_GREEN = "#2F855A"
C_RED = "#C53030"
C_ORANGE = "#DD6B20"
C_GRAY = "#718096"
C_PURPLE = "#6B46C1"
C_LIGHT = "#E2E8F0"
PALETTE = [C_NAVY, C_TEAL, C_GOLD, C_GREEN, C_ORANGE, C_PURPLE, C_RED, C_GRAY]

KIND_CN = {
    "page": "页面",
    "api": "API",
    "heartbeat": "心跳",
    "static": "静态资源",
    "probe": "扫描探测",
    "admin": "管理",
    "login": "登录",
    "other": "其他",
}
NET_CN = {
    "Office A": "办公室出口 A",
    "Network B": "晚间网络 B",
    "Network C": "网络 C",
    "Network D": "网络 D",
    "Carrier 39.144": "运营商 39.144",
    "Localhost (dev)": "本机开发",
    "Other": "其他",
}
DEV_CN = {
    "Edge": "Edge",
    "Mobile": "手机",
    "Chrome": "Chrome",
    "Quark": "夸克",
    "bot/script": "爬虫/脚本",
    "Safari": "Safari",
    "Other": "其他",
}
WD_CN = {0: "一", 1: "二", 2: "三", 3: "四", 4: "五", 5: "六", 6: "日"}
SELFTEST = "__login_history_selftest__"
_CN_FONT: FontProperties | None = None


def configure_matplotlib() -> None:
    global _CN_FONT
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.facecolor"] = "white"
    plt.rcParams["axes.facecolor"] = "white"
    plt.rcParams["axes.edgecolor"] = "#CBD5E0"
    plt.rcParams["axes.grid"] = False
    for path in [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]:
        if os.path.isfile(path):
            try:
                fontManager.addfont(path)
                _CN_FONT = FontProperties(fname=path)
                plt.rcParams["font.family"] = "sans-serif"
                plt.rcParams["font.sans-serif"] = [_CN_FONT.get_name(), "Microsoft YaHei", "SimHei"]
                return
            except Exception:
                continue
    _CN_FONT = FontProperties(family="Microsoft YaHei")


def fp() -> dict:
    return {"fontproperties": _CN_FONT} if _CN_FONT is not None else {}


def set_run_font(run, *, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_text(p, text, *, size=11, bold=False, color=None):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def para(doc, text="", *, size=11, bold=False, color=None, align=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.first_line_indent = Cm(0)
    if align:
        p.alignment = align
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    return p


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E0")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, size=9, bold=False, color=None, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, "" if text is None else str(text), size=size, bold=bold, color=color)
    set_cell_border(cell)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=9, bold=True, color=WHITE)
        shade(table.rows[0].cells[i], "1A365D")
    for r_i, row in enumerate(rows):
        fill = "F7FAFC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(
                table.rows[r_i + 1].cells[c_i],
                val,
                size=8,
                color=TEXT,
                align="center" if c_i else "left",
            )
            shade(table.rows[r_i + 1].cells[c_i], fill)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    return table


def add_picture(doc, path: Path, width_in=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    return p


def caption(doc, text: str):
    para(doc, text, size=8, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def style_ax(ax, title: str, xlabel: str, ylabel: str):
    ax.set_title(title, **fp(), fontsize=12, color=C_NAVY, pad=10)
    ax.set_xlabel(xlabel, **fp(), fontsize=9, color="#4A5568")
    ax.set_ylabel(ylabel, **fp(), fontsize=9, color="#4A5568")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(colors="#4A5568", labelsize=8)
    for lab in ax.get_xticklabels() + ax.get_yticklabels():
        lab.set_fontproperties(_CN_FONT)


def save_fig(fig, name: str) -> Path:
    path = CHART_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fmt_int(n) -> str:
    return f"{int(n):,}"


def fmt_pct(n, d) -> str:
    if d == 0:
        return "—"
    return f"{100.0 * n / d:.1f}%"


def classify_login_device(ua: str | None) -> str:
    s = (ua or "").lower()
    if "iphone" in s or "android" in s or "mobile" in s or "huawei" in s:
        return "手机"
    if "edg/" in s:
        return "Edge"
    if "quark" in s:
        return "夸克"
    if "chrome" in s:
        return "Chrome"
    return "其他"


def classify_login_net(ip: str | None) -> str:
    ip = (ip or "").strip()
    if ip in ("::1", "127.0.0.1"):
        return "本机开发"
    if ip.startswith("116.237.193."):
        return "办公室出口 A"
    if ip.startswith("116.234.199."):
        return "晚间网络 B"
    if ip.startswith("116.234.86."):
        return "网络 C"
    if ip.startswith("111.187."):
        return "网络 D"
    if ip.startswith("39.144."):
        return "运营商 39.144"
    return "其他"


def date_label(day: str) -> str:
    dt = pd.to_datetime(day)
    return f"{dt.month}/{dt.day}({WD_CN[dt.weekday()]})"


def load() -> dict:
    summary = json.loads((RAW / "summary.json").read_text(encoding="utf-8"))
    daily = pd.read_csv(RAW / "daily.csv")
    hourly = pd.read_csv(RAW / "hourly.csv")
    paths = pd.read_csv(RAW / "top_paths.csv")
    pages = pd.read_csv(RAW / "top_pages.csv")
    status = pd.read_csv(RAW / "status.csv")
    devices = pd.read_csv(RAW / "devices.csv")
    networks = pd.read_csv(RAW / "networks.csv")
    methods = pd.read_csv(RAW / "methods.csv")
    ips = pd.read_csv(RAW / "ips.csv")
    login = pd.read_csv(RAW / "login.csv")
    users = pd.read_csv(RAW / "users.csv")
    login["ts"] = pd.to_datetime(login["ts"])
    login["success"] = login["success"].astype(str).str.lower().isin(["t", "true", "1"])
    login["who"] = login["who"].str.strip()
    login["who_norm"] = login["who"].str.lower()
    login["date"] = login["ts"].dt.strftime("%Y-%m-%d")
    login["hour"] = login["ts"].dt.hour
    login["device"] = login["user_agent"].map(classify_login_device)
    login["network"] = login["ip"].map(classify_login_net)
    daily["date"] = daily["date"].astype(str)
    daily["gb"] = daily["bytes"] / (1024 ** 3)
    daily["weekday"] = pd.to_datetime(daily["date"]).dt.weekday
    hourly["date"] = hourly["date"].astype(str)
    return {
        "summary": summary,
        "daily": daily,
        "hourly": hourly,
        "paths": paths,
        "pages": pages,
        "status": status,
        "devices": devices,
        "networks": networks,
        "methods": methods,
        "ips": ips,
        "login": login,
        "users": users,
    }


def feature_rows(paths: pd.DataFrame) -> list[tuple[str, int, str]]:
    rules = [
        ("心跳 /api/presence", "/api/presence", "polling"),
        ("跟踪基金 API", "/ma/api/tracking-funds", "api"),
        ("CTP 行情 live", "/ma/api/ctp-market/live", "polling"),
        ("CTP 行情其它", "/ma/api/ctp-market", "api"),
        ("实时行情 API", "/ma/api/realtime-quotes", "api"),
        ("投资笔记 API", "/ma/api/investment-notes", "api"),
        ("私募产品 API", "/ma/api/private-funds", "api"),
        ("尽调表 API", "/ma/api/due-diligence-table", "api"),
        ("最近页面 API", "/ma/api/recent-pages", "polling"),
        ("运营/产品要素 API", "/ma/api/ops/", "api"),
        ("MOM API", "/ma/api/mom-analysis", "api"),
        ("FOF overview API", "/ma/api/:id/fof-overview", "api"),
        ("知识库 API", "/api/knowledge-base", "api"),
        ("鉴权 /api/auth/me", "/api/auth/me", "polling"),
        ("管理 API", "/api/admin/", "admin"),
        ("私募产品页", "/ma/dashboard/private-funds", "page"),
        ("MOM 分析页", "/ma/dashboard/mom-analysis", "page"),
        ("全天候", "/ma/dashboard/all-weather", "prefetch"),
        ("期货市场", "/ma/dashboard/futures-market", "prefetch"),
        ("期权市场", "/ma/dashboard/options-market", "prefetch"),
        ("宏观市场", "/ma/dashboard/macro-market", "prefetch"),
        ("AI 知识", "/ma/dashboard/ai-knowledge", "prefetch"),
        ("AI 投研", "/ma/dashboard/ai-researcher", "prefetch"),
        ("实时行情页", "/ma/dashboard/realtime-quotes", "prefetch"),
        ("股票市场", "/ma/dashboard/stock-market", "prefetch"),
        ("工具页", "/ma/dashboard/tools", "prefetch"),
        ("设置页", "/ma/dashboard/settings", "page"),
        ("投研看板首页", "/ma/dashboard", "page"),
        ("登录页", "/login", "page"),
        ("Vercel insights 404", "/_vercel/insights", "noise"),
        ("图标 / 静态", "/icon", "static"),
        ("Next 静态资源", "/_next/", "static"),
    ]
    used = set()
    out = []
    for name, prefix, kind in rules:
        mask = paths["path"].astype(str).str.startswith(prefix)
        if prefix == "/ma/dashboard":
            mask = paths["path"].eq("/ma/dashboard") | paths["path"].eq("/ma/dashboard/:id")
        hits = int(paths.loc[mask & ~paths["path"].isin(used), "hits"].sum())
        for p in paths.loc[mask, "path"]:
            used.add(p)
        if hits:
            out.append((name, hits, kind))
    rest = int(paths.loc[~paths["path"].isin(used), "hits"].sum())
    if rest:
        out.append(("其它（未归入上表的路径）", rest, "other"))
    return out


def make_charts(d: dict) -> dict[str, Path]:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    charts = {}
    daily = d["daily"]
    hourly = d["hourly"]
    login = d["login"]
    status = d["status"]
    networks = d["networks"]
    devices = d["devices"]
    feats = feature_rows(d["paths"])
    labels = [date_label(x) for x in daily["date"]]
    x = np.arange(len(daily))

    kinds = ["page", "api", "heartbeat", "static", "probe", "admin", "login", "other"]
    colors = {
        "page": C_NAVY, "api": C_TEAL, "heartbeat": C_GOLD, "static": C_GRAY,
        "probe": C_RED, "admin": C_PURPLE, "login": C_GREEN, "other": C_ORANGE,
    }
    fig, ax = plt.subplots(figsize=(8.6, 3.9))
    bottom = np.zeros(len(daily))
    for k in kinds:
        vals = daily[k].to_numpy(dtype=float)
        ax.bar(x, vals, bottom=bottom, color=colors[k], width=0.72, label=KIND_CN[k])
        bottom += vals
    ax.set_xticks(x)
    ax.set_xticklabels(labels, **fp(), rotation=45, ha="right")
    ax.legend(prop=_CN_FONT, fontsize=7, frameon=False, ncol=4, loc="upper right")
    style_ax(ax, "按日请求量（按类型堆叠）", "日期（北京时间）", "请求数")
    charts["daily_kind"] = save_fig(fig, "cn_daily_kind.png")

    fig, ax = plt.subplots(figsize=(8.6, 3.6))
    product = daily["product"].to_numpy()
    rest = daily["all"].to_numpy() - product
    ax.bar(x - 0.18, product, width=0.36, color=C_NAVY, label="产品流量（页面+API+管理+登录）")
    ax.bar(x + 0.18, rest, width=0.36, color=C_LIGHT, label="其余（静态 / 心跳 / 扫描）")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, **fp(), rotation=45, ha="right")
    ax.legend(prop=_CN_FONT, fontsize=8, frameon=False)
    style_ax(ax, "产品流量 vs 其余流量", "日期（北京时间）", "请求数")
    charts["daily_product"] = save_fig(fig, "cn_daily_product.png")

    hours = np.arange(24)
    prod_h = hourly.groupby("hour")["product"].sum().reindex(hours, fill_value=0)
    fig, ax = plt.subplots(figsize=(8.4, 3.5))
    bar_colors = [C_RED if v >= 15000 else (C_GOLD if v >= 8000 else C_TEAL) for v in prod_h]
    ax.bar(hours, prod_h.to_numpy(), color=bar_colors, width=0.72)
    ax.set_xticks(hours)
    style_ax(ax, "产品请求按钟点合计（观察期内）", "小时（北京时间）", "产品请求数")
    charts["hourly_product"] = save_fig(fig, "cn_hourly_product.png")

    day_order = list(daily["date"])
    mat = (
        hourly.pivot_table(index="date", columns="hour", values="product", aggfunc="sum")
        .reindex(index=day_order, columns=hours, fill_value=0)
        .fillna(0)
    )
    fig, ax = plt.subplots(figsize=(8.6, 5.0))
    im = ax.imshow(mat.to_numpy(), aspect="auto", cmap="Blues")
    ax.set_xticks(hours)
    ax.set_yticks(range(len(day_order)))
    ax.set_yticklabels(labels, **fp())
    ax.set_xlabel("小时（北京时间）", **fp(), fontsize=9, color="#4A5568")
    ax.set_title("产品请求热力图（页面 + API + 管理 + 登录）", **fp(), fontsize=12, color=C_NAVY, pad=10)
    ax.tick_params(colors="#4A5568", labelsize=8)
    cbar = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02)
    cbar.ax.tick_params(labelsize=7)
    cbar.set_label("请求数", **fp(), fontsize=8)
    fig.tight_layout()
    charts["heatmap"] = save_fig(fig, "cn_heatmap.png")

    top = sorted(feats, key=lambda r: r[1], reverse=True)[:14]
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    names = [r[0] for r in top][::-1]
    vals = [r[1] for r in top][::-1]
    kinds_l = [r[2] for r in top][::-1]
    kind_color = {
        "polling": C_GOLD, "api": C_TEAL, "page": C_NAVY, "prefetch": C_PURPLE,
        "admin": C_ORANGE, "noise": C_RED, "static": C_GRAY, "other": C_ORANGE,
    }
    y = np.arange(len(names))
    ax.barh(y, vals, color=[kind_color.get(k, C_GRAY) for k in kinds_l], height=0.62)
    ax.set_yticks(y)
    ax.set_yticklabels(names, **fp())
    style_ax(ax, "高频路径分组（颜色=流量类型）", "", "命中次数")
    charts["features"] = save_fig(fig, "cn_features.png")

    st = status.sort_values("hits", ascending=False)
    fig, ax = plt.subplots(figsize=(8.4, 3.4))
    sc = [C_GREEN if s < 400 else (C_GOLD if s < 500 else C_RED) for s in st["status"]]
    ax.bar([str(s) for s in st["status"]], st["hits"], color=sc, width=0.7)
    style_ax(ax, "HTTP 状态码分布", "状态码", "命中次数")
    charts["status"] = save_fig(fig, "cn_status.png")

    fig, ax = plt.subplots(figsize=(8.4, 3.6))
    net = networks.copy()
    net["label"] = net["network"].map(lambda v: NET_CN.get(v, v))
    net = net.sort_values("hits", ascending=True)
    y = np.arange(len(net))
    ax.barh(y, net["hits"], color=C_NAVY, height=0.55)
    ax.set_yticks(y)
    ax.set_yticklabels(list(net["label"]), **fp())
    style_ax(ax, "按客户端网段的请求量", "", "命中次数")
    charts["networks"] = save_fig(fig, "cn_networks.png")

    fig, ax = plt.subplots(figsize=(8.4, 3.5))
    dev = devices.copy()
    dev["label"] = dev["device"].map(lambda v: DEV_CN.get(v, v))
    ax.bar(dev["label"], dev["hits"], color=PALETTE[: len(dev)], width=0.62)
    style_ax(ax, "按终端类型的请求量", "", "命中次数")
    charts["devices"] = save_fig(fig, "cn_devices.png")

    period = login.loc[(~login["who_norm"].eq(SELFTEST.lower())) & (login["date"].isin(set(daily["date"])))]
    daily_login = period.groupby(["date", "who"]).size().unstack(fill_value=0).reindex(daily["date"], fill_value=0)
    fig, ax = plt.subplots(figsize=(8.6, 3.7))
    x2 = np.arange(len(daily_login))
    bottom = np.zeros(len(daily_login))
    for i, name in enumerate(daily_login.columns):
        vals = daily_login[name].to_numpy()
        ax.bar(x2, vals, bottom=bottom, color=PALETTE[i % len(PALETTE)], width=0.62, label=name)
        bottom += vals
    ax.set_xticks(x2)
    ax.set_xticklabels(labels, **fp(), rotation=45, ha="right")
    ax.legend(prop=_CN_FONT, fontsize=8, frameon=False, ncol=4, loc="upper left")
    style_ax(ax, "按人统计的登录事件", "日期（北京时间）", "登录次数")
    charts["logins"] = save_fig(fig, "cn_logins.png")
    return charts


def status_hits(status: pd.DataFrame, code: int) -> int:
    hit = status.loc[status["status"] == code, "hits"]
    return int(hit.sum()) if len(hit) else 0


def build_doc(d: dict, charts: dict[str, Path]) -> Path:
    s = d["summary"]
    daily = d["daily"]
    hourly = d["hourly"]
    login_all = d["login"]
    users = d["users"]
    ips = d["ips"]
    status = d["status"]
    devices = d["devices"]
    methods = d["methods"]
    pages = d["pages"]
    feats = feature_rows(d["paths"])

    n = int(s["hits_in_week"])
    n_prod = int(s["product_hits"])
    n_ip = int(s["unique_ips"])
    gb = s["bytes_in_week"] / (1024 ** 3)
    kind = s["kind"]
    first_ts = datetime.fromisoformat(s["first_ts"])
    last_ts = datetime.fromisoformat(s["last_ts"])
    window_cn = (
        f"{first_ts.strftime('%Y年%m月%d日 %H:%M')} – {last_ts.strftime('%Y年%m月%d日 %H:%M')}（北京时间）"
    )
    report_path = OUT_DIR / f"网站流量分析报告_{first_ts.strftime('%Y%m%d')}-{last_ts.strftime('%Y%m%d')}.docx"

    ok_2xx = int(status.loc[status["status"].between(200, 299), "hits"].sum())
    n_304 = status_hits(status, 304)
    n_404 = status_hits(status, 404)
    n_499 = status_hits(status, 499)
    n_502 = status_hits(status, 502)
    n_500 = status_hits(status, 500)
    n_403 = status_hits(status, 403)
    n_400 = status_hits(status, 400)
    n_probe = int(kind["probe"])
    n_hb = int(kind["heartbeat"])

    real_ips = ips[(ips["page"] > 0) | (ips["api"] > 0)]
    scanner_ips = ips[(ips["page"] == 0) & (ips["api"] == 0)]
    office_ip = "116.237.193.158"
    office_hits = float(ips.loc[ips["ip"] == office_ip, "hits"].sum())
    office_share = office_hits / n if n else 0

    prefetch_pages = [
        "/ma/dashboard/all-weather", "/ma/dashboard/futures-market", "/ma/dashboard/options-market",
        "/ma/dashboard/macro-market", "/ma/dashboard/ai-knowledge", "/ma/dashboard/ai-researcher",
        "/ma/dashboard/realtime-quotes", "/ma/dashboard/stock-market", "/ma/dashboard/tools",
    ]
    prefetch_hits = int(pages.loc[pages["path"].isin(prefetch_pages), "hits"].sum())
    private_page = int(pages.loc[pages["path"].str.startswith("/ma/dashboard/private-funds"), "hits"].sum())
    mom_page = int(pages.loc[pages["path"].str.startswith("/ma/dashboard/mom-analysis"), "hits"].sum())
    home_page = int(pages.loc[pages["path"].isin(["/ma/dashboard", "/ma/dashboard/:id"]), "hits"].sum())

    peak = hourly.loc[hourly["product"].idxmax()]
    peak_label = f"{date_label(str(peak['date']))} {int(peak['hour']):02d}:00–{int(peak['hour']):02d}:59"

    weekday = daily[daily["weekday"] < 5]
    weekend = daily[daily["weekday"] >= 5]
    busiest = daily.loc[daily["all"].idxmax()]
    quietest = daily.loc[daily["product"].idxmin()]

    hour_sum = hourly.groupby("hour")["product"].sum()
    top_hours = hour_sum.sort_values(ascending=False).head(3)
    lunch = int(hour_sum.get(12, 0))
    night = int(hour_sum.loc[hour_sum.index.isin(range(0, 7))].sum())

    login = login_all.loc[~login_all["who_norm"].eq(SELFTEST.lower())].copy()
    in_win = login.loc[login["date"].isin(set(daily["date"]))].copy()
    login_ok = int(in_win["success"].sum())
    login_fail = int((~in_win["success"]).sum())
    people = sorted(in_win["who_norm"].unique())
    named = set(users["name"].str.lower())
    inactive = [n_ for n_ in users["name"] if n_.lower() not in set(people)]
    new_users = users.loc[
        (users["created"] >= first_ts.strftime("%Y-%m-%d"))
        & (users["created"] <= last_ts.strftime("%Y-%m-%d 23:59:59")),
        "name",
    ].tolist()

    by_who = (
        in_win.assign(who=np.where(in_win["who_norm"] == "musheng", "musheng", in_win["who"]))
        .groupby("who")
        .agg(events=("ts", "size"), ok=("success", "sum"), first=("ts", "min"), last=("ts", "max"))
        .sort_values("events", ascending=False)
    )

    fail_rows = in_win.loc[~in_win["success"]]
    vercel_404 = int(d["paths"].loc[d["paths"]["path"].astype(str).str.startswith("/_vercel/insights"), "hits"].sum())

    presence = next((h for name, h, _ in feats if name.startswith("心跳")), 0)
    tracking = next((h for name, h, _ in feats if name.startswith("跟踪基金")), 0)
    ctp_live = next((h for name, h, _ in feats if "live" in name.lower() or "CTP 行情 live" in name), 0)

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(1.8)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)

    para(doc, "投研看板  ·  内部运营报告", size=10, color=GOLD, space_after=2)
    p = para(doc, "网站流量分析报告", size=22, bold=True, color=NAVY, space_after=4)
    p.runs[0].font.size = Pt(22)
    para(
        doc,
        f"数据来源：nginx access.log（含轮转文件）与 public.auth_login_history  ·  "
        f"时区：Asia/Shanghai (UTC+8)  ·  观察窗口：{window_cn}  ·  "
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        size=9,
        color=MUTED,
        space_after=14,
    )

    heading(doc, "一、执行摘要", 1)
    para(
        doc,
        f"本报告依据生产服务器已落盘的访问记录，覆盖 {window_cn}，共 "
        f"{fmt_int(n)} 条 HTTP 请求、传输 {gb:.2f} GB，来自 {fmt_int(n_ip)} 个独立客户端 IP。"
        f"其中产品侧请求（页面、业务 API、管理、登录）{fmt_int(n_prod)} 条，占 {fmt_pct(n_prod, n)}。"
        f"最后一天 {date_label(daily['date'].iloc[-1])} 只统计到 {last_ts.strftime('%H:%M')}，"
        f"与完整工作日对比时不宜直接当作需求下滑。",
    )
    para(
        doc,
        f"办公室出口 A（{office_ip}）贡献了全部请求的 {office_share:.0%}。"
        f"同一 NAT 后面是多人，因此「独立 IP」不能当作用户数。"
        f"人头应以登录历史为准：观察期内 {len(in_win)} 次登录事件，"
        f"成功 {login_ok} / 失败 {login_fail}，出现 {len(people)} 个账号。"
        + (f"本期新开账号：{'、'.join(new_users)}。" if new_users else ""),
        space_after=8,
    )
    para(doc, "这 15 天的核心判断：", size=11, bold=True, color=NAVY, space_after=4)
    para(
        doc,
        f"1）工作日主导、周末几乎空窗。工作日日均全部请求 {fmt_int(weekday['all'].mean())}、"
        f"产品请求 {fmt_int(weekday['product'].mean())}；周末日均分别只有 "
        f"{fmt_int(weekend['all'].mean())} 和 {fmt_int(weekend['product'].mean())}。"
        f"最忙的一天是 {date_label(busiest['date'])}（{fmt_int(busiest['all'])}），"
        f"产品请求最少的一天是 {date_label(quietest['date'])}（仅 {fmt_int(quietest['product'])}）。",
        space_after=3,
    )
    para(
        doc,
        f"2）页面计数会夸大浏览。侧栏行情/AI/工具共 9 条路由合计 {fmt_int(prefetch_hits)} 次，"
        f"条数几乎一样，是 Next.js 预取/导航预热，不是九个同等受欢迎的页面。"
        f"真正有意图的落点是私募产品（{fmt_int(private_page)} 次页面）、"
        f"看板首页（{fmt_int(home_page)}）和 MOM 分析（{fmt_int(mom_page)}）。",
        space_after=3,
    )
    para(
        doc,
        f"3）最重的接口是轮询而不是点击：心跳 {fmt_int(presence)}、跟踪基金 {fmt_int(tracking)}、"
        f"CTP live {fmt_int(ctp_live)}。行情页开着不关，部署检查器也会看到「很忙」。",
        space_after=3,
    )
    para(
        doc,
        f"4）可用性总体可接受：2xx 占 {fmt_pct(ok_2xx, n)}。"
        f"404 共 {fmt_int(n_404)}（含 Vercel insights 脚本 {fmt_int(vercel_404)} 次及扫描路径），"
        f"499 {fmt_int(n_499)}，502 {fmt_int(n_502)}，500 {fmt_int(n_500)}。",
        space_after=3,
    )
    para(
        doc,
        f"5）公网噪声真实存在：归类为扫描探测的请求 {fmt_int(n_probe)} 条，"
        f"多出现在凌晨，打 /.env、WordPress、广告探测等，基本未形成产品 API 流量。",
        space_after=12,
    )

    heading(doc, "二、数据口径与注意点", 1)
    add_table(
        doc,
        ["项目", "取值"],
        [
            ["HTTP 来源", s.get("http_source", "/var/log/nginx/access.log + 轮转文件")],
            ["登录来源", s.get("login_source", "public.auth_login_history")],
            ["时区", "Asia/Shanghai（UTC+8）"],
            ["窗口起", first_ts.strftime("%Y-%m-%d %H:%M:%S")],
            ["窗口止", last_ts.strftime("%Y-%m-%d %H:%M:%S") + "（末日为部分日）"],
            ["扫描行数 / 窗口内命中", f"{fmt_int(s['lines_scanned'])} / {fmt_int(n)}"],
            ["传输体积", f"{gb:.2f} GB"],
            ["独立客户端 IP", fmt_int(n_ip)],
            ["产生页面或 API 的 IP", fmt_int(len(real_ips))],
            ["产品请求", f"{fmt_int(n_prod)}（{fmt_pct(n_prod, n)}）"],
            ["页面 / API / 心跳 / 静态 / 扫描",
             f"{fmt_int(kind['page'])} / {fmt_int(kind['api'])} / {fmt_int(n_hb)} / "
             f"{fmt_int(kind['static'])} / {fmt_int(n_probe)}"],
            ["登录事件（成功 / 失败）", f"{len(in_win)}（{login_ok} / {login_fail}）"],
            ["注册账号 / 本期出现", f"{len(users)} / {len(people)}"],
            ["本期新开账号", "、".join(new_users) if new_users else "无"],
        ],
        col_widths=[6.6, 10.4],
    )
    para(doc, "", space_after=6)
    para(
        doc,
        "一条命中是一次 HTTP 请求，既不是一个人，也不是一个会话。"
        "心跳、CTP live、最近页面写入、Next.js 路由预取都会放大请求量。"
        "办公室同事共享 116.237.193.158，IP 去重会低估工位人数、高估公网扫描器。"
        "静态资源和 /_next 资产仍计入总量，便于看缓存/304，但不代表产品使用。"
        "登录表从 8 月 23 日起有记录；正文人头统计只取与 nginx 窗口重叠的部分。",
        size=10,
        color=MUTED,
        space_after=12,
    )

    heading(doc, "三、按日流量", 1)
    para(
        doc,
        "周一最重：8 月 31 日与 9 月 7 日都在 4.0 万请求上下，产品请求约 3.1–3.2 万。"
        "周三、周五明显更轻。周末几乎没有真人："
        f"{date_label(quietest['date'])} 产品请求只有 {fmt_int(quietest['product'])}，"
        "更像扫描和残留心跳。9 月 6 日周日中午有一次明显的管理端使用，属于例外，不是周末常态。",
    )
    add_table(
        doc,
        ["日期", "全部", "产品", "页面", "API", "心跳", "扫描", "IP", "GB"],
        [
            [
                date_label(row["date"]),
                fmt_int(row["all"]),
                fmt_int(row["product"]),
                fmt_int(row["page"]),
                fmt_int(row["api"]),
                fmt_int(row["heartbeat"]),
                fmt_int(row["probe"]),
                fmt_int(row["unique_ips"]),
                f"{row['gb']:.2f}",
            ]
            for _, row in daily.iterrows()
        ],
        col_widths=[2.4, 1.7, 1.7, 1.5, 1.5, 1.5, 1.5, 1.3, 1.5],
    )
    para(doc, "", space_after=6)
    add_picture(doc, charts["daily_kind"], 6.4)
    caption(doc, "图 1. 按日 nginx 请求量，按类型堆叠。9 月 13 日只统计到 20:13。")
    add_picture(doc, charts["daily_product"], 6.4)
    caption(doc, "图 2. 产品侧流量与静态 / 心跳 / 扫描等其余流量对比。")

    heading(doc, "四、何时最忙", 1)
    top_h_txt = "、".join([f"{int(h):02d} 时（{fmt_int(v)}）" for h, v in top_hours.items()])
    para(
        doc,
        f"最忙的一个产品小时是 {peak_label}，共 {fmt_int(peak['product'])} 次产品请求"
        f"（页面 {fmt_int(peak['page'])}，API {fmt_int(peak['api'])}）。"
        f"按钟点合计，最密的三个小时是 {top_h_txt}。"
        f"12 时合计 {fmt_int(lunch)}，工作日普遍有午休回落。"
        f"0–6 时产品请求合计仅 {fmt_int(night)}，几乎没有真人。",
    )
    para(
        doc,
        "热力图上，工作日 09:00–11:30、13:00–17:30 是核心色块；"
        "18:00 之后主要是管理端在晚间网络 B 继续工作，并夹少量手机访问。"
        "凌晨偶发的「很忙」多半是扫描，产品热力图里是空的。",
    )
    add_picture(doc, charts["hourly_product"], 6.4)
    caption(doc, "图 3. 产品请求按钟点合计。红色 ≥ 15,000，金色 ≥ 8,000。")
    add_picture(doc, charts["heatmap"], 6.4)
    caption(doc, "图 4. 产品请求热力图。隔夜空白是真实空窗；扫描高峰不会出现在这张图里。")

    heading(doc, "五、大家在用什么", 1)
    para(
        doc,
        "路径要分四层看，否则排行会骗人。第一层是轮询（心跳、CTP live、auth/me、最近页面）；"
        "第二层是页面触发的业务 API（跟踪基金、私募产品、投资笔记、MOM）；"
        "第三层是 HTML 页面；第四层是侧栏兄弟路由预取——"
        "行情/AI/工具约 7,650–7,800 次、条数几乎相同，是 <Link prefetch> 的典型指纹。",
    )
    add_picture(doc, charts["features"], 6.4)
    caption(doc, "图 5. 高频路径分组。深蓝=页面，青绿=API，金色=轮询，紫色=疑似预取，红色=噪声。")
    add_table(
        doc,
        ["路径分组", "命中", "类型"],
        [
            [
                name,
                fmt_int(hits),
                {"polling": "轮询", "api": "API", "page": "页面", "prefetch": "预取",
                 "admin": "管理", "noise": "噪声", "static": "静态", "other": "其他"}.get(k, k),
            ]
            for name, hits, k in feats[:18]
        ],
        col_widths=[7.0, 3.5, 6.5],
    )
    para(doc, "", space_after=8)
    para(doc, "有意图的页面落点（合并 :id 后）：", size=11, bold=True, color=NAVY, space_after=4)
    page_rows = []
    for _, r in pages.head(16).iterrows():
        tag = "预取簇" if r["path"] in prefetch_pages else "页面"
        page_rows.append([r["path"], fmt_int(r["hits"]), tag])
    add_table(doc, ["页面", "命中", "读法"], page_rows, col_widths=[9.5, 2.5, 5.0])
    para(doc, "", space_after=6)
    para(
        doc,
        "从产品经理视角：私募产品是本期主工作面（列表 + 产品代码页 + 笔记/标签/尽调/份额接口）。"
        "MOM 分析是第二工作面（风控报告、交易员分析、数据导入、单账户风控、carry）。"
        "跟踪基金和 CTP live 是常开组件。设置页次数偏高，部分来自导航预热，不宜直接当成「大家都在改设置」。",
        space_after=12,
    )

    heading(doc, "六、谁在用、从哪来", 1)
    para(
        doc,
        f"登录历史是人头普查。观察期内 {len(in_win)} 次鉴权，成功 {login_ok} 次。"
        + (
            f"失败 {login_fail} 次"
            + (
                f"，涉及：{'、'.join(sorted(fail_rows['who'].unique()))}。"
                if len(fail_rows)
                else "。"
            )
            if login_fail
            else "无失败记录。"
        )
        + " 9 月 2 日晚 Musheng（大写）在 iPhone 上因密码错误失败，13 秒后 musheng（新账号）登录成功，是首次登录大小写问题，不是攻击。",
    )
    add_picture(doc, charts["logins"], 6.4)
    caption(doc, "图 6. 按人统计的登录事件。cshen 含本机开发服务器上的登录。")
    who_rows = []
    for name, r in by_who.iterrows():
        who_rows.append([
            name,
            int(r.events),
            int(r.ok),
            r["first"].strftime("%m-%d %H:%M"),
            r["last"].strftime("%m-%d %H:%M"),
        ])
    add_table(
        doc,
        ["账号", "次数", "成功", "首次（北京时间）", "末次（北京时间）"],
        who_rows,
        col_widths=[3.4, 2.4, 2.4, 4.4, 4.4],
    )
    para(doc, "", space_after=8)

    profiles = []
    for name, grp in in_win.groupby("who"):
        nets = "、".join(sorted(grp["network"].unique()))
        devices_u = "、".join(sorted(grp["device"].unique()))
        profiles.append(
            f"{name}：{len(grp)} 次，{devices_u}，网络 {nets}，"
            f"{grp['ts'].min().strftime('%m-%d %H:%M')} 至 {grp['ts'].max().strftime('%m-%d %H:%M')}。"
        )
    para(doc, " ".join(profiles), space_after=8)
    para(
        doc,
        "本期未登录账号："
        + ("、".join(inactive) if inactive else "无")
        + "。只对这 15 天窗口有效——有的同事隔周才出现。",
        space_after=8,
    )
    para(doc, "网段构成（nginx 全部请求）：", size=11, bold=True, color=NAVY, space_after=4)
    add_picture(doc, charts["networks"], 6.4)
    caption(doc, "图 7. 客户端网段。办公室出口 A 是共享 NAT，多人共用一个 IP。")
    add_picture(doc, charts["devices"], 6.4)
    caption(doc, "图 8. 终端类型。手机流量真实存在，不全是 UA 伪造。")

    top_ip_rows = []
    for _, r in ips.head(12).iterrows():
        if str(r.ip).endswith(".158"):
            role = "办公室 NAT / 多人"
        elif str(r.ip).endswith(".216"):
            role = "晚间 / 管理端"
        elif r.page == 0 and r.api == 0:
            role = "扫描器"
        else:
            role = "用户 / 手机 / 其他出口"
        top_ip_rows.append([
            r.ip, NET_CN.get(r.network, r.network), fmt_int(r.hits),
            fmt_int(r.page), fmt_int(r.api), role,
        ])
    add_table(
        doc,
        ["IP", "网段", "命中", "页面", "API", "读法"],
        top_ip_rows,
        col_widths=[3.6, 3.4, 1.6, 1.6, 1.6, 5.2],
    )
    para(doc, "", space_after=6)
    device_txt = "、".join(
        f"{DEV_CN.get(r.device, r.device)} {fmt_int(r.hits)}" for _, r in devices.iterrows()
    )
    para(
        doc,
        f"产生过页面或 API 的 IP 有 {fmt_int(len(real_ips))} 个；"
        f"列表里纯扫描、无产品路径的 IP 有 {fmt_int(len(scanner_ips))} 个。"
        f"终端构成：{device_txt}。",
        space_after=12,
    )

    heading(doc, "七、可靠性与扫描噪声", 1)
    para(
        doc,
        f"2xx 占比 {fmt_pct(ok_2xx, n)}（{fmt_int(ok_2xx)} / {fmt_int(n)}）。"
        f"304 Not Modified {fmt_int(n_304)}，说明图标和部分 HTML 被浏览器缓存，这是健康信号。"
        "需要单独说一句的错误如下。",
    )
    add_picture(doc, charts["status"], 6.4)
    caption(doc, "图 9. HTTP 状态码。绿色=2xx/3xx，金色=4xx，红色=5xx。")
    add_table(
        doc,
        ["状态码", "次数", "占比", "读法"],
        [
            ["200", fmt_int(status_hits(status, 200)), fmt_pct(status_hits(status, 200), n), "正常成功"],
            ["304", fmt_int(n_304), fmt_pct(n_304, n), "缓存命中"],
            ["404", fmt_int(n_404), fmt_pct(n_404, n),
             f"缺失 /_vercel/insights/script.js（约 {fmt_int(vercel_404)}）+ 扫描路径"],
            ["400", fmt_int(n_400), fmt_pct(n_400, n), "畸形请求 / TLS 握手被当成 HTTP"],
            ["403", fmt_int(n_403), fmt_pct(n_403, n), "鉴权拒绝"],
            ["499", fmt_int(n_499), fmt_pct(n_499, n), "客户端断开——发布、刷新或处理过慢"],
            ["500", fmt_int(n_500), fmt_pct(n_500, n), "应用错误，量不大，仍值得 grep 日志"],
            ["502", fmt_int(n_502), fmt_pct(n_502, n), "nginx → Next 上游短暂不可用"],
        ],
        col_widths=[2.0, 2.2, 2.2, 10.6],
    )
    para(doc, "", space_after=8)
    get_n = int(methods.loc[methods["method"] == "GET", "hits"].sum())
    post_n = int(methods.loc[methods["method"] == "POST", "hits"].sum())
    put_n = int(methods.loc[methods["method"] == "PUT", "hits"].sum())
    patch_n = int(methods.loc[methods["method"] == "PATCH", "hits"].sum())
    head_n = int(methods.loc[methods["method"] == "HEAD", "hits"].sum())
    para(
        doc,
        f"方法：GET {fmt_int(get_n)}，POST {fmt_int(post_n)}，PUT {fmt_int(put_n)}，"
        f"PATCH {fmt_int(patch_n)}，HEAD {fmt_int(head_n)}。"
        "PRI / CONNECT / 乱码动词来自扫描器的 HTTP/2 或 TLS 探测，不是应用本身。"
        "PUT 主要来自 /ma/api/recent-pages。",
        space_after=8,
    )
    para(
        doc,
        "404 列表里有两件卫生事项。第一，前端仍在这台自建主机上请求 "
        "/_vercel/insights/script.js，每次打开页面都付一次 404，应去掉这段脚本。"
        "第二，主机在公网，/.env 和 WordPress 探测不会停；目前结果是 404/400，应继续保持这个结果。",
        space_after=12,
    )

    heading(doc, "八、部署窗口（本期证据）", 1)
    para(doc, "用产品请求密度加登录时钟，不要用独立 IP。北京时间窗口排序：", space_after=6)
    para(doc, "1）最佳：00:30–07:00，以及周六大部分时间。没有产品用户。接受扫描 404，不要当成负载。", space_after=3)
    para(doc, "2）当日低风险：12:00–12:40（本期工作日普遍有午休低谷）。", space_after=3)
    para(
        doc,
        "3）可以但需打招呼：18:30–20:00，且当晚管理端不在晚间网络 B。"
        "本期多数工作日晚上 cshen 都在线；9 月 2 日 20:00 还赶上 zhougang / musheng 首次登录。",
        space_after=3,
    )
    para(
        doc,
        "4）不要突然发布：工作日 09:00–11:30 与 13:00–17:30。"
        "这是热力图核心、CTP live 轮询、以及同事重叠时段（cshen + sunjie，部分上午还有 luoshuang）。",
        space_after=12,
    )

    heading(doc, "九、建议", 1)
    para(doc, "产品 / 运营", size=11, bold=True, color=NAVY, space_after=4)
    para(doc, "• 本期真正被使用的页面是私募产品和 MOM 分析；工作日 09:00–17:30 要保证列表和详情接口快。", space_after=3)
    para(doc, "• CTP live 与跟踪基金轮询是背景负载。行情页挂着不关，部署检查器会显示「很忙」。", space_after=3)
    para(doc, "• 9 月 2 日晚两位同事用手机完成首次登录。确认他们能自己找到私募产品，不必管理员在旁边。", space_after=8)
    para(doc, "工程卫生", size=11, bold=True, color=NAVY, space_after=4)
    para(doc, f"• 去掉 Vercel insights 脚本（本期约 {fmt_int(vercel_404)} 次重复 404）。", space_after=3)
    para(doc, f"• 排查 502（{fmt_int(n_502)}）与 499（{fmt_int(n_499)}）：进程重启或 nginx 上游超时。", space_after=3)
    para(doc, "• 可选：若侧栏 9 条预取带来的重复 SSR 有成本，可关掉兄弟行情/AI 链接的 Next.js prefetch。", space_after=3)
    para(doc, "• 部署就绪继续看已登录用户存在，不要看独立 IP——办公室 NAT 会把 sunjie 藏在 cshen 后面。", space_after=8)
    para(doc, "度量", size=11, bold=True, color=NAVY, space_after=4)
    para(
        doc,
        "本期覆盖 nginx 轮转窗口内全部约两周记录，末日为部分日。"
        "请继续保留 access.log.*.gz（约 14 天），以便下周仍能复盘。"
        "人头看登录历史，功能看 nginx。",
        space_after=14,
    )
    para(
        doc,
        "密级：内部。流量日志含客户端 IP 与 User-Agent，请勿转发到团队以外。",
        size=8,
        color=MUTED,
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(report_path))
    return report_path


def main() -> None:
    configure_matplotlib()
    d = load()
    charts = make_charts(d)
    path = build_doc(d, charts)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
