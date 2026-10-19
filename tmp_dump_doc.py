from pathlib import Path
from docx import Document
docs = Path(r"d:\coding\market_website\docs")
target = [p for p in docs.glob("*.docx") if "\u7b56\u7565" in p.name][0]
doc = Document(str(target))
out = []
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else ""
    out.append(f"[{i}|{style}] {p.text}")
for ti, table in enumerate(doc.tables):
    out.append(f"\n=== TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} ===")
    for ri, row in enumerate(table.rows):
        cells = [c.text.replace("\n", " ") for c in row.cells]
        out.append(f"R{ri}: " + " | ".join(cells))
Path(r"d:\coding\market_website\tmp_doc_text.txt").write_text("\n".join(out), encoding="utf-8")
print("paras", len(doc.paragraphs), "tables", len(doc.tables), "target", target.name.encode("unicode_escape").decode())
