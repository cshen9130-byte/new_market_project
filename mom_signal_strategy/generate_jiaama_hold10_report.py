# -*- coding: utf-8 -*-
"""Deep-dive: 只做加码同向, enter and hold 10 trading days (20M account)."""
from __future__ import annotations

import math
import sys
import traceback
from datetime import date
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "report_output" / "jiaama_hold10"
CHART_DIR = OUT_DIR / "charts"
REPORT_PATH = BASE_DIR / "MOM决策信号_只做加码同向_持有10日_回测报告.docx"
REPORT_PATH_ASCII = BASE_DIR / "MOM_signal_jiaama_only_hold10_backtest.docx"

sys.path.insert(0, str(BASE_DIR))
from _mom_20m_account import (  # noqa: E402
    COMM_RATE,
    MAX_MARGIN_UTIL,
    SLIP_RATE,
    START_EQUITY,
    account_stats,
    fmt_yuan,
    run_account,
    run_account_fixed_hold,
)
from generate_jiaama_crowd_report import (  # noqa: E402
    add_strategy_alignment,
    bucket_stats,
    drawdown_episodes,
    half_sample,
    streak_stats,
    tstat,
    year_month_tables,
)
from generate_jiaama_only_report import (  # noqa: E402
    draw_event_horizon_charts,
    event_horizon_stats,
    group_event_table,
)
from generate_mom_signal_backtest_report import (  # noqa: E402
    GOLD,
    MUTED,
    NAVY,
    TEXT,
    add_picture,
    add_table,
    add_text,
    apply_font,
    build_returns,
    build_signals,
    caption,
    configure_matplotlib,
    event_study,
    fmt_num,
    fmt_pct,
    fmt_t,
    fp,
    get_conn,
    heading,
    load_data,
    load_env,
    para,
    set_run_font,
)

C_NAVY = "#1A365D"
C_GOLD = "#C9A227"
C_RED = "#C53030"
C_GREEN = "#2F855A"
C_BLUE = "#2B6CB0"
C_ORANGE = "#DD6B20"
C_TEAL = "#0F766E"
HOLD_DAYS = 10


def save_fig(fig, name: str) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def entry_hold_stats(holds: pd.DataFrame) -> pd.DataFrame:
    if holds.empty or "entry_date" not in holds.columns:
        return pd.DataFrame()
    g = holds.groupby(["product", "name", "entry_date", "dir", "action"], as_index=False).agg(
        days=("hold_date", "nunique"),
        pnl=("pnl", "sum"),
        sector=("sector", "first"),
    )
    return g


def draw_charts(acct, acct_daily, ev, ev_row) -> dict:
    daily = acct["daily"]
    holds = acct["holds"]
    trades = acct["trades"]
    out = {}
    if daily.empty:
        return out
    x = pd.to_datetime(daily["return_date"])

    fig, ax = plt.subplots(figsize=(11.2, 5.2), dpi=160)
    ax.plot(x, daily["equity"] / 1e4, color=C_NAVY, lw=1.95, label="加码 · 持有 10 日")
    if not acct_daily["daily"].empty:
        ax.plot(pd.to_datetime(acct_daily["daily"]["return_date"]), acct_daily["daily"]["equity"] / 1e4,
                color=C_RED, lw=1.25, alpha=0.85, label="加码 · 信号在才持有（原回测）")
    ax.axhline(START_EQUITY / 1e4, color="#A0AEC0", ls="--", lw=1, label="起始 2,000 万")
    ax.set_title("2,000 万账户权益（扣手续费与滑点）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("权益（万元）", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["equity"] = save_fig(fig, "equity.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
    ax.fill_between(x, daily["dd"] * 100, 0, color=C_RED, alpha=0.35)
    ax.set_title("持有 10 日：账户回撤", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("回撤（%）", **fp())
    apply_font(ax)
    out["dd"] = save_fig(fig, "dd.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
    ax.plot(x, daily["cum_gross"] / 1e4, color=C_BLUE, lw=1.5, label="累计毛盈亏")
    ax.plot(x, daily["cum_cost"] / 1e4, color=C_ORANGE, lw=1.5, label="累计费用")
    ax.plot(x, daily["cum_net"] / 1e4, color=C_NAVY, lw=1.8, label="累计净盈亏")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("毛盈亏、费用与净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["cost"] = save_fig(fig, "cost.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.2), dpi=160)
    ax.plot(x, daily["margin"] / 1e4, color=C_NAVY, lw=1.4, label="保证金")
    ax.plot(x, daily["gross_notional"] / 1e4, color=C_ORANGE, lw=1.2, label="名义本金")
    ax.set_title("保证金占用与名义本金", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["margin"] = save_fig(fig, "margin.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.6), dpi=160)
    ax.plot(x, daily["n"], color=C_NAVY, lw=1.3)
    ax.set_title("每日持仓品种数（上限 8，持有期内不因信号消失而平）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("个数", **fp())
    apply_font(ax)
    out["n"] = save_fig(fig, "n.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    s = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("ME").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in s.values]
    ax.bar(s.index, s.values / 1e4, width=20, color=colors, align="center")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("月度净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    apply_font(ax)
    out["monthly"] = save_fig(fig, "monthly.png")

    fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=160)
    y = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("YE").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in y.values]
    ax.bar([str(i.year) for i in y.index], y.values / 1e4, color=colors, width=0.55)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("分年净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    apply_font(ax)
    out["yearly"] = save_fig(fig, "yearly.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
    roll = daily["ret"].rolling(60, min_periods=30)
    sharpe = roll.mean() / roll.std(ddof=1) * math.sqrt(252)
    ax.plot(x, sharpe, color=C_NAVY, lw=1.4)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.axhline(1.0, color=C_GOLD, ls="--", lw=1)
    ax.set_title("滚动 60 日年化夏普", fontsize=13, color=C_NAVY, **fp())
    apply_font(ax)
    out["roll_sharpe"] = save_fig(fig, "roll_sharpe.png")

    fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
    ax.hist(daily["ret"] * 100, bins=36, color=C_NAVY, alpha=0.82, edgecolor="white")
    ax.axvline(0, color=C_ORANGE, lw=1.2)
    ax.set_title("日收益率分布", fontsize=13, color=C_NAVY, **fp())
    ax.set_xlabel("日收益（%）", **fp())
    apply_font(ax)
    out["hist"] = save_fig(fig, "hist.png")

    out.update(draw_event_horizon_charts(ev_row, save_fig))

    entries = entry_hold_stats(holds)
    if not entries.empty:
        fig, ax = plt.subplots(figsize=(8.6, 4.4), dpi=160)
        ax.hist(entries["days"], bins=np.arange(0.5, HOLD_DAYS + 1.5, 1), color=C_NAVY, edgecolor="white")
        ax.set_title("每笔开仓的实际持有交易日", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("交易日", **fp())
        apply_font(ax)
        out["hold_hist"] = save_fig(fig, "hold_hist.png")

        fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
        by_dir = entries.groupby("dir")["pnl"].sum()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_dir.values]
        ax.bar(by_dir.index, by_dir.values / 1e4, color=colors, width=0.45)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("每笔 10 日持仓：多 vs 空累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["ls"] = save_fig(fig, "ls.png")

    if not holds.empty:
        fig, ax = plt.subplots(figsize=(10.4, 5.0), dpi=160)
        by_sec = holds.groupby("sector")["pnl"].sum().sort_values()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_sec.values]
        ax.barh(by_sec.index, by_sec.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("板块累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["sector"] = save_fig(fig, "sector.png")

        fig, ax = plt.subplots(figsize=(10.6, 6.0), dpi=160)
        by_p = holds.groupby("name")["pnl"].sum().sort_values()
        top = pd.concat([by_p.head(10), by_p.tail(10)]).drop_duplicates()
        colors = [C_RED if v >= 0 else C_GREEN for v in top.values]
        ax.barh(top.index, top.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("品种累计毛盈亏：最亏 / 最赚各 10 个", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["product"] = save_fig(fig, "product.png")

    if not trades.empty:
        fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
        t = trades.copy()
        t["dt"] = pd.to_datetime(t["trade_date"])
        turn = t.groupby("dt")["notional"].sum()
        ax.plot(turn.index, turn.rolling(10, min_periods=3).mean() / 1e4, color=C_NAVY, lw=1.5)
        ax.set_title("10 日平均成交名义（换手应低于逐日调仓）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["turnover"] = save_fig(fig, "turnover.png")

    return out


def set_core_header(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "MOM 每日风控  ·  只做加码同向 · 持有 10 日", size=8, color=MUTED)
    fp_ = section.footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp_, "内部研究  ·  加码开仓后持有 10 个交易日  ·  模拟账户 2,000 万  ·  ", size=8, color=MUTED)
    run = fp_.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color=MUTED)


def build_report(ctx: dict) -> Path:
    configure_matplotlib()
    daily = ctx["acct"]["daily"]
    holds = ctx["acct"]["holds"]
    trades = ctx["acct"]["trades"]
    st, st_d = ctx["stats"], ctx["stats_daily"]
    ch = ctx["charts"]
    ev_row = ctx["ev_row"]
    sig = ctx["sig"]
    start, end = ctx["start"], ctx["end"]
    episodes, monthly, yearly = ctx["episodes"], ctx["monthly"], ctx["yearly"]
    halves, streaks = ctx["halves"], ctx["streaks"]
    entries = ctx["entries"]
    ev = ctx["ev"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    set_core_header(doc)

    para(doc, "MOM 每日风控", size=12, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "只做加码同向 · 持有 10 个交易日", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "开仓后手数冻结  ·  不因信号消失提前平仓  ·  2,000 万模拟账户", size=13, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=10)
    para(
        doc,
        f"样本 {start} 至 {end}　　账户日 {st.get('n')}　　"
        f"手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp　　"
        f"生成于 {date.today().isoformat()}",
        size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=16,
    )

    heading(doc, "一、和原回测差在哪", 1)
    para(
        doc,
        "原「只做加码」账户不是严格的隔夜 1 日策略，而是每个信号日收盘后按当天还在的加码名单重建持仓："
        "品种明天还是加码就留下并可能改手数，不再是加码就下一交易日平掉。"
        "加码经常只亮一两天，所以实际持仓中位数大约 2 个交易日，约一半笔数只持有 1 日。"
        "本报告改成：加码一旦开仓，手数冻结，持有 10 个交易日再平；"
        "持有期内同一品种不再加仓、不重置计时；若共识方向反过来则先平再按新方向开。"
        "控拥挤、补风格、观望、暂缓、减码仍然不做。",
    )
    add_table(
        doc,
        ["规则", "原回测（信号在才持有）", "本报告（持有 10 日）"],
        [
            ["开仓", "当天加码、按当日权益定手数", "当天加码、按开仓日权益定手数"],
            ["持有中", "每日按新名单改手数", "手数冻结"],
            ["平仓", "加码消失的下一交易日", "满 10 个交易日后平"],
            ["同一品种再出现加码", "继续持有并调仓", "持有期内忽略（不叠仓）"],
            ["费用", "名单一变就双边开平", "通常只有开仓+到期平仓"],
        ],
        col_widths=[3.4, 6.6, 6.4],
    )
    caption(doc, "表 1  两种持有口径。原回测更像滚动 overlay；本报告更像事件驱动的 10 日持有。")
    para(
        doc,
        f"账户初始权益 {fmt_yuan(START_EQUITY)}。新开仓按当时已有持仓数 + 新开个数分摊约 2.2 倍名义，"
        f"最多 8 个品种，单品种一日 1σ 不超过权益 1.2%，保证金不超过 {MAX_MARGIN_UTIL:.0%}。"
        f"费用单边手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp，开平都收。",
    )

    heading(doc, "二、核心结论", 1)
    if st:
        para(
            doc,
            f"持有 10 日口径：期末权益 {fmt_yuan(st.get('end'))}，净盈亏 {fmt_yuan(st.get('pnl'), signed=True)}，"
            f"相对起始 {fmt_pct((st.get('end') - START_EQUITY) / START_EQUITY)}。"
            f"年化 {fmt_pct(st.get('cagr'))}，波动 {fmt_pct(st.get('vol'))}，夏普 {fmt_num(st.get('sharpe'))}，"
            f"最大回撤 {fmt_pct(st.get('maxdd'))}（约 {fmt_yuan(st.get('maxdd_yuan'))}），"
            f"日胜率 {fmt_pct(st.get('hit'), already=False)}。"
            f"费用合计 {fmt_yuan(st.get('total_cost'))}，毛盈亏 {fmt_yuan(st.get('total_gross'), signed=True)}。"
            f"日均持仓 {fmt_num(st.get('avg_names'), 1)} 个，日均名义 {fmt_yuan(st.get('avg_notional'))}。",
        )
    add_table(
        doc,
        ["口径", "期末权益", "净盈亏", "年化", "波动", "夏普", "最大回撤", "费用"],
        [
            ["加码 · 持有 10 日", fmt_yuan(st.get("end")), fmt_yuan(st.get("pnl"), signed=True),
             fmt_pct(st.get("cagr")), fmt_pct(st.get("vol")), fmt_num(st.get("sharpe")),
             fmt_pct(st.get("maxdd")), fmt_yuan(st.get("total_cost"))],
            ["加码 · 信号在才持有（原）", fmt_yuan(st_d.get("end")), fmt_yuan(st_d.get("pnl"), signed=True),
             fmt_pct(st_d.get("cagr")), fmt_pct(st_d.get("vol")), fmt_num(st_d.get("sharpe")),
             fmt_pct(st_d.get("maxdd")), fmt_yuan(st_d.get("total_cost"))],
        ],
        signed_cols={2, 3, 5, 6},
    )
    caption(doc, "表 2  同一套加码信号、同一套 2,000 万仓位与费率，只改持有期。")

    n_add = int((sig["action"] == "加码").sum())
    para(
        doc,
        f"产品级加码 {n_add:,} 条。事件研究里 10 日累计对齐 {fmt_pct(ev_row.get('mu10'))}，"
        f"t={fmt_t(ev_row.get('t10'))}，折合每日 {fmt_pct(ev_row.get('mu10_pd'))}；"
        f"次日 {fmt_pct(ev_row.get('mu1'))}（折合每日 {fmt_pct(ev_row.get('mu1_pd'))}），"
        f"第6–10日增量 {fmt_pct(ev_row.get('mu_inc6_10'))}。"
        "10 日累计柱更高，不表示持有 10 日更好——要比折合每日、增量，以及表 2 的账户对照。"
        "若 10 日账户明显差于原回测，说明加码的边更短，拉长持有是在吃均值回归；"
        "若更好，说明原回测过早平掉了还在走的共识。",
    )
    if not entries.empty:
        para(
            doc,
            f"实际开仓 {len(entries)} 笔，平均持有 {entries['days'].mean():.1f} 个交易日，"
            f"中位数 {entries['days'].median():.0f}，满 10 日的 {int((entries['days'] >= HOLD_DAYS).sum())} 笔。"
            "样本末未满期的持仓会短于 10 日。持有期内再出现的加码被忽略，所以开仓笔数少于信号条数。",
        )

    heading(doc, "三、账户路径", 1)
    for key, text in (
        ("equity", "图 1  蓝线=持有 10 日；红线=原「信号在才持有」。"),
        ("dd", "图 2  持有 10 日的回撤。拉长持有通常让回撤更连续。"),
        ("cost", "图 3  毛盈亏、费用与净盈亏。换手应低于逐日调仓。"),
        ("margin", "图 4  保证金与名义。持有期内名义不随名单每天归零。"),
        ("n", "图 5  每日持仓品种数。"),
        ("hold_hist", "图 6  每笔开仓的持有天数。目标是 10。"),
        ("monthly", "图 7  月度净盈亏。"),
        ("yearly", "图 8  分年净盈亏。"),
        ("roll_sharpe", "图 9  滚动 60 日年化夏普。"),
        ("hist", "图 10  日收益分布。"),
    ):
        if key in ch:
            add_picture(doc, ch[key])
            caption(doc, text)

    heading(doc, "四、事件研究：10 日到底有没有边", 1)
    para(
        doc,
        "事件研究按共识方向对齐品种收益，不问账户仓位、费用和名额。"
        "累计 10 日均值高于次日，不能当成「本账户应该更好」："
        "窗口更长、重叠样本会把同一段行情数好几遍，而且持有期内新加码进不去。"
        "要看增量（第6–10日还有没有边）和折合每日，再对照表 2 的账户路径。",
    )
    add_table(
        doc,
        ["口径", "样本", "均值", "折合每日", "中位数", "胜率", "t"],
        [
            ["累计·次日", f"{ev_row.get('n1')}", fmt_pct(ev_row.get("mu1")), fmt_pct(ev_row.get("mu1_pd")),
             fmt_pct(ev_row.get("med1")), fmt_pct(ev_row.get("hit1"), already=False), fmt_t(ev_row.get("t1"))],
            ["累计·5 日", f"{ev_row.get('n5')}", fmt_pct(ev_row.get("mu5")), fmt_pct(ev_row.get("mu5_pd")),
             fmt_pct(ev_row.get("med5")), fmt_pct(ev_row.get("hit5"), already=False), fmt_t(ev_row.get("t5"))],
            ["累计·10 日", f"{ev_row.get('n10')}", fmt_pct(ev_row.get("mu10")), fmt_pct(ev_row.get("mu10_pd")),
             fmt_pct(ev_row.get("med10")), fmt_pct(ev_row.get("hit10"), already=False), fmt_t(ev_row.get("t10"))],
            ["累计·20 日", f"{ev_row.get('n20')}", fmt_pct(ev_row.get("mu20")), fmt_pct(ev_row.get("mu20_pd")),
             fmt_pct(ev_row.get("med20")), fmt_pct(ev_row.get("hit20"), already=False), fmt_t(ev_row.get("t20"))],
            ["增量·第1日", f"{ev_row.get('n_inc1')}", fmt_pct(ev_row.get("mu_inc1")), "—",
             fmt_pct(ev_row.get("med_inc1")), fmt_pct(ev_row.get("hit_inc1"), already=False), fmt_t(ev_row.get("t_inc1"))],
            ["增量·第2–5日", f"{ev_row.get('n_inc2_5')}", fmt_pct(ev_row.get("mu_inc2_5")), "—",
             fmt_pct(ev_row.get("med_inc2_5")), fmt_pct(ev_row.get("hit_inc2_5"), already=False), fmt_t(ev_row.get("t_inc2_5"))],
            ["增量·第6–10日", f"{ev_row.get('n_inc6_10')}", fmt_pct(ev_row.get("mu_inc6_10")), "—",
             fmt_pct(ev_row.get("med_inc6_10")), fmt_pct(ev_row.get("hit_inc6_10"), already=False), fmt_t(ev_row.get("t_inc6_10"))],
            ["增量·第11–20日", f"{ev_row.get('n_inc11_20')}", fmt_pct(ev_row.get("mu_inc11_20")), "—",
             fmt_pct(ev_row.get("med_inc11_20")), fmt_pct(ev_row.get("hit_inc11_20"), already=False), fmt_t(ev_row.get("t_inc11_20"))],
        ],
        signed_cols={2, 3, 4, 5},
    )
    caption(doc, "表 3  累计窗口不能横比持有期。本账户对应「累计·10 日」，但决策应看增量和表 2。")
    if "event" in ch:
        add_picture(doc, ch["event"])
        caption(doc, "图 11  左：增量窗口。右：累计÷天数。10 日累计柱更高，不表示持有 10 日更好。")
    if "hit" in ch:
        add_picture(doc, ch["hit"])
        caption(doc, "图 11b  增量窗口胜率。累计胜率会把第 1 日的赢面带到 10 日，不能横比。")
    if ev_row.get("n_610_on") or ev_row.get("n_610_off"):
        para(
            doc,
            "第2–5日加码没了之后增量转负；第6–10日事件上「没了还拿」可以更高，"
            "但仍是重叠样本，且占住名额会挤掉新加码。以表 2 账户对照为准。"
            f"第2–5日还在/没了：{fmt_pct(ev_row.get('mu_25_on'))} vs {fmt_pct(ev_row.get('mu_25_off'))}；"
            f"第6–10日还在/没了：{fmt_pct(ev_row.get('mu_610_on'))} vs {fmt_pct(ev_row.get('mu_610_off'))}。",
        )
        add_table(
            doc,
            ["窗口", "加码是否还在", "样本", "增量均值", "胜率", "t"],
            [
                ["第2–5日", "还在", f"{ev_row.get('n_25_on')}", fmt_pct(ev_row.get("mu_25_on")),
                 fmt_pct(ev_row.get("hit_25_on"), already=False), fmt_t(ev_row.get("t_25_on"))],
                ["第2–5日", "已经没了", f"{ev_row.get('n_25_off')}", fmt_pct(ev_row.get("mu_25_off")),
                 fmt_pct(ev_row.get("hit_25_off"), already=False), fmt_t(ev_row.get("t_25_off"))],
                ["第6–10日", "还在", f"{ev_row.get('n_610_on')}", fmt_pct(ev_row.get("mu_610_on")),
                 fmt_pct(ev_row.get("hit_610_on"), already=False), fmt_t(ev_row.get("t_610_on"))],
                ["第6–10日", "已经没了", f"{ev_row.get('n_610_off')}", fmt_pct(ev_row.get("mu_610_off")),
                 fmt_pct(ev_row.get("hit_610_off"), already=False), fmt_t(ev_row.get("t_610_off"))],
            ],
            signed_cols={3, 4},
        )
        caption(doc, "表 3b  按加码还在不在拆开。绿柱高也不能推翻表 2 的账户结论。")
    if "persist" in ch:
        add_picture(doc, ch["persist"])
        caption(doc, "图 11c  红=信号还在；绿=已经没了。")

    add_ev = ev[ev["action"] == "加码"].copy()
    if not add_ev.empty:
        add_ev["strength"] = add_ev["q_pct"].abs() + add_ev["s_pct"].abs()
        b = bucket_stats(add_ev.assign(strat_1=add_ev["strat_10"]), "加码", "strength",
                         [0, 10, 16, 24, 80], ["<10", "10–16", "16–24", "≥24"])
        if not b.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in b.itertuples(index=False)]
            add_table(doc, ["|q%|+|s%|", "样本", "10日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 4  加码强度分层的 10 日对齐（把 strat_10 临时当作分层收益）。")
        add_ev["side"] = np.where(add_ev["kind"] == "consensus_short", "空头共识", "多头共识")
        # group_event_table uses strat_1; build a view with 10d in strat_1
        view = add_ev.copy()
        view["strat_1"] = view["strat_10"]
        view["strat_5"] = view["strat_20"]
        side = group_event_table(view, "side")
        if not side.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in side.itertuples(index=False)]
            add_table(doc, ["方向", "样本", "10日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 5  多头共识 vs 空头共识的 10 日对齐。")

    heading(doc, "五、多空、板块、品种与每笔持仓", 1)
    for key, text in (
        ("ls", "图 12  每笔 10 日持仓的多空毛盈亏。"),
        ("sector", "图 13  板块累计毛盈亏。"),
        ("product", "图 14  品种两端。"),
    ):
        if key in ch:
            add_picture(doc, ch[key])
            caption(doc, text)

    if not entries.empty:
        rows = []
        for d, r in entries.groupby("dir").agg(n=("pnl", "size"), pnl=("pnl", "sum"), days=("days", "mean")).iterrows():
            rows.append([d, f"{int(r['n'])}", f"{r['days']:.1f}", fmt_yuan(r["pnl"], signed=True)])
        add_table(doc, ["方向", "开仓笔数", "平均持有天", "累计毛盈亏"], rows, signed_cols={3})
        caption(doc, "表 6  按开仓方向。")

        best = entries.sort_values("pnl", ascending=False)
        show = pd.concat([best.head(12), best.tail(8)]).drop_duplicates()
        rows = []
        for r in show.itertuples(index=False):
            rows.append([
                f"{r.name}({r.product})", r.sector, r.dir, str(r.entry_date),
                str(int(r.days)), fmt_yuan(r.pnl, signed=True),
            ])
        add_table(doc, ["品种", "板块", "方向", "开仓日", "持有天", "该笔毛盈亏"], rows, signed_cols={5})
        caption(doc, "表 7  单笔 10 日持仓头部 12 + 尾部 8。")

    if not holds.empty:
        sec = holds.groupby("sector").agg(
            days=("hold_date", "nunique"), names=("product", "nunique"),
            pnl=("pnl", "sum"), hit=("pnl", lambda s: float((s > 0).mean())),
        ).sort_values("pnl", ascending=False)
        rows = [[i, f"{int(r.names)}", f"{int(r.days)}", fmt_yuan(r.pnl, signed=True), fmt_pct(r.hit, already=False)]
                for i, r in sec.iterrows()]
        add_table(doc, ["板块", "品种数", "有仓天数", "累计毛盈亏", "条数胜率"], rows, signed_cols={3, 4})
        caption(doc, "表 8  板块归因。")

    heading(doc, "六、分年、分月与回撤", 1)
    if not yearly.empty:
        rows = []
        for r in yearly.itertuples(index=False):
            rows.append([
                str(r.year), fmt_yuan(r.pnl, signed=True), fmt_yuan(r.cost),
                fmt_pct(r.cagr), fmt_pct(r.vol), fmt_num(r.sharpe), fmt_pct(r.maxdd),
                fmt_pct(r.hit, already=False), str(int(r.days)),
            ])
        add_table(doc, ["年", "净盈亏", "费用", "年化", "波动", "夏普", "最大回撤", "日胜率", "天数"], rows, signed_cols={1, 3, 5, 6})
        caption(doc, "表 9  分年绩效。")
    if not monthly.empty:
        rows = []
        for ym, r in monthly.iterrows():
            rows.append([
                ym, fmt_yuan(r["pnl"], signed=True), fmt_yuan(r["gross"], signed=True),
                fmt_yuan(r["cost"]), fmt_num(r["n"], 1), fmt_pct(r["hit"], already=False), fmt_yuan(r["eq"]),
            ])
        add_table(doc, ["月份", "净盈亏", "毛盈亏", "费用", "日均品种", "日胜率", "月末权益"], rows, signed_cols={1, 2})
        caption(doc, "表 10  分月结果。")

    if not episodes.empty:
        heading(doc, "6.1 深度回撤段", 2)
        rows = []
        for r in episodes.head(8).itertuples(index=False):
            rows.append([
                r.start, r.trough, r.end, str(int(r.days)), str(int(r.trough_days)),
                fmt_pct(r.depth), fmt_yuan(r.lost, signed=True),
                "未修复" if getattr(r, "open", False) else "已修复",
            ])
        add_table(doc, ["开始", "谷底", "结束", "持续天", "到谷底", "深度", "权益损失", "状态"], rows, signed_cols={5, 6})
        caption(doc, "表 11  深度至少 3% 的回撤段。")
        worst = episodes.iloc[0]
        para(
            doc,
            f"最深一段从 {worst['start']} 到谷底 {worst['trough']}，深度 {fmt_pct(worst['depth'])}，"
            f"权益最多少了 {fmt_yuan(worst['lost'])}。"
            "固定持有 10 日意味着共识反转后还要再拿一段时间，回撤往往比「信号没了就平」更深。",
        )

    heading(doc, "七、成交与费用", 1)
    if "turnover" in ch:
        add_picture(doc, ch["turnover"])
        caption(doc, "图 15  10 日平均成交名义。固定持有应明显低于逐日重建。")
    if not trades.empty:
        para(
            doc,
            f"成交 {len(trades):,} 笔，名义合计 {fmt_yuan(trades['notional'].sum())}，"
            f"费用 {fmt_yuan(trades['cost'].sum())}。"
            f"开/加 {(trades['side'] == '开/加').sum():,}，平仓 {(trades['side'] == '平仓').sum():,}，"
            f"反手 {(trades['side'] == '反手/平').sum():,}。",
        )
        add_table(
            doc,
            ["费用项", "持有10日", "原逐日调仓", "说明"],
            [
                ["手续费", fmt_yuan(st.get("total_comm")), fmt_yuan(st_d.get("total_comm")), f"单边 {COMM_RATE * 1e4:.1f}bp"],
                ["滑点", fmt_yuan(st.get("total_slip")), fmt_yuan(st_d.get("total_slip")), f"单边 {SLIP_RATE * 1e4:.1f}bp"],
                ["费用合计", fmt_yuan(st.get("total_cost")), fmt_yuan(st_d.get("total_cost")), "开平都收"],
                ["毛盈亏", fmt_yuan(st.get("total_gross"), signed=True), fmt_yuan(st_d.get("total_gross"), signed=True), "持仓×涨跌"],
                ["净盈亏", fmt_yuan(st.get("pnl"), signed=True), fmt_yuan(st_d.get("pnl"), signed=True), "毛−费"],
            ],
            signed_cols={1, 2},
        )
        caption(doc, "表 12  费用对照。持有 10 日的主要好处通常是少换手。")
        t20 = trades.sort_values("notional", ascending=False).head(12)
        trows = [[r.trade_date, f"{r.name}({r.product})", r.side,
                  f"{int(r.old_lots)}→{int(r.new_lots)}", fmt_yuan(r.notional), fmt_yuan(r.cost)]
                 for r in t20.itertuples(index=False)]
        add_table(doc, ["成交日", "品种", "动作", "手数", "名义", "费用"], trows)
        caption(doc, "表 13  名义最大的 12 笔。")

    heading(doc, "八、稳健性", 1)
    if halves:
        rows = [[h["name"], f"{h['start']} ~ {h['end']}", fmt_yuan(h.get("pnl"), signed=True),
                 fmt_pct(h.get("cagr")), fmt_num(h.get("sharpe")), fmt_pct(h.get("maxdd")),
                 fmt_pct(h.get("hit"), already=False)] for h in halves]
        add_table(doc, ["分段", "区间", "净盈亏", "年化", "夏普", "最大回撤", "日胜率"], rows, signed_cols={2, 3, 4, 5})
        caption(doc, "表 14  前后半样本。")
    if streaks:
        para(doc, f"赚钱日 {streaks['win_days']}、亏钱日 {streaks['loss_days']}；"
             f"最长连赢 {streaks['max_win_streak']} 日，最长连亏 {streaks['max_loss_streak']} 日。")
    if not daily.empty:
        rows = []
        for label, df in (("最亏", daily.nsmallest(8, "pnl_net")), ("最赚", daily.nlargest(8, "pnl_net"))):
            for r in df.itertuples(index=False):
                rows.append([label, r.return_date, fmt_yuan(r.pnl_net, signed=True),
                             fmt_yuan(r.cost), f"{int(r.n)}", fmt_pct(r.ret)])
        add_table(doc, ["类型", "日期", "净盈亏", "费用", "持仓数", "日收益"], rows, signed_cols={2, 5})
        caption(doc, "表 15  单日两端。")

    heading(doc, "九、持仓快照", 1)
    para(doc, "手数在开仓日冻结。下表是样本末日仍未到期的 10 日持仓（若有）。")

    def hold_table(df: pd.DataFrame, title: str):
        if df is None or df.empty:
            para(doc, f"{title}：当日无持仓。", first_line=False)
            return
        rows = []
        for r in df.sort_values("notional", ascending=False).itertuples(index=False):
            left = getattr(r, "remaining_after", "")
            rows.append([
                f"{r.name}({r.product})", r.sector, r.dir, str(int(r.lots)),
                getattr(r, "entry_date", ""), str(left),
                fmt_yuan(r.notional), fmt_yuan(r.pnl, signed=True),
            ])
        add_table(doc, ["品种", "板块", "方向", "手数", "开仓日", "剩余天", "名义", "当日盈亏"], rows, signed_cols={7})
        caption(doc, title)

    if not holds.empty:
        last_d = holds["hold_date"].max()
        hold_table(holds[holds["hold_date"] == last_d], f"表 16  样本末日 {last_d} 持仓。")

    heading(doc, "十、结论与局限", 1)
    para(
        doc,
        "原只做加码回测是「信号还在才拿」，不是固定隔夜 1 日，但加码名单换得快，很多笔事实上只拿 1～2 日。"
        "本报告把持有期钉死为 10 个交易日，用来回答：加码的边是脉冲还是能抱十天。"
        "对照表 2：账户路径才是持有期的答案。表 3 的 10 日累计均值高于次日，只说明窗口更长，"
        "不说明该抱十天；看增量和折合每日。若账户又差于原回测，就不要为了「少交易」去拉长持有。",
    )
    para(
        doc,
        "局限与原报告相同：主力连续价、近似保证金、无涨跌停与换月；"
        "持有期内忽略重复加码，会漏掉「加码持续很久本该加仓」的路径；"
        "最多 8 个名额时，后到的加码进不去。手数冻结也不再做波动目标再平衡。",
    )
    para(
        doc,
        f"图表在 mom_signal_strategy/report_output/jiaama_hold10/charts/。生成于 {date.today().isoformat()}。",
        size=10, color=MUTED, first_line=False,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(REPORT_PATH))
    try:
        doc.save(str(REPORT_PATH_ASCII))
    except Exception:
        pass
    return REPORT_PATH


def main():
    load_env()
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_conn()
    try:
        pos, px = load_data(conn)
    finally:
        conn.close()
    if pos.empty or px.empty:
        raise SystemExit("empty market data")

    wide, clean, close, _ = build_returns(px)
    sig = build_signals(pos, clean)
    if sig.empty:
        raise SystemExit("no signals")

    print("Event study…")
    ev = add_strategy_alignment(event_study(sig, wide))
    ev_row = event_horizon_stats(ev, "加码", sig=sig)

    print("20M: 加码 hold 10 days…")
    acct = run_account_fixed_hold(
        sig, close, wide, clean, hold_days=HOLD_DAYS,
        allowed_actions={"加码"}, include_bufengge=False,
    )
    print("20M: 加码 daily rebuild (original)…")
    acct_daily = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码"})

    stats = account_stats(acct["daily"])
    stats_daily = account_stats(acct_daily["daily"])
    print("Drawing charts…")
    charts = draw_charts(acct, acct_daily, ev, ev_row)

    if not acct["daily"].empty:
        acct["daily"].to_csv(OUT_DIR / "acct_daily.csv", index=False, encoding="utf-8-sig")
    if not acct["holds"].empty:
        acct["holds"].to_csv(OUT_DIR / "acct_holdings.csv", index=False, encoding="utf-8-sig")
    if not acct["trades"].empty:
        acct["trades"].to_csv(OUT_DIR / "acct_trades.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([ev_row]).to_csv(OUT_DIR / "event_stats.csv", index=False, encoding="utf-8-sig")

    monthly, yearly = year_month_tables(acct["daily"]) if not acct["daily"].empty else (pd.DataFrame(), pd.DataFrame())
    ctx = {
        "sig": sig, "ev": ev, "ev_row": ev_row, "acct": acct,
        "stats": stats, "stats_daily": stats_daily, "charts": charts,
        "start": str(sig["date"].min()), "end": str(sig["date"].max()),
        "episodes": drawdown_episodes(acct["daily"]), "monthly": monthly, "yearly": yearly,
        "halves": half_sample(acct["daily"]),
        "streaks": streak_stats(acct["daily"]["pnl_net"]) if not acct["daily"].empty else {},
        "entries": entry_hold_stats(acct["holds"]),
    }
    print("Writing Word report…")
    path = build_report(ctx)
    print(f"Wrote {path}")
    print(f"Also {REPORT_PATH_ASCII}")
    if stats:
        print(f"Hold10 end {stats.get('end'):,.0f} Sharpe {stats.get('sharpe')} MaxDD {stats.get('maxdd')}")
        print(f"Daily  end {stats_daily.get('end'):,.0f} Sharpe {stats_daily.get('sharpe')} MaxDD {stats_daily.get('maxdd')}")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        raise
