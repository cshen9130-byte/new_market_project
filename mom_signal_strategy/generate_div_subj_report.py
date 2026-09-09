# -*- coding: utf-8 -*-
"""Deep-dive Word report: 分歧跟主观（加码跟共识 + 观望跟主观）, 20M account."""
from __future__ import annotations

import math
import sys
import traceback
from datetime import date
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "report_output" / "div_subj"
CHART_DIR = OUT_DIR / "charts"
REPORT_PATH = BASE_DIR / "MOM决策信号_分歧跟主观_回测报告.docx"
REPORT_PATH_ASCII = BASE_DIR / "MOM_signal_divergence_follow_subjective_backtest.docx"

sys.path.insert(0, str(BASE_DIR))
from _mom_20m_account import (  # noqa: E402
    COMM_RATE,
    MAX_MARGIN_UTIL,
    SLIP_RATE,
    START_EQUITY,
    account_stats,
    fmt_yuan,
    run_account,
)
from generate_jiaama_crowd_report import (  # noqa: E402
    bucket_stats,
    drawdown_episodes,
    half_sample,
    streak_stats,
    tstat,
    year_month_tables,
)
from generate_mom_signal_backtest_report import (  # noqa: E402
    GOLD,
    MUTED,
    NAVY,
    TEXT,
    add_picture,
    add_table,
    add_text,
    apply_font,
    build_returns,
    build_signals,
    caption,
    configure_matplotlib,
    event_study,
    fmt_num,
    fmt_pct,
    fmt_t,
    fp,
    get_conn,
    heading,
    load_data,
    load_env,
    para,
    set_run_font,
)

C_NAVY = "#1A365D"
C_GOLD = "#C9A227"
C_RED = "#C53030"
C_GREEN = "#2F855A"
C_BLUE = "#2B6CB0"
C_ORANGE = "#DD6B20"
C_VIOLET = "#6D28D9"
C_TEAL = "#0F766E"


def save_fig(fig, name: str) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def align_div_subj(ev: pd.DataFrame) -> pd.DataFrame:
    out = ev.copy()
    dirs = []
    qdirs = []
    for r in out.itertuples(index=False):
        if r.action == "加码":
            d = float(r.dir) if r.dir else (1.0 if (r.q_pct + r.s_pct) >= 0 else -1.0)
        elif r.action == "观望":
            d = 1.0 if r.s_pct > 0 else (-1.0 if r.s_pct < 0 else 0.0)
        else:
            d = 0.0
        dirs.append(d)
        qdirs.append(1.0 if r.q_pct > 0 else (-1.0 if r.q_pct < 0 else 0.0))
    out["strat_dir"] = dirs
    out["quant_dir"] = qdirs
    for src, dest, dcol in (
        ("r1", "strat_1", "strat_dir"),
        ("r5", "strat_5", "strat_dir"),
        ("r10", "strat_10", "strat_dir"),
        ("r20", "strat_20", "strat_dir"),
        ("r1", "quant_1", "quant_dir"),
        ("r5", "quant_5", "quant_dir"),
        ("r10", "quant_10", "quant_dir"),
        ("r20", "quant_20", "quant_dir"),
    ):
        aligned = []
        for r in out.itertuples(index=False):
            raw = getattr(r, src)
            d = getattr(r, dcol)
            if raw is None or (isinstance(raw, float) and math.isnan(raw)) or d == 0:
                aligned.append(np.nan)
            else:
                aligned.append(float(d) * float(raw))
        out[dest] = aligned
    return out


def event_horizon_stats(ev: pd.DataFrame, action: str, col_prefix="strat") -> dict:
    g = ev[ev["action"] == action]
    row = {"action": action, "n": int(len(g))}
    for h in (1, 5, 10, 20):
        col = f"{col_prefix}_{h}"
        if col not in g.columns:
            row[f"n{h}"] = 0
            row[f"mu{h}"] = row[f"hit{h}"] = row[f"t{h}"] = row[f"med{h}"] = None
            continue
        s = pd.to_numeric(g[col], errors="coerce").dropna()
        row[f"n{h}"] = int(len(s))
        row[f"mu{h}"] = float(s.mean()) if len(s) else None
        row[f"hit{h}"] = float((s > 0).mean()) if len(s) else None
        row[f"t{h}"] = tstat(s)
        row[f"med{h}"] = float(s.median()) if len(s) else None
    return row


def group_event_table(ev: pd.DataFrame, action: str, col: str) -> pd.DataFrame:
    g = ev[ev["action"] == action].copy()
    if g.empty or col not in g.columns:
        return pd.DataFrame()
    rows = []
    for lab, sub in g.groupby(col, dropna=False):
        s = pd.to_numeric(sub["strat_1"], errors="coerce").dropna()
        rows.append({
            "bucket": "（空）" if pd.isna(lab) else str(lab),
            "n": int(len(s)),
            "mu1": float(s.mean()) if len(s) else None,
            "hit1": float((s > 0).mean()) if len(s) else None,
            "t1": tstat(s),
            "mu5": float(pd.to_numeric(sub["strat_5"], errors="coerce").dropna().mean()) if len(sub) else None,
        })
    return pd.DataFrame(rows).sort_values("n", ascending=False)


def draw_charts(acct, acct_add, acct_div, acct_q, ev, ev_add, ev_div) -> dict:
    daily = acct["daily"]
    holds = acct["holds"]
    trades = acct["trades"]
    out = {}
    if daily.empty:
        return out
    x = pd.to_datetime(daily["return_date"])

    fig, ax = plt.subplots(figsize=(11.2, 5.2), dpi=160)
    ax.plot(x, daily["equity"] / 1e4, color=C_NAVY, lw=1.95, label="分歧跟主观")
    if not acct_add["daily"].empty:
        ax.plot(pd.to_datetime(acct_add["daily"]["return_date"]), acct_add["daily"]["equity"] / 1e4,
                color=C_RED, lw=1.25, alpha=0.85, label="只做加码")
    if not acct_div["daily"].empty:
        ax.plot(pd.to_datetime(acct_div["daily"]["return_date"]), acct_div["daily"]["equity"] / 1e4,
                color=C_VIOLET, lw=1.25, alpha=0.9, label="只做观望跟主观")
    if not acct_q["daily"].empty:
        ax.plot(pd.to_datetime(acct_q["daily"]["return_date"]), acct_q["daily"]["equity"] / 1e4,
                color=C_GOLD, lw=1.2, alpha=0.85, label="分歧跟量化（对照）")
    ax.axhline(START_EQUITY / 1e4, color="#A0AEC0", ls="--", lw=1, label="起始 2,000 万")
    ax.set_title("2,000 万账户权益（扣手续费与滑点）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("权益（万元）", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["equity"] = save_fig(fig, "equity.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
    ax.fill_between(x, daily["dd"] * 100, 0, color=C_RED, alpha=0.35)
    ax.set_title("分歧跟主观：账户回撤", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("回撤（%）", **fp())
    apply_font(ax)
    out["dd"] = save_fig(fig, "dd.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
    ax.plot(x, daily["cum_gross"] / 1e4, color=C_BLUE, lw=1.5, label="累计毛盈亏")
    ax.plot(x, daily["cum_cost"] / 1e4, color=C_ORANGE, lw=1.5, label="累计费用")
    ax.plot(x, daily["cum_net"] / 1e4, color=C_NAVY, lw=1.8, label="累计净盈亏")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("毛盈亏、费用与净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["cost"] = save_fig(fig, "cost.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.2), dpi=160)
    ax.plot(x, daily["margin"] / 1e4, color=C_NAVY, lw=1.4, label="保证金")
    ax.plot(x, daily["gross_notional"] / 1e4, color=C_ORANGE, lw=1.2, label="名义本金")
    ax.set_title("保证金占用与名义本金", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["margin"] = save_fig(fig, "margin.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.6), dpi=160)
    ax.plot(x, daily["n"], color=C_NAVY, lw=1.3)
    ax.set_title("每日持仓品种数（上限 8）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("个数", **fp())
    apply_font(ax)
    out["n"] = save_fig(fig, "n.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    s = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("ME").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in s.values]
    ax.bar(s.index, s.values / 1e4, width=20, color=colors, align="center")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("月度净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    apply_font(ax)
    out["monthly"] = save_fig(fig, "monthly.png")

    fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=160)
    y = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("YE").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in y.values]
    ax.bar([str(i.year) for i in y.index], y.values / 1e4, color=colors, width=0.55)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("分年净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    apply_font(ax)
    out["yearly"] = save_fig(fig, "yearly.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
    roll = daily["ret"].rolling(60, min_periods=30)
    sharpe = roll.mean() / roll.std(ddof=1) * math.sqrt(252)
    ax.plot(x, sharpe, color=C_NAVY, lw=1.4)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.axhline(1.0, color=C_GOLD, ls="--", lw=1)
    ax.set_title("滚动 60 日年化夏普", fontsize=13, color=C_NAVY, **fp())
    apply_font(ax)
    out["roll_sharpe"] = save_fig(fig, "roll_sharpe.png")

    fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
    ax.hist(daily["ret"] * 100, bins=36, color=C_NAVY, alpha=0.82, edgecolor="white")
    ax.axvline(0, color=C_ORANGE, lw=1.2)
    ax.set_title("日收益率分布", fontsize=13, color=C_NAVY, **fp())
    ax.set_xlabel("日收益（%）", **fp())
    apply_font(ax)
    out["hist"] = save_fig(fig, "hist.png")

    if not holds.empty:
        fig, ax = plt.subplots(figsize=(10.6, 4.6), dpi=160)
        for action, color in (("加码", C_RED), ("观望", C_VIOLET)):
            h = holds[holds["action"] == action]
            if h.empty:
                continue
            c = h.groupby("hold_date")["pnl"].sum()
            ax.plot(pd.to_datetime(c.index), c.cumsum() / 1e4, color=color, lw=1.8, label=f"{action} 累计毛盈亏")
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("持仓毛盈亏：加码 vs 观望跟主观", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["by_action"] = save_fig(fig, "by_action.png")

        fig, ax = plt.subplots(figsize=(8.4, 4.2), dpi=160)
        by_act = holds.groupby("action")["pnl"].sum().reindex(["加码", "观望"]).dropna()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_act.values]
        ax.bar(by_act.index.map(lambda a: "加码" if a == "加码" else "观望跟主观"), by_act.values / 1e4,
               color=colors, width=0.5)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("两种信号的累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["action_bar"] = save_fig(fig, "action_bar.png")

        mix = holds.groupby(["hold_date", "action"]).size().unstack(fill_value=0)
        mix = mix.reindex(columns=["加码", "观望"], fill_value=0)
        fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
        ax.stackplot(pd.to_datetime(mix.index), mix["加码"], mix["观望"],
                     labels=["加码", "观望跟主观"], colors=[C_RED, C_VIOLET], alpha=0.85)
        ax.set_title("每日持仓中两种信号的个数", fontsize=13, color=C_NAVY, **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["mix"] = save_fig(fig, "mix.png")

        fig, ax = plt.subplots(figsize=(8.4, 4.2), dpi=160)
        by_dir = holds.groupby("dir")["pnl"].sum()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_dir.values]
        ax.bar(by_dir.index, by_dir.values / 1e4, color=colors, width=0.45)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("多头 vs 空头累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["ls"] = save_fig(fig, "ls.png")

        fig, ax = plt.subplots(figsize=(10.4, 5.0), dpi=160)
        by_sec = holds.groupby("sector")["pnl"].sum().sort_values()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_sec.values]
        ax.barh(by_sec.index, by_sec.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("板块累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["sector"] = save_fig(fig, "sector.png")

        fig, ax = plt.subplots(figsize=(10.6, 6.0), dpi=160)
        by_p = holds.groupby("name")["pnl"].sum().sort_values()
        top = pd.concat([by_p.head(10), by_p.tail(10)]).drop_duplicates()
        colors = [C_RED if v >= 0 else C_GREEN for v in top.values]
        ax.barh(top.index, top.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("品种累计毛盈亏：最亏 / 最赚各 10 个", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["product"] = save_fig(fig, "product.png")

    fig, ax = plt.subplots(figsize=(10.2, 5.0), dpi=160)
    horizons = [1, 5, 10, 20]
    x_pos = np.arange(len(horizons))
    width = 0.36
    for i, (row, label, color) in enumerate(((ev_add, "加码跟共识", C_RED), (ev_div, "观望跟主观", C_VIOLET))):
        vals = [row.get(f"mu{h}") or 0 for h in horizons]
        ax.bar(x_pos + (i - 0.5) * width, [v * 100 for v in vals], width=width, color=color, label=label)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_xticks(x_pos)
    ax.set_xticklabels(["次日", "5日", "10日", "20日"])
    ax.set_ylabel("策略方向对齐平均收益（%）", **fp())
    ax.set_title("事件研究：按本策略方向对齐", fontsize=13, color=C_NAVY, **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["event"] = save_fig(fig, "event.png")

    fig, ax = plt.subplots(figsize=(9.2, 4.6), dpi=160)
    labels, hits, colors = [], [], []
    for row, name, color in ((ev_add, "加码", C_RED), (ev_div, "观望跟主观", C_VIOLET)):
        for h, lab in ((1, "次日"), (5, "5日"), (20, "20日")):
            v = row.get(f"hit{h}")
            if v is None:
                continue
            labels.append(f"{name}\n{lab}")
            hits.append(v * 100)
            colors.append(color)
    ax.bar(range(len(hits)), hits, color=colors, width=0.62)
    ax.axhline(50, color="#A0AEC0", ls="--", lw=1)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_ylabel("胜率（%）", **fp())
    ax.set_title("策略方向对齐胜率", fontsize=13, color=C_NAVY, **fp())
    apply_font(ax)
    out["hit"] = save_fig(fig, "hit.png")

    div = ev[ev["action"] == "观望"]
    if not div.empty:
        cnt = div.groupby("date").size()
        fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
        ax.bar(pd.to_datetime(cnt.index), cnt.values, color=C_VIOLET, width=1.2, alpha=0.85)
        ax.set_title("每日观望（分歧）信号条数", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("条数", **fp())
        apply_font(ax)
        out["sig_n"] = save_fig(fig, "sig_n.png")

        sub = div.copy()
        sub["sabs"] = sub["s_pct"].abs()
        b = bucket_stats(sub, "观望", "sabs", [0, 4, 8, 14, 80], ["<4", "4–8", "8–14", "≥14"])
        if not b.empty:
            fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
            vals = [(v or 0) * 100 for v in b["mu1"]]
            colors = [C_RED if v >= 0 else C_GREEN for v in vals]
            ax.bar(b["bucket"].astype(str), vals, color=colors, width=0.55)
            ax.axhline(0, color="#4A5568", lw=0.8)
            ax.set_title("观望跟主观：按 |主观风险%| 分层的次日对齐", fontsize=13, color=C_NAVY, **fp())
            ax.set_ylabel("%", **fp())
            apply_font(ax)
            out["s_strength"] = save_fig(fig, "s_strength.png")

    if not trades.empty:
        fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
        t = trades.copy()
        t["dt"] = pd.to_datetime(t["trade_date"])
        turn = t.groupby("dt")["notional"].sum()
        ax.plot(turn.index, turn.rolling(10, min_periods=3).mean() / 1e4, color=C_NAVY, lw=1.5)
        ax.set_title("10 日平均成交名义（换手）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["turnover"] = save_fig(fig, "turnover.png")

    return out


def set_core_header(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "MOM 每日风控  ·  分歧跟主观 深度回测", size=8, color=MUTED)
    fp_ = section.footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp_, "内部研究  ·  加码跟共识 + 观望跟主观  ·  模拟账户 2,000 万  ·  ", size=8, color=MUTED)
    run = fp_.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color=MUTED)


def ev_rows(row: dict) -> list:
    return [
        ["次日", f"{row.get('n1')}", fmt_pct(row.get("mu1")), fmt_pct(row.get("med1")),
         fmt_pct(row.get("hit1"), already=False), fmt_t(row.get("t1"))],
        ["5 日", f"{row.get('n5')}", fmt_pct(row.get("mu5")), fmt_pct(row.get("med5")),
         fmt_pct(row.get("hit5"), already=False), fmt_t(row.get("t5"))],
        ["10 日", f"{row.get('n10')}", fmt_pct(row.get("mu10")), fmt_pct(row.get("med10")),
         fmt_pct(row.get("hit10"), already=False), fmt_t(row.get("t10"))],
        ["20 日", f"{row.get('n20')}", fmt_pct(row.get("mu20")), fmt_pct(row.get("med20")),
         fmt_pct(row.get("hit20"), already=False), fmt_t(row.get("t20"))],
    ]


def build_report(ctx: dict) -> Path:
    configure_matplotlib()
    daily = ctx["acct"]["daily"]
    holds = ctx["acct"]["holds"]
    trades = ctx["acct"]["trades"]
    st, st_add, st_div, st_q = ctx["stats"], ctx["stats_add"], ctx["stats_div"], ctx["stats_q"]
    ch = ctx["charts"]
    ev_add, ev_div, ev_div_q = ctx["ev_add"], ctx["ev_div"], ctx["ev_div_q"]
    sig, ev = ctx["sig"], ctx["ev"]
    start, end = ctx["start"], ctx["end"]
    episodes, monthly, yearly = ctx["episodes"], ctx["monthly"], ctx["yearly"]
    halves, streaks = ctx["halves"], ctx["streaks"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    set_core_header(doc)

    para(doc, "MOM 每日风控", size=12, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "只做分歧跟主观", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "加码跟共识  ·  观望日跟主观方向  ·  2,000 万模拟账户深度回测", size=13, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=10)
    para(
        doc,
        f"样本 {start} 至 {end}　　交易日 {st.get('n')}　　"
        f"手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp（开平都收）　　"
        f"生成于 {date.today().isoformat()}",
        size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=16,
    )

    heading(doc, "一、策略定义", 1)
    para(
        doc,
        "总报告里「分歧跟主观」是样本内 overlay 夏普最高的一条研究线索。"
        "本报告把它落到 2,000 万可交易账户：加码日与量化/主观共识同向；"
        "观望日（两边都≥3% 但方向相反）改跟主观风险占比的方向；"
        "控拥挤、补风格、暂缓加码、减码准备一律空仓。"
        "对照账户分别拿掉观望腿、只留观望腿、以及把观望改成跟量化，用来回答「这一腿值不值得做、该不该跟主观」。",
    )
    add_table(
        doc,
        ["信号", "本账户怎么做", "经济含义"],
        [
            ["加码", "与量化/主观共识同向", "两边已经确认，复制这份 beta"],
            ["观望", "跟主观方向（主观多为买、主观空为卖）", "两边冲突时，站主观一侧"],
            ["控拥挤 / 补风格 / 暂缓 / 减码", "空仓并平掉旧仓", "不在本策略里"],
        ],
        col_widths=[4.4, 6.2, 5.8],
    )
    caption(doc, "表 1  分歧跟主观的下单规则。信号不再是加码或观望时，下一交易日平仓。")
    para(
        doc,
        f"账户初始权益 {fmt_yuan(START_EQUITY)}。仓位按当日权益约 2.2 倍名义、最多 8 个品种等权分配，"
        f"单品种一日 1σ 亏损不超过权益的 1.2%，保证金占用不超过权益的 {MAX_MARGIN_UTIL:.0%}。"
        f"手数取整。费用按成交名义单边手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp，"
        "每手手续费不低于 3 元，开平都收。行情用主力连续合约。",
    )

    heading(doc, "二、核心结论", 1)
    if st:
        para(
            doc,
            f"样本内期末权益 {fmt_yuan(st.get('end'))}，净盈亏 {fmt_yuan(st.get('pnl'), signed=True)}，"
            f"相对起始资金 {fmt_pct((st.get('end') - START_EQUITY) / START_EQUITY)}。"
            f"年化 {fmt_pct(st.get('cagr'))}，波动 {fmt_pct(st.get('vol'))}，夏普 {fmt_num(st.get('sharpe'))}，"
            f"卡玛 {fmt_num(st.get('calmar'))}，最大回撤 {fmt_pct(st.get('maxdd'))}"
            f"（约 {fmt_yuan(st.get('maxdd_yuan'))}），日胜率 {fmt_pct(st.get('hit'), already=False)}。"
            f"手续费 {fmt_yuan(st.get('total_comm'))}、滑点 {fmt_yuan(st.get('total_slip'))}，"
            f"费用合计 {fmt_yuan(st.get('total_cost'))}，毛盈亏 {fmt_yuan(st.get('total_gross'), signed=True)}。"
            f"日均持仓 {fmt_num(st.get('avg_names'), 1)} 个品种，日均保证金 {fmt_yuan(st.get('avg_margin'))}，"
            f"占用率 {fmt_pct(st.get('avg_util'))}，日均名义 {fmt_yuan(st.get('avg_notional'))}。",
        )
    add_table(
        doc,
        ["口径", "期末权益", "净盈亏", "年化", "波动", "夏普", "最大回撤", "费用"],
        [
            ["分歧跟主观（本报告）", fmt_yuan(st.get("end")), fmt_yuan(st.get("pnl"), signed=True),
             fmt_pct(st.get("cagr")), fmt_pct(st.get("vol")), fmt_num(st.get("sharpe")),
             fmt_pct(st.get("maxdd")), fmt_yuan(st.get("total_cost"))],
            ["只做加码（对照）", fmt_yuan(st_add.get("end")), fmt_yuan(st_add.get("pnl"), signed=True),
             fmt_pct(st_add.get("cagr")), fmt_pct(st_add.get("vol")), fmt_num(st_add.get("sharpe")),
             fmt_pct(st_add.get("maxdd")), fmt_yuan(st_add.get("total_cost"))],
            ["只做观望跟主观（对照）", fmt_yuan(st_div.get("end")), fmt_yuan(st_div.get("pnl"), signed=True),
             fmt_pct(st_div.get("cagr")), fmt_pct(st_div.get("vol")), fmt_num(st_div.get("sharpe")),
             fmt_pct(st_div.get("maxdd")), fmt_yuan(st_div.get("total_cost"))],
            ["分歧跟量化（对照）", fmt_yuan(st_q.get("end")), fmt_yuan(st_q.get("pnl"), signed=True),
             fmt_pct(st_q.get("cagr")), fmt_pct(st_q.get("vol")), fmt_num(st_q.get("sharpe")),
             fmt_pct(st_q.get("maxdd")), fmt_yuan(st_q.get("total_cost"))],
        ],
        signed_cols={2, 3, 5, 6},
    )
    caption(doc, "表 2  主策略与三套对照。四套账户共用仓位和费用规则，只改允许的信号和观望跟哪一侧。")

    n_add = int((sig["action"] == "加码").sum())
    n_div = int((sig["action"] == "观望").sum())
    n_all = int(len(sig))
    add_pnl = float(holds.loc[holds["action"] == "加码", "pnl"].sum()) if not holds.empty else 0.0
    div_pnl = float(holds.loc[holds["action"] == "观望", "pnl"].sum()) if not holds.empty else 0.0
    para(
        doc,
        f"全样本产品级信号 {n_all:,} 条，加码 {n_add:,}（{n_add / n_all:.1%}），"
        f"观望 {n_div:,}（{n_div / n_all:.1%}）。"
        f"持仓层累计毛盈亏：加码 {fmt_yuan(add_pnl, signed=True)}，观望跟主观 {fmt_yuan(div_pnl, signed=True)}。"
        f"观望事件按主观方向对齐，次日均值 {fmt_pct(ev_div.get('mu1'))}，胜率 {fmt_pct(ev_div.get('hit1'), already=False)}，"
        f"t={fmt_t(ev_div.get('t1'))}；若改按量化方向对齐，次日 {fmt_pct(ev_div_q.get('mu1'))}，t={fmt_t(ev_div_q.get('t1'))}。"
        "总报告里这条 overlay 好看，账户层要看观望腿是否真的贡献了可交易的盈亏，还是只是换手和波动。",
    )

    heading(doc, "三、账户路径", 1)
    caps = [
        ("equity", "图 1  蓝线=分歧跟主观；红=只做加码；紫=只做观望跟主观；金=观望改跟量化。"),
        ("dd", "图 2  本策略回撤。"),
        ("cost", "图 3  毛盈亏、费用与净盈亏。观望腿会增加换手，费用通常高于只做加码。"),
        ("margin", "图 4  保证金与名义本金。观望日也开仓，占用会比只做加码更满。"),
        ("n", "图 5  每日持仓品种数。"),
        ("sig_n", "图 6  每日观望信号条数。"),
        ("monthly", "图 7  月度净盈亏。"),
        ("yearly", "图 8  分年净盈亏。"),
        ("roll_sharpe", "图 9  滚动 60 日年化夏普，金虚线=1。"),
        ("hist", "图 10  日收益直方图。"),
    ]
    for key, text in caps:
        if key in ch:
            add_picture(doc, ch[key])
            caption(doc, text)

    heading(doc, "四、加码 vs 观望跟主观", 1)
    para(
        doc,
        "把持仓毛盈亏按信号拆开，看观望腿是增强还是拖累。"
        "对照账户把全部仓位预算给其中一腿，数字会和「拆贡献」不同："
        "组合里观望和加码抢 8 个名额，单腿账户则把 2.2 倍名义都给自己。",
    )
    if "by_action" in ch:
        add_picture(doc, ch["by_action"])
        caption(doc, "图 11  加码、观望跟主观的累计毛盈亏路径。")
    if "action_bar" in ch:
        add_picture(doc, ch["action_bar"])
        caption(doc, "图 12  两种信号累计毛盈亏。")
    if "mix" in ch:
        add_picture(doc, ch["mix"])
        caption(doc, "图 13  每日持仓里加码与观望的个数。")

    if not holds.empty:
        rows = []
        for action, label in (("加码", "加码"), ("观望", "观望跟主观")):
            h = holds[holds["action"] == action]
            if h.empty:
                continue
            s = h.groupby("hold_date")["pnl"].sum()
            rows.append([
                label, f"{len(h):,}", f"{h['hold_date'].nunique()}", f"{h['product'].nunique()}",
                fmt_yuan(h["pnl"].sum(), signed=True), fmt_yuan(s.mean(), signed=True),
                fmt_pct(float((s > 0).mean()), already=False), fmt_t(tstat(s)),
            ])
        add_table(doc, ["信号", "持仓条数", "有仓天数", "品种数", "累计毛盈亏", "日均毛盈亏", "日胜率", "t"], rows, signed_cols={4, 5, 6})
        caption(doc, "表 3  持仓层按信号处理。日胜率按「当天该信号持仓合计是否赚钱」计。")

    heading(doc, "五、事件研究（按本策略方向对齐）", 1)
    para(
        doc,
        "加码按共识方向对齐；观望按主观方向对齐。这与总报告里观望没有签名方向不同——"
        "这里检验的是「如果按本账户规则去交易，品种本身有没有边」。"
        "观望再按量化方向对齐一次，用来对照该不该跟主观。",
    )
    add_table(doc, ["持有期", "样本", "对齐均值", "中位数", "胜率", "t"], ev_rows(ev_add), signed_cols={2, 3, 4})
    caption(doc, "表 4  加码 · 共识方向对齐。")
    add_table(doc, ["持有期", "样本", "对齐均值", "中位数", "胜率", "t"], ev_rows(ev_div), signed_cols={2, 3, 4})
    caption(doc, "表 5  观望 · 主观方向对齐（本策略）。")
    add_table(doc, ["持有期", "样本", "对齐均值", "中位数", "胜率", "t"], ev_rows(ev_div_q), signed_cols={2, 3, 4})
    caption(doc, "表 6  同一批观望样本，改按量化方向对齐。若表 5 明显好于表 6，跟主观才有依据。")
    if "event" in ch:
        add_picture(doc, ch["event"])
        caption(doc, "图 14  加码 vs 观望跟主观的持有期对齐收益。")
    if "hit" in ch:
        add_picture(doc, ch["hit"])
        caption(doc, "图 15  对齐胜率，虚线 50%。")

    div_ev = ev[ev["action"] == "观望"].copy()
    if not div_ev.empty:
        heading(doc, "5.1 观望腿：什么时候跟主观更干净", 2)
        div_ev["sabs"] = div_ev["s_pct"].abs()
        div_ev["qabs"] = div_ev["q_pct"].abs()
        div_ev["s_lead"] = div_ev["sabs"] - div_ev["qabs"]
        div_ev["side"] = np.where(div_ev["s_pct"] > 0, "主观偏多", "主观偏空")
        b = bucket_stats(div_ev, "观望", "sabs", [0, 4, 8, 14, 80], ["<4", "4–8", "8–14", "≥14"])
        if not b.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in b.itertuples(index=False)]
            add_table(doc, ["|主观风险%|", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 7  观望按主观仓位强度分层。主观自己都不重时，跟它更像噪声。")
        if "s_strength" in ch:
            add_picture(doc, ch["s_strength"])
            caption(doc, "图 16  观望跟主观按 |s%| 分层的次日对齐。")
        lead = bucket_stats(div_ev, "观望", "s_lead", [-80, -4, 0, 4, 80], ["主观更轻≥4", "主观略轻", "主观略重", "主观更重≥4"])
        if not lead.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in lead.itertuples(index=False)]
            add_table(doc, ["主观相对量化", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 8  |s%|−|q%|。主观比量化更重时，跟主观是否更稳。")
        side = group_event_table(div_ev, "观望", "side")
        if not side.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1), fmt_pct(r.mu5)]
                    for r in side.itertuples(index=False)]
            add_table(doc, ["主观方向", "样本", "次日对齐", "胜率", "t", "5日对齐"], rows, signed_cols={2, 3, 5})
            caption(doc, "表 9  观望日主观偏多 vs 偏空。")

        sec_ev = div_ev.dropna(subset=["strat_1"]).groupby("sector")["strat_1"]
        rows = []
        for sec_name, s in sec_ev:
            if len(s) < 8:
                continue
            rows.append([sec_name, f"{len(s)}", fmt_pct(float(s.mean())),
                         fmt_pct(float((s > 0).mean()), already=False), fmt_t(tstat(s))])
        if rows:
            add_table(doc, ["板块", "观望样本(≥8)", "次日对齐(跟主观)", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 10  观望跟主观的板块事件。")

    heading(doc, "六、多空、板块与品种", 1)
    for key, text in (
        ("ls", "图 17  账户多头 vs 空头累计毛盈亏。"),
        ("sector", "图 18  板块累计毛盈亏。"),
        ("product", "图 19  品种两端。"),
    ):
        if key in ch:
            add_picture(doc, ch[key])
            caption(doc, text)

    if not holds.empty:
        by_dir = holds.groupby("dir").agg(days=("hold_date", "nunique"), pnl=("pnl", "sum"), n=("product", "size"))
        rows = [[d, f"{int(r['n']):,}", f"{int(r['days'])}", fmt_yuan(r["pnl"], signed=True)] for d, r in by_dir.iterrows()]
        add_table(doc, ["方向", "持仓条数", "有仓天数", "累计毛盈亏"], rows, signed_cols={3})
        caption(doc, "表 11  多空拆解。")

        sec = holds.groupby("sector").agg(
            days=("hold_date", "nunique"), names=("product", "nunique"),
            pnl=("pnl", "sum"), hit=("pnl", lambda s: float((s > 0).mean())),
        ).sort_values("pnl", ascending=False)
        rows = [[i, f"{int(r.names)}", f"{int(r.days)}", fmt_yuan(r.pnl, signed=True), fmt_pct(r.hit, already=False)]
                for i, r in sec.iterrows()]
        add_table(doc, ["板块", "品种数", "有仓天数", "累计毛盈亏", "条数胜率"], rows, signed_cols={3, 4})
        caption(doc, "表 12  板块归因。")

        prod = holds.groupby(["name", "product", "sector"]).agg(
            days=("hold_date", "nunique"),
            long_days=("dir", lambda s: int((s == "多").sum())),
            short_days=("dir", lambda s: int((s == "空").sum())),
            lots=("lots", lambda s: s.abs().mean()),
            notional=("notional", "mean"),
            pnl=("pnl", "sum"),
        ).sort_values("pnl", ascending=False)
        top_bottom = pd.concat([prod.head(12), prod.tail(8)]).drop_duplicates()
        rows = []
        for (nme, p, sec_), r in top_bottom.iterrows():
            rows.append([
                f"{nme}({p})", sec_, f"{int(r['days'])}", f"{int(r['long_days'])}/{int(r['short_days'])}",
                f"{r['lots']:.1f}", fmt_yuan(r["notional"]), fmt_yuan(r["pnl"], signed=True),
            ])
        add_table(doc, ["品种", "板块", "持仓天", "多/空条数", "平均|手数|", "平均名义", "累计毛盈亏"], rows, signed_cols={6})
        caption(doc, "表 13  品种层头部 12 + 尾部 8。")
        pos = prod[prod["pnl"] > 0]["pnl"]
        if pos.sum() > 0:
            para(doc, f"正贡献品种里，前 5 名约占全部正贡献的 {fmt_pct(float(pos.head(5).sum() / pos.sum()), already=False)}。")

    heading(doc, "七、分年、分月与回撤段", 1)
    if not yearly.empty:
        rows = []
        for r in yearly.itertuples(index=False):
            rows.append([
                str(r.year), fmt_yuan(r.pnl, signed=True), fmt_yuan(r.cost),
                fmt_pct(r.cagr), fmt_pct(r.vol), fmt_num(r.sharpe), fmt_pct(r.maxdd),
                fmt_pct(r.hit, already=False), str(int(r.days)),
            ])
        add_table(doc, ["年", "净盈亏", "费用", "年化", "波动", "夏普", "最大回撤", "日胜率", "天数"], rows, signed_cols={1, 3, 5, 6})
        caption(doc, "表 14  分年账户绩效。")
    if not monthly.empty:
        rows = []
        for ym, r in monthly.iterrows():
            rows.append([
                ym, fmt_yuan(r["pnl"], signed=True), fmt_yuan(r["gross"], signed=True),
                fmt_yuan(r["cost"]), fmt_num(r["n"], 1), fmt_pct(r["hit"], already=False), fmt_yuan(r["eq"]),
            ])
        add_table(doc, ["月份", "净盈亏", "毛盈亏", "费用", "日均品种", "日胜率", "月末权益"], rows, signed_cols={1, 2})
        caption(doc, "表 15  分月账户结果。")
        para(doc, f"有数据的月份共 {len(monthly)} 个，月度赚钱比例 {fmt_pct(float((monthly['pnl'] > 0).mean()), already=False)}。")

    if not episodes.empty:
        heading(doc, "7.1 深度回撤段", 2)
        rows = []
        for r in episodes.head(8).itertuples(index=False):
            rows.append([
                r.start, r.trough, r.end, str(int(r.days)), str(int(r.trough_days)),
                fmt_pct(r.depth), fmt_yuan(r.lost, signed=True),
                "未修复" if getattr(r, "open", False) else "已修复",
            ])
        add_table(doc, ["开始", "谷底", "结束", "持续天", "到谷底", "深度", "权益损失", "状态"], rows, signed_cols={5, 6})
        caption(doc, "表 16  深度至少 3% 的回撤段。")
        worst = episodes.iloc[0]
        para(
            doc,
            f"最深一段从 {worst['start']} 到谷底 {worst['trough']}，深度 {fmt_pct(worst['depth'])}，"
            f"权益最多少了 {fmt_yuan(worst['lost'])}，持续 {int(worst['days'])} 个交易日。"
            "观望日两边本来就在打架，跟主观等于主动选边，选错时回撤会快于只做加码。",
        )

    heading(doc, "八、成交、换手与费用", 1)
    if "turnover" in ch:
        add_picture(doc, ch["turnover"])
        caption(doc, "图 20  10 日平均成交名义。观望名单和加码名单切换时，换手会上去。")
    if not trades.empty:
        para(
            doc,
            f"全样本成交 {len(trades):,} 笔，成交名义合计 {fmt_yuan(trades['notional'].sum())}，"
            f"费用合计 {fmt_yuan(trades['cost'].sum())}，平均每笔 {fmt_yuan(trades['cost'].mean())}。"
            f"开/加 {(trades['side'] == '开/加').sum():,} 笔，平仓 {(trades['side'] == '平仓').sum():,} 笔，"
            f"反手/平 {(trades['side'] == '反手/平').sum():,} 笔。",
        )
        t20 = trades.sort_values("notional", ascending=False).head(15)
        trows = [[r.trade_date, f"{r.name}({r.product})", r.action, r.side,
                  f"{int(r.old_lots)}→{int(r.new_lots)}", fmt_yuan(r.notional), fmt_yuan(r.cost)]
                 for r in t20.itertuples(index=False)]
        add_table(doc, ["成交日", "品种", "信号", "动作", "手数", "成交名义", "费用"], trows)
        caption(doc, "表 17  名义最大的 15 笔调仓。")
        add_table(
            doc,
            ["费用项", "金额", "说明"],
            [
                ["手续费", fmt_yuan(st.get("total_comm")), f"成交名义 × {COMM_RATE * 1e4:.1f}bp，每手不低于 3 元"],
                ["滑点", fmt_yuan(st.get("total_slip")), f"成交名义 × {SLIP_RATE * 1e4:.1f}bp"],
                ["费用合计", fmt_yuan(st.get("total_cost")), "上面两项之和"],
                ["毛盈亏", fmt_yuan(st.get("total_gross"), signed=True), "持仓名义 × 次日涨跌"],
                ["净盈亏", fmt_yuan(st.get("pnl"), signed=True), "毛盈亏 − 费用"],
                ["费用/|毛盈亏|", fmt_pct(abs(st.get("total_cost") or 0) / abs(st.get("total_gross") or 1), already=False), "摩擦侵蚀"],
            ],
            signed_cols={1},
        )
        caption(doc, "表 18  费用拆解。主力连续没有换月移仓。")
        if st.get("total_cost") and not daily.empty:
            rows = []
            for k, label in ((0.5, "半价摩擦"), (1.0, "基准（本报告）"), (2.0, "双倍摩擦")):
                extra = (k - 1.0) * daily["cost"]
                adj = daily["pnl_net"] - extra
                eq = START_EQUITY + adj.cumsum()
                tmp = daily.copy()
                tmp["pnl_net"] = adj
                tmp["equity"] = eq
                tmp["ret"] = adj / (eq - adj).replace(0, np.nan)
                tmp["dd"] = eq / eq.cummax() - 1.0
                st_k = account_stats(tmp)
                rows.append([label, fmt_yuan(st_k.get("end")), fmt_yuan(st_k.get("pnl"), signed=True),
                             fmt_pct(st_k.get("cagr")), fmt_num(st_k.get("sharpe")), fmt_pct(st_k.get("maxdd"))])
            add_table(doc, ["摩擦假设", "期末权益", "净盈亏", "年化", "夏普", "最大回撤"], rows, signed_cols={2, 3, 4, 5})
            caption(doc, "表 19  费用按比例缩放后的近似敏感性（仓位不重算）。观望腿对费率更敏感。")

    heading(doc, "九、稳健性", 1)
    if halves:
        rows = [[h["name"], f"{h['start']} ~ {h['end']}", fmt_yuan(h.get("pnl"), signed=True),
                 fmt_pct(h.get("cagr")), fmt_pct(h.get("vol")), fmt_num(h.get("sharpe")),
                 fmt_pct(h.get("maxdd")), fmt_pct(h.get("hit"), already=False)] for h in halves]
        add_table(doc, ["分段", "区间", "净盈亏", "年化", "波动", "夏普", "最大回撤", "日胜率"], rows, signed_cols={2, 3, 5, 6})
        caption(doc, "表 20  前后半样本。观望跟主观若只在半段有效，就不能当成稳定规则。")
    if streaks:
        para(doc, f"赚钱日 {streaks['win_days']}、亏钱日 {streaks['loss_days']}；"
             f"最长连赢 {streaks['max_win_streak']} 日，最长连亏 {streaks['max_loss_streak']} 日。")
    if not daily.empty:
        heading(doc, "9.1 最赚与最亏的交易日", 2)
        rows = []
        for label, df in (("最亏", daily.nsmallest(8, "pnl_net")), ("最赚", daily.nlargest(8, "pnl_net"))):
            for r in df.itertuples(index=False):
                rows.append([label, r.return_date, fmt_yuan(r.pnl_net, signed=True), fmt_yuan(r.pnl_gross, signed=True),
                             fmt_yuan(r.cost), f"{int(r.n)}", fmt_pct(r.ret)])
        add_table(doc, ["类型", "日期", "净盈亏", "毛盈亏", "费用", "持仓数", "日收益"], rows, signed_cols={2, 3, 6})
        caption(doc, "表 21  单日两端。")

    heading(doc, "十、持仓快照", 1)
    para(doc, "持仓在信号日收盘后生成，下一交易日生效。加码与共识同向，观望与主观同向。")

    def hold_table(df: pd.DataFrame, title: str):
        if df is None or df.empty:
            para(doc, f"{title}：当日无持仓。", first_line=False)
            return
        rows = []
        for r in df.sort_values("notional", ascending=False).itertuples(index=False):
            rows.append([
                f"{r.name}({r.product})", r.sector, r.action, r.dir, str(int(r.lots)),
                f"{r.price:,.2f}" if pd.notna(r.price) else "—",
                fmt_yuan(r.notional), fmt_yuan(r.margin), fmt_yuan(r.pnl, signed=True),
            ])
        rows.append(["合计", "", "", "", str(int(df["lots"].abs().sum())), "",
                     fmt_yuan(df["notional"].sum()), fmt_yuan(df["margin"].sum()), fmt_yuan(df["pnl"].sum(), signed=True)])
        add_table(doc, ["品种", "板块", "信号", "方向", "手数", "价格", "名义", "保证金", "当日盈亏"], rows, signed_cols={8})
        caption(doc, title)

    if not holds.empty:
        last_d = holds["hold_date"].max()
        hold_table(holds[holds["hold_date"] == last_d], f"表 22  样本末日 {last_d} 持仓。")
        snap = holds[holds["hold_date"] == "2026-09-01"]
        if snap.empty:
            snap = holds[holds["signal_date"] == "2026-09-01"]
        if not snap.empty:
            hold_table(snap, "表 23  2026-09-01 持仓（与风控页同一天）。")

    heading(doc, "十一、结论与局限", 1)
    para(
        doc,
        "分歧跟主观 = 加码复制共识 + 观望日主动站主观。它比只做加码更满仓、换手更高，"
        "样本内 overlay 好看，不等于 2,000 万账户里观望腿一定赚钱。"
        "要以表 2 的单腿对照和表 5/表 6 的事件 t 为准：若观望跟主观不显著、或跟量化差不多，"
        "就不要因为总报告里夏普最高就放大这条规则。",
    )
    para(
        doc,
        "主要局限。第一，主力连续价，没有换月移仓。"
        "第二，保证金率为近似值；无涨跌停、强平。"
        "第三，观望是两边冲突，选边本身就是在赌哪一类投顾这段更对，样本期的主观账户集合一变，边可能消失。"
        "第四，仓位规则是研究设定。第五，观望事件 t 在总报告里就不强，账户层更要防过拟合。",
    )
    para(
        doc,
        "使用建议。先把加码当成主腿。观望跟主观只宜当研究线索："
        "仅在主观仓位明显更重、且后半样本和双倍摩擦后仍然不塌时，才考虑小仓位试。"
        "不要把六种信号捆成一篮子，也不要把 overlay 夏普直接当成账户承诺。",
    )
    para(
        doc,
        f"图表原文件在 mom_signal_strategy/report_output/div_subj/charts/。报告生成于 {date.today().isoformat()}。",
        size=10, color=MUTED, first_line=False,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(REPORT_PATH))
    try:
        doc.save(str(REPORT_PATH_ASCII))
    except Exception:
        pass
    return REPORT_PATH


def main():
    load_env()
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_conn()
    try:
        pos, px = load_data(conn)
    finally:
        conn.close()
    if pos.empty:
        raise SystemExit("mom_position_details is empty")
    if px.empty:
        raise SystemExit("raw_akshare_futures_daily is empty")

    wide, clean, close, _ = build_returns(px)
    sig = build_signals(pos, clean)
    if sig.empty:
        raise SystemExit("no signals reconstructed")

    print("Event study…")
    ev = align_div_subj(event_study(sig, wide))
    ev_add = event_horizon_stats(ev, "加码")
    ev_div = event_horizon_stats(ev, "观望")
    ev_div_q = event_horizon_stats(ev, "观望", col_prefix="quant")
    ev_div_q["action"] = "观望跟量化"

    print("20M account: 分歧跟主观…")
    acct = run_account(sig, close, wide, clean, include_bufengge=False,
                       allowed_actions={"加码", "观望"}, divergence_follow="subj")
    print("20M account: 只做加码…")
    acct_add = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码"})
    print("20M account: 只做观望跟主观…")
    acct_div = run_account(sig, close, wide, clean, include_bufengge=False,
                           allowed_actions={"观望"}, divergence_follow="subj")
    print("20M account: 分歧跟量化…")
    acct_q = run_account(sig, close, wide, clean, include_bufengge=False,
                         allowed_actions={"加码", "观望"}, divergence_follow="quant")

    stats = account_stats(acct["daily"])
    stats_add = account_stats(acct_add["daily"])
    stats_div = account_stats(acct_div["daily"])
    stats_q = account_stats(acct_q["daily"])

    print("Drawing charts…")
    charts = draw_charts(acct, acct_add, acct_div, acct_q, ev, ev_add, ev_div)

    if not acct["daily"].empty:
        acct["daily"].to_csv(OUT_DIR / "acct_daily.csv", index=False, encoding="utf-8-sig")
    if not acct["holds"].empty:
        acct["holds"].to_csv(OUT_DIR / "acct_holdings.csv", index=False, encoding="utf-8-sig")
    if not acct["trades"].empty:
        acct["trades"].to_csv(OUT_DIR / "acct_trades.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([ev_add, ev_div, ev_div_q]).to_csv(OUT_DIR / "event_stats.csv", index=False, encoding="utf-8-sig")

    monthly, yearly = year_month_tables(acct["daily"]) if not acct["daily"].empty else (pd.DataFrame(), pd.DataFrame())
    ctx = {
        "sig": sig, "ev": ev, "ev_add": ev_add, "ev_div": ev_div, "ev_div_q": ev_div_q,
        "acct": acct, "stats": stats, "stats_add": stats_add, "stats_div": stats_div, "stats_q": stats_q,
        "charts": charts, "start": str(sig["date"].min()), "end": str(sig["date"].max()),
        "episodes": drawdown_episodes(acct["daily"]), "monthly": monthly, "yearly": yearly,
        "halves": half_sample(acct["daily"]),
        "streaks": streak_stats(acct["daily"]["pnl_net"]) if not acct["daily"].empty else {},
    }
    print("Writing Word report…")
    path = build_report(ctx)
    print(f"Wrote {path}")
    print(f"Also {REPORT_PATH_ASCII}")
    if stats:
        print(f"End equity {stats.get('end'):,.0f}  Sharpe {stats.get('sharpe')}  MaxDD {stats.get('maxdd')}")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        raise
