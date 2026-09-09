# -*- coding: utf-8 -*-
"""Deep-dive Word report: 只做加码同向（其余空仓）, 20M account."""
from __future__ import annotations

import math
import sys
import traceback
from datetime import date
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "report_output" / "jiaama_only"
CHART_DIR = OUT_DIR / "charts"
REPORT_PATH = BASE_DIR / "MOM决策信号_只做加码同向_回测报告.docx"
REPORT_PATH_ASCII = BASE_DIR / "MOM_signal_jiaama_only_backtest.docx"

sys.path.insert(0, str(BASE_DIR))
from _mom_20m_account import (  # noqa: E402
    COMM_RATE,
    MAX_MARGIN_UTIL,
    SLIP_RATE,
    START_EQUITY,
    account_stats,
    fmt_yuan,
    run_account,
)
from _mom_roll import load_roll_context  # noqa: E402
from generate_jiaama_crowd_report import (  # noqa: E402
    add_strategy_alignment,
    bucket_stats,
    drawdown_episodes,
    half_sample,
    slice_stats,
    streak_stats,
    tstat,
    year_month_tables,
)
from generate_mom_signal_backtest_report import (  # noqa: E402
    GOLD,
    MUTED,
    NAVY,
    TEXT,
    add_picture,
    add_table,
    add_text,
    apply_font,
    build_returns,
    build_signals,
    caption,
    configure_matplotlib,
    event_study,
    fmt_num,
    fmt_pct,
    fmt_t,
    fp,
    get_conn,
    heading,
    load_data,
    load_env,
    para,
    set_run_font,
)

C_NAVY = "#1A365D"
C_GOLD = "#C9A227"
C_RED = "#C53030"
C_GREEN = "#2F855A"
C_BLUE = "#2B6CB0"
C_ORANGE = "#DD6B20"
C_VIOLET = "#6D28D9"
C_TEAL = "#0F766E"


def save_fig(fig, name: str) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def mask_keep(sig: pd.DataFrame, keep: pd.Series) -> pd.DataFrame:
    out = sig.copy()
    out.loc[~keep.reindex(out.index, fill_value=False), "action"] = "观望"
    return out


def event_horizon_stats(ev: pd.DataFrame, action="加码") -> dict:
    g = ev[ev["action"] == action]
    row = {"action": action, "n": int(len(g))}
    for h, col in ((1, "strat_1"), (5, "strat_5"), (10, "strat_10"), (20, "strat_20")):
        s = pd.to_numeric(g[col], errors="coerce").dropna()
        row[f"n{h}"] = int(len(s))
        row[f"mu{h}"] = float(s.mean()) if len(s) else None
        row[f"hit{h}"] = float((s > 0).mean()) if len(s) else None
        row[f"t{h}"] = tstat(s)
        row[f"med{h}"] = float(s.median()) if len(s) else None
    return row


def group_event_table(ev: pd.DataFrame, col: str) -> pd.DataFrame:
    g = ev[ev["action"] == "加码"].copy()
    if g.empty or col not in g.columns:
        return pd.DataFrame()
    rows = []
    for lab, sub in g.groupby(col, dropna=False):
        s = pd.to_numeric(sub["strat_1"], errors="coerce").dropna()
        rows.append({
            "bucket": "（空）" if pd.isna(lab) else str(lab),
            "n": int(len(s)),
            "mu1": float(s.mean()) if len(s) else None,
            "hit1": float((s > 0).mean()) if len(s) else None,
            "t1": tstat(s),
            "mu5": float(pd.to_numeric(sub["strat_5"], errors="coerce").dropna().mean()) if len(sub) else None,
        })
    return pd.DataFrame(rows).sort_values("n", ascending=False)


def draw_charts(acct, acct_combo, acct_long, acct_short, ev, ev_row) -> dict:
    daily = acct["daily"]
    holds = acct["holds"]
    trades = acct["trades"]
    out = {}
    if daily.empty:
        return out
    x = pd.to_datetime(daily["return_date"])

    fig, ax = plt.subplots(figsize=(11.2, 5.2), dpi=160)
    ax.plot(x, daily["equity"] / 1e4, color=C_NAVY, lw=1.95, label="只做加码同向")
    if not acct_combo["daily"].empty:
        ax.plot(pd.to_datetime(acct_combo["daily"]["return_date"]), acct_combo["daily"]["equity"] / 1e4,
                color=C_GOLD, lw=1.2, alpha=0.85, label="加码 + 控拥挤（对照）")
    if not acct_long["daily"].empty:
        ax.plot(pd.to_datetime(acct_long["daily"]["return_date"]), acct_long["daily"]["equity"] / 1e4,
                color=C_RED, lw=1.2, alpha=0.85, label="只做加码多头")
    if not acct_short["daily"].empty:
        ax.plot(pd.to_datetime(acct_short["daily"]["return_date"]), acct_short["daily"]["equity"] / 1e4,
                color=C_TEAL, lw=1.2, alpha=0.9, label="只做加码空头")
    ax.axhline(START_EQUITY / 1e4, color="#A0AEC0", ls="--", lw=1, label="起始 2,000 万")
    ax.set_title("2,000 万账户权益（扣手续费与滑点）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("权益（万元）", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["equity"] = save_fig(fig, "equity.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
    ax.fill_between(x, daily["dd"] * 100, 0, color=C_RED, alpha=0.35)
    ax.set_title("只做加码：账户回撤", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("回撤（%）", **fp())
    apply_font(ax)
    out["dd"] = save_fig(fig, "dd.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
    ax.plot(x, daily["cum_gross"] / 1e4, color=C_BLUE, lw=1.5, label="累计毛盈亏")
    ax.plot(x, daily["cum_cost"] / 1e4, color=C_ORANGE, lw=1.5, label="累计费用")
    ax.plot(x, daily["cum_net"] / 1e4, color=C_NAVY, lw=1.8, label="累计净盈亏")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("毛盈亏、费用与净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["cost"] = save_fig(fig, "cost.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.2), dpi=160)
    ax.plot(x, daily["margin"] / 1e4, color=C_NAVY, lw=1.4, label="保证金")
    ax.plot(x, daily["gross_notional"] / 1e4, color=C_ORANGE, lw=1.2, label="名义本金")
    ax.set_title("保证金占用与名义本金", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["margin"] = save_fig(fig, "margin.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.6), dpi=160)
    ax.plot(x, daily["n"], color=C_NAVY, lw=1.3)
    ax.set_title("每日持仓品种数（上限 8，只交易加码）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("个数", **fp())
    apply_font(ax)
    out["n"] = save_fig(fig, "n.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    s = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("ME").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in s.values]
    ax.bar(s.index, s.values / 1e4, width=20, color=colors, align="center")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("月度净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    apply_font(ax)
    out["monthly"] = save_fig(fig, "monthly.png")

    fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=160)
    y = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("YE").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in y.values]
    ax.bar([str(i.year) for i in y.index], y.values / 1e4, color=colors, width=0.55)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("分年净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    apply_font(ax)
    out["yearly"] = save_fig(fig, "yearly.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
    roll = daily["ret"].rolling(60, min_periods=30)
    sharpe = roll.mean() / roll.std(ddof=1) * math.sqrt(252)
    ax.plot(x, sharpe, color=C_NAVY, lw=1.4)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.axhline(1.0, color=C_GOLD, ls="--", lw=1)
    ax.set_title("滚动 60 日年化夏普", fontsize=13, color=C_NAVY, **fp())
    apply_font(ax)
    out["roll_sharpe"] = save_fig(fig, "roll_sharpe.png")

    fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
    ax.hist(daily["ret"] * 100, bins=36, color=C_NAVY, alpha=0.82, edgecolor="white")
    ax.axvline(0, color=C_ORANGE, lw=1.2)
    ax.set_title("日收益率分布", fontsize=13, color=C_NAVY, **fp())
    ax.set_xlabel("日收益（%）", **fp())
    apply_font(ax)
    out["hist"] = save_fig(fig, "hist.png")

    add_ev = ev[ev["action"] == "加码"]
    if not add_ev.empty:
        fig, ax = plt.subplots(figsize=(9.4, 4.6), dpi=160)
        hs = [1, 5, 10, 20]
        vals = [ev_row.get(f"mu{h}") or 0 for h in hs]
        colors = [C_RED if v >= 0 else C_GREEN for v in vals]
        ax.bar(["次日", "5日", "10日", "20日"], [v * 100 for v in vals], color=colors, width=0.55)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_ylabel("共识方向对齐平均收益（%）", **fp())
        ax.set_title("加码事件研究：持有期", fontsize=13, color=C_NAVY, **fp())
        apply_font(ax)
        out["event"] = save_fig(fig, "event.png")

        fig, ax = plt.subplots(figsize=(8.6, 4.4), dpi=160)
        hits = [ev_row.get(f"hit{h}") for h in (1, 5, 10, 20)]
        ax.bar(["次日", "5日", "10日", "20日"], [h * 100 if h is not None else 0 for h in hits],
               color=C_NAVY, width=0.55)
        ax.axhline(50, color="#A0AEC0", ls="--", lw=1)
        ax.set_ylabel("胜率（%）", **fp())
        ax.set_title("加码对齐胜率", fontsize=13, color=C_NAVY, **fp())
        apply_font(ax)
        out["hit"] = save_fig(fig, "hit.png")

        cnt = add_ev.groupby("date").size()
        fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
        ax.bar(pd.to_datetime(cnt.index), cnt.values, color=C_NAVY, width=1.2, alpha=0.85)
        ax.set_title("每日加码信号条数（产品级）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("条数", **fp())
        apply_font(ax)
        out["sig_n"] = save_fig(fig, "sig_n.png")

        add_ev = add_ev.copy()
        add_ev["strength"] = add_ev["q_pct"].abs() + add_ev["s_pct"].abs()
        b = bucket_stats(add_ev, "加码", "strength", [0, 10, 16, 24, 80], ["<10", "10–16", "16–24", "≥24"])
        if not b.empty:
            fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
            vals = [(v or 0) * 100 for v in b["mu1"]]
            colors = [C_RED if v >= 0 else C_GREEN for v in vals]
            ax.bar(b["bucket"].astype(str), vals, color=colors, width=0.55)
            ax.axhline(0, color="#4A5568", lw=0.8)
            ax.set_title("加码强度分层：次日对齐收益", fontsize=13, color=C_NAVY, **fp())
            ax.set_ylabel("%", **fp())
            apply_font(ax)
            out["strength"] = save_fig(fig, "strength.png")

        kind = add_ev.copy()
        kind["side"] = np.where(kind["kind"] == "consensus_short", "空头共识", "多头共识")
        side = group_event_table(kind, "side")
        if not side.empty:
            fig, ax = plt.subplots(figsize=(7.6, 4.4), dpi=160)
            vals = [(v or 0) * 100 for v in side["mu1"]]
            colors = [C_RED if v >= 0 else C_GREEN for v in vals]
            ax.bar(side["bucket"], vals, color=colors, width=0.45)
            ax.axhline(0, color="#4A5568", lw=0.8)
            ax.set_title("加码多头共识 vs 空头共识：次日对齐", fontsize=13, color=C_NAVY, **fp())
            ax.set_ylabel("%", **fp())
            apply_font(ax)
            out["side_event"] = save_fig(fig, "side_event.png")

    if not holds.empty:
        fig, ax = plt.subplots(figsize=(8.4, 4.2), dpi=160)
        by_dir = holds.groupby("dir")["pnl"].sum()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_dir.values]
        ax.bar(by_dir.index, by_dir.values / 1e4, color=colors, width=0.45)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("账户持仓：多头 vs 空头累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["ls"] = save_fig(fig, "ls.png")

        fig, ax = plt.subplots(figsize=(10.4, 5.0), dpi=160)
        by_sec = holds.groupby("sector")["pnl"].sum().sort_values()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_sec.values]
        ax.barh(by_sec.index, by_sec.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("板块累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["sector"] = save_fig(fig, "sector.png")

        fig, ax = plt.subplots(figsize=(10.6, 6.0), dpi=160)
        by_p = holds.groupby("name")["pnl"].sum().sort_values()
        top = pd.concat([by_p.head(10), by_p.tail(10)]).drop_duplicates()
        colors = [C_RED if v >= 0 else C_GREEN for v in top.values]
        ax.barh(top.index, top.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("品种累计毛盈亏：最亏 / 最赚各 10 个", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["product"] = save_fig(fig, "product.png")

        mix = holds.groupby(["hold_date", "dir"]).size().unstack(fill_value=0)
        fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
        cols = [c for c in ("多", "空") if c in mix.columns]
        colors = {"多": C_RED, "空": C_TEAL}
        ax.stackplot(pd.to_datetime(mix.index), *[mix[c] for c in cols],
                     labels=cols, colors=[colors[c] for c in cols], alpha=0.85)
        ax.set_title("每日持仓中的多/空个数", fontsize=13, color=C_NAVY, **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["mix"] = save_fig(fig, "mix.png")

    if not trades.empty:
        fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
        t = trades.copy()
        t["dt"] = pd.to_datetime(t["trade_date"])
        turn = t.groupby("dt")["notional"].sum()
        ax.plot(turn.index, turn.rolling(10, min_periods=3).mean() / 1e4, color=C_NAVY, lw=1.5)
        ax.set_title("10 日平均成交名义（换手）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["turnover"] = save_fig(fig, "turnover.png")

    if "roll_cost" in daily.columns and daily["roll_cost"].sum() > 0:
        fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
        ax.plot(x, daily["roll_cost"].cumsum() / 1e4, color=C_ORANGE, lw=1.6)
        ax.set_title("累计换月/展期费用（旧约平 + 新约开）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["roll"] = save_fig(fig, "roll_cost.png")

    return out


def set_core_header(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "MOM 每日风控  ·  只做加码同向 深度回测", size=8, color=MUTED)
    fp_ = section.footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp_, "内部研究  ·  加码跟共识  ·  其余空仓  ·  模拟账户 2,000 万  ·  ", size=8, color=MUTED)
    run = fp_.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color=MUTED)


def build_report(ctx: dict) -> Path:
    configure_matplotlib()
    daily = ctx["acct"]["daily"]
    holds = ctx["acct"]["holds"]
    trades = ctx["acct"]["trades"]
    st = ctx["stats"]
    st_combo = ctx["stats_combo"]
    st_long = ctx["stats_long"]
    st_short = ctx["stats_short"]
    ch = ctx["charts"]
    ev_row = ctx["ev_row"]
    sig = ctx["sig"]
    start, end = ctx["start"], ctx["end"]
    episodes = ctx["episodes"]
    monthly, yearly = ctx["monthly"], ctx["yearly"]
    halves = ctx["halves"]
    streaks = ctx["streaks"]
    ev = ctx["ev"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    set_core_header(doc)

    para(doc, "MOM 每日风控", size=12, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "只做加码同向", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "其余信号空仓  ·  2,000 万模拟账户深度回测", size=13, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=10)
    para(
        doc,
        f"样本 {start} 至 {end}　　交易日 {st.get('n')}　　"
        f"手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp（开平都收）　　"
        f"生成于 {date.today().isoformat()}",
        size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=16,
    )

    heading(doc, "一、策略定义", 1)
    para(
        doc,
        "本报告只检验一条规则：每个交易日收盘后读取 MOM「量化 vs 主观」决策信号，"
        "下一交易日只交易「加码」——与量化/主观存量共识同向开仓。"
        "控拥挤、补风格、观望、暂缓加码、减码准备全部空仓；信号消失或不再是加码时，下一交易日平掉。"
        "这是六种动作里方向最干净的一条，也是上一份「加码+控拥挤」报告里的主腿。",
    )
    add_table(
        doc,
        ["信号", "本账户怎么做", "经济含义"],
        [
            ["加码", "与量化/主观共识同向（共识多则买、共识空则卖）", "两边已经站在同一边，账户复制这份 beta"],
            ["控拥挤", "空仓并平掉旧仓", "拥挤反向不在本策略里"],
            ["观望 / 暂缓加码 / 减码准备", "空仓并平掉旧仓", "没有共识加码，就不下注"],
            ["补风格", "不做", "覆盖提示，不译成期货仓位"],
        ],
        col_widths=[4.2, 6.4, 5.8],
    )
    caption(doc, "表 1  只做加码同向的下单规则。")

    para(
        doc,
        f"账户初始权益 {fmt_yuan(START_EQUITY)}。仓位按当日权益约 2.2 倍名义、最多 8 个品种等权分配，"
        f"单品种一日 1σ 亏损不超过权益的 1.2%，保证金占用不超过权益的 {MAX_MARGIN_UTIL:.0%}（券商系数 1.1）。"
        f"手数取整。费用按成交名义单边手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp，"
        "每手手续费不低于 3 元，开平都收。"
        "持仓盯指定主力合约的收盘价，不用主力连续的拼接涨跌；"
        "持有期内主力换月时，按旧约平仓、新约开仓做展期，双边各收一次手续费和滑点。"
        "换月跳空不再计入当日盈亏（那不是实盘的钱）。",
    )

    heading(doc, "二、核心结论", 1)
    if st:
        para(
            doc,
            f"样本内期末权益 {fmt_yuan(st.get('end'))}，净盈亏 {fmt_yuan(st.get('pnl'), signed=True)}，"
            f"相对起始资金 {fmt_pct((st.get('end') - START_EQUITY) / START_EQUITY)}。"
            f"年化 {fmt_pct(st.get('cagr'))}，波动 {fmt_pct(st.get('vol'))}，夏普 {fmt_num(st.get('sharpe'))}，"
            f"卡玛 {fmt_num(st.get('calmar'))}，最大回撤 {fmt_pct(st.get('maxdd'))}"
            f"（约 {fmt_yuan(st.get('maxdd_yuan'))}），日胜率 {fmt_pct(st.get('hit'), already=False)}。"
            f"手续费 {fmt_yuan(st.get('total_comm'))}、滑点 {fmt_yuan(st.get('total_slip'))}，"
            f"费用合计 {fmt_yuan(st.get('total_cost'))}，毛盈亏 {fmt_yuan(st.get('total_gross'), signed=True)}。"
            f"日均持仓 {fmt_num(st.get('avg_names'), 1)} 个品种，日均保证金 {fmt_yuan(st.get('avg_margin'))}，"
            f"占用率 {fmt_pct(st.get('avg_util'))}，日均名义 {fmt_yuan(st.get('avg_notional'))}。",
        )

    add_table(
        doc,
        ["口径", "期末权益", "净盈亏", "年化", "波动", "夏普", "最大回撤", "费用"],
        [
            [
                "只做加码同向（本报告）",
                fmt_yuan(st.get("end")),
                fmt_yuan(st.get("pnl"), signed=True),
                fmt_pct(st.get("cagr")),
                fmt_pct(st.get("vol")),
                fmt_num(st.get("sharpe")),
                fmt_pct(st.get("maxdd")),
                fmt_yuan(st.get("total_cost")),
            ],
            [
                "加码 + 控拥挤（对照）",
                fmt_yuan(st_combo.get("end")),
                fmt_yuan(st_combo.get("pnl"), signed=True),
                fmt_pct(st_combo.get("cagr")),
                fmt_pct(st_combo.get("vol")),
                fmt_num(st_combo.get("sharpe")),
                fmt_pct(st_combo.get("maxdd")),
                fmt_yuan(st_combo.get("total_cost")),
            ],
            [
                "只做加码多头共识",
                fmt_yuan(st_long.get("end")),
                fmt_yuan(st_long.get("pnl"), signed=True),
                fmt_pct(st_long.get("cagr")),
                fmt_pct(st_long.get("vol")),
                fmt_num(st_long.get("sharpe")),
                fmt_pct(st_long.get("maxdd")),
                fmt_yuan(st_long.get("total_cost")),
            ],
            [
                "只做加码空头共识",
                fmt_yuan(st_short.get("end")),
                fmt_yuan(st_short.get("pnl"), signed=True),
                fmt_pct(st_short.get("cagr")),
                fmt_pct(st_short.get("vol")),
                fmt_num(st_short.get("sharpe")),
                fmt_pct(st_short.get("maxdd")),
                fmt_yuan(st_short.get("total_cost")),
            ],
        ],
        signed_cols={2, 3, 5, 6},
    )
    caption(doc, "表 2  本策略与三套对照。均已计入指定月换月/展期。多/空对照把仓位预算全部给一侧。")

    n_add = int((sig["action"] == "加码").sum())
    n_all = int(len(sig))
    n_long = int(((sig["action"] == "加码") & (sig["kind"] == "consensus_long")).sum())
    n_short = int(((sig["action"] == "加码") & (sig["kind"] == "consensus_short")).sum())
    flat_days = int((daily["n"] == 0).sum()) if not daily.empty else 0
    para(
        doc,
        f"全样本产品级信号 {n_all:,} 条，加码 {n_add:,} 条（{n_add / n_all:.1%}），"
        f"其中多头共识 {n_long:,}、空头共识 {n_short:,}。"
        f"账户 {st.get('n')} 个交易日里，有 {flat_days} 日完全空仓——加码名单不是每天都有。"
        f"事件上次日对齐均值 {fmt_pct(ev_row.get('mu1'))}，胜率 {fmt_pct(ev_row.get('hit1'), already=False)}，"
        f"t={fmt_t(ev_row.get('t1'))}。"
        "对照「加码+控拥挤」用来回答：拿掉拥挤反向之后，成绩会变好还是变差；"
        "多/空对照用来回答：加码是不是只在单边行情里好看。",
    )

    heading(doc, "三、账户路径", 1)
    if "equity" in ch:
        add_picture(doc, ch["equity"])
        caption(doc, "图 1  蓝线=只做加码；金线=加上控拥挤的对照；红/青绿=只做多头或只做空头共识。")
    if "dd" in ch:
        add_picture(doc, ch["dd"])
        caption(doc, "图 2  回撤。深度和持续时间比期末盈亏更能说明能不能拿得住。")
    if "cost" in ch:
        add_picture(doc, ch["cost"])
        caption(doc, "图 3  毛盈亏、累计费用与净盈亏。")
    if "margin" in ch:
        add_picture(doc, ch["margin"])
        caption(doc, "图 4  保证金与名义本金。没有加码日名义会掉到接近零。")
    if "n" in ch:
        add_picture(doc, ch["n"])
        caption(doc, "图 5  每日持仓品种数。经常低于 8，说明同时满足加码的品种并不密。")
    if "sig_n" in ch:
        add_picture(doc, ch["sig_n"])
        caption(doc, "图 6  每日产品级加码条数。账户最多只取强度最高的 8 个。")
    if "monthly" in ch:
        add_picture(doc, ch["monthly"])
        caption(doc, "图 7  月度净盈亏（万元，红赚绿亏）。")
    if "yearly" in ch:
        add_picture(doc, ch["yearly"])
        caption(doc, "图 8  分年净盈亏。样本短的年份只作参考。")
    if "roll_sharpe" in ch:
        add_picture(doc, ch["roll_sharpe"])
        caption(doc, "图 9  滚动 60 日年化夏普。金虚线=1。")
    if "hist" in ch:
        add_picture(doc, ch["hist"])
        caption(doc, "图 10  日收益直方图。左尾厚度决定回撤体验。")

    heading(doc, "四、事件研究：加码本身有没有边", 1)
    para(
        doc,
        "事件研究不问账户仓位，只问：出现加码后，品种收益按共识方向对齐，均值是否显著不为零。"
        "这是账户赚钱的微观基础。若次日 t 不强、拉长持有期后又衰减，账户成绩就更依赖仓位规则和少数品种路径。",
    )
    add_table(
        doc,
        ["持有期", "样本", "对齐均值", "中位数", "胜率", "t"],
        [
            ["次日", f"{ev_row.get('n1')}", fmt_pct(ev_row.get("mu1")), fmt_pct(ev_row.get("med1")),
             fmt_pct(ev_row.get("hit1"), already=False), fmt_t(ev_row.get("t1"))],
            ["5 日", f"{ev_row.get('n5')}", fmt_pct(ev_row.get("mu5")), fmt_pct(ev_row.get("med5")),
             fmt_pct(ev_row.get("hit5"), already=False), fmt_t(ev_row.get("t5"))],
            ["10 日", f"{ev_row.get('n10')}", fmt_pct(ev_row.get("mu10")), fmt_pct(ev_row.get("med10")),
             fmt_pct(ev_row.get("hit10"), already=False), fmt_t(ev_row.get("t10"))],
            ["20 日", f"{ev_row.get('n20')}", fmt_pct(ev_row.get("mu20")), fmt_pct(ev_row.get("med20")),
             fmt_pct(ev_row.get("hit20"), already=False), fmt_t(ev_row.get("t20"))],
        ],
        signed_cols={2, 3, 4},
    )
    caption(doc, "表 3  加码按共识方向对齐。|t|<2 时不宜把均值当成稳定边。")
    if "event" in ch:
        add_picture(doc, ch["event"])
        caption(doc, "图 11  加码在次日 / 5 / 10 / 20 日的对齐收益。")
    if "hit" in ch:
        add_picture(doc, ch["hit"])
        caption(doc, "图 12  对齐胜率，虚线 50%。")

    add_ev = ev[ev["action"] == "加码"].copy()
    if not add_ev.empty:
        add_ev["strength"] = add_ev["q_pct"].abs() + add_ev["s_pct"].abs()
        b = bucket_stats(add_ev, "加码", "strength", [0, 10, 16, 24, 80], ["<10", "10–16", "16–24", "≥24"])
        if not b.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in b.itertuples(index=False)]
            add_table(doc, ["|q%|+|s%|", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 4  按两边风险占比之和分层。更「重」的共识是否更干净。")
        if "strength" in ch:
            add_picture(doc, ch["strength"])
            caption(doc, "图 13  加码强度分层的次日对齐收益。")

        add_ev["side"] = np.where(add_ev["kind"] == "consensus_short", "空头共识", "多头共识")
        side = group_event_table(add_ev, "side")
        if not side.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1), fmt_pct(r.mu5)]
                    for r in side.itertuples(index=False)]
            add_table(doc, ["方向", "样本", "次日对齐", "胜率", "t", "5日对齐"], rows, signed_cols={2, 3, 5})
            caption(doc, "表 5  多头共识 vs 空头共识。若一侧显著、一侧没有边，账户就不是市场中性。")
        if "side_event" in ch:
            add_picture(doc, ch["side_event"])
            caption(doc, "图 14  多/空共识的次日对齐。")

        flow = group_event_table(add_ev, "kind1d")
        if not flow.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in flow.itertuples(index=False)]
            add_table(doc, ["1 日调仓形态", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 6  加码按 1 日资金流形态。both_add=两边同日加；diverge=存量同向但当日调仓反向（5 日仍同加才维持加码）。")

        flow5 = group_event_table(add_ev, "kind5d")
        if not flow5.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in flow5.itertuples(index=False)]
            add_table(doc, ["5 日调仓形态", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 7  加码按 5 日资金流形态。")

        add_ev["gap"] = (add_ev["q_pct"].abs() - add_ev["s_pct"].abs()).abs()
        gap = bucket_stats(add_ev, "加码", "gap", [-0.1, 2, 5, 12, 80], ["≤2", "2–5", "5–12", "≥12"])
        if not gap.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)]
                    for r in gap.itertuples(index=False)]
            add_table(doc, ["| |q%|−|s%| |", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 8  量化与主观风险占比是否接近。差距很大时，共识更像「一侧主导、另一侧勉强同向」。")

    heading(doc, "五、多空、板块与品种", 1)
    if "ls" in ch:
        add_picture(doc, ch["ls"])
        caption(doc, "图 15  账户持仓层多头 vs 空头累计毛盈亏。")
    if "mix" in ch:
        add_picture(doc, ch["mix"])
        caption(doc, "图 16  每日持仓中的多/空个数。")
    if "sector" in ch:
        add_picture(doc, ch["sector"])
        caption(doc, "图 17  板块累计毛盈亏。")
    if "product" in ch:
        add_picture(doc, ch["product"])
        caption(doc, "图 18  品种两端。用来看集中度，不是推荐交易名单。")

    if not holds.empty:
        by_dir = holds.groupby("dir").agg(days=("hold_date", "nunique"), pnl=("pnl", "sum"), n=("product", "size"))
        rows = []
        for d, r in by_dir.iterrows():
            rows.append([d, f"{int(r['n']):,}", f"{int(r['days'])}", fmt_yuan(r["pnl"], signed=True)])
        add_table(doc, ["方向", "持仓条数", "有仓天数", "累计毛盈亏"], rows, signed_cols={3})
        caption(doc, "表 9  账户多空拆解。")

        sec = holds.groupby("sector").agg(
            days=("hold_date", "nunique"),
            names=("product", "nunique"),
            pnl=("pnl", "sum"),
            hit=("pnl", lambda s: float((s > 0).mean())),
        ).sort_values("pnl", ascending=False)
        rows = [[i, f"{int(r.names)}", f"{int(r.days)}", fmt_yuan(r.pnl, signed=True), fmt_pct(r.hit, already=False)]
                for i, r in sec.iterrows()]
        add_table(doc, ["板块", "品种数", "有仓天数", "累计毛盈亏", "条数胜率"], rows, signed_cols={3, 4})
        caption(doc, "表 10  板块归因。条数胜率按持仓行计。")

        if not add_ev.empty:
            sec_ev = add_ev.dropna(subset=["strat_1"]).groupby("sector")["strat_1"]
            rows = []
            for sec_name, s in sec_ev:
                if len(s) < 8:
                    continue
                rows.append([sec_name, f"{len(s)}", fmt_pct(float(s.mean())), fmt_pct(float((s > 0).mean()), already=False), fmt_t(tstat(s))])
            rows.sort(key=lambda r: -len(add_ev[add_ev["sector"] == r[0]]))
            if rows:
                add_table(doc, ["板块", "加码样本(≥8)", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
                caption(doc, "表 11  加码事件的板块分解。样本很少的板块不列入。")

        prod = holds.groupby(["name", "product", "sector"]).agg(
            days=("hold_date", "nunique"),
            long_days=("dir", lambda s: int((s == "多").sum())),
            short_days=("dir", lambda s: int((s == "空").sum())),
            lots=("lots", lambda s: s.abs().mean()),
            notional=("notional", "mean"),
            pnl=("pnl", "sum"),
        ).sort_values("pnl", ascending=False)
        top_bottom = pd.concat([prod.head(12), prod.tail(8)]).drop_duplicates()
        rows = []
        for (nme, p, sec_), r in top_bottom.iterrows():
            rows.append([
                f"{nme}({p})", sec_, f"{int(r['days'])}", f"{int(r['long_days'])}/{int(r['short_days'])}",
                f"{r['lots']:.1f}", fmt_yuan(r["notional"]), fmt_yuan(r["pnl"], signed=True),
            ])
        add_table(doc, ["品种", "板块", "持仓天", "多/空条数", "平均|手数|", "平均名义", "累计毛盈亏"], rows, signed_cols={6})
        caption(doc, "表 12  品种层头部 12 + 尾部 8。")

        if len(prod) >= 5:
            pos = prod[prod["pnl"] > 0]["pnl"]
            conc = float(pos.head(5).sum() / pos.sum()) if pos.sum() > 0 else None
            para(
                doc,
                f"正贡献品种里，前 5 名约占全部正贡献的 {fmt_pct(conc, already=False)}。"
                "集中度高时，样本内成绩更容易被少数品种的行情路径带走。",
            )

    heading(doc, "六、分年、分月与回撤段", 1)
    if not yearly.empty:
        rows = []
        for r in yearly.itertuples(index=False):
            rows.append([
                str(r.year), fmt_yuan(r.pnl, signed=True), fmt_yuan(r.cost),
                fmt_pct(r.cagr), fmt_pct(r.vol), fmt_num(r.sharpe), fmt_pct(r.maxdd),
                fmt_pct(r.hit, already=False), str(int(r.days)),
            ])
        add_table(doc, ["年", "净盈亏", "费用", "年化", "波动", "夏普", "最大回撤", "日胜率", "天数"], rows, signed_cols={1, 3, 5, 6})
        caption(doc, "表 13  分年账户绩效。年内年化按该年交易日年化。")

    if not monthly.empty:
        rows = []
        for ym, r in monthly.iterrows():
            rows.append([
                ym, fmt_yuan(r["pnl"], signed=True), fmt_yuan(r["gross"], signed=True),
                fmt_yuan(r["cost"]), fmt_num(r["n"], 1), fmt_pct(r["hit"], already=False),
                fmt_yuan(r["eq"]),
            ])
        add_table(doc, ["月份", "净盈亏", "毛盈亏", "费用", "日均品种", "日胜率", "月末权益"], rows, signed_cols={1, 2})
        caption(doc, "表 14  分月账户结果。")
        win_m = float((monthly["pnl"] > 0).mean())
        para(doc, f"有数据的月份共 {len(monthly)} 个，月度赚钱比例 {fmt_pct(win_m, already=False)}。")

    if not episodes.empty:
        heading(doc, "6.1 深度回撤段", 2)
        rows = []
        for r in episodes.head(8).itertuples(index=False):
            rows.append([
                r.start, r.trough, r.end, str(int(r.days)), str(int(r.trough_days)),
                fmt_pct(r.depth), fmt_yuan(r.lost, signed=True),
                "未修复" if getattr(r, "open", False) else "已修复",
            ])
        add_table(doc, ["开始", "谷底", "结束", "持续天", "到谷底", "深度", "权益损失", "状态"], rows, signed_cols={5, 6})
        caption(doc, "表 15  深度至少 3% 的回撤段，按深度排序。")
        worst = episodes.iloc[0]
        para(
            doc,
            f"最深一段从 {worst['start']} 到谷底 {worst['trough']}，深度 {fmt_pct(worst['depth'])}，"
            f"权益最多少了 {fmt_yuan(worst['lost'])}，持续 {int(worst['days'])} 个交易日。"
            "加码跟的是存量共识，共识拥挤后若趋势反转，回撤会来得很快。",
        )

    heading(doc, "七、成交、换手与费用", 1)
    if "turnover" in ch:
        add_picture(doc, ch["turnover"])
        caption(doc, "图 19  10 日平均成交名义。加码名单日度变化大时，费用会明显吃收益。")
    if not trades.empty:
        para(
            doc,
            f"全样本成交 {len(trades):,} 笔，成交名义合计 {fmt_yuan(trades['notional'].sum())}，"
            f"费用合计 {fmt_yuan(trades['cost'].sum())}，平均每笔 {fmt_yuan(trades['cost'].mean())}。"
            f"开/加 {(trades['side'] == '开/加').sum():,} 笔，平仓 {(trades['side'] == '平仓').sum():,} 笔，"
            f"反手/平 {(trades['side'] == '反手/平').sum():,} 笔。",
        )
        t20 = trades.sort_values("notional", ascending=False).head(15)
        trows = []
        for r in t20.itertuples(index=False):
            trows.append([
                r.trade_date, f"{r.name}({r.product})", r.action, r.side,
                f"{int(r.old_lots)}→{int(r.new_lots)}", fmt_yuan(r.notional), fmt_yuan(r.cost),
            ])
        add_table(doc, ["成交日", "品种", "信号", "动作", "手数", "成交名义", "费用"], trows)
        caption(doc, "表 16  名义最大的 15 笔调仓。")

        add_table(
            doc,
            ["费用项", "金额", "说明"],
            [
                ["手续费", fmt_yuan(st.get("total_comm")), f"成交名义 × {COMM_RATE * 1e4:.1f}bp，每手不低于 3 元，开平都收"],
                ["滑点", fmt_yuan(st.get("total_slip")), f"成交名义 × {SLIP_RATE * 1e4:.1f}bp"],
                ["费用合计", fmt_yuan(st.get("total_cost")), "上面两项之和"],
                ["毛盈亏", fmt_yuan(st.get("total_gross"), signed=True), "持仓名义 × 次日涨跌"],
                ["净盈亏", fmt_yuan(st.get("pnl"), signed=True), "毛盈亏 − 费用"],
                ["费用/|毛盈亏|", fmt_pct(abs(st.get("total_cost") or 0) / abs(st.get("total_gross") or 1), already=False), "摩擦侵蚀比例"],
            ],
            signed_cols={1},
        )
        caption(doc, "表 17  费用拆解。含换月展期的开平双边手续费与滑点。")

        if st.get("total_cost") and not daily.empty:
            rows = []
            for k, label in ((0.5, "半价摩擦"), (1.0, "基准（本报告）"), (2.0, "双倍摩擦")):
                extra = (k - 1.0) * daily["cost"]
                adj = daily["pnl_net"] - extra
                eq = START_EQUITY + adj.cumsum()
                tmp = daily.copy()
                tmp["pnl_net"] = adj
                tmp["equity"] = eq
                tmp["ret"] = adj / (eq - adj).replace(0, np.nan)
                tmp["dd"] = eq / eq.cummax() - 1.0
                st_k = account_stats(tmp)
                rows.append([
                    label, fmt_yuan(st_k.get("end")), fmt_yuan(st_k.get("pnl"), signed=True),
                    fmt_pct(st_k.get("cagr")), fmt_num(st_k.get("sharpe")), fmt_pct(st_k.get("maxdd")),
                ])
            add_table(doc, ["摩擦假设", "期末权益", "净盈亏", "年化", "夏普", "最大回撤"], rows, signed_cols={2, 3, 4, 5})
            caption(doc, "表 18  把已发生费用按比例缩放后的近似敏感性（仓位不重算）。")

    heading(doc, "7.1 换月、展期与移仓价差", 2)
    para(
        doc,
        "实盘拿的是指定交割月，不是交易所拼好的主力连续。"
        "本账户每个品种盯当日持仓主力合约的收盘价做盈亏；"
        "当持仓主力从旧约切到新约（OI 最大合约变化，或 rollover 日历），"
        "视为展期：旧约全部平掉、新约按当时手数开仓，各收一次手续费和滑点。"
        "连续价在换月日的跳空不再记进盈亏——那一跳不是结算单上的钱。"
        "移仓价差通过「旧约自己的涨跌 + 换成新约之后走新约」进入路径，而不是把新约减旧约当成当天亏损。",
    )
    st_nr = ctx.get("stats_noroll") or {}
    daily_r = daily
    n_roll = int(daily_r["n_rolls"].sum()) if "n_rolls" in daily_r.columns else 0
    roll_cost = float(daily_r["roll_cost"].sum()) if "roll_cost" in daily_r.columns else 0.0
    rolls = trades[trades["side"] == "移仓"] if (not trades.empty and "side" in trades.columns) else pd.DataFrame()
    para(
        doc,
        f"样本内持仓期发生移仓 {n_roll} 次，展期费用（手续费+滑点）{fmt_yuan(roll_cost)}。"
        + (
            f"若不做换月、仍用主力连续涨跌（旧口径），期末权益 {fmt_yuan(st_nr.get('end'))}，"
            f"净盈亏 {fmt_yuan(st_nr.get('pnl'), signed=True)}，夏普 {fmt_num(st_nr.get('sharpe'))}，"
            f"回撤 {fmt_pct(st_nr.get('maxdd'))}。差额主要来自：去掉虚假跳空、加上真实展期开平、以及新约之后的价格路径。"
            if st_nr else ""
        ),
    )
    if st_nr:
        add_table(
            doc,
            ["口径", "期末权益", "净盈亏", "年化", "夏普", "最大回撤", "费用"],
            [
                ["指定月 + 换月展期（本报告）", fmt_yuan(st.get("end")), fmt_yuan(st.get("pnl"), signed=True),
                 fmt_pct(st.get("cagr")), fmt_num(st.get("sharpe")), fmt_pct(st.get("maxdd")),
                 fmt_yuan(st.get("total_cost"))],
                ["主力连续、不换月（旧）", fmt_yuan(st_nr.get("end")), fmt_yuan(st_nr.get("pnl"), signed=True),
                 fmt_pct(st_nr.get("cagr")), fmt_num(st_nr.get("sharpe")), fmt_pct(st_nr.get("maxdd")),
                 fmt_yuan(st_nr.get("total_cost"))],
            ],
            signed_cols={2, 3, 4, 5},
        )
        caption(doc, "表 18b  同一套加码规则：计入换月 vs 主力连续。实盘应对齐本报告口径。")
    if "roll" in ch:
        add_picture(doc, ch["roll"])
        caption(doc, "图 19b  累计展期费用（换月双边开平）。")
    if not rolls.empty:
        r20 = rolls.sort_values("notional", ascending=False).head(12)
        rrows = []
        for r in r20.itertuples(index=False):
            rrows.append([
                r.trade_date, f"{r.name}({r.product})",
                getattr(r, "from_contract", "") or "",
                getattr(r, "to_contract", "") or "",
                f"{int(r.old_lots)}→{int(r.new_lots)}",
                fmt_yuan(r.notional), fmt_yuan(r.cost),
            ])
        add_table(doc, ["日期", "品种", "旧约", "新约", "手数", "开平名义", "展期费用"], rrows)
        caption(doc, "表 18c  名义最大的 12 笔移仓。")

    heading(doc, "八、稳健性", 1)
    if halves:
        rows = []
        for h in halves:
            rows.append([
                h["name"], f"{h['start']} ~ {h['end']}", fmt_yuan(h.get("pnl"), signed=True),
                fmt_pct(h.get("cagr")), fmt_pct(h.get("vol")), fmt_num(h.get("sharpe")),
                fmt_pct(h.get("maxdd")), fmt_pct(h.get("hit"), already=False),
            ])
        add_table(doc, ["分段", "区间", "净盈亏", "年化", "波动", "夏普", "最大回撤", "日胜率"], rows, signed_cols={2, 3, 5, 6})
        caption(doc, "表 19  前后半样本。若半段反号，全样本结论不能外推。")

    if streaks:
        para(
            doc,
            f"赚钱日 {streaks['win_days']}、亏钱日 {streaks['loss_days']}；"
            f"最长连赢 {streaks['max_win_streak']} 日，最长连亏 {streaks['max_loss_streak']} 日。",
        )

    if not daily.empty:
        worst = daily.nsmallest(8, "pnl_net")
        best = daily.nlargest(8, "pnl_net")
        heading(doc, "8.1 最赚与最亏的交易日", 2)
        rows = []
        for label, df in (("最亏", worst), ("最赚", best)):
            for r in df.itertuples(index=False):
                rows.append([
                    label, r.return_date, fmt_yuan(r.pnl_net, signed=True), fmt_yuan(r.pnl_gross, signed=True),
                    fmt_yuan(r.cost), f"{int(r.n)}", fmt_pct(r.ret),
                ])
        add_table(doc, ["类型", "日期", "净盈亏", "毛盈亏", "费用", "持仓数", "日收益"], rows, signed_cols={2, 3, 6})
        caption(doc, "表 20  单日两端。最亏日若集中在同一板块，尾部风险并不分散。")

    heading(doc, "九、持仓快照", 1)
    para(doc, "持仓在信号日收盘后生成，下一交易日生效。方向与加码共识同向。")

    def hold_table(df: pd.DataFrame, title: str):
        if df is None or df.empty:
            para(doc, f"{title}：当日无持仓。", first_line=False)
            return
        rows = []
        for r in df.sort_values("notional", ascending=False).itertuples(index=False):
            rows.append([
                f"{r.name}({r.product})", r.sector, r.action, r.dir, str(int(r.lots)),
                f"{r.price:,.2f}" if pd.notna(r.price) else "—",
                fmt_yuan(r.notional), fmt_yuan(r.margin), fmt_yuan(r.pnl, signed=True),
            ])
        rows.append([
            "合计", "", "", "", str(int(df["lots"].abs().sum())), "",
            fmt_yuan(df["notional"].sum()), fmt_yuan(df["margin"].sum()), fmt_yuan(df["pnl"].sum(), signed=True),
        ])
        add_table(doc, ["品种", "板块", "信号", "方向", "手数", "价格", "名义", "保证金", "当日盈亏"], rows, signed_cols={8})
        caption(doc, title)

    if not holds.empty:
        last_d = holds["hold_date"].max()
        hold_table(holds[holds["hold_date"] == last_d], f"表 21  样本末日 {last_d} 持仓。")
        snap = holds[holds["hold_date"] == "2026-09-01"]
        if snap.empty:
            snap = holds[holds["signal_date"] == "2026-09-01"]
        if not snap.empty:
            hold_table(snap, "表 22  2026-09-01 持仓（与风控页同一天）。")

        tenure = holds.groupby(["product", "name"]).agg(days=("hold_date", "nunique"), pnl=("pnl", "sum"))
        para(
            doc,
            f"全程出现过 {holds['product'].nunique()} 个品种；"
            f"单品种平均持仓 {tenure['days'].mean():.1f} 个交易日，中位数 {tenure['days'].median():.0f} 日，"
            f"最长 {int(tenure['days'].max())} 日。"
            "加码名单换得快时，账户更像滚动共识 overlay，而不是中长期配置。",
        )

    heading(doc, "十、结论与局限", 1)
    para(
        doc,
        "只做加码同向，是把页面上的「共识加码」直接译成期货账户："
        "量化与主观存量已经站在同一边，才复制这份方向；其余动作一律空仓。"
        "账户数字回答的是「若用 2,000 万只跟加码」，不是单位净值 overlay，也不是对未来收益的承诺。",
    )
    para(
        doc,
        "和「加码+控拥挤」比，本策略更干净、交易日更稀疏，也少了一条样本很少的反向腿。"
        "若对照账户明显更差，说明控拥挤在本样本里是拖累；若更好，说明拥挤反向有时能对冲加码的拥挤回撤。"
        "多/空对照若一侧贡献了绝大部分盈亏，加码就不是市场中性，外推时要按当时的多空结构打折。",
    )
    para(
        doc,
        "主要局限。第一，换月已按持仓主力合约盯市，并在主力切换日双边开平；"
        "移仓价差体现在「旧约真实涨跌 + 新约之后的路径」，连续合约跳空不计入盈亏。"
        "仍没有盘口排队、涨跌停、强平和真实换月滑点加宽。"
        "第二，保证金率是品种近似值。"
        "第三，信号来自 MOM 存量仓位，账户集合变化后边可能消失。"
        "第四，2.2 倍名义、最多 8 个、1.2% 风险帽是研究设定。"
        "第五，加码次日事件 t 往往只是弱显著。",
    )
    para(
        doc,
        "使用建议。把加码当成「两边已经确认才加 beta」的配置提示，而不是高频预测模型。"
        "暂缓加码和减码准备继续排除：存量同向不等于现在还该加。"
        "若后半样本或双倍摩擦后夏普明显塌掉，就把它当成信号设计的检验，而不是可以放大的期货策略。",
    )
    para(
        doc,
        f"图表原文件在 mom_signal_strategy/report_output/jiaama_only/charts/。报告生成于 {date.today().isoformat()}。",
        size=10, color=MUTED, first_line=False,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(REPORT_PATH))
    try:
        doc.save(str(REPORT_PATH_ASCII))
    except Exception:
        pass
    return REPORT_PATH


def main():
    load_env()
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_conn()
    try:
        pos, px = load_data(conn)
        if pos.empty:
            raise SystemExit("mom_position_details is empty")
        if px.empty:
            raise SystemExit("raw_akshare_futures_daily is empty")
        wide, clean, close, _ = build_returns(px)
        sig = build_signals(pos, clean)
        if sig.empty:
            raise SystemExit("no signals reconstructed")
        roll = load_roll_context(conn, "2025-01-01", str(sig["date"].max()))
    finally:
        conn.close()

    print("Event study…")
    ev = add_strategy_alignment(event_study(sig, wide))
    extra = sig[["date", "product", "kind1d", "kind5d"]].drop_duplicates()
    ev = ev.merge(extra, on=["date", "product"], how="left")
    ev_row = event_horizon_stats(ev, "加码")

    is_add = sig["action"] == "加码"
    is_long = is_add & (sig["kind"] == "consensus_long")
    is_short = is_add & (sig["kind"] == "consensus_short")

    print("20M account: 只做加码（指定月+换月）…")
    acct = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码"}, roll=roll)
    print("20M account: 只做加码（主力连续、不换月）…")
    acct_noroll = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码"})
    print("20M account: 加码+控拥挤 对照…")
    acct_combo = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码", "控拥挤"}, roll=roll)
    print("20M account: 只做加码多头…")
    acct_long = run_account(mask_keep(sig, is_long), close, wide, clean, include_bufengge=False, allowed_actions={"加码"}, roll=roll)
    print("20M account: 只做加码空头…")
    acct_short = run_account(mask_keep(sig, is_short), close, wide, clean, include_bufengge=False, allowed_actions={"加码"}, roll=roll)

    stats = account_stats(acct["daily"])
    stats_noroll = account_stats(acct_noroll["daily"])
    stats_combo = account_stats(acct_combo["daily"])
    stats_long = account_stats(acct_long["daily"])
    stats_short = account_stats(acct_short["daily"])

    print("Drawing charts…")
    charts = draw_charts(acct, acct_combo, acct_long, acct_short, ev, ev_row)

    if not acct["daily"].empty:
        acct["daily"].to_csv(OUT_DIR / "acct_daily.csv", index=False, encoding="utf-8-sig")
    if not acct["holds"].empty:
        acct["holds"].to_csv(OUT_DIR / "acct_holdings.csv", index=False, encoding="utf-8-sig")
    if not acct["trades"].empty:
        acct["trades"].to_csv(OUT_DIR / "acct_trades.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([ev_row]).to_csv(OUT_DIR / "event_stats.csv", index=False, encoding="utf-8-sig")

    monthly, yearly = year_month_tables(acct["daily"]) if not acct["daily"].empty else (pd.DataFrame(), pd.DataFrame())
    ctx = {
        "sig": sig,
        "ev": ev,
        "ev_row": ev_row,
        "acct": acct,
        "stats": stats,
        "stats_noroll": stats_noroll,
        "stats_combo": stats_combo,
        "stats_long": stats_long,
        "stats_short": stats_short,
        "charts": charts,
        "start": str(sig["date"].min()),
        "end": str(sig["date"].max()),
        "episodes": drawdown_episodes(acct["daily"]),
        "monthly": monthly,
        "yearly": yearly,
        "halves": half_sample(acct["daily"]),
        "streaks": streak_stats(acct["daily"]["pnl_net"]) if not acct["daily"].empty else {},
    }
    print("Writing Word report…")
    path = build_report(ctx)
    print(f"Wrote {path}")
    print(f"Also {REPORT_PATH_ASCII}")
    if stats:
        print(f"End equity {stats.get('end'):,.0f}  Sharpe {stats.get('sharpe')}  MaxDD {stats.get('maxdd')}")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        raise
