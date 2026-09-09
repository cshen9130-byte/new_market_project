# -*- coding: utf-8 -*-
"""Deep-dive Word report: 只做加码同向 + 控拥挤反向（不做补风格）, 20M account."""
from __future__ import annotations

import math
import sys
import traceback
from datetime import date
from pathlib import Path

import matplotlib
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "report_output" / "jiaama_crowd"
CHART_DIR = OUT_DIR / "charts"
REPORT_PATH = BASE_DIR / "MOM决策信号_加码同向控拥挤反向_回测报告.docx"
REPORT_PATH_ASCII = BASE_DIR / "MOM_signal_jiaama_fade_crowd_backtest.docx"

sys.path.insert(0, str(BASE_DIR))
from _mom_20m_account import (  # noqa: E402
    COMM_RATE,
    MAX_MARGIN_UTIL,
    SLIP_RATE,
    START_EQUITY,
    account_stats,
    fmt_yuan,
    run_account,
)
from generate_mom_signal_backtest_report import (  # noqa: E402
    GOLD,
    GREEN,
    MUTED,
    NAVY,
    RED,
    TEXT,
    WHITE,
    add_picture,
    add_table,
    add_text,
    apply_font,
    build_returns,
    build_signals,
    caption,
    configure_matplotlib,
    event_study,
    fmt_num,
    fmt_pct,
    fmt_t,
    fp,
    get_conn,
    heading,
    load_data,
    load_env,
    para,
    set_run_font,
)

C_NAVY = "#1A365D"
C_GOLD = "#C9A227"
C_RED = "#C53030"
C_GREEN = "#2F855A"
C_BLUE = "#2B6CB0"
C_ORANGE = "#DD6B20"
C_VIOLET = "#6D28D9"
C_TEAL = "#0F766E"
C_GRAY = "#718096"

CORE_ACTIONS = {"加码", "控拥挤"}


def save_fig(fig, name: str) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def tstat(s: pd.Series) -> float | None:
    s = pd.to_numeric(s, errors="coerce").dropna()
    if len(s) < 3:
        return None
    se = float(s.std(ddof=1) / math.sqrt(len(s)))
    if se <= 0:
        return None
    return float(s.mean() / se)


def drawdown_episodes(daily: pd.DataFrame, min_depth=-0.03) -> pd.DataFrame:
    if daily.empty:
        return pd.DataFrame()
    eq = daily["equity"].to_numpy(dtype=float)
    dates = daily["return_date"].astype(str).to_list()
    peak = -np.inf
    peak_i = 0
    in_dd = False
    start_i = 0
    trough_i = 0
    trough_v = 0.0
    rows = []
    for i, v in enumerate(eq):
        if v >= peak:
            if in_dd:
                depth = trough_v / eq[start_i] - 1.0 if eq[start_i] else 0.0
                if depth <= min_depth:
                    rows.append({
                        "start": dates[start_i],
                        "trough": dates[trough_i],
                        "end": dates[i],
                        "days": i - start_i,
                        "trough_days": trough_i - start_i,
                        "depth": depth,
                        "peak_equity": eq[start_i],
                        "trough_equity": trough_v,
                        "lost": trough_v - eq[start_i],
                    })
                in_dd = False
            peak = v
            peak_i = i
        elif v < peak:
            if not in_dd:
                in_dd = True
                start_i = peak_i
                trough_i = i
                trough_v = v
            elif v < trough_v:
                trough_i = i
                trough_v = v
    if in_dd:
        depth = trough_v / eq[start_i] - 1.0 if eq[start_i] else 0.0
        if depth <= min_depth:
            rows.append({
                "start": dates[start_i],
                "trough": dates[trough_i],
                "end": dates[-1],
                "days": len(eq) - 1 - start_i,
                "trough_days": trough_i - start_i,
                "depth": depth,
                "peak_equity": eq[start_i],
                "trough_equity": trough_v,
                "lost": trough_v - eq[start_i],
                "open": True,
            })
    return pd.DataFrame(rows).sort_values("depth") if rows else pd.DataFrame()


def streak_stats(pnl: pd.Series) -> dict:
    wins = losses = max_w = max_l = cur_w = cur_l = 0
    for v in pnl.fillna(0.0):
        if v > 0:
            cur_w += 1
            cur_l = 0
            max_w = max(max_w, cur_w)
            wins += 1
        elif v < 0:
            cur_l += 1
            cur_w = 0
            max_l = max(max_l, cur_l)
            losses += 1
        else:
            cur_w = cur_l = 0
    return {"win_days": wins, "loss_days": losses, "max_win_streak": max_w, "max_loss_streak": max_l}


def add_strategy_alignment(ev: pd.DataFrame) -> pd.DataFrame:
    out = ev.copy()

    def sdir(r):
        if r.action == "加码":
            return float(r.dir) if r.dir else (1.0 if (r.q_pct + r.s_pct) >= 0 else -1.0)
        if r.action == "控拥挤":
            d = float(r.dir) if r.dir else (1.0 if r.q_pct > 0 else -1.0)
            return -d
        return 0.0

    out["strat_dir"] = [sdir(r) for r in out.itertuples(index=False)]
    for src, dest in (("r1", "strat_1"), ("r5", "strat_5"), ("r10", "strat_10"), ("r20", "strat_20")):
        aligned = []
        for r in out.itertuples(index=False):
            raw = getattr(r, src)
            if raw is None or (isinstance(raw, float) and math.isnan(raw)) or r.strat_dir == 0:
                aligned.append(np.nan)
            else:
                aligned.append(float(r.strat_dir) * float(raw))
        out[dest] = aligned
    return out


def action_event_stats(ev: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for action in ("加码", "控拥挤"):
        g = ev[ev["action"] == action]
        row = {"action": action, "n": int(len(g))}
        for h, col in ((1, "strat_1"), (5, "strat_5"), (10, "strat_10"), (20, "strat_20")):
            s = pd.to_numeric(g[col], errors="coerce").dropna()
            row[f"n{h}"] = int(len(s))
            row[f"mu{h}"] = float(s.mean()) if len(s) else None
            row[f"hit{h}"] = float((s > 0).mean()) if len(s) else None
            row[f"t{h}"] = tstat(s)
            row[f"med{h}"] = float(s.median()) if len(s) else None
        rows.append(row)
    return pd.DataFrame(rows)


def bucket_stats(ev: pd.DataFrame, action: str, col: str, bins, labels) -> pd.DataFrame:
    g = ev[ev["action"] == action].copy()
    if g.empty:
        return pd.DataFrame()
    g["bucket"] = pd.cut(g[col], bins=bins, labels=labels, include_lowest=True)
    rows = []
    for lab, sub in g.groupby("bucket", observed=False):
        s = pd.to_numeric(sub["strat_1"], errors="coerce").dropna()
        rows.append({
            "bucket": str(lab),
            "n": int(len(s)),
            "mu1": float(s.mean()) if len(s) else None,
            "hit1": float((s > 0).mean()) if len(s) else None,
            "t1": tstat(s),
        })
    return pd.DataFrame(rows)


def slice_stats(daily: pd.DataFrame) -> dict:
    """Performance of a date slice using that slice's own starting equity and peak."""
    if daily.empty:
        return {}
    d = daily.reset_index(drop=True)
    start = float(d["equity"].iloc[0] - d["pnl_net"].iloc[0])
    end = float(d["equity"].iloc[-1])
    r = d["ret"]
    n = len(r)
    years = n / 252.0
    cagr = (end / start) ** (1 / years) - 1 if years > 0 and start > 0 and end > 0 else None
    vol = float(r.std(ddof=1) * math.sqrt(252)) if n > 1 else None
    sharpe = (float(r.mean()) / float(r.std(ddof=1)) * math.sqrt(252)) if n > 1 and r.std(ddof=1) > 0 else None
    path = start + d["pnl_net"].cumsum()
    maxdd = float((path / path.cummax() - 1.0).min()) if len(path) else None
    return {
        "start": start,
        "end": end,
        "pnl": end - start,
        "cagr": cagr,
        "vol": vol,
        "sharpe": sharpe,
        "maxdd": maxdd,
        "hit": float((d["pnl_net"] > 0).mean()),
        "n": n,
        "total_cost": float(d["cost"].sum()),
    }


def year_month_tables(daily: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    d = daily.copy()
    d["dt"] = pd.to_datetime(d["return_date"])
    d["year"] = d["dt"].dt.year
    d["ym"] = d["dt"].dt.strftime("%Y-%m")
    monthly = d.groupby("ym").agg(
        pnl=("pnl_net", "sum"),
        gross=("pnl_gross", "sum"),
        cost=("cost", "sum"),
        n=("n", "mean"),
        eq=("equity", "last"),
        hit=("pnl_net", lambda s: float((s > 0).mean())),
        days=("pnl_net", "size"),
    )
    yearly = []
    for y, g in d.groupby("year"):
        st = slice_stats(g)
        yearly.append({
            "year": int(y),
            "start": st.get("start"),
            "end": st.get("end"),
            "pnl": st.get("pnl"),
            "cost": st.get("total_cost"),
            "cagr": st.get("cagr"),
            "vol": st.get("vol"),
            "sharpe": st.get("sharpe"),
            "maxdd": st.get("maxdd"),
            "hit": st.get("hit"),
            "days": int(len(g)),
        })
    return monthly, pd.DataFrame(yearly)


def half_sample(daily: pd.DataFrame) -> list[dict]:
    if daily.empty or len(daily) < 40:
        return []
    mid = len(daily) // 2
    out = []
    for name, part in (("前半样本", daily.iloc[:mid].copy()), ("后半样本", daily.iloc[mid:].copy())):
        st = slice_stats(part)
        out.append({
            "name": name,
            "start": str(part["return_date"].iloc[0]),
            "end": str(part["return_date"].iloc[-1]),
            **st,
        })
    return out


def draw_charts(acct: dict, add_acct: dict, fade_acct: dict, ev: pd.DataFrame, ev_stats: pd.DataFrame) -> dict:
    daily = acct["daily"]
    holds = acct["holds"]
    trades = acct["trades"]
    out = {}
    if daily.empty:
        return out
    x = pd.to_datetime(daily["return_date"])

    fig, ax = plt.subplots(figsize=(11.2, 5.2), dpi=160)
    ax.plot(x, daily["equity"] / 1e4, color=C_NAVY, lw=1.9, label="加码同向 + 控拥挤反向")
    if not add_acct["daily"].empty:
        ax.plot(pd.to_datetime(add_acct["daily"]["return_date"]), add_acct["daily"]["equity"] / 1e4,
                color=C_RED, lw=1.3, alpha=0.85, label="只做加码")
    if not fade_acct["daily"].empty:
        ax.plot(pd.to_datetime(fade_acct["daily"]["return_date"]), fade_acct["daily"]["equity"] / 1e4,
                color=C_GOLD, lw=1.3, alpha=0.9, label="只做控拥挤反向")
    ax.axhline(START_EQUITY / 1e4, color="#A0AEC0", ls="--", lw=1, label="起始 2,000 万")
    ax.set_title("2,000 万账户权益（扣手续费与滑点）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("权益（万元）", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["equity"] = save_fig(fig, "equity.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.8), dpi=160)
    ax.fill_between(x, daily["dd"] * 100, 0, color=C_RED, alpha=0.35)
    ax.set_title("组合回撤", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("回撤（%）", **fp())
    apply_font(ax)
    out["dd"] = save_fig(fig, "dd.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.4), dpi=160)
    ax.plot(x, daily["cum_gross"] / 1e4, color=C_BLUE, lw=1.5, label="累计毛盈亏")
    ax.plot(x, daily["cum_cost"] / 1e4, color=C_ORANGE, lw=1.5, label="累计费用")
    ax.plot(x, daily["cum_net"] / 1e4, color=C_NAVY, lw=1.8, label="累计净盈亏")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("毛盈亏、费用与净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["cost"] = save_fig(fig, "cost.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.2), dpi=160)
    ax.plot(x, daily["margin"] / 1e4, color=C_NAVY, lw=1.4, label="保证金")
    ax.plot(x, daily["gross_notional"] / 1e4, color=C_ORANGE, lw=1.2, label="名义本金")
    ax.set_title("保证金占用与名义本金", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.legend(frameon=False, fontsize=8)
    apply_font(ax)
    out["margin"] = save_fig(fig, "margin.png")

    fig, ax = plt.subplots(figsize=(11.2, 3.6), dpi=160)
    ax.plot(x, daily["n"], color=C_NAVY, lw=1.3)
    ax.set_title("每日持仓品种数（上限 8）", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("个数", **fp())
    apply_font(ax)
    out["n"] = save_fig(fig, "n.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    s = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("ME").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in s.values]
    ax.bar(s.index, s.values / 1e4, width=20, color=colors, align="center")
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("月度净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    apply_font(ax)
    out["monthly"] = save_fig(fig, "monthly.png")

    fig, ax = plt.subplots(figsize=(8.6, 4.8), dpi=160)
    y = daily.set_index(pd.to_datetime(daily["return_date"]))["pnl_net"].resample("YE").sum()
    colors = [C_RED if v >= 0 else C_GREEN for v in y.values]
    labels = [str(i.year) for i in y.index]
    ax.bar(labels, y.values / 1e4, color=colors, width=0.55)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.set_title("分年净盈亏", fontsize=13, color=C_NAVY, **fp())
    ax.set_ylabel("万元", **fp())
    apply_font(ax)
    out["yearly"] = save_fig(fig, "yearly.png")

    fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
    roll = daily["ret"].rolling(60, min_periods=30)
    sharpe = roll.mean() / roll.std(ddof=1) * math.sqrt(252)
    ax.plot(x, sharpe, color=C_NAVY, lw=1.4)
    ax.axhline(0, color="#4A5568", lw=0.8)
    ax.axhline(1.0, color=C_GOLD, ls="--", lw=1)
    ax.set_title("滚动 60 日年化夏普", fontsize=13, color=C_NAVY, **fp())
    apply_font(ax)
    out["roll_sharpe"] = save_fig(fig, "roll_sharpe.png")

    fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
    ax.hist(daily["ret"] * 100, bins=36, color=C_NAVY, alpha=0.82, edgecolor="white")
    ax.axvline(0, color=C_ORANGE, lw=1.2)
    ax.set_title("日收益率分布", fontsize=13, color=C_NAVY, **fp())
    ax.set_xlabel("日收益（%）", **fp())
    apply_font(ax)
    out["hist"] = save_fig(fig, "hist.png")

    if not holds.empty:
        fig, ax = plt.subplots(figsize=(10.6, 4.6), dpi=160)
        for action, color in (("加码", C_RED), ("控拥挤", C_GOLD)):
            h = holds[holds["action"] == action]
            if h.empty:
                continue
            c = h.groupby("hold_date")["pnl"].sum()
            ax.plot(pd.to_datetime(c.index), c.cumsum() / 1e4, color=color, lw=1.8, label=f"{action} 累计毛盈亏")
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("持仓毛盈亏：加码 vs 控拥挤", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["by_action"] = save_fig(fig, "by_action.png")

        fig, ax = plt.subplots(figsize=(8.4, 4.2), dpi=160)
        by_act = holds.groupby("action")["pnl"].sum().reindex(["加码", "控拥挤"]).dropna()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_act.values]
        ax.bar(by_act.index, by_act.values / 1e4, color=colors, width=0.5)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("两种信号的累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["action_bar"] = save_fig(fig, "action_bar.png")

        fig, ax = plt.subplots(figsize=(8.4, 4.2), dpi=160)
        by_dir = holds.groupby("dir")["pnl"].sum()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_dir.values]
        ax.bar(by_dir.index, by_dir.values / 1e4, color=colors, width=0.45)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_title("多头 vs 空头累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["ls"] = save_fig(fig, "ls.png")

        fig, ax = plt.subplots(figsize=(10.4, 5.0), dpi=160)
        by_sec = holds.groupby("sector")["pnl"].sum().sort_values()
        colors = [C_RED if v >= 0 else C_GREEN for v in by_sec.values]
        ax.barh(by_sec.index, by_sec.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("板块累计毛盈亏", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["sector"] = save_fig(fig, "sector.png")

        fig, ax = plt.subplots(figsize=(10.6, 6.0), dpi=160)
        by_p = holds.groupby("name")["pnl"].sum().sort_values()
        top = pd.concat([by_p.head(10), by_p.tail(10)]).drop_duplicates()
        colors = [C_RED if v >= 0 else C_GREEN for v in top.values]
        ax.barh(top.index, top.values / 1e4, color=colors)
        ax.axvline(0, color="#4A5568", lw=0.8)
        ax.set_title("品种累计毛盈亏：最亏 / 最赚各 10 个", fontsize=13, color=C_NAVY, **fp())
        ax.set_xlabel("万元", **fp())
        apply_font(ax)
        out["product"] = save_fig(fig, "product.png")

        fig, ax = plt.subplots(figsize=(10.6, 4.8), dpi=160)
        mix = holds.groupby(["hold_date", "action"]).size().unstack(fill_value=0)
        mix = mix.reindex(columns=["加码", "控拥挤"], fill_value=0)
        ax.stackplot(pd.to_datetime(mix.index), mix["加码"], mix["控拥挤"],
                     labels=["加码", "控拥挤"], colors=[C_RED, C_GOLD], alpha=0.85)
        ax.set_title("每日持仓中两种信号的个数", fontsize=13, color=C_NAVY, **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["mix"] = save_fig(fig, "mix.png")

    core = ev[ev["action"].isin(CORE_ACTIONS)]
    if not core.empty and not ev_stats.empty:
        fig, ax = plt.subplots(figsize=(10.2, 5.0), dpi=160)
        horizons = [1, 5, 10, 20]
        x_pos = np.arange(len(horizons))
        width = 0.36
        for i, action in enumerate(("加码", "控拥挤")):
            row = ev_stats[ev_stats["action"] == action]
            if row.empty:
                continue
            vals = [row.iloc[0].get(f"mu{h}") or 0 for h in horizons]
            ax.bar(x_pos + (i - 0.5) * width, [v * 100 for v in vals], width=width,
                   color=C_RED if action == "加码" else C_GOLD, label=action)
        ax.axhline(0, color="#4A5568", lw=0.8)
        ax.set_xticks(x_pos)
        ax.set_xticklabels(["次日", "5日", "10日", "20日"])
        ax.set_ylabel("策略方向对齐后的平均收益（%）", **fp())
        ax.set_title("事件研究：按本策略方向对齐", fontsize=13, color=C_NAVY, **fp())
        ax.legend(frameon=False, fontsize=8)
        apply_font(ax)
        out["event"] = save_fig(fig, "event.png")

        fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=160)
        hits = []
        labels = []
        colors = []
        for action, color in (("加码", C_RED), ("控拥挤", C_GOLD)):
            row = ev_stats[ev_stats["action"] == action]
            if row.empty:
                continue
            for h, lab in ((1, "次日"), (5, "5日"), (20, "20日")):
                v = row.iloc[0].get(f"hit{h}")
                if v is None:
                    continue
                labels.append(f"{action}\n{lab}")
                hits.append(v * 100)
                colors.append(color)
        ax.bar(range(len(hits)), hits, color=colors, width=0.62)
        ax.axhline(50, color="#A0AEC0", ls="--", lw=1)
        ax.set_xticks(range(len(labels)))
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_ylabel("胜率（%）", **fp())
        ax.set_title("策略方向对齐胜率", fontsize=13, color=C_NAVY, **fp())
        apply_font(ax)
        out["hit"] = save_fig(fig, "hit.png")

    if not trades.empty:
        fig, ax = plt.subplots(figsize=(11.2, 4.0), dpi=160)
        t = trades.copy()
        t["dt"] = pd.to_datetime(t["trade_date"])
        turn = t.groupby("dt")["notional"].sum()
        ax.plot(turn.index, turn.rolling(10, min_periods=3).mean() / 1e4, color=C_NAVY, lw=1.5)
        ax.set_title("10 日平均成交名义（换手）", fontsize=13, color=C_NAVY, **fp())
        ax.set_ylabel("万元", **fp())
        apply_font(ax)
        out["turnover"] = save_fig(fig, "turnover.png")

    return out


def set_core_header(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "MOM 每日风控  ·  加码同向 + 控拥挤反向 深度回测", size=8, color=MUTED)
    fp_ = section.footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp_, "内部研究  ·  不做补风格  ·  模拟账户 2,000 万  ·  ", size=8, color=MUTED)
    run = fp_.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color=MUTED)


def build_report(ctx: dict) -> Path:
    configure_matplotlib()
    daily = ctx["acct"]["daily"]
    holds = ctx["acct"]["holds"]
    trades = ctx["acct"]["trades"]
    st = ctx["stats"]
    st_add = ctx["stats_add"]
    st_fade = ctx["stats_fade"]
    ch = ctx["charts"]
    ev_stats = ctx["ev_stats"]
    sig = ctx["sig"]
    start, end = ctx["start"], ctx["end"]
    episodes = ctx["episodes"]
    monthly, yearly = ctx["monthly"], ctx["yearly"]
    halves = ctx["halves"]
    streaks = ctx["streaks"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    set_core_header(doc)

    para(doc, "MOM 每日风控", size=12, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "只做加码同向 + 控拥挤反向", size=22, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=4)
    para(doc, "不做补风格  ·  2,000 万模拟账户深度回测", size=13, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=10)
    para(
        doc,
        f"样本 {start} 至 {end}　　交易日 {st.get('n')}　　"
        f"手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp（开平都收）　　"
        f"生成于 {date.today().isoformat()}",
        size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, space_after=16,
    )

    heading(doc, "一、策略定义", 1)
    para(
        doc,
        "本报告只检验一条可交易规则：每个交易日收盘后读取 MOM「量化 vs 主观」决策信号，"
        "下一交易日按信号调仓。加码与量化/主观存量共识同向开仓；控拥挤与共识反向开仓；"
        "观望、暂缓加码、减码准备一律空仓；补风格一律忽略。"
        "这是上一份总报告里的「核心口径」，这里把它单独做成账户深度回测，不再混进补风格。",
    )
    add_table(
        doc,
        ["信号", "本账户怎么做", "经济含义"],
        [
            ["加码", "与量化/主观共识同向", "两边已经站在同一边，账户复制这份 beta"],
            ["控拥挤", "与共识反向", "两边都已很重，账户做拥挤反向/对冲"],
            ["观望", "空仓；若昨有仓则下一交易日平掉", "两边方向冲突，不下注"],
            ["暂缓加码", "空仓并平掉旧仓", "存量仍同向，但边际不再同加"],
            ["减码准备", "空仓并平掉旧仓", "两边同时在减，不再跟"],
            ["补风格", "不做", "覆盖/引进投顾提示，不译成期货仓位"],
        ],
        col_widths=[3.0, 6.2, 7.2],
    )
    caption(doc, "表 1  下单规则。信号消失或改成空仓类动作时，下一交易日平仓。")

    para(
        doc,
        f"账户初始权益 {fmt_yuan(START_EQUITY)}。仓位按当日权益约 2.2 倍名义、最多 8 个品种等权分配，"
        f"单品种一日 1σ 亏损不超过权益的 1.2%，保证金占用不超过权益的 {MAX_MARGIN_UTIL:.0%}（券商系数 1.1）。"
        f"手数取整。费用按成交名义单边手续费 {COMM_RATE * 1e4:.1f}bp + 滑点 {SLIP_RATE * 1e4:.1f}bp，"
        "每手手续费不低于 3 元，开平都收。行情用主力连续合约，没有真实换月盈亏。",
    )

    heading(doc, "二、核心结论", 1)
    if st:
        para(
            doc,
            f"样本内期末权益 {fmt_yuan(st.get('end'))}，净盈亏 {fmt_yuan(st.get('pnl'), signed=True)}，"
            f"相对起始资金 {fmt_pct((st.get('end') - START_EQUITY) / START_EQUITY)}。"
            f"年化 {fmt_pct(st.get('cagr'))}，波动 {fmt_pct(st.get('vol'))}，夏普 {fmt_num(st.get('sharpe'))}，"
            f"卡玛 {fmt_num(st.get('calmar'))}，最大回撤 {fmt_pct(st.get('maxdd'))}"
            f"（约 {fmt_yuan(st.get('maxdd_yuan'))}），日胜率 {fmt_pct(st.get('hit'), already=False)}。"
            f"手续费 {fmt_yuan(st.get('total_comm'))}、滑点 {fmt_yuan(st.get('total_slip'))}，"
            f"费用合计 {fmt_yuan(st.get('total_cost'))}，毛盈亏 {fmt_yuan(st.get('total_gross'), signed=True)}。"
            f"日均持仓 {fmt_num(st.get('avg_names'), 1)} 个品种，日均保证金 {fmt_yuan(st.get('avg_margin'))}，"
            f"占用率 {fmt_pct(st.get('avg_util'))}，日均名义 {fmt_yuan(st.get('avg_notional'))}。",
        )
    add_table(
        doc,
        ["口径", "期末权益", "净盈亏", "年化", "波动", "夏普", "最大回撤", "费用"],
        [
            [
                "加码同向 + 控拥挤反向",
                fmt_yuan(st.get("end")),
                fmt_yuan(st.get("pnl"), signed=True),
                fmt_pct(st.get("cagr")),
                fmt_pct(st.get("vol")),
                fmt_num(st.get("sharpe")),
                fmt_pct(st.get("maxdd")),
                fmt_yuan(st.get("total_cost")),
            ],
            [
                "只做加码（对照）",
                fmt_yuan(st_add.get("end")),
                fmt_yuan(st_add.get("pnl"), signed=True),
                fmt_pct(st_add.get("cagr")),
                fmt_pct(st_add.get("vol")),
                fmt_num(st_add.get("sharpe")),
                fmt_pct(st_add.get("maxdd")),
                fmt_yuan(st_add.get("total_cost")),
            ],
            [
                "只做控拥挤反向（对照）",
                fmt_yuan(st_fade.get("end")),
                fmt_yuan(st_fade.get("pnl"), signed=True),
                fmt_pct(st_fade.get("cagr")),
                fmt_pct(st_fade.get("vol")),
                fmt_num(st_fade.get("sharpe")),
                fmt_pct(st_fade.get("maxdd")),
                fmt_yuan(st_fade.get("total_cost")),
            ],
        ],
        signed_cols={2, 3, 5, 6},
    )
    caption(doc, "表 2  主策略与两个单腿对照。三套账户共用同一套仓位、保证金和费用规则，只是允许的信号不同。")

    add_pnl = float((holds[holds["action"] == "加码"]["pnl"].sum()) if not holds.empty else 0)
    fade_pnl = float((holds[holds["action"] == "控拥挤"]["pnl"].sum()) if not holds.empty else 0)
    n_add = int((sig["action"] == "加码").sum())
    n_fade = int((sig["action"] == "控拥挤").sum())
    n_all = int(len(sig))
    para(
        doc,
        f"全样本产品级信号 {n_all:,} 条，其中加码 {n_add:,} 条（{n_add / n_all:.1%}），"
        f"控拥挤 {n_fade:,} 条（{n_fade / n_all:.1%}）。"
        f"账户持仓层累计毛盈亏：加码 {fmt_yuan(add_pnl, signed=True)}，"
        f"控拥挤 {fmt_yuan(fade_pnl, signed=True)}。"
        + (
            "加码是收益主力；控拥挤条数少，更多是间歇性的反向对冲，贡献可正可负，不能事先当成稳赚的保险。"
            if abs(add_pnl) >= abs(fade_pnl)
            else "控拥挤在本样本里贡献并不小，说明拥挤反向有时能单独赚钱，但不能外推成稳定 alpha。"
        )
        + " 只做加码的对照账户用来回答「控拥挤这一腿值不值得做」；只做控拥挤的对照用来看反向腿本身有没有边。",
    )

    heading(doc, "三、账户路径", 1)
    if "equity" in ch:
        add_picture(doc, ch["equity"])
        caption(doc, "图 1  三套口径的权益。蓝线是本策略；红线只跟加码；金线只做拥挤反向。")
    if "dd" in ch:
        add_picture(doc, ch["dd"])
        caption(doc, "图 2  本策略回撤。深度和持续时间比期末盈亏更能说明能不能拿得住。")
    if "cost" in ch:
        add_picture(doc, ch["cost"])
        caption(doc, "图 3  毛盈亏、累计费用与净盈亏。费用按成交名义逐笔计提。")
    if "margin" in ch:
        add_picture(doc, ch["margin"])
        caption(doc, "图 4  保证金与名义本金。名义随信号个数和权益复利变化。")
    if "n" in ch:
        add_picture(doc, ch["n"])
        caption(doc, "图 5  每日持仓品种数。经常低于 8，说明可交易的加码/控拥挤名单并不密。")
    if "monthly" in ch:
        add_picture(doc, ch["monthly"])
        caption(doc, "图 6  月度净盈亏（万元，红赚绿亏）。")
    if "yearly" in ch:
        add_picture(doc, ch["yearly"])
        caption(doc, "图 7  分年净盈亏。样本短的年份只作参考。")
    if "roll_sharpe" in ch:
        add_picture(doc, ch["roll_sharpe"])
        caption(doc, "图 8  滚动 60 日年化夏普。金虚线=1。长时间在零轴下说明策略有失效段。")
    if "hist" in ch:
        add_picture(doc, ch["hist"])
        caption(doc, "图 9  日收益直方图。左尾厚度决定回撤体验。")

    heading(doc, "四、加码 vs 控拥挤", 1)
    para(
        doc,
        "把持仓毛盈亏按信号动作拆开，能直接看到两只腿各自赚不赚钱。"
        "注意：这是持仓市值×当日涨跌的毛贡献，还没有把手续费摊回到每一条信号；"
        "费用在账户层已经扣过。对照账户（只开其中一腿）则把仓位预算全部给这一腿，数字会和「拆贡献」不同。",
    )
    if "by_action" in ch:
        add_picture(doc, ch["by_action"])
        caption(doc, "图 10  加码、控拥挤持仓毛盈亏的累计路径。")
    if "action_bar" in ch:
        add_picture(doc, ch["action_bar"])
        caption(doc, "图 11  两种信号累计毛盈亏柱状图。")
    if "mix" in ch:
        add_picture(doc, ch["mix"])
        caption(doc, "图 12  每日持仓里加码与控拥挤的个数。多数日子是加码占主导。")

    if not holds.empty:
        rows = []
        for action in ("加码", "控拥挤"):
            h = holds[holds["action"] == action]
            if h.empty:
                continue
            s = h.groupby("hold_date")["pnl"].sum()
            rows.append([
                action,
                f"{len(h):,}",
                f"{h['hold_date'].nunique()}",
                f"{h['product'].nunique()}",
                fmt_yuan(h["pnl"].sum(), signed=True),
                fmt_yuan(s.mean(), signed=True),
                fmt_pct(float((s > 0).mean()), already=False),
                fmt_t(tstat(s)),
            ])
        add_table(doc, ["信号", "持仓条数", "有仓天数", "品种数", "累计毛盈亏", "日均毛盈亏", "日胜率", "t"], rows, signed_cols={4, 5, 6})
        caption(doc, "表 3  持仓层按信号处理。日胜率按「当天该信号持仓合计是否赚钱」计。")

    heading(doc, "五、事件研究（按本策略方向对齐）", 1)
    para(
        doc,
        "事件研究不问账户仓位，只问：信号出现后，品种收益按本策略方向对齐，均值是否显著不为零。"
        "加码对齐共识方向；控拥挤对齐反向。这与总报告里「控拥挤仍按共识方向签名」不同——"
        "这里检验的是「如果按本账户规则去交易，品种本身有没有边」。",
    )
    if not ev_stats.empty:
        ev_rows = []
        for r in ev_stats.itertuples(index=False):
            ev_rows.append([
                r.action, f"{int(r.n):,}",
                fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1),
                fmt_pct(r.mu5), fmt_t(r.t5),
                fmt_pct(r.mu10), fmt_t(r.t10),
                fmt_pct(r.mu20), fmt_t(r.t20),
            ])
        add_table(
            doc,
            ["信号", "样本", "次日均值", "次日胜率", "次日t", "5日均值", "5日t", "10日均值", "10日t", "20日均值", "20日t"],
            ev_rows,
            signed_cols={2, 3, 5, 7, 9},
        )
        caption(doc, "表 4  策略方向对齐的事件研究。t 的绝对值小于 2 时，不宜把均值当成稳定边。")
    if "event" in ch:
        add_picture(doc, ch["event"])
        caption(doc, "图 13  次日 / 5 / 10 / 20 日对齐收益。")
    if "hit" in ch:
        add_picture(doc, ch["hit"])
        caption(doc, "图 14  对齐胜率，虚线 50%。")

    add_row = ev_stats[ev_stats["action"] == "加码"] if not ev_stats.empty else pd.DataFrame()
    fade_row = ev_stats[ev_stats["action"] == "控拥挤"] if not ev_stats.empty else pd.DataFrame()
    bits = []
    if not add_row.empty:
        bits.append(
            f"加码 {int(add_row.iloc[0]['n']):,} 条，次日对齐均值 {fmt_pct(add_row.iloc[0]['mu1'])}，"
            f"胜率 {fmt_pct(add_row.iloc[0]['hit1'], already=False)}，t={fmt_t(add_row.iloc[0]['t1'])}。"
        )
    if not fade_row.empty:
        bits.append(
            f"控拥挤 {int(fade_row.iloc[0]['n']):,} 条，次日对齐均值 {fmt_pct(fade_row.iloc[0]['mu1'])}，"
            f"胜率 {fmt_pct(fade_row.iloc[0]['hit1'], already=False)}，t={fmt_t(fade_row.iloc[0]['t1'])}。"
        )
    if bits:
        para(doc, " ".join(bits) + " 持有期拉长后若均值衰减或 t 变弱，说明边更像短持有的拥挤/共识脉冲，而不是可以抱很久的趋势。")

    ev = ctx["ev"]
    add_ev = ev[ev["action"] == "加码"].copy()
    if not add_ev.empty:
        add_ev["strength"] = add_ev["q_pct"].abs() + add_ev["s_pct"].abs()
        b = bucket_stats(add_ev, "加码", "strength", [0, 10, 16, 24, 80], ["<10", "10–16", "16–24", "≥24"])
        if not b.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)] for r in b.itertuples(index=False)]
            add_table(doc, ["加码 |q%|+|s%|", "样本", "次日对齐", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 5  加码按两边风险占比之和分层。更「重」的共识是否更干净。")
    fade_ev = ev[ev["action"] == "控拥挤"].copy()
    if not fade_ev.empty:
        fade_ev["crowd"] = fade_ev["q_pct"].abs() + fade_ev["s_pct"].abs()
        b = bucket_stats(fade_ev, "控拥挤", "crowd", [24, 30, 40, 80], ["25–30", "30–40", "≥40"])
        if not b.empty:
            rows = [[r.bucket, f"{int(r.n)}", fmt_pct(r.mu1), fmt_pct(r.hit1, already=False), fmt_t(r.t1)] for r in b.itertuples(index=False)]
            add_table(doc, ["控拥挤 |q%|+|s%|", "样本", "次日对齐（反向）", "胜率", "t"], rows, signed_cols={2, 3})
            caption(doc, "表 6  控拥挤按拥挤程度分层。若更拥挤反而对齐收益更差，反向腿就要更小心。")

    heading(doc, "六、多空、板块与品种", 1)
    if "ls" in ch:
        add_picture(doc, ch["ls"])
        caption(doc, "图 15  多头与空头累计毛盈亏。一边独大时，策略更像单边 beta 而不是市场中性。")
    if "sector" in ch:
        add_picture(doc, ch["sector"])
        caption(doc, "图 16  板块累计毛盈亏。")
    if "product" in ch:
        add_picture(doc, ch["product"])
        caption(doc, "图 17  品种两端。用来看集中度，不是推荐交易名单。")

    if not holds.empty:
        by_dir = holds.groupby("dir").agg(days=("hold_date", "nunique"), pnl=("pnl", "sum"), n=("product", "size"))
        rows = []
        for d, r in by_dir.iterrows():
            rows.append([d, f"{int(r['n']):,}", f"{int(r['days'])}", fmt_yuan(r["pnl"], signed=True)])
        add_table(doc, ["方向", "持仓条数", "有仓天数", "累计毛盈亏"], rows, signed_cols={3})
        caption(doc, "表 7  多空拆解。")

        sec = holds.groupby("sector").agg(
            days=("hold_date", "nunique"),
            names=("product", "nunique"),
            pnl=("pnl", "sum"),
            hit=("pnl", lambda s: float((s > 0).mean())),
        ).sort_values("pnl", ascending=False)
        rows = [[i, f"{int(r.names)}", f"{int(r.days)}", fmt_yuan(r.pnl, signed=True), fmt_pct(r.hit, already=False)] for i, r in sec.iterrows()]
        add_table(doc, ["板块", "品种数", "有仓天数", "累计毛盈亏", "条数胜率"], rows, signed_cols={3, 4})
        caption(doc, "表 8  板块归因。条数胜率按持仓行计，不是按日。")

        prod = holds.groupby(["name", "product", "sector"]).agg(
            days=("hold_date", "nunique"),
            long_days=("dir", lambda s: int((s == "多").sum())),
            short_days=("dir", lambda s: int((s == "空").sum())),
            lots=("lots", lambda s: s.abs().mean()),
            notional=("notional", "mean"),
            pnl=("pnl", "sum"),
        ).sort_values("pnl", ascending=False)
        top_bottom = pd.concat([prod.head(12), prod.tail(8)]).drop_duplicates()
        rows = []
        for (nme, p, sec_), r in top_bottom.iterrows():
            rows.append([
                f"{nme}({p})", sec_, f"{int(r['days'])}", f"{int(r['long_days'])}/{int(r['short_days'])}",
                f"{r['lots']:.1f}", fmt_yuan(r["notional"]), fmt_yuan(r["pnl"], signed=True),
            ])
        add_table(doc, ["品种", "板块", "持仓天", "多/空条数", "平均|手数|", "平均名义", "累计毛盈亏"], rows, signed_cols={6})
        caption(doc, "表 9  品种层头部 12 + 尾部 8。")

        if len(prod) >= 5:
            total = float(prod["pnl"].sum())
            pos = prod[prod["pnl"] > 0]["pnl"]
            conc = float(pos.head(5).sum() / pos.sum()) if pos.sum() > 0 else None
            para(
                doc,
                f"正贡献品种里，前 5 名约占全部正贡献的 {fmt_pct(conc, already=False)}。"
                "若集中度很高，样本内成绩更容易被少数品种的行情路径带走，外推时要打折。",
            )

    heading(doc, "七、分年、分月与回撤段", 1)
    if not yearly.empty:
        rows = []
        for r in yearly.itertuples(index=False):
            rows.append([
                str(r.year), fmt_yuan(r.pnl, signed=True), fmt_yuan(r.cost),
                fmt_pct(r.cagr), fmt_pct(r.vol), fmt_num(r.sharpe), fmt_pct(r.maxdd),
                fmt_pct(r.hit, already=False), str(int(r.days)),
            ])
        add_table(doc, ["年", "净盈亏", "费用", "年化", "波动", "夏普", "最大回撤", "日胜率", "天数"], rows, signed_cols={1, 3, 5, 6})
        caption(doc, "表 10  分年账户绩效。年内年化按该年交易日年化，短年不要和全年比。")

    if not monthly.empty:
        rows = []
        for ym, r in monthly.iterrows():
            rows.append([
                ym, fmt_yuan(r["pnl"], signed=True), fmt_yuan(r["gross"], signed=True),
                fmt_yuan(r["cost"]), fmt_num(r["n"], 1), fmt_pct(r["hit"], already=False),
                fmt_yuan(r["eq"]),
            ])
        add_table(doc, ["月份", "净盈亏", "毛盈亏", "费用", "日均品种", "日胜率", "月末权益"], rows, signed_cols={1, 2})
        caption(doc, "表 11  分月账户结果。")
        win_m = float((monthly["pnl"] > 0).mean())
        para(doc, f"有数据的月份共 {len(monthly)} 个，月度赚钱比例 {fmt_pct(win_m, already=False)}。连续亏月比单月亏损更能说明策略阶段性失效。")

    if not episodes.empty:
        heading(doc, "7.1 深度回撤段", 2)
        rows = []
        for r in episodes.head(8).itertuples(index=False):
            rows.append([
                r.start, r.trough, r.end, str(int(r.days)), str(int(r.trough_days)),
                fmt_pct(r.depth), fmt_yuan(r.lost, signed=True),
                "未修复" if getattr(r, "open", False) else "已修复",
            ])
        add_table(doc, ["开始", "谷底", "结束", "持续天", "到谷底", "深度", "权益损失", "状态"], rows, signed_cols={5, 6})
        caption(doc, "表 12  深度至少 3% 的回撤段，按深度排序。未修复表示样本结束时仍在水下。")
        worst = episodes.iloc[0]
        para(
            doc,
            f"最深一段从 {worst['start']} 到谷底 {worst['trough']}，深度 {fmt_pct(worst['depth'])}，"
            f"权益最多少了 {fmt_yuan(worst['lost'])}，持续 {int(worst['days'])} 个交易日。"
            "若这段正好叠在某个板块单边或信号密度骤变上，账户体验会明显差于年化数字。",
        )

    heading(doc, "八、成交、换手与费用", 1)
    if "turnover" in ch:
        add_picture(doc, ch["turnover"])
        caption(doc, "图 18  10 日平均成交名义。名单日度变化大时，费用会明显吃收益。")
    if not trades.empty:
        para(
            doc,
            f"全样本成交 {len(trades):,} 笔，成交名义合计 {fmt_yuan(trades['notional'].sum())}，"
            f"费用合计 {fmt_yuan(trades['cost'].sum())}，平均每笔 {fmt_yuan(trades['cost'].mean())}。"
            f"开/加 {(trades['side'] == '开/加').sum():,} 笔，平仓 {(trades['side'] == '平仓').sum():,} 笔，"
            f"反手/平 {(trades['side'] == '反手/平').sum():,} 笔。",
        )
        t20 = trades.sort_values("notional", ascending=False).head(15)
        trows = []
        for r in t20.itertuples(index=False):
            trows.append([
                r.trade_date, f"{r.name}({r.product})", r.action, r.side,
                f"{int(r.old_lots)}→{int(r.new_lots)}", fmt_yuan(r.notional), fmt_yuan(r.cost),
            ])
        add_table(doc, ["成交日", "品种", "信号", "动作", "手数", "成交名义", "费用"], trows)
        caption(doc, "表 13  名义最大的 15 笔调仓。")

        add_table(
            doc,
            ["费用项", "金额", "说明"],
            [
                ["手续费", fmt_yuan(st.get("total_comm")), f"成交名义 × {COMM_RATE * 1e4:.1f}bp，每手不低于 3 元，开平都收"],
                ["滑点", fmt_yuan(st.get("total_slip")), f"成交名义 × {SLIP_RATE * 1e4:.1f}bp"],
                ["费用合计", fmt_yuan(st.get("total_cost")), "上面两项之和"],
                ["毛盈亏", fmt_yuan(st.get("total_gross"), signed=True), "持仓名义 × 次日涨跌"],
                ["净盈亏", fmt_yuan(st.get("pnl"), signed=True), "毛盈亏 − 费用"],
                ["费用/|毛盈亏|", fmt_pct(abs(st.get("total_cost") or 0) / abs(st.get("total_gross") or 1), already=False), "摩擦侵蚀比例"],
            ],
            signed_cols={1},
        )
        caption(doc, "表 14  费用拆解。主力连续没有换月移仓，换月成本未另计。")

        if st.get("total_cost") and not daily.empty:
            rows = []
            for k, label in ((0.5, "半价摩擦"), (1.0, "基准（本报告）"), (2.0, "双倍摩擦")):
                extra = (k - 1.0) * daily["cost"]
                adj = daily["pnl_net"] - extra
                eq = START_EQUITY + adj.cumsum()
                tmp = daily.copy()
                tmp["pnl_net"] = adj
                tmp["equity"] = eq
                tmp["ret"] = adj / (eq - adj).replace(0, np.nan)
                tmp["dd"] = eq / eq.cummax() - 1.0
                st_k = account_stats(tmp)
                rows.append([
                    label, fmt_yuan(st_k.get("end")), fmt_yuan(st_k.get("pnl"), signed=True),
                    fmt_pct(st_k.get("cagr")), fmt_num(st_k.get("sharpe")), fmt_pct(st_k.get("maxdd")),
                ])
            add_table(doc, ["摩擦假设", "期末权益", "净盈亏", "年化", "夏普", "最大回撤"], rows, signed_cols={2, 3, 4, 5})
            caption(doc, "表 15  把已发生费用按比例缩放后的近似敏感性（仓位不重算）。若双倍摩擦后夏普仍明显为正，结论对费率更稳健。")

    heading(doc, "九、稳健性", 1)
    if halves:
        rows = []
        for h in halves:
            rows.append([
                h["name"], f"{h['start']} ~ {h['end']}", fmt_yuan(h.get("pnl"), signed=True),
                fmt_pct(h.get("cagr")), fmt_pct(h.get("vol")), fmt_num(h.get("sharpe")),
                fmt_pct(h.get("maxdd")), fmt_pct(h.get("hit"), already=False),
            ])
        add_table(doc, ["分段", "区间", "净盈亏", "年化", "波动", "夏普", "最大回撤", "日胜率"], rows, signed_cols={2, 3, 5, 6})
        caption(doc, "表 16  前后半样本。若半段反号，全样本结论不能外推。")

    if streaks:
        para(
            doc,
            f"赚钱日 {streaks['win_days']}、亏钱日 {streaks['loss_days']}；"
            f"最长连赢 {streaks['max_win_streak']} 日，最长连亏 {streaks['max_loss_streak']} 日。"
            "连亏长度决定会不会在谷底砍掉规则。",
        )

    if not daily.empty:
        worst = daily.nsmallest(8, "pnl_net")
        best = daily.nlargest(8, "pnl_net")
        heading(doc, "9.1 最赚与最亏的交易日", 2)
        rows = []
        for label, df in (("最亏", worst), ("最赚", best)):
            for r in df.itertuples(index=False):
                rows.append([
                    label, r.return_date, fmt_yuan(r.pnl_net, signed=True), fmt_yuan(r.pnl_gross, signed=True),
                    fmt_yuan(r.cost), f"{int(r.n)}", fmt_pct(r.ret),
                ])
        add_table(doc, ["类型", "日期", "净盈亏", "毛盈亏", "费用", "持仓数", "日收益"], rows, signed_cols={2, 3, 6})
        caption(doc, "表 17  单日两端。若最亏日集中在同一板块或同一类信号，说明尾部风险并不分散。")

    heading(doc, "十、持仓快照", 1)
    para(
        doc,
        "持仓在信号日收盘后生成，下一交易日生效。方向已按规则处理：加码与共识同向，控拥挤与共识反向。",
    )

    def hold_table(df: pd.DataFrame, title: str):
        if df is None or df.empty:
            para(doc, f"{title}：当日无持仓。", first_line=False)
            return
        rows = []
        for r in df.sort_values("notional", ascending=False).itertuples(index=False):
            rows.append([
                f"{r.name}({r.product})", r.sector, r.action, r.dir, str(int(r.lots)),
                f"{r.price:,.2f}" if pd.notna(r.price) else "—",
                fmt_yuan(r.notional), fmt_yuan(r.margin), fmt_yuan(r.pnl, signed=True),
            ])
        rows.append([
            "合计", "", "", "", str(int(df["lots"].abs().sum())), "",
            fmt_yuan(df["notional"].sum()), fmt_yuan(df["margin"].sum()), fmt_yuan(df["pnl"].sum(), signed=True),
        ])
        add_table(doc, ["品种", "板块", "信号", "方向", "手数", "价格", "名义", "保证金", "当日盈亏"], rows, signed_cols={8})
        caption(doc, title)

    if not holds.empty:
        last_d = holds["hold_date"].max()
        hold_table(holds[holds["hold_date"] == last_d], f"表 18  样本末日 {last_d} 持仓。")
        snap = holds[holds["hold_date"] == "2026-09-01"]
        if snap.empty:
            snap = holds[holds["signal_date"] == "2026-09-01"]
        if not snap.empty:
            hold_table(snap, "表 19  2026-09-01 持仓（与风控页同一天）。")

    heading(doc, "十一、结论与局限", 1)
    para(
        doc,
        "这条规则的优点是干净：只在「两边同向加码」或「两边已经过重」时下单，不把覆盖提示（补风格）译成期货，"
        "也不在观望/暂缓/减码日强行找方向。账户数字回答的是「若用 2,000 万按规则跟信号」，"
        "不是单位净值 overlay，也不是对未来收益的承诺。",
    )
    para(
        doc,
        "主要局限。第一，行情是主力连续价，没有指定月换月、展期和移仓价差。"
        "第二，保证金率是品种近似值，不是当日交易所公布值；没有涨跌停、强平、资金划转。"
        "第三，信号来自 MOM 存量仓位，样本期的品种结构、量化/主观账户集合一旦变化，边可能消失。"
        "第四，仓位规则（2.2 倍名义、最多 8 个、1.2% 风险帽）是研究设定，换成别的杠杆，回撤和年化会一起变。"
        "第五，控拥挤样本显著少于加码，反向腿的统计更不稳。",
    )
    para(
        doc,
        "使用建议。把加码当成主腿、控拥挤当成有条件的拥挤反向，而不是六种信号捆成一篮子。"
        "补风格继续留在页面上做覆盖提示，不要进这个账户。"
        "若后半样本或双倍摩擦后夏普明显塌掉，就把它当成信号设计的检验，而不是一组可以放大的期货策略。",
    )
    para(
        doc,
        f"图表原文件在 mom_signal_strategy/report_output/jiaama_crowd/charts/。报告生成于 {date.today().isoformat()}。",
        size=10, color=MUTED, first_line=False,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(REPORT_PATH))
    try:
        doc.save(str(REPORT_PATH_ASCII))
    except Exception:
        pass
    return REPORT_PATH


def main():
    load_env()
    configure_matplotlib()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_conn()
    try:
        pos, px = load_data(conn)
    finally:
        conn.close()
    if pos.empty:
        raise SystemExit("mom_position_details is empty")
    if px.empty:
        raise SystemExit("raw_akshare_futures_daily is empty")

    wide, clean, close, _ = build_returns(px)
    sig = build_signals(pos, clean)
    if sig.empty:
        raise SystemExit("no signals reconstructed")

    print("Event study…")
    ev = add_strategy_alignment(event_study(sig, wide))
    ev_stats = action_event_stats(ev)

    print("20M account: 加码 + 控拥挤…")
    acct = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions=CORE_ACTIONS)
    print("20M account: 只做加码…")
    acct_add = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"加码"})
    print("20M account: 只做控拥挤…")
    acct_fade = run_account(sig, close, wide, clean, include_bufengge=False, allowed_actions={"控拥挤"})

    stats = account_stats(acct["daily"])
    stats_add = account_stats(acct_add["daily"])
    stats_fade = account_stats(acct_fade["daily"])

    print("Drawing charts…")
    charts = draw_charts(acct, acct_add, acct_fade, ev, ev_stats)

    if not acct["daily"].empty:
        acct["daily"].to_csv(OUT_DIR / "acct_daily.csv", index=False, encoding="utf-8-sig")
    if not acct["holds"].empty:
        acct["holds"].to_csv(OUT_DIR / "acct_holdings.csv", index=False, encoding="utf-8-sig")
    if not acct["trades"].empty:
        acct["trades"].to_csv(OUT_DIR / "acct_trades.csv", index=False, encoding="utf-8-sig")
    ev_stats.to_csv(OUT_DIR / "event_stats.csv", index=False, encoding="utf-8-sig")

    monthly, yearly = year_month_tables(acct["daily"]) if not acct["daily"].empty else (pd.DataFrame(), pd.DataFrame())
    ctx = {
        "sig": sig,
        "ev": ev,
        "ev_stats": ev_stats,
        "acct": acct,
        "acct_add": acct_add,
        "acct_fade": acct_fade,
        "stats": stats,
        "stats_add": stats_add,
        "stats_fade": stats_fade,
        "charts": charts,
        "start": str(sig["date"].min()),
        "end": str(sig["date"].max()),
        "episodes": drawdown_episodes(acct["daily"]),
        "monthly": monthly,
        "yearly": yearly,
        "halves": half_sample(acct["daily"]),
        "streaks": streak_stats(acct["daily"]["pnl_net"]) if not acct["daily"].empty else {},
    }
    print("Writing Word report…")
    path = build_report(ctx)
    print(f"Wrote {path}")
    print(f"Also {REPORT_PATH_ASCII}")
    if stats:
        print(f"End equity {stats.get('end'):,.0f}  Sharpe {stats.get('sharpe')}  MaxDD {stats.get('maxdd')}")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        raise
